"""Turning comments and markup: what the returned file did, as evidence.

"Turn the comments" and "turn their markup" hand the model a document that
other people have been working in, and ask it to act on their work: make the
change a comment asks for, reply in the comment's thread, resolve it or leave
it open with a question; accept some of the other side's tracked changes,
reject others, and write our counter-proposals as tracked changes of our own.

The model decides all of that (DECISIONS 30). This module is the evidence on
the other side of the decision: given the file we were sent and the file that
came back, what happened to every comment and every tracked change, and
whether anything happened that nobody decided. It runs twice on the same
pair, in the sandbox before the file is written to the outputs (the skill's
finish.py) and on the host before it is attached, and the report is what the
model is shown when something is wrong.

What it holds the file to:

- **No comment is deleted or rewritten.** Every comment in the original is
  in the result with the same author and words. New comments are ours.
- **Other people's tracked changes are untouched** unless accepted or
  rejected by id, and the ones that were are gone.
- **The decisions are what they say.** The original with exactly those
  accepts and rejects applied, and the result with our own new tracked
  changes rejected, read the same, paragraph for paragraph. So nothing
  changed except by a decision on a named change or by a tracked change of
  ours, which the lawyer can see and reject.
- **The file is sound**: redline.verify, the same gate as every attachment.
"""

from __future__ import annotations

import difflib
import re
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from secondeye.pipeline import ooxml, redline, wordcomments


@dataclass
class CommentOutcome:
    id: str
    author: str
    clause: str
    anchored: str
    text: str
    # resolved: done in Word. open: replied to and left open (a question).
    # untouched: neither.
    status: str
    reply: str = ""

    def as_dict(self) -> dict:
        return dict(self.__dict__)


@dataclass
class Examination:
    verified: bool
    verification: str
    problems: list[str] = field(default_factory=list)
    comments: list[CommentOutcome] = field(default_factory=list)
    accepted: list[dict] = field(default_factory=list)
    rejected: list[dict] = field(default_factory=list)
    still_tracked: list[dict] = field(default_factory=list)
    ours: list[dict] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.verified and not self.problems

    def report(self) -> str:
        """The verification report, in full, in words the model can act on."""
        state = "passed" if self.verified else "FAILED"
        lines = [f"File check: {state} ({self.verification})."]
        if self.problems:
            lines.append("Problems with the returned file:")
            lines += [f"- {p}" for p in self.problems]
        else:
            lines.append("No comment was deleted or changed, and nothing changed except "
                         "by the decisions and tracked changes listed.")
        done = [c for c in self.comments if c.status == "resolved"]
        open_ = [c for c in self.comments if c.status == "open"]
        untouched = [c for c in self.comments if c.status == "untouched"]
        lines.append(f"Comments: {len(done)} resolved, {len(open_)} open with a reply, "
                     f"{len(untouched)} untouched.")
        if self.accepted or self.rejected or self.still_tracked:
            lines.append(f"Their tracked changes: {len(self.accepted)} accepted, "
                         f"{len(self.rejected)} rejected, {len(self.still_tracked)} "
                         "left tracked.")
        lines.append(f"Tracked changes of ours: {len(self.ours)}.")
        return "\n".join(lines)

    def as_dict(self) -> dict:
        return {
            "ok": self.ok, "verified": self.verified, "verification": self.verification,
            "problems": self.problems,
            "comments": [c.as_dict() for c in self.comments],
            "accepted": self.accepted, "rejected": self.rejected,
            "still_tracked": self.still_tracked, "ours": self.ours,
        }


def _short(text: str, limit: int = 60) -> str:
    text = " ".join((text or "").split())
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _changes(content: bytes) -> list[ooxml.TrackedChange]:
    return ooxml.tracked_changes(Document(BytesIO(content)))


def _describe(change: ooxml.TrackedChange) -> dict:
    return {"id": change.id, "kind": change.kind, "author": change.author,
            "text": change.text, "part": change.part}


def examine(original: bytes, turned: bytes, author: str,
            accepted: list[str] | tuple = (), rejected: list[str] | tuple = ()
            ) -> Examination:
    """Everything the returned file did to the one we were sent."""
    verified, why = redline.verify(turned)
    exam = Examination(verified=verified, verification=why)
    if not verified:
        exam.problems.append(f"the file does not pass the Word checks: {why}")
        return exam

    _examine_comments(original, turned, author, exam)
    _examine_changes(original, turned, author, list(accepted), list(rejected), exam)
    return exam


def _examine_comments(original: bytes, turned: bytes, author: str, exam: Examination) -> None:
    before = {c.id: c for c in wordcomments.read(original)}
    after = {c.id: c for c in wordcomments.read(turned)}

    for cid, c in before.items():
        now = after.get(cid)
        where = f" on cl. {c.clause}" if c.clause else ""
        if now is None:
            exam.problems.append(f'comment {cid} by {c.author}{where} ("{_short(c.text)}") '
                                 "was deleted; comments are never deleted")
        elif now.author != c.author or now.text != c.text:
            exam.problems.append(f"comment {cid} by {c.author}{where} was changed; other "
                                 "people's comments are never edited")
    for cid, c in after.items():
        if cid not in before and c.author != author:
            exam.problems.append(f'comment {cid} ("{_short(c.text)}") is signed "{c.author}", '
                                 f'not "{author}"')

    for cid, c in before.items():
        if c.parent or cid not in after or c.author == author or c.done:
            # A reply already in the file is part of its thread; a comment
            # of our own (an earlier review's question) or one already
            # resolved is not one to turn.
            continue
        now = after[cid]
        ours = [after[r] for r in now.replies if r not in before and after[r].author == author]
        status = "resolved" if now.done else ("open" if ours else "untouched")
        exam.comments.append(CommentOutcome(
            id=cid, author=c.author, clause=c.clause, anchored=c.anchored, text=c.text,
            status=status, reply=ours[-1].text if ours else ""))


def _examine_changes(original: bytes, turned: bytes, author: str, accepted: list[str],
                     rejected: list[str], exam: Examination) -> None:
    before = _changes(original)
    after = _changes(turned)
    after_keys = {(c.part, c.id): c for c in after}
    before_keys = {(c.part, c.id) for c in before}

    def named(change, ids) -> bool:
        return change.id in ids or f"{change.part}:{change.id}" in ids or any(
            e.get(qn("w:id")) in ids for e in change.elements)

    for change in before:
        now = after_keys.get((change.part, change.id))
        label = f'{change.author}\'s change {change.id} ({change.kind} "{_short(change.text)}")'
        if now is not None:
            if (now.kind, now.author, now.text) != (change.kind, change.author, change.text):
                exam.problems.append(f"{label} was altered; decide it by id or leave it alone")
            elif named(change, accepted) or named(change, rejected):
                exam.problems.append(f"{label} is recorded as decided but is still tracked")
            else:
                exam.still_tracked.append(_describe(change))
            continue
        if named(change, accepted):
            exam.accepted.append(_describe(change))
        elif named(change, rejected):
            exam.rejected.append(_describe(change))
        else:
            exam.problems.append(f"{label} is gone, but was neither accepted nor rejected "
                                 "with accept_reject.py")

    new_ids: set[str] = set()
    for change in after:
        if (change.part, change.id) in before_keys:
            continue
        new_ids.update(e.get(qn("w:id")) for e in change.elements)
        if change.author != author:
            exam.problems.append(f'a new tracked change ({change.kind} "{_short(change.text)}") '
                                 f'is signed "{change.author}", not "{author}"')
        else:
            exam.ours.append(_describe(change))

    if exam.problems:
        return
    # The proof: the decisions replayed on the original, against the result
    # with our own new changes taken back out.
    try:
        replay = Document(BytesIO(original))
        ooxml.decide(replay, accept=accepted, reject=rejected)
        buffer = BytesIO()
        replay.save(buffer)
    except (ooxml.RevisionNotFound, ooxml.CannotDecide) as e:
        exam.problems.append(f"the recorded decisions cannot be replayed: {e}")
        return
    expected = _reading(buffer.getvalue(), set())
    actual = _reading(turned, new_ids)
    if expected != actual:
        exam.problems.append("the text differs from the original with only the recorded "
                             "decisions and your tracked changes applied: "
                             + _first_difference(expected, actual))


def _reading(content: bytes, hide: set[str]) -> list[str]:
    """Each paragraph as it reads with every tracked change shown as made,
    except those with an id in `hide`, which are shown as rejected."""
    document = Document(BytesIO(content))
    out: list[str] = []
    for story in ooxml._story_elements(document):
        joined = ""
        for p in story.iter(qn("w:p")):
            if _ancestor(p, qn("w:p")) is not None:
                continue            # a text box's paragraphs are read in their own right
            text = "".join(_visible(p, hide))
            mark = p.find(f"{qn('w:pPr')}/{qn('w:rPr')}")
            merged = False
            if mark is not None:
                ins, dele = mark.find(qn("w:ins")), mark.find(qn("w:del"))
                if ins is not None and ins.get(qn("w:id")) in hide:
                    merged = True
                if dele is not None and dele.get(qn("w:id")) not in hide:
                    merged = True
            joined += text
            if not merged:
                out.append(" ".join(joined.split()))
                joined = ""
        if joined:
            out.append(" ".join(joined.split()))
    return [line for line in out if line]


def _visible(node, hide: set[str]):
    for child in node:
        if not isinstance(child.tag, str):
            continue
        tag = child.tag
        if tag in (qn("w:pPr"), qn("w:rPr")):
            continue
        if tag in (qn("w:ins"), qn("w:moveTo")):
            if child.get(qn("w:id")) in hide:
                continue
            yield from _visible(child, hide)
        elif tag in (qn("w:del"), qn("w:moveFrom")):
            if child.get(qn("w:id")) in hide:
                yield from _deleted_text(child)
        elif tag == qn("w:t"):
            yield child.text or ""
        elif tag == qn("w:tab"):
            yield " "
        elif tag in (qn("w:p"), qn("w:txbxContent")):
            continue
        else:
            yield from _visible(child, hide)


def _deleted_text(node):
    for child in node.iter(qn("w:delText"), qn("w:t")):
        yield child.text or ""


def _ancestor(node, tag):
    parent = node.getparent()
    while parent is not None:
        if parent.tag == tag:
            return parent
        parent = parent.getparent()
    return None


def _first_difference(expected: list[str], actual: list[str]) -> str:
    for line in difflib.unified_diff(expected, actual, lineterm="", n=0):
        if line.startswith(("---", "+++", "@@")):
            continue
        sign = "expected" if line.startswith("-") else "found"
        return f'{sign} "{_short(line[1:], 120)}"'
    return "the paragraphs are in a different order"


# --------------------------------------------------------------------------
# The blackline: our turn against their version
# --------------------------------------------------------------------------


def blackline(theirs: bytes, ours: bytes, author: str) -> tuple[bytes | None, str]:
    """Their version (their markup all accepted) compared with ours (all of
    it accepted, their decided changes and our counter-proposals), as
    tracked changes on theirs. Returns (file or None, a note)."""
    from secondeye.pipeline import clean, compare

    try:
        their_clean = clean.clean(theirs).content
        our_clean = clean.clean(ours).content
    except clean.CannotClean as e:
        return None, str(e)
    result = compare.compare(their_clean, our_clean, author=author)
    if result.identical or not result.content:
        return None, "our turn reads the same as their version"
    ok, why = redline.verify(result.content, expect_revisions=True)
    if not ok:
        return None, f"the blackline did not verify ({why})"
    note = "" if result.proven else "the comparison could not prove every change"
    return result.content, note


def stem(filename: str) -> str:
    base = re.sub(r"\.docx?$", "", filename or "document", flags=re.IGNORECASE)
    return re.sub(r"\s*\((?:redline|comments turned|our turn|clean)\)\s*$", "", base) or "document"


__all__ = ["CommentOutcome", "Examination", "blackline", "examine", "stem"]
