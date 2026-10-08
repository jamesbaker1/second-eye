"""Document -> reviewable text, with stable anchors back into the source.

The review model never sees raw OOXML. It sees numbered paragraphs, so a finding
can name the paragraph it belongs to and the redliner can find it again.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from io import BytesIO

from docx import Document

from secondeye.models import Attachment
from secondeye.pipeline.filetype import (
    SANDBOX_CONVERTIBLE,
    Kind,
    identify,
    message,
)


@dataclass
class Block:
    index: int
    text: str
    style: str
    kind: str  # "paragraph" | "heading" | "table-cell" | "list-item"
    # The number Word draws in front of the paragraph from its list definition:
    # "2.2", "(a)", "Schedule 3". It is not in `text`, because it is not in the
    # runs and an anchor that included it could never be found again; the
    # checks that reason about structure read `display` instead.
    number: str = ""
    # For a table cell: (table, row, column), counted from 0 in document
    # order, so arithmetic over a schedule can read the grid the text
    # flattened. None for everything that is not a cell.
    cell: tuple[int, int, int] | None = None

    @property
    def display(self) -> str:
        """The paragraph as the reader sees it, automatic number included."""
        return f"{self.number} {self.text}" if self.number else self.text


@dataclass
class ExtractedDoc:
    blocks: list[Block]
    docx: Document | None
    filename: str

    # The original file, when the model should see it directly rather than
    # only as extracted text. Claude reads PDFs natively, which preserves
    # layout, tables and signature pages that a text layer flattens, and lets
    # it read a scan that has no text layer at all. The extracted blocks are
    # still produced, because the deterministic checks need text and because
    # findings have to anchor to something.
    native: tuple[bytes, str] | None = None
    # The text was transcribed by the model from the page images rather than
    # read from a text layer. A transcription can be wrong in ways a text layer
    # cannot, so the deterministic checks do not run on it (checks.run_all) and
    # the reply says the copy was built from a scan. What it buys is a Word
    # copy to write tracked changes into, and a comparison, for a document that
    # used to get neither.
    transcribed: bool = False
    # Some paragraph is numbered by Word through a list definition this reader
    # could not follow. Its clause numbers are then unknown, and a check that
    # compared references against the numbers it did find would report every
    # reference to the unknown ones as broken.
    numbering_unresolved: bool = False

    def as_prompt(self, limit: int | None = None) -> str:
        """The document as the agent sees it.

        `limit` caps the characters rendered. A hundred-page credit agreement
        does not fit one review comfortably, and silently truncating it would
        produce a confident review of the first third. When the cap bites, the
        caller is told so it can say so in the reply.
        """
        # Reset first. This used to persist from an earlier render with a
        # smaller limit, so a document well inside the budget was reviewed in
        # full while the agent was told it had been truncated and instructed to
        # confine its findings.
        self.truncated_at = None
        lines, used = [], 0
        for b in self.blocks:
            if not b.text.strip():
                continue
            # The automatic number is labelled as such rather than run into the
            # text, so it is never mistaken for words an anchor can quote.
            kind = f"{b.kind}, numbered {b.number}" if b.number else b.kind
            line = f"[{b.index}] ({kind}) {b.text}"
            if limit is not None and used + len(line) > limit:
                self.truncated_at = b.index
                break
            lines.append(line)
            used += len(line) + 1
        return "\n".join(lines)

    truncated_at: int | None = None

    @property
    def characters(self) -> int:
        return sum(len(b.text) for b in self.blocks)


log = logging.getLogger(__name__)


class Unreadable(Exception):
    """Carries a sentence we are willing to email back verbatim."""


def extract(att: Attachment) -> ExtractedDoc:
    """Attachment to reviewable blocks, dispatching on content not extension."""
    kind = identify(att.content, att.filename)

    if kind is Kind.DOCX:
        try:
            return _extract_docx(att)
        except Unreadable:
            raise
        except Exception as e:
            raise Unreadable(message(Kind.DAMAGED)) from e

    if kind is Kind.TEXT:
        text = att.content.decode("utf-8", errors="replace")
        blocks = [
            Block(i, line, "Normal", "paragraph")
            for i, line in enumerate(text.splitlines())
        ]
        return ExtractedDoc(blocks=blocks, docx=None, filename=att.filename)

    if kind is Kind.PDF:
        return _extract_pdf(att)

    if kind in SANDBOX_CONVERTIBLE:
        # A legacy .doc is the case that matters. Telling a lawyer to open it
        # and re-save is exactly the friction this product exists to remove.
        converted = _via_sandbox_convert(att, kind)
        if converted is not None:
            return converted

    if kind is Kind.PPTX:
        return _extract_pptx(att)

    if kind is Kind.XLSX:
        return _extract_xlsx(att)

    raise Unreadable(message(kind))


def _story_blocks(doc, start: int) -> list[Block]:
    """Header, footer and note text, which the writer can edit.

    The writer walks every story part, so an anchor in a header is editable.
    The extractor used to read only the body, so the agent could not see what
    it was allowed to change, the ambiguity guard was computed over different
    text than the writer used, and a change could land somewhere the review had
    never looked. Same parts, same rule: existing ones only, never the
    get-or-create accessors that would materialise them.
    """
    from secondeye.pipeline.ooxml import _paragraph_text, _story_parts

    blocks: list[Block] = []
    index = start
    for part in _story_parts(doc):
        name = str(getattr(part, "partname", "")).rsplit("/", 1)[-1]
        kind = ("header" if "header" in name else
                "footer" if "footer" in name else "note")
        element = getattr(part, "element", None)
        if element is None:
            continue
        texts = [
            _paragraph_text(p) for p in _paragraphs_of(element, part)
        ]
        for text in texts:
            if not text.strip():
                continue
            blocks.append(Block(index, f"[{kind}] {text}", "Normal", kind))
            index += 1
    return blocks


def _paragraphs_of(element, part):
    from secondeye.pipeline.ooxml import _paragraphs_under

    return _paragraphs_under(element, part)


def _extract_docx(att: Attachment) -> ExtractedDoc:
    """The body in document order, then the headers, footers and notes.

    In order, and into block-level content controls. Tables used to be read
    after every paragraph, so a definitions table sat at the end of the
    document as far as the checks knew; and a paragraph inside a content
    control was not read at all, although the writer can find and edit it, so
    an "[INSERT AMOUNT]" in a template's form field went unreported.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(BytesIO(att.content))
    numbers = _ListNumbers(doc)
    blocks: list[Block] = []
    unresolved = False

    # Style names by style id. python-docx resolves `paragraph.style` by
    # walking every style in the document, twice per paragraph as this was
    # written, which was nine tenths of the time taken to read a long
    # agreement, and a review reads the same document several times over.
    style_names: dict[str | None, str] = {}

    def style_of(para) -> str:
        style_id = para._p.style
        if style_id not in style_names:
            try:
                style = para.style
                style_names[style_id] = style.name if style is not None else "Normal"
            except Exception:  # noqa: BLE001 - a dangling style id is not fatal
                style_names[style_id] = "Normal"
        return style_names[style_id]

    def paragraph(p_element) -> None:
        nonlocal unresolved
        para = Paragraph(p_element, doc)
        style = style_of(para)
        kind = "heading" if style.startswith(("Heading", "Title")) else (
            "list-item" if "List" in style else "paragraph"
        )
        number, known = numbers.label(p_element)
        unresolved = unresolved or not known
        blocks.append(Block(len(blocks), visible_text(p_element), style, kind, number))

    tables = 0

    def table(tbl_element) -> None:
        nonlocal unresolved, tables
        seen: list = []     # the elements themselves, so their ids stay theirs
        number = tables
        tables += 1
        for r, row in enumerate(Table(tbl_element, doc).rows):
            for c, cell in enumerate(row.cells):
                # A merged cell is returned once per grid column it spans.
                if any(tc is cell._tc for tc in seen):
                    continue
                seen.append(cell._tc)
                paras = cell.paragraphs
                text = "\n".join(visible_text(p._p) for p in paras)
                label = ""
                for p in paras:
                    label, known = numbers.label(p._p)
                    unresolved = unresolved or not known
                    if label:
                        break
                blocks.append(Block(len(blocks), text, "TableCell", "table-cell", label,
                                    cell=(number, r, c)))

    def walk(container) -> None:
        for child in container:
            if not isinstance(child.tag, str):
                continue
            if child.tag == _W + "p":
                paragraph(child)
            elif child.tag == _W + "tbl":
                table(child)
            elif child.tag in (_W + "sdt", _W + "sdtContent", _W + "customXml"):
                walk(child)

    walk(doc.element.body)

    # Headers, footers and notes. The writer can edit them, so the agent has to
    # be able to see them: otherwise it cannot review what it is allowed to
    # change, and the two disagree about how often a phrase occurs, which is
    # what the ambiguity guard is computed from.
    blocks.extend(_story_blocks(doc, len(blocks)))
    return ExtractedDoc(blocks=blocks, docx=doc, filename=att.filename,
                        numbering_unresolved=unresolved)


# --------------------------------------------------------------------------
# Word's automatic numbering
# --------------------------------------------------------------------------
#
# Almost every agreement a firm produces numbers its clauses with Word's list
# numbering, and that number is drawn by Word, not stored in the text. Read
# naively, such a document has no clause 2.2 at all, and every "Section 2.2"
# in it points at nothing. This renders the number the way Word draws it, from
# the list definition in numbering.xml: the level's format and text pattern,
# its start value, a start override on the list instance, and the rule that a
# level restarts when a level above it advances.


_ROMAN = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
          (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]


def _roman(n: int) -> str:
    out = []
    for value, letters in _ROMAN:
        while n >= value:
            out.append(letters)
            n -= value
    return "".join(out)


def _letters(n: int) -> str:
    """Word's letter numbering: a..z, then aa..zz, then aaa."""
    if n <= 0:
        return ""
    return chr(ord("a") + (n - 1) % 26) * ((n - 1) // 26 + 1)


def _format_number(n: int, fmt: str) -> str:
    if fmt == "lowerLetter":
        return _letters(n)
    if fmt == "upperLetter":
        return _letters(n).upper()
    if fmt == "lowerRoman":
        return _roman(n)
    if fmt == "upperRoman":
        return _roman(n).upper()
    if fmt == "decimalZero":
        return f"{n:02d}"
    if fmt == "ordinal":
        suffix = "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
        return f"{n}{suffix}"
    return str(n)


@dataclass
class _Level:
    start: int = 1
    fmt: str = "decimal"
    text: str = ""
    restart: int | None = None    # w:lvlRestart: restart after this (1-based) level only
    legal: bool = False           # w:isLgl: show every level as a decimal


class _ListNumbers:
    """The label Word draws for each numbered paragraph, in document order.

    Counters are kept per abstract list, which is how Word continues a list
    across two list instances that share one definition. A start override on
    an instance restarts the list the first time that instance is used.
    """

    def __init__(self, doc) -> None:
        self.abstract: dict[str, dict[int, _Level]] = {}
        self.abstract_link: dict[str, str] = {}      # abstractNumId -> numStyleLink style
        self.nums: dict[str, tuple[str, dict[int, int]]] = {}
        self.style_numpr: dict[str, tuple[str | None, int | None]] = {}
        self.style_based_on: dict[str, str] = {}
        self.counters: dict[str, dict[int, int]] = {}
        self.started: set[str] = set()
        try:
            numbering = doc.part.numbering_part.element
        except Exception:  # noqa: BLE001 - no numbering part: nothing is auto-numbered
            numbering = None
        if numbering is not None:
            self._read_numbering(numbering)
        try:
            self._read_styles(doc.styles.element)
        except Exception:  # noqa: BLE001 - numbering by style is then unknown
            log.info("could not read paragraph styles for list numbering")

    @staticmethod
    def _val(node, tag: str) -> str | None:
        child = node.find(_W + tag) if node is not None else None
        return child.get(_W + "val") if child is not None else None

    def _read_numbering(self, numbering) -> None:
        for absn in numbering.findall(_W + "abstractNum"):
            aid = absn.get(_W + "abstractNumId")
            link = self._val(absn, "numStyleLink")
            if link:
                self.abstract_link[aid] = link
            levels: dict[int, _Level] = {}
            for lvl in absn.findall(_W + "lvl"):
                levels[int(lvl.get(_W + "ilvl", "0"))] = self._level(lvl)
            self.abstract[aid] = levels
        for num in numbering.findall(_W + "num"):
            nid = num.get(_W + "numId")
            aid = self._val(num, "abstractNumId")
            overrides: dict[int, int] = {}
            for over in num.findall(_W + "lvlOverride"):
                ilvl = int(over.get(_W + "ilvl", "0"))
                start = self._val(over, "startOverride")
                if start is not None and start.lstrip("-").isdigit():
                    overrides[ilvl] = int(start)
                lvl = over.find(_W + "lvl")
                if lvl is not None and aid in self.abstract:
                    # A level redefined on the instance. Rare; the abstract list
                    # is copied so the override does not leak to its siblings.
                    self.abstract[f"{aid}/{nid}"] = dict(self.abstract[aid])
                    self.abstract[f"{aid}/{nid}"][ilvl] = self._level(lvl)
                    aid = f"{aid}/{nid}"
            if aid is not None:
                self.nums[nid] = (aid, overrides)

    def _level(self, lvl) -> _Level:
        start = self._val(lvl, "start")
        restart = self._val(lvl, "lvlRestart")
        return _Level(
            start=int(start) if start and start.lstrip("-").isdigit() else 1,
            fmt=self._val(lvl, "numFmt") or "decimal",
            text=self._val(lvl, "lvlText") or "",
            restart=int(restart) if restart and restart.isdigit() else None,
            legal=lvl.find(_W + "isLgl") is not None,
        )

    def _read_styles(self, styles) -> None:
        for style in styles.findall(_W + "style"):
            sid = style.get(_W + "styleId")
            based = self._val(style, "basedOn")
            if based:
                self.style_based_on[sid] = based
            numpr = style.find(f"{_W}pPr/{_W}numPr")
            if numpr is not None:
                ilvl = self._val(numpr, "ilvl")
                self.style_numpr[sid] = (self._val(numpr, "numId"),
                                         int(ilvl) if ilvl and ilvl.isdigit() else None)

    def _from_style(self, sid: str | None) -> tuple[str | None, int | None]:
        seen: set[str] = set()
        while sid and sid not in seen:
            seen.add(sid)
            if sid in self.style_numpr:
                return self.style_numpr[sid]
            sid = self.style_based_on.get(sid)
        return None, None

    def _abstract_for(self, nid: str) -> tuple[str | None, dict[int, int]]:
        if nid not in self.nums:
            return None, {}
        aid, overrides = self.nums[nid]
        # A list defined by a list style points at the style, which points at
        # the instance that carries the real definition.
        hops = 0
        while aid in self.abstract_link and hops < 4:
            linked, _ = self._from_style(self.abstract_link[aid])
            if not linked or linked not in self.nums:
                break
            aid = self.nums[linked][0]
            hops += 1
        return aid, overrides

    def label(self, p_element) -> tuple[str, bool]:
        """(the drawn number or "", whether it could be worked out)."""
        ppr = p_element.find(_W + "pPr")
        numpr = ppr.find(_W + "numPr") if ppr is not None else None
        nid = ilvl = None
        if numpr is not None:
            nid = self._val(numpr, "numId")
            raw = self._val(numpr, "ilvl")
            ilvl = int(raw) if raw and raw.isdigit() else None
        if nid is None:
            sid = self._val(ppr, "pStyle") if ppr is not None else None
            style_nid, style_ilvl = self._from_style(sid)
            nid = style_nid
            if ilvl is None:
                ilvl = style_ilvl
        if nid is None or nid == "0":
            return "", True
        ilvl = ilvl or 0
        aid, overrides = self._abstract_for(nid)
        if aid is None or aid not in self.abstract:
            return "", False
        levels = self.abstract[aid]
        level = levels.get(ilvl)
        if level is None:
            return "", False

        counters = self.counters.setdefault(aid, {})
        if nid not in self.started:
            self.started.add(nid)
            for lvl_index, start in overrides.items():
                # The override restarts the list: the level is set so that
                # the next paragraph at it takes the override's value.
                counters[lvl_index] = start - 1
        current = counters.get(ilvl)
        counters[ilvl] = (current + 1) if current is not None else levels[ilvl].start
        for deeper in [k for k in counters if k > ilvl]:
            restart = levels.get(deeper, _Level()).restart
            if restart == 0:
                continue
            if restart is None or restart - 1 >= ilvl:
                del counters[deeper]

        if level.fmt in ("bullet", "none") or not level.text:
            return ("" if level.fmt == "bullet" else level.text.strip()), True

        def value(m: re.Match) -> str:
            k = int(m.group(1)) - 1
            lvl = levels.get(k, _Level())
            n = counters.get(k, lvl.start)
            fmt = "decimal" if level.legal and k < ilvl else lvl.fmt
            return _format_number(n, fmt)

        return re.sub(r"%([1-9])", value, level.text).strip(), True


# Word namespace, and the elements it renders as ordinary paragraph text.
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_TRANSPARENT = {_W + "ins", _W + "hyperlink", _W + "smartTag", _W + "sdtContent"}


def visible_text(paragraph_element) -> str:
    """Everything Word shows in this paragraph, in document order.

    Not `python-docx`'s `paragraph.text`, which reads only the direct `w:r`
    children of `w:p`. Any text an author typed with track changes on lives
    inside a `w:ins`, and hyperlink and smart-tag text is nested too, so all of
    it was invisible to `Block.text` and therefore to every deterministic
    check. The commonest state of a document in a live negotiation is "carries
    somebody's tracked changes", and in that state a `[PRICE TO BE CONFIRMED]`
    left in an insertion was missed entirely.

    Runs inside a `w:del` are skipped: that text has been deleted, so it is not
    in the document and a finding must never anchor to it. This is the same set
    of containers `ooxml._runs` treats as visible, deliberately, so that a
    finding anchored to text found here can be located again by the redliner.
    """
    out: list[str] = []

    def walk(node) -> None:
        for child in node:
            if not isinstance(child.tag, str):
                continue  # a comment or processing instruction
            if child.tag == _W + "r":
                out.append(_run_text(child))
            elif child.tag in _TRANSPARENT:
                walk(child)

    walk(paragraph_element)
    return "".join(out)


def _run_text(run) -> str:
    """One run's text, matching what python-docx would have produced for it.

    Tabs and line breaks are kept as characters: clause numbering arrives as
    "1.<tab>Definitions" and the numbering and cross-reference checks read the
    start of the block. `w:delText` and `w:instrText` are excluded, being
    deleted text and field code rather than anything the reader sees.
    """
    parts: list[str] = []
    for node in run.iter():
        if node.tag == _W + "t":
            parts.append(node.text or "")
        elif node.tag == _W + "tab":
            parts.append("\t")
        elif node.tag in (_W + "br", _W + "cr"):
            parts.append("\n")
    return "".join(parts)


def _extract_pdf(att: Attachment) -> ExtractedDoc:
    """PDF to reviewable blocks.

    Memo-only by nature: there is no tracked change to write into a PDF, so
    intake.mode_for downgrades the request. But a litigator forwarding a brief
    should still get a review, and previously this branch raised, told them to
    reply "review it anyway", and then raised again on that reply. There was no
    sequence of replies that produced a review.

    pypdfium2 is used rather than PyMuPDF, which is the more obvious choice and
    is AGPL. That licence is a real problem for a commercial product and it is
    not worth the better table handling here.
    """
    try:
        import pypdfium2
    except ImportError as e:  # pragma: no cover - dependency is declared
        raise Unreadable(
            "I could not read that PDF. Send me the Word version and I will "
            "review it, with tracked changes."
        ) from e

    try:
        pdf = pypdfium2.PdfDocument(BytesIO(att.content))
        pages = [page.get_textpage().get_text_bounded() or "" for page in pdf]
    except Exception as e:
        raise Unreadable(
            "That PDF would not open. It may be damaged or password protected."
        ) from e

    # A scan has no text layer. The model reads the pages and writes the text
    # down, verbatim, so that a Word copy can be built and tracked changes
    # written into it. Not an OCR engine: the owner's decision is that reading
    # goes through the model, and the same model is about to read the pages for
    # the review anyway. A transcription can be wrong in ways a text layer
    # cannot, which is why checks.run_all does nothing on it.
    transcribed = False
    if not any(t.strip() for t in pages) and pages and len(pages) <= TRANSCRIBE_MAX_PAGES \
            and transcription_available():
        try:
            read = _transcribe(att.content, len(pages))
        except Exception:
            # No credit, a network fault, a refusal: the scan is then what it
            # was before, a document the model reads in the review and nothing
            # can be marked up in. Logged, not raised.
            log.exception("could not transcribe a scanned PDF; reviewing the pages only")
            read = []
        if read and any(t.strip() for t in read):
            pages, transcribed = read, True

    blocks: list[Block] = []
    index = 0
    for page_number, text in enumerate(pages, start=1):
        for paragraph in _pdf_paragraphs(text):
            blocks.append(Block(index, paragraph, "Normal", "paragraph"))
            index += 1
        blocks.append(Block(index, f"[page {page_number}]", "Normal", "page-break"))
        index += 1

    # Without a transcription a scan has no text, and the deterministic checks
    # cannot run on it. That is no longer a reason to refuse: the model reads
    # the pages directly in the review, so it degrades to a substantive review
    # rather than nothing.
    return ExtractedDoc(
        blocks=blocks,
        docx=None,
        filename=att.filename,
        native=(att.content, "application/pdf"),
        transcribed=transcribed,
    )


# A scan longer than this is not transcribed: the transcription is one model
# call carrying every page, and a 300-page scanned data-room bundle is not
# what this product is for.
TRANSCRIBE_MAX_PAGES = 60
_PAGE_MARK = "=== PAGE {n} ==="


def transcription_available() -> bool:
    """A model can be called: an API key is configured."""
    from secondeye import policy
    from secondeye.config import settings

    # Not for a no-AI client (policy.py): a scan of theirs is unreadable.
    return bool(settings().anthropic_api_key) and not policy.ai_forbidden()


def _transcribe(pdf: bytes, page_count: int) -> list[str]:
    """The pages' text, one string per page, as the model reads them.

    One call, the whole document as a native PDF block, so the model sees each
    page as printed. The instruction asks for the words and nothing else: no
    summary, no correction of what looks like a typo, a page marker between
    pages so the text can be laid back onto the pages it came from.
    """
    import base64

    from secondeye.config import anthropic_client, refusal_fallback, settings

    cfg = settings()
    client = anthropic_client()
    instruction = (
        f"This PDF is a scan of {page_count} page{'s' if page_count != 1 else ''}. "
        "Transcribe every page, word for word, as printed: keep the paragraph "
        "breaks, clause numbers, headings and signature blocks; do not correct, "
        "summarise, translate or comment; write nothing that is not on the page. "
        f'Begin each page with a line reading exactly "{_PAGE_MARK.format(n="N")}" '
        "where N is the page number, starting at 1. If a page is blank or has "
        "no readable text, write the marker and nothing else."
    )
    # Copying, not reasoning: low effort keeps a sixty-page transcription
    # quick, and Haiku takes no effort setting at all.
    options = {} if "haiku" in cfg.effective_model.lower() else {"output_config": {"effort": "low"}}
    # Beta, for the refusal fallback (config.refusal_fallback). Streamed,
    # because a sixty-page transcription is a long turn.
    with client.beta.messages.stream(
        model=cfg.effective_model,
        max_tokens=64000,
        **options,
        **refusal_fallback(cfg.effective_model),
        messages=[{"role": "user", "content": [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                            "data": base64.standard_b64encode(pdf).decode()}},
            {"type": "text", "text": instruction},
        ]}],
    ) as stream:
        message = stream.get_final_message()
    # A refusal can come mid-stream, after part of the text. That part is not
    # a transcription of the whole scan and must not be used as one.
    if message.stop_reason == "refusal":
        raise RuntimeError("the model declined to transcribe this scan")
    text = "".join(b.text for b in message.content if b.type == "text")
    return _split_pages(text, page_count)


def _split_pages(text: str, page_count: int) -> list[str]:
    """Lay the transcription back onto its pages by the markers, tolerating a
    missing or extra one: the count of pages in the result is always the
    count in the PDF."""
    import re

    parts = re.split(r"^=== PAGE (\d+) ===\s*$", text, flags=re.MULTILINE)
    pages = [""] * page_count
    if len(parts) <= 1:
        # No markers at all. Everything onto the first page: the text is still
        # right, only its page is not.
        pages[0] = text.strip()
        return pages
    for number, body in zip(parts[1::2], parts[2::2], strict=False):
        n = int(number)
        if 1 <= n <= page_count:
            pages[n - 1] = (pages[n - 1] + "\n" + body).strip()
    return pages


def _pdf_paragraphs(text: str) -> list[str]:
    """Join the hard-wrapped lines a PDF text layer produces back into clauses.

    Extracted PDF text breaks at the visual line, so a sentence arrives as five
    fragments. Left alone, every cross-reference and defined-term check reads
    across a break that is not there in the document.
    """
    lines = [line.rstrip() for line in text.splitlines()]
    paragraphs: list[str] = []
    current: list[str] = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if current:
                paragraphs.append(" ".join(current))
                current = []
            continue
        # A new numbered clause always starts a paragraph.
        if current and re.match(r"^\s*(?:\d+(?:\.\d+)*[.)]|\([a-z0-9]+\))\s", line):
            paragraphs.append(" ".join(current))
            current = []
        current.append(stripped)
        # A line ending in a full stop that is not an abbreviation ends a clause.
        if stripped.endswith((".", ";", ":")) and len(stripped) > 40:
            paragraphs.append(" ".join(current))
            current = []
    if current:
        paragraphs.append(" ".join(current))
    return [p for p in paragraphs if p.strip()]


def _via_sandbox_convert(att: Attachment, kind: Kind) -> ExtractedDoc | None:
    """Convert an unreadable format to .docx.

    Locally first. LibreOffice does this as a subprocess with the document
    never leaving the machine, which is strictly better than a session for
    privileged material and needs no data-handling story at all. The associate
    agent's sandbox is the fallback for deployments without LibreOffice.
    """
    from secondeye import convert, managed

    converted = None
    if convert.available():
        try:
            converted = convert.to_docx(att.content, att.filename)
        except convert.ConversionUnavailable as e:
            log.info("local conversion failed for %s: %s", att.filename, e)

    if converted is None:
        if not managed.configured():
            return None
        try:
            converted = managed.convert_to_docx(att.content, att.filename)
        except managed.JobFailed as e:
            log.info("session conversion unavailable for %s: %s", att.filename, e)
            return None
    if identify(converted, att.filename) is not Kind.DOCX:
        log.warning("the conversion returned something that is not a .docx")
        return None
    return _extract_docx(
        Attachment(filename=_as_docx(att.filename), content_type=att.content_type,
                   size_bytes=len(converted), content=converted)
    )


def _as_docx(filename: str) -> str:
    stem = filename.rsplit(".", 1)[0] if "." in filename else filename
    return f"{stem}.docx"


def _extract_pptx(att: Attachment) -> ExtractedDoc:
    """A deck as readable text, slide by slide.

    Decks arrive attached to deals constantly. They cannot carry a tracked
    change, so the review is memo-only, and plain text is enough.

    Read locally with python-pptx rather than in Anthropic's code-execution
    container, which has the same library. The container is not eligible for
    zero data retention and keeps uploaded files for up to thirty days, which
    is not a trade worth making for privileged material when the same work runs
    in-process. See docs/sandbox.md.
    """
    try:
        from pptx import Presentation
    except ImportError as e:  # pragma: no cover - dependency is declared
        # NOT message(Kind.PPTX): that says "I can read it and send you notes",
        # which is exactly what a missing library means we cannot do.
        raise Unreadable(
            "I could not read that PowerPoint file. Send me the text you want "
            "checked and I will review that."
        ) from e

    blocks: list[Block] = []
    index = 0
    try:
        deck = Presentation(BytesIO(att.content))
        for number, slide in enumerate(deck.slides, start=1):
            blocks.append(Block(index, f"[slide {number}]", "Normal", "page-break"))
            index += 1
            for shape in slide.shapes:
                for line in _shape_text(shape):
                    blocks.append(Block(index, line, "Normal", "paragraph"))
                    index += 1
            notes = getattr(slide, "notes_slide", None) if slide.has_notes_slide else None
            if notes is not None and notes.notes_text_frame is not None:
                text = (notes.notes_text_frame.text or "").strip()
                if text:
                    blocks.append(Block(index, f"Speaker notes: {text}", "Normal", "paragraph"))
                    index += 1
    except Unreadable:
        raise
    except Exception as e:
        raise Unreadable(
            "That PowerPoint file would not open. It may be damaged."
        ) from e

    return ExtractedDoc(blocks=blocks, docx=None, filename=att.filename)


def _shape_text(shape) -> list[str]:
    """Text from one shape, including tables and grouped shapes."""
    out: list[str] = []
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
        return out
    if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
        for inner in shape.shapes:            # a grouped shape
            out.extend(_shape_text(inner))
        return out
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                out.append(text)
    return out


def _extract_xlsx(att: Attachment) -> ExtractedDoc:
    """A workbook as readable text, sheet by sheet.

    Values rather than formulas: the review is about what the document says,
    and a lawyer reading a schedule of payments cares about the numbers, not
    how they were computed.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover - dependency is declared
        raise Unreadable(
            "I could not read that Excel file. Send me the text you want "
            "checked and I will review that."
        ) from e

    blocks: list[Block] = []
    index = 0
    try:
        book = load_workbook(BytesIO(att.content), data_only=True, read_only=True)
        for sheet in book.worksheets:
            blocks.append(Block(index, f"[sheet {sheet.title}]", "Normal", "page-break"))
            index += 1
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if not any(cells):
                    continue
                # Trailing empties are padding, not content.
                while cells and not cells[-1]:
                    cells.pop()
                blocks.append(Block(index, " | ".join(cells), "Normal", "table-cell"))
                index += 1
        book.close()
    except Unreadable:
        raise
    except Exception as e:
        raise Unreadable("That Excel file would not open. It may be damaged.") from e

    return ExtractedDoc(blocks=blocks, docx=None, filename=att.filename)
