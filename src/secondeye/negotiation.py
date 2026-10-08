"""The negotiation ledger and the covering-note lie detector.

Opposing counsel sends v2 with "We've accepted all your points except the
liability cap." Usually that is nearly true, and the lawyer who takes it on
trust signs a document in which the fraud carve-out quietly went missing. The
associate the lawyer wants reads v2 against every point we raised and answers
first with whether their email is true.

Built the way DECISIONS 29-30 build everything now. The model decides: what
happened to each point, which changes nobody asked for and how much each one
matters, and whether each thing their email claims is true. It does so in the
review session, with the `lra-negotiation` skill, and records
`negotiation.json` through that skill's validator. What this module keeps is
what the model cannot hold between emails, and the evidence it is given:

- **The ledger.** Every point we raised on a document, per round: the issues
  list's rows on their paper, the substantive tracked changes we wrote, and the
  asks in a covering email of ours that we were copied on. Each with its
  clause, what we asked, and its status after the last version they sent.
- **The version we sent.** The copy we were BCC'd on when it went out, or else
  the document as we returned it: the conversation's original with our active
  changes written in (thread.rebuild).
- **The evidence.** The exact comparison of that version against theirs
  (pipeline/compare.py), labelled by clause; the ledger; and their covering
  email with the quoted history taken off. Handed to the session as a fenced
  block and as a file the skill's scripts read.
- **The reply's first lines,** from the validated file: the verdict on their
  claims, the tally, what was not accepted and what they changed unasked,
  clause first; and the issues list again with a status column.

The ledger follows the document. A forwarded v2 usually arrives as a new mail
thread, so the ledger is found by the conversation that holds the document
(by name, then by how alike the texts are, as reconcile.py finds one) and
moved to the conversation holding the latest version, where the next round
will look for it.
"""

from __future__ import annotations

import importlib.util
import json
import logging
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from functools import cache
from html import escape
from io import BytesIO
from types import ModuleType

from secondeye import blobs
from secondeye.models import Attachment, Finding, InboundEmail, OutboundEmail, Severity
from secondeye.store import connect

log = logging.getLogger(__name__)

EVIDENCE_FILE = "negotiation-evidence.json"
OUTPUT_FILE = "negotiation.json"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SCHEMA = """
CREATE TABLE IF NOT EXISTS negotiation_points (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    thread_id TEXT NOT NULL,
    round INTEGER NOT NULL,           -- the round we raised it in
    source TEXT NOT NULL,             -- issue | change | email
    clause TEXT,
    asked TEXT NOT NULL,
    anchor TEXT,
    status TEXT NOT NULL DEFAULT 'open',  -- open | accepted | rejected | partial | changed
    status_round INTEGER,             -- the round whose version gave it that status
    evidence TEXT,                    -- the sentence that says why
    created_at TEXT,
    updated_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_negotiation_points_thread ON negotiation_points(thread_id);

-- The version that went to the other side, when we were copied on it.
CREATE TABLE IF NOT EXISTS negotiation_sent (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    thread_id TEXT NOT NULL,
    filename TEXT,
    content BLOB,
    created_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_negotiation_sent_thread ON negotiation_sent(thread_id);

-- Each version of theirs read against the ledger, as the model recorded it.
CREATE TABLE IF NOT EXISTS negotiation_rounds (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    thread_id TEXT NOT NULL,
    round INTEGER NOT NULL,
    filename TEXT,
    outcome TEXT NOT NULL,            -- json: the validated negotiation.json
    points TEXT NOT NULL,             -- json: the ledger it answered, by point id
    cover INTEGER NOT NULL DEFAULT 0, -- whether they sent a covering email
    created_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_negotiation_rounds_thread ON negotiation_rounds(thread_id);

-- The evidence a review session was given, until its reply is composed: the
-- start and the finish of a review can be different processes (flow.py).
-- Sealed, because it quotes the document.
CREATE TABLE IF NOT EXISTS negotiation_evidence (
    job_id TEXT PRIMARY KEY,
    thread_id TEXT NOT NULL,
    evidence BLOB NOT NULL,
    created_at TEXT
);
"""

# Columns added after a table first shipped: (table, column, declaration).
# init() adds any a database lacks, once per database per process, the way
# memory.py and thread.py bring an older database up to date.
LATER_COLUMNS: tuple[tuple[str, str, str], ...] = ()

_migrated: set[str] = set()


def init() -> None:
    from secondeye.config import settings

    url = settings().database_url
    with connect() as c:
        c.executescript(SCHEMA)
        if url in _migrated:
            return
        for table, column, declaration in LATER_COLUMNS:
            have = {r[1] for r in c.execute(f"PRAGMA table_info({table})")}
            if column not in have:
                c.execute(f"ALTER TABLE {table} ADD COLUMN {column} {declaration}")
        _migrated.add(url)


# --------------------------------------------------------------------------
# The ledger
# --------------------------------------------------------------------------

SOURCES = {"issue": "issues list", "change": "tracked change", "email": "our email"}


@dataclass
class Point:
    id: int
    thread_id: str
    round: int
    source: str
    clause: str
    asked: str
    anchor: str = ""
    status: str = "open"
    status_round: int | None = None
    evidence: str = ""

    @property
    def pid(self) -> str:
        return f"P{self.id}"

    def as_evidence(self) -> dict:
        out = {"id": self.pid, "clause": self.clause, "asked": self.asked,
               "source": self.source, "round": self.round}
        if self.anchor:
            out["their_text_when_raised"] = self.anchor
        if self.status != "open":
            out["status_last_round"] = self.status
        return out


_POINT_COLUMNS = ("id, thread_id, round, source, clause, asked, anchor, status, "
                  "status_round, evidence")


def points(thread_id: str) -> list[Point]:
    init()
    with connect() as c:
        rows = c.execute(f"SELECT {_POINT_COLUMNS} FROM negotiation_points "
                         "WHERE thread_id = ? ORDER BY id", (thread_id,)).fetchall()
    return [Point(id=r[0], thread_id=r[1], round=r[2], source=r[3], clause=r[4] or "",
                  asked=r[5], anchor=r[6] or "", status=r[7] or "open",
                  status_round=r[8], evidence=r[9] or "") for r in rows]


def _key(text: str) -> str:
    return re.sub(r"\W+", " ", (text or "").lower()).strip()


def record_points(thread_id: str, round_number: int, raised: list[dict]) -> int:
    """Add points to the ledger; returns how many were new. A point already in
    the ledger and not yet accepted (same clause and ask, or same anchor) is
    the same point still open, not a second one: the review of their next
    version finds the same off-playbook cap again, and it is one point."""
    init()
    existing = [p for p in points(thread_id) if p.status != "accepted"]
    seen = {(p.source, _key(p.clause), _key(p.asked)) for p in existing}
    anchors = {_key(p.anchor) for p in existing if p.anchor}
    now = datetime.now(UTC).isoformat()
    added = 0
    with connect() as c:
        for r in raised:
            asked = (r.get("asked") or "").strip()
            if not asked:
                continue
            key = (r["source"], _key(r.get("clause", "")), _key(asked))
            anchor = _key(r.get("anchor", ""))
            if key in seen or (anchor and anchor in anchors):
                continue
            c.execute(
                "INSERT INTO negotiation_points (thread_id, round, source, clause, asked, "
                "anchor, status, created_at, updated_at) VALUES (?,?,?,?,?,?,'open',?,?)",
                (thread_id, round_number, r["source"], r.get("clause", ""), asked[:2000],
                 (r.get("anchor") or "")[:2000], now, now))
            seen.add(key)
            if anchor:
                anchors.add(anchor)
            added += 1
    return added


def record_sent(thread_id: str, filename: str, content: bytes) -> None:
    init()
    with connect() as c:
        c.execute("INSERT INTO negotiation_sent (thread_id, filename, content, created_at) "
                  "VALUES (?,?,?,?)", (thread_id, filename, blobs.stash(content, "doc"),
                                       datetime.now(UTC).isoformat()))


def latest_sent(thread_id: str) -> tuple[str, bytes, str] | None:
    """(filename, content, day) of the last version we were copied on sending."""
    init()
    with connect() as c:
        row = c.execute("SELECT filename, content, created_at FROM negotiation_sent "
                        "WHERE thread_id = ? ORDER BY id DESC LIMIT 1", (thread_id,)).fetchone()
    if not row:
        return None
    return row[0] or "", blobs.fetch(row[1]), _day(row[2])


def rounds_done(thread_id: str) -> int:
    init()
    with connect() as c:
        row = c.execute("SELECT COUNT(*) FROM negotiation_rounds WHERE thread_id = ?",
                        (thread_id,)).fetchone()
    return row[0] if row else 0


def move(old: str, new: str) -> None:
    """The ledger follows the document to the conversation holding its
    latest version."""
    if not old or old == new:
        return
    init()
    with connect() as c:
        for table in ("negotiation_points", "negotiation_sent", "negotiation_rounds"):
            c.execute(f"UPDATE {table} SET thread_id = ? WHERE thread_id = ?", (new, old))
    # The versions the conversation held go with it ("blackline against v1").
    from secondeye import blackline

    blackline.move(old, new)


def sweep() -> None:
    """Rows whose conversation has gone (thread.purge, retention). Held a day,
    because the ledger is moved to a conversation a moment before that
    conversation's row is written. Never raises."""
    try:
        init()
        cutoff = (datetime.now(UTC) - timedelta(days=1)).isoformat()
        with connect() as c:
            gone = "thread_id NOT IN (SELECT id FROM threads)"
            for (stored,) in c.execute(
                    f"SELECT content FROM negotiation_sent WHERE {gone} AND created_at < ?",
                    (cutoff,)).fetchall():
                blobs.discard(stored)
            c.execute(f"DELETE FROM negotiation_sent WHERE {gone} AND created_at < ?", (cutoff,))
            c.execute(f"DELETE FROM negotiation_points WHERE {gone} AND updated_at < ?",
                      (cutoff,))
            c.execute(f"DELETE FROM negotiation_rounds WHERE {gone} AND created_at < ?",
                      (cutoff,))
            for (stored,) in c.execute("SELECT evidence FROM negotiation_evidence "
                                       "WHERE created_at < ?", (cutoff,)).fetchall():
                blobs.discard(stored)
            c.execute("DELETE FROM negotiation_evidence WHERE created_at < ?", (cutoff,))
    except Exception:
        log.warning("negotiation retention sweep failed", exc_info=True)


def _day(stamp: str | None) -> str:
    try:
        day = datetime.fromisoformat((stamp or "")[:10])
    except ValueError:
        return ""
    return f"{day.day} {day.strftime('%b')}"


# --------------------------------------------------------------------------
# Which conversation holds the ledger for this document
# --------------------------------------------------------------------------


def ledger_for(user: str, thread_key: str, att: Attachment, text: str) -> str | None:
    """This conversation, if it has a ledger; otherwise the lawyer's
    conversation with a ledger whose document this one is a version of: by
    filename, held to the similarity bar, because two clients' NDAs are both
    "NDA.docx"; then by the words alone."""
    from secondeye import reconcile, thread

    init()
    if points(thread_key):
        return thread_key
    with connect() as c:
        rows = c.execute(
            "SELECT t.id, t.filename FROM threads t WHERE t.owner = ? AND EXISTS "
            "(SELECT 1 FROM negotiation_points p WHERE p.thread_id = t.id) "
            "ORDER BY t.updated_at DESC LIMIT 20", (user.lower(),)).fetchall()
    if not rows or not text:
        return None
    stem = reconcile._stem(att.filename)
    ordered = sorted(rows, key=lambda r: reconcile._stem(r[1] or "") != stem)
    for key, _ in ordered:
        held = thread.original_of(key)
        if held is None or not held[1]:
            continue
        earlier = reconcile._text_of(Attachment(filename=held[0], content_type="",
                                                size_bytes=len(held[1]), content=held[1]))
        if earlier and reconcile._similar(earlier, text):
            return key
    return None


# --------------------------------------------------------------------------
# Their covering email, and "did they accept our changes?"
# --------------------------------------------------------------------------

_HEADER = re.compile(r"^\s*[*>]*\s*(?:from|sent|date|to|cc|subject)\s*:", re.IGNORECASE)
_FROM = re.compile(r"^\s*[*>]*\s*From:\*?\s*(.+)$", re.IGNORECASE | re.MULTILINE)
_ADDRESS = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+")
COVER_LIMIT = 4000


def cover_note(email: InboundEmail, owner: str = "") -> tuple[str, str]:
    """(who wrote it, what it says) for the other side's covering email, the
    quoted history below it taken off. The counterparty's message in the
    chain the lawyer forwarded, when they forwarded one; the email itself when
    the other side wrote to us; ("", "") when there is none.

    Never our own side's words: anyone inside the firm, anyone on the
    allowlist and the conversation's owner (`owner`). The allowlisted lawyer
    often writes from a mailbox outside FIRM_DOMAINS (Gmail, in the Falcon
    rehearsal), and reading "Final check, please" as the other side's
    covering email made the reply say it could not settle which of our
    points they had accepted."""
    from secondeye.pipeline.router import strip_reply

    body = email.body or ""
    found = _FROM.search(body)
    if found and _forwarded(email, body[:found.start()]):
        # The first quoted From: may be the lawyer's own message, or a
        # colleague's, with the counterparty's further down the chain.
        for header in _FROM.finditer(body, found.start()):
            address = _ADDRESS.search(header.group(1))
            author = address.group(0).lower() if address else ""
            if author and _ours(author, owner):
                continue
            lines = body[header.start():].splitlines()
            at = 0
            while at < len(lines) and (_HEADER.match(lines[at]) or not lines[at].strip()):
                at += 1
            rest = lines[at:]
            if rest and all(line.lstrip().startswith(">") or not line.strip() for line in rest):
                rest = [re.sub(r"^\s*>\s?", "", line) for line in rest]
            text = strip_reply("\n" + "\n".join(rest)).strip()
            return author, text[:COVER_LIMIT]
        return "", ""
    if not _ours(email.from_address, owner):
        return email.from_address.lower(), strip_reply(body).strip()[:COVER_LIMIT]
    return "", ""


def _ours(address: str, owner: str = "") -> bool:
    """Whether an address is on our side: the firm's, allowlisted (by address
    or domain), or the conversation's owner."""
    from secondeye.config import settings
    from secondeye.models import bare_address

    cfg = settings()
    lowered = bare_address(address or "").lower()
    if not lowered:
        return False
    return (cfg.is_internal(lowered)
            or lowered in cfg.allowlist
            or lowered.split("@")[-1] in cfg.allowlist
            or (bool(owner) and lowered == bare_address(owner).lower()))


def _forwarded(email: InboundEmail, above: str) -> bool:
    """Whether the first "From:" header in the body opens a forwarded message:
    a forward's subject, or a forward or Outlook separator above it."""
    from secondeye.pipeline import their_paper

    return bool(their_paper._FORWARD_SUBJECT.match(email.subject or "")
                or their_paper._FORWARD_MARKER.search(above)
                or re.search(r"_{5,}|-{2,}\s*Original Message", above, re.IGNORECASE))


_ASKS = re.compile(
    r"\b(?:did|have|has|do|does)\s+(?:they|the other side|opposing counsel|the counterparty|"
    r"[a-z]+)\s+(?:\w+\s+){0,2}?(?:accept|agree|take|took|reject|push(?:ed)? back on|"
    r"come back on)\w*\s+(?:to\s+)?(?:(?:all|any|each)\s+(?:of\s+)?)?(?:our|my|the)\s+"
    r"(?:changes|points|comments|mark-?ups?|asks|amendments|redlines?|edits|issues(?: list)?)"
    r"|\bwhat\s+did\s+they\s+(?:accept|agree\s+to|take|reject)\b"
    r"|\bwhich\s+of\s+(?:our|my)\s+(?:changes|points|comments|asks)\b"
    r"|\b(?:status|scorecard)\s+of\s+(?:our|my)\s+(?:changes|points|asks)\b",
    re.IGNORECASE,
)


def asks_status(email: InboundEmail) -> bool:
    """The lawyer asking what the other side did with our points, in their
    own words (never the forwarded message's)."""
    from secondeye.pipeline.router import strip_reply

    return bool(_ASKS.search(strip_reply(email.body or "")))


# --------------------------------------------------------------------------
# The evidence the review session is given
# --------------------------------------------------------------------------

BLOCK_CHARS = 12000
CHANGE_CHARS = 1500

TASK = (
    "The block above is a new version of a document we are negotiating, from the "
    "other side: the points we raised, the complete comparison against the version "
    "we sent, and their covering email. The same evidence is mounted at "
    f"`/workspace/{EVIDENCE_FILE}`. Alongside the review, use the lra-negotiation "
    "skill: decide what they did with every point, which changes nobody asked for, "
    "and whether each thing their email claims is true, and record it with "
    "validate_negotiation.py."
)


@dataclass
class Evidence:
    thread_id: str
    data: dict

    def file(self) -> bytes:
        return json.dumps(self.data, ensure_ascii=False, indent=1).encode()

    def block(self) -> str:
        d = self.data
        lines = [(f"Their version: {d['theirs']['filename']}. The version we sent: "
                  f"{d['ours']['filename']} ({d['ours']['how']})."), "",
                 f"Points we raised ({len(d['points'])}):"]
        for p in d["points"]:
            was = f"; last round: {p['status_last_round']}" if p.get("status_last_round") else ""
            lines.append(f"{p['id']} | {p['clause'] or '(no clause)'} | "
                         f"{SOURCES.get(p['source'], p['source'])}, round {p['round']}{was} | "
                         f"asked: {_one_line(p['asked'])}")
        changes = d["changes"]
        lines += ["", (f"Changes from the version we sent to theirs ({len(changes)}, "
                       "complete; a clause not listed did not change):")]
        used = 0
        for c in changes:
            entry = f"[{c['index']}] {c['kind']} at {c['where']}"
            if c.get("before"):
                entry += "\nBEFORE: " + c["before"]
            if c.get("after"):
                entry += "\nAFTER: " + c["after"]
            if used + len(entry) > BLOCK_CHARS:
                lines.append(f"({len(changes) - c['index']} further changes: read them in "
                             f"/workspace/{EVIDENCE_FILE})")
                break
            lines.append(entry)
            used += len(entry)
        cover = d["cover"]
        if cover["text"]:
            lines += ["", f"Their covering email (from {cover['from'] or 'the other side'}):",
                      cover["text"]]
        else:
            lines += ["", "They sent no covering email with it."]
        return "\n".join(lines)


def _one_line(text: str, limit: int = 400) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    return text if len(text) <= limit else text[: limit - 1].rsplit(" ", 1)[0] + "…"


def evidence_for(job) -> Evidence | None:
    """The evidence for a review of `job` (handler._Review), or None when this
    is not a version from the other side of a document with a ledger."""
    from secondeye import reconcile, thread
    from secondeye.models import Mode
    from secondeye.pipeline import compare, extract, identity, their_paper

    if job.doc is None or job.entry is identity.EntryMode.BCC_SENT:
        return None
    text = reconcile._norm(" ".join(b.text for b in job.doc.blocks))
    tid = ledger_for(job.user, job.thread_key, job.att, text)
    if tid is None:
        return None
    raised = points(tid)
    if not raised:
        return None
    theirs = (asks_status(job.email)
              or their_paper.detect(job.email, job.att.content, job.entry, Mode.REDLINE).theirs
              or bool(cover_note(job.email, job.user)[1]))
    if not theirs:
        return None

    sent = latest_sent(tid)
    if sent is not None:
        base_name, base, day = sent
        how = f"the copy you sent{' on ' + day if day else ''}"
    else:
        state = thread.load(tid)
        if state is None or not state.original:
            return None
        base_name = state.filename
        try:
            base = thread.rebuild(state)[0]
        except Exception:
            log.exception("could not rebuild the version we returned; using the original")
            base = state.original
        how = f"as we returned it on {state.seen_on}"
    if base in (job.att.content, job.working):
        return None     # our own version back: nothing of theirs to read

    try:
        if job.working is None:
            raise ValueError("no Word copy of their version")
        comparison = compare.compare(base, job.working)
    except Exception as e:  # noqa: BLE001 - the text comparison is the fallback
        log.info("comparing as text (%s)", e)
        earlier = extract.extract(Attachment(filename=base_name, content_type="",
                                             size_bytes=len(base), content=base))
        comparison = compare.compare_text(earlier, job.doc)

    changes = [{"index": i, "kind": c.kind, "where": c.where,
                "before": _one_line(c.before, CHANGE_CHARS),
                "after": _one_line(c.after, CHANGE_CHARS)}
               for i, c in enumerate(comparison.changes)]
    author, cover = cover_note(job.email, job.user)
    return Evidence(tid, {
        "ours": {"filename": base_name, "how": how},
        "theirs": {"filename": job.att.filename, "label": version_label(job.att.filename)},
        "points": [p.as_evidence() for p in raised],
        "changes": changes,
        "cover": {"from": author, "text": cover},
    })


def evidence_block(job) -> tuple[str, bytes | None]:
    """What review.first_message and the mounted file carry, and the evidence
    kept for the reply. ("", None) when there is none. Never raises: the
    review goes ahead without it."""
    from secondeye.config import settings

    if not settings().sandbox_negotiation_skill_id.strip():
        return "", None     # the session could not record anything
    try:
        evidence = evidence_for(job)
        if evidence is None:
            return "", None
        _keep(job.job_id, evidence)
        return evidence.block(), evidence.file()
    except Exception:
        log.exception("could not build the negotiation evidence")
        return "", None


def _keep(job_id: str, evidence: Evidence) -> None:
    init()
    with connect() as c:
        old = c.execute("SELECT evidence FROM negotiation_evidence WHERE job_id = ?",
                        (job_id,)).fetchone()
        if old:
            blobs.discard(old[0])
        c.execute("INSERT OR REPLACE INTO negotiation_evidence (job_id, thread_id, evidence, "
                  "created_at) VALUES (?,?,?,?)",
                  (job_id, evidence.thread_id, blobs.stash(evidence.file(), "doc"),
                   datetime.now(UTC).isoformat()))


def _kept(job_id: str) -> Evidence | None:
    init()
    with connect() as c:
        row = c.execute("SELECT thread_id, evidence FROM negotiation_evidence WHERE job_id = ?",
                        (job_id,)).fetchone()
    if not row:
        return None
    return Evidence(row[0], json.loads(blobs.fetch(row[1])))


def _forget(job_id: str) -> None:
    with connect() as c:
        row = c.execute("SELECT evidence FROM negotiation_evidence WHERE job_id = ?",
                        (job_id,)).fetchone()
        if row:
            blobs.discard(row[0])
        c.execute("DELETE FROM negotiation_evidence WHERE job_id = ?", (job_id,))


_VERSION = re.compile(r"\b(v\s?\d+|version\s+\d+|draft\s+\d+|turn\s+\d+)\b", re.IGNORECASE)


def version_label(filename: str) -> str:
    """"v2" from "Falcon SPA v2.docx"; "their version" when the name says nothing."""
    m = _VERSION.search((filename or "").rsplit(".", 1)[0])
    return m.group(1).replace(" ", "").lower() if m else "their version"


# --------------------------------------------------------------------------
# The skill's rules, loaded from the skill so both sides check the same things
# --------------------------------------------------------------------------


@cache
def rules() -> ModuleType:
    from secondeye import skillsync

    path = skillsync.SKILLS / "lra-negotiation" / "scripts" / "rules.py"
    spec = importlib.util.spec_from_file_location("secondeye_negotiation_rules", path)
    if spec is None or spec.loader is None:  # pragma: no cover - the file ships in the image
        raise ImportError(f"cannot load {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def read_output(data: bytes | None) -> dict | None:
    """negotiation.json as the session left it, or None. Checked properly
    once the evidence is to hand (settle)."""
    if not data:
        return None
    try:
        value = json.loads(data)
    except ValueError:
        log.warning("%s is not JSON", OUTPUT_FILE)
        return None
    return value if isinstance(value, dict) else None


# --------------------------------------------------------------------------
# The reply
# --------------------------------------------------------------------------

STATUS_WORDS = {"accepted": "Accepted", "rejected": "Rejected",
                "partial": "Partly accepted", "changed": "Changed another way"}
VERDICT_WORDS = {"true": "True", "false": "Not true", "partly": "Partly true",
                 "unverifiable": "Can't tell from the documents"}
_RISK = {"high": 0, "medium": 1, "low": 2}
_NUMBERS = ("none", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine")


@dataclass
class Settled:
    """What the reply says about the negotiation."""

    outcome: dict | None = None
    ledger: dict[str, dict] = field(default_factory=dict)   # point id -> clause, asked, source
    version: str = "their version"
    cover: bool = False
    filename: str = ""
    notes: list[str] = field(default_factory=list)
    attachment: Attachment | None = None

    # -- the words ------------------------------------------------------------

    def _entries(self) -> list[dict]:
        return list(self.outcome["points"]) if self.outcome else []

    def _clause(self, entry: dict) -> str:
        raised = self.ledger.get(entry["id"].partition(".")[0], {})
        clause = raised.get("clause") or ""
        if entry.get("ask"):
            return f"{clause}: {entry['ask']}" if clause else f"Your email: {entry['ask']}"
        return clause or _one_line(raised.get("asked", ""), 80)

    def verdict(self) -> str:
        o = self.outcome or {}
        claims = [c for c in o.get("claims", []) if c["verdict"] != "unverifiable"]
        hidden = sorted((u for u in o.get("unrequested", [])
                         if not u["mentioned"] and u["risk"] in ("high", "medium")),
                        key=lambda u: _RISK[u["risk"]])[:3]
        doing = _join([u["summary"].strip().rstrip(".") for u in hidden])
        if any(c["verdict"] in ("false", "partly") for c in claims):
            head = "Not what they said."
            if doing:
                admitted = [e["label"] for e in self._entries()
                            if e["status"] != "accepted" and e["mentioned"]]
                besides = f"Besides {_join(admitted)}, " if admitted else ""
                head += f" {besides}{self.version} {doing}."
                head = head.replace(f". {self.version}", f". {_cap(self.version)}")
            return head
        if claims:
            return (f"Matches what they said, but {self.version} also {doing}." if doing
                    else "Matches what they said.")
        head = ("Their email doesn't say what they changed." if self.cover
                else "They sent no covering note.")
        return f"{head} {_cap(self.version)} {doing}." if doing else head

    def tally(self) -> str:
        entries = self._entries()
        n = len(entries)
        if not n:
            return ""
        accepted = sum(1 for e in entries if e["status"] == "accepted")
        missing = n - accepted
        silent = sum(1 for e in entries if e["status"] != "accepted" and not e["mentioned"])
        noun = "point" if n == 1 else "points"
        if not missing:
            return "Your one point was accepted." if n == 1 else f"All {n} of your points were accepted."
        if not accepted:
            line = ("Your one point was not accepted" if n == 1
                    else f"None of your {n} {noun} were accepted")
        else:
            line = (f"{accepted} of your {n} {noun} {'was' if accepted == 1 else 'were'} "
                    f"accepted; {missing} {'was' if missing == 1 else 'were'} not")
        if self.cover and silent:
            if silent == missing == 1:
                line += " — and it isn't mentioned in their email"
            elif silent == missing:
                line += " — and their email mentions none of them"
            else:
                word = _NUMBERS[silent] if silent < len(_NUMBERS) else str(silent)
                line += (f" — and {word} of those {'isn' if silent == 1 else 'aren'}'t "
                         "mentioned in their email")
        return line + "."

    def sections(self) -> list[tuple[str, list[str]]]:
        o = self.outcome or {}
        out: list[tuple[str, list[str]]] = []
        missed = [e for e in self._entries() if e["status"] != "accepted"]
        if missed:
            items = []
            for e in missed:
                said = "" if e["mentioned"] or not self.cover else ", not mentioned in their email"
                items.append(f"{self._clause(e)}. {STATUS_WORDS[e['status']]}{said}: "
                             f"{_sentence(e['evidence'])}")
            out.append(("Not accepted", items))
        unasked = sorted(o.get("unrequested", []), key=lambda u: _RISK[u["risk"]])
        if unasked:
            items = []
            for u in unasked:
                said = ", not mentioned in their email" if self.cover and not u["mentioned"] else ""
                items.append(f"{_clause_label(u['clause'])}: {u['summary'].strip().rstrip('.')}. "
                             f"{u['risk'].capitalize()} risk{said}.")
            out.append(("Changes you didn't ask for", items))
        claims = o.get("claims", [])
        if claims:
            out.append(("What their email says", [
                f"“{_one_line(c['quote'], 200)}” {VERDICT_WORDS[c['verdict']]}: "
                f"{_sentence(c['evidence'])}" for c in claims]))
        kept = [self._clause(e) for e in self._entries() if e["status"] == "accepted"]
        if kept:
            out.append(("Accepted", ["; ".join(kept) + "."]))
        return out

    def text(self) -> str:
        lines = [self.verdict()]
        tally = self.tally()
        if tally:
            lines.append(tally)
        lines.append("")
        for heading, items in self.sections():
            lines.append(heading + ":")
            lines += [f"  - {i}" for i in items]
            lines.append("")
        if self.attachment is not None:
            lines += [f"{self.attachment.filename} has every point with its status.", ""]
        return "\n".join(lines)

    def html(self) -> str:
        parts = [(f'<p style="font-size:17px;font-weight:600;margin:0 0 .4em">'
                  f"{escape(self.verdict())}</p>")]
        tally = self.tally()
        if tally:
            parts.append(f'<p style="margin:0 0 1em">{escape(tally)}</p>')
        for heading, items in self.sections():
            parts.append(f'<p style="margin:0 0 .3em;font-weight:600">{escape(heading)}</p>'
                         '<ul style="margin:0 0 1em;padding-left:1.2em">'
                         + "".join(f"<li>{escape(i)}</li>" for i in items) + "</ul>")
        if self.attachment is not None:
            parts.append(f'<p style="margin:0 0 1em">{escape(self.attachment.filename)} has '
                         "every point with its status.</p>")
        parts.append('<hr style="border:0;border-top:1px solid #ddd;margin:1.2em 0">')
        return "".join(parts)


def _cap(text: str) -> str:
    return text[:1].upper() + text[1:]


def _join(items: list[str]) -> str:
    items = [i for i in items if i]
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


def _sentence(text: str) -> str:
    text = _one_line(text, 400)
    return text if text[-1:] in ".?!”\"')" else text + "."


def _clause_label(clause: str) -> str:
    clause = (clause or "").strip()
    return f"Clause {clause}" if re.match(r"^\d", clause) else _cap(clause) or "Elsewhere"


def lead(out: OutboundEmail, settled: Settled) -> None:
    """Put the negotiation first in the reply: the verdict on their claims is
    the first line a phone shows."""
    if settled.outcome is None:
        return
    out.text_body = settled.text().rstrip("\n") + "\n\n" + out.text_body
    if out.html_body:
        cut = out.html_body.find(">") + 1
        out.html_body = out.html_body[:cut] + settled.html() + out.html_body[cut:]
    if settled.attachment is not None:
        out.attachments.append(settled.attachment)


# --------------------------------------------------------------------------
# Settling a round: check the file, keep the result, build the reply's part
# --------------------------------------------------------------------------


def settle(job, result) -> Settled | None:
    """The negotiation part of the reply to `job`'s review, or None when the
    review was given no negotiation evidence. Updates the ledger with what the
    model decided, and moves the ledger to this conversation. The file is
    checked against the evidence first, as the sandbox did: it came out of a
    sandbox that runs model-written code."""
    evidence = _kept(job.job_id)
    if evidence is None:
        return None
    ledger = {p["id"]: p for p in evidence.data["points"]}
    settled = Settled(ledger=ledger, version=evidence.data["theirs"]["label"],
                      cover=bool(evidence.data["cover"]["text"]),
                      filename=job.att.filename)
    raw = getattr(result, "negotiation", None)
    if raw is None:
        settled.notes.append(
            f"I couldn't settle this time which of the {len(ledger)} points we raised "
            "they accepted. Ask me \"did they accept our changes?\" to try again.")
        _forget(job.job_id)
        return settled
    outcome, problems = rules().check(raw, evidence.data)
    if outcome is None:
        log.warning("negotiation.json refused on the host: %s", "; ".join(problems[:5]))
        settled.notes.append(
            f"I read their version against the {len(ledger)} points we raised, but my "
            "account of it did not hold together, so I have left it out rather than guess.")
        _forget(job.job_id)
        return settled

    settled.outcome = outcome
    record_round(evidence.thread_id, job.thread_key, job.att.filename, outcome,
                 ledger, settled.cover)
    settled.attachment = status_list(settled)
    _forget(job.job_id)
    return settled


def record_round(ledger_thread: str, thread_key: str, filename: str, outcome: dict,
                 ledger: dict[str, dict], cover: bool) -> None:
    """Keep the round and each point's new status, and move the ledger to the
    conversation now holding the latest version."""
    init()
    round_number = rounds_done(ledger_thread) + 1
    now = datetime.now(UTC).isoformat()
    by_point: dict[str, list[dict]] = {}
    for entry in outcome["points"]:
        by_point.setdefault(entry["id"].partition(".")[0], []).append(entry)
    with connect() as c:
        for pid, entries in by_point.items():
            statuses = {e["status"] for e in entries}
            status = statuses.pop() if len(statuses) == 1 else "partial"
            evidence = " ".join(e["evidence"] for e in entries)[:2000]
            c.execute("UPDATE negotiation_points SET status = ?, status_round = ?, "
                      "evidence = ?, updated_at = ? WHERE id = ? AND thread_id = ?",
                      (status, round_number, evidence, now, int(pid[1:]), ledger_thread))
        c.execute("INSERT INTO negotiation_rounds (thread_id, round, filename, outcome, "
                  "points, cover, created_at) VALUES (?,?,?,?,?,?,?)",
                  (ledger_thread, round_number, filename, json.dumps(outcome),
                   json.dumps(ledger), 1 if cover else 0, now))
    move(ledger_thread, thread_key)


def last_round(thread_id: str) -> Settled | None:
    init()
    with connect() as c:
        row = c.execute("SELECT filename, outcome, points, cover FROM negotiation_rounds "
                        "WHERE thread_id = ? ORDER BY id DESC LIMIT 1", (thread_id,)).fetchone()
    if not row:
        return None
    settled = Settled(outcome=json.loads(row[1]), ledger=json.loads(row[2]),
                      version=version_label(row[0] or ""), cover=bool(row[3]),
                      filename=row[0] or "")
    settled.attachment = status_list(settled)
    return settled


# --------------------------------------------------------------------------
# The issues list again, with a status column
# --------------------------------------------------------------------------

COLUMNS = ("Clause", "What we asked", "Status", "Their version")


def status_name(filename: str) -> str:
    stem, _, _ = (filename or "document").rpartition(".")
    return f"{stem or filename} (issues list, updated).docx"


def status_list(settled: Settled) -> Attachment | None:
    """Every point with its status after their version, and what they changed
    unasked. Verified like every file we send; None, logged, if it fails."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    from secondeye.pipeline import issues_list, redline

    if not settled.outcome:
        return None
    try:
        document = Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2))
        props = document.core_properties
        props.author = props.last_modified_by = redline.configured_author()
        props.title = f"Issues list: {settled.filename}"
        props.comments = props.keywords = ""
        document.add_heading(f"Issues list: {settled.filename}", level=1)
        document.add_paragraph(settled.verdict())
        tally = settled.tally()
        if tally:
            document.add_paragraph(tally)

        def table(headings: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
            grid = document.add_table(rows=1, cols=len(headings))
            grid.style = "Table Grid"
            issues_list._repeat_as_header(grid.rows[0])
            for cell, heading in zip(grid.rows[0].cells, headings, strict=True):
                cell.text = ""
                cell.paragraphs[0].add_run(heading).bold = True
            for values in rows:
                for cell, value in zip(grid.add_row().cells, values, strict=True):
                    cell.text = value
            for row in grid.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(10)

        rows = []
        for e in settled.outcome["points"]:
            raised = settled.ledger.get(e["id"].partition(".")[0], {})
            status = STATUS_WORDS[e["status"]]
            if settled.cover and e["status"] != "accepted" and not e["mentioned"]:
                status += " (not mentioned in their email)"
            rows.append((settled._clause(e), _one_line(e.get("ask") or raised.get("asked", ""),
                                                       issues_list.QUOTE_LIMIT),
                         status, _one_line(e["evidence"], issues_list.QUOTE_LIMIT)))
        table(COLUMNS, rows)
        unasked = sorted(settled.outcome.get("unrequested", []), key=lambda u: _RISK[u["risk"]])
        if unasked:
            document.add_heading("Changes we did not ask for", level=2)
            table(("Clause", "What it does", "Risk", "In their email"),
                  [(_clause_label(u["clause"]), _one_line(u["summary"], 300),
                    u["risk"].capitalize(), "Yes" if u["mentioned"] else "No")
                   for u in unasked])
        buffer = BytesIO()
        document.save(buffer)
        content = buffer.getvalue()
    except Exception:
        log.exception("could not build the updated issues list")
        return None
    ok, why = redline.verify(content)
    if not ok:
        log.error("the updated issues list failed verification (%s); not attaching it", why)
        return None
    return Attachment(filename=status_name(settled.filename), content_type=DOCX_TYPE,
                      size_bytes=len(content), content=content)


# --------------------------------------------------------------------------
# What we raised this round, recorded once the reply is composed
# --------------------------------------------------------------------------


def raised_in(result, applied: list[Finding]) -> list[dict]:
    """The points a review puts to the other side: the issues list's rows on
    their paper, and the substantive changes we wrote as tracked changes."""
    from secondeye.pipeline import issues_list, their_paper

    out: list[dict] = []
    if result.their_paper:
        for f in their_paper.points(result.findings):
            out.append({"source": "issue", "clause": issues_list._clause(f),
                        "asked": f.our_position or f.response or f.suggested_text
                        or f.explanation or f.title,
                        "anchor": f.anchor})
    for f in applied:
        if f.severity not in (Severity.BLOCKER, Severity.SUBSTANTIVE):
            continue
        clause = issues_list._clause(f) or f.title
        if f.edit_kind == "insert_after":
            asked = f"Add: “{f.suggested_text}”"
        else:
            asked = f"Replace “{f.anchor}” with “{f.suggested_text}”"
        out.append({"source": "change", "clause": clause, "asked": asked,
                    "anchor": f.suggested_text or ""})
    return out


def record_review(job, result, applied: list[Finding]) -> None:
    """After a review's reply is composed: bring the document's ledger to
    this conversation, and add what this review raised. On a document the
    lawyer sent with us copied, that copy is the version we sent and their
    covering words are asks of their own. Never raises."""
    from secondeye import reconcile
    from secondeye.pipeline import identity
    from secondeye.pipeline.router import strip_reply

    try:
        sweep()
        text = reconcile._norm(" ".join(b.text for b in job.doc.blocks)) if job.doc else ""
        other = ledger_for(job.user, job.thread_key, job.att, text)
        if other and other != job.thread_key:
            move(other, job.thread_key)
        round_number = rounds_done(job.thread_key) + 1
        raised = raised_in(result, applied)
        if identity.external_recipients(job.email):
            # Written to the other side with us copied: its asks are points too.
            words = strip_reply(job.email.body or "").strip()
            if len(words) > 20:
                raised.append({"source": "email", "clause": "", "asked": words[:2000]})
        if job.entry is identity.EntryMode.BCC_SENT and (raised or points(job.thread_key)):
            record_sent(job.thread_key, job.att.filename, job.att.content)
        record_points(job.thread_key, round_number, raised)
    except Exception:
        log.exception("could not record the negotiation points")


# --------------------------------------------------------------------------
# "Did they accept our changes?", with nothing attached
# --------------------------------------------------------------------------


def answer(email: InboundEmail, thread_id: str | None) -> OutboundEmail | None:
    """The reply to the question on a conversation with a ledger, from the
    last version of theirs we read; None when there is no ledger here, and
    the message is handled as it would have been."""
    from secondeye.pipeline import reply

    if not thread_id or not points(thread_id):
        return None
    settled = last_round(thread_id)
    if settled is None:
        n = len(points(thread_id))
        text = (f"Nothing from them yet: I haven't seen a version of theirs since we raised "
                f"our {n} point{'s' if n != 1 else ''}. Forward it with their covering email "
                "when it comes and I'll check each point, and what their email says, against it.")
        body, html_body, attachments = text + "\n", reply.text_as_html(text), []
    else:
        body = settled.text().rstrip("\n") + "\n"
        html_body = ('<div style="font-family:-apple-system,Segoe UI,Roboto,sans-serif;'
                     'font-size:15px;line-height:1.5;color:#111;max-width:42em">'
                     + settled.html().rsplit("<hr", 1)[0] + "</div>")
        attachments = [settled.attachment] if settled.attachment else []
    return OutboundEmail(
        to=[email.from_address], subject=reply._subject(email.subject or "Negotiation"),
        text_body=body, html_body=html_body, in_reply_to=email.message_id,
        thread_id=email.thread_id, attachments=attachments)


__all__ = [
    "EVIDENCE_FILE", "OUTPUT_FILE", "TASK", "Evidence", "Point", "Settled", "answer",
    "asks_status", "cover_note", "evidence_block", "evidence_for", "lead", "ledger_for",
    "points", "read_output", "record_points", "record_review", "record_round",
    "record_sent", "rules", "settle", "status_list", "version_label",
]
