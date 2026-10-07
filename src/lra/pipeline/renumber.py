"""Renumbering typed clause numbers, and every reference to them.

The most-asked-for DocXtools feature, done narrowly enough to be provable.

Scope is the clause numbers a drafter typed as text: "4. Payment", "4.3 The
Supplier shall". Word's automatic numbering is left alone, because Word
renumbers that itself and a tracked change cannot touch a number that is not in
the text. Schedules and their Parts number their paragraphs in a space of their
own and are never renumbered; references in them to the body's clauses are.

What gets written, each as its own tracked change:

- every typed clause number that is out of sequence, top level and sub-clause,
  with the sub-clauses of a renumbered clause following it ("5.2" -> "4.2");
- every reference to a renumbered clause: "clause 5", "Clause 5.2", "clause
  5(b)", "Sections 5 and 6", "clauses 5 to 7".

What is never touched: amounts, dates, "Schedule 3", "paragraph 2",
"section 1.1502-6 of the Treasury Regulations", "section 5 of the Supply
Agreement", "Article 6 GDPR". The reference patterns and the external-instrument
test are the cross-reference check's own, so the two can never disagree about
what counts as a reference.

It refuses, with one sentence, rather than guess:

- a reference to a clause number that appears twice (which of them is meant?);
- a reference to a missing clause whose number another clause would take after
  renumbering (it would silently start pointing somewhere);
- a sub-clause that does not sit under its parent, or a family that does not
  start at .1 ("2.5 million shares" is not clause 2.5);
- a number that would need changing but is drawn by Word;
- any edit the writer declines, since half a renumber is worse than none.

And nothing is returned unless the output passes the proof gate: it opens by
the same test as every redline, its text is exactly the original with the
planned numbers swapped, and on a fresh extraction the numbering and
cross-reference checks are silent about everything this repaired.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document

from lra.models import Attachment
from lra.pipeline import checks, extract, redline
from lra.pipeline.ooxml import (
    AnchorNotSafelyEditable,
    RevisionWriter,
    _paragraph_text,
    _paragraphs_under,
    _story_parts,
)

# A top-level clause, exactly as checks.numbering reads one: a number, then a
# full stop or bracket, then text. Sub-clauses are "4.3 Text" or "4.3. Text",
# and never a figure with its unit after it.
_TOP = re.compile(r"^\s*([0-9]+)[.)]\s+\S")
_SUB = re.compile(
    r"^\s*([0-9]+(?:\.[0-9]+)+)\.?(?=\s+\S)"
    r"(?!\.?\s+(?:%|per\b|percent|million|billion|thousand|times|x\b))"
)
# "Article 6(1)(f) GDPR", "section 4 UCC": an instrument named without "of".
_ACRONYM_AFTER = re.compile(r"^(?:\([0-9A-Za-z]{1,4}\))*\s+(?:UK\s+|EU\s+)?[A-Z]{2,}\b")

_REFUSED = "so I have not sent a renumbered copy"


class CannotRenumber(Exception):
    """Carries a sentence we are willing to email back verbatim."""


@dataclass
class ReferenceChange:
    where: str        # "Clause 7", "Schedule 2", "Header"
    before: str       # "clause 5.2"
    after: str        # "clause 4.2"


@dataclass
class Renumbered:
    content: bytes
    filename: str
    clauses: list[tuple[str, str]] = field(default_factory=list)  # (old, new)
    references: list[ReferenceChange] = field(default_factory=list)
    word: str = "clause"
    was_wrong: list[str] = field(default_factory=list)

    @property
    def headline(self) -> str:
        """ "Numbering fixed: clauses 4–12 renumbered and 9 references updated,
        as tracked changes." """
        top = [new for _, new in self.clauses if "." not in new]
        shown = top or [new for _, new in self.clauses]
        if len(shown) == 1:
            span = f"{self.word} {shown[0]}"
        else:
            span = f"{self.word}s {shown[0]}–{shown[-1]}"
        n = sum(_count_numbers(r.before, r.after) for r in self.references)
        refs = ("no references needed updating" if n == 0 else
                f"{n} reference{'s' if n != 1 else ''} updated")
        return f"Numbering fixed: {span} renumbered and {refs}, as tracked changes."


def output_name(filename: str) -> str:
    stem, _, ext = filename.rpartition(".")
    stem = (stem or filename).replace(" (redline)", "")
    return f"{stem} (renumbered).{ext or 'docx'}"


# --------------------------------------------------------------------------
# Reading the structure
# --------------------------------------------------------------------------


@dataclass
class _Clause:
    para: int                      # index into the paragraph list
    old: tuple[int, ...]
    new: tuple[int, ...] = ()
    auto: bool = False
    span: tuple[int, int] = (0, 0)  # the typed number's characters


def _label(path: tuple[int, ...]) -> str:
    return ".".join(str(n) for n in path)


def _paragraphs(document) -> list[tuple[object, str]]:
    """Every paragraph in document order, with the story it belongs to: the
    body first (tables included), then headers, footers and notes."""
    out = [(p, "body") for p in _paragraphs_under(document.element.body, document)]
    for part in _story_parts(document):
        element = getattr(part, "element", None)
        if element is None:
            continue
        name = str(getattr(part, "partname", "")).rsplit("/", 1)[-1]
        story = ("Header" if "header" in name else
                 "Footer" if "footer" in name else "Footnote")
        out.extend((p, story) for p in _paragraphs_under(element, part))
    return out


def _in_table(para) -> bool:
    node = para._p.getparent()
    while node is not None:
        if node.tag.endswith("}tc"):
            return True
        node = node.getparent()
    return False


def _style(para) -> str:
    try:
        style = para.style
        return style.name if style is not None else ""
    except Exception:  # noqa: BLE001 - a dangling style id is not fatal
        return ""


def _structure(document, paras, word: str = "clause") -> tuple[list[_Clause], int]:
    """The body's clause numbers, old and new, and where the body ends.

    New numbers are sequential in document order: top-level from 1, each
    sub-clause family from .1 under its parent's new number.
    """
    numbers = extract._ListNumbers(document)
    clauses: list[_Clause] = []
    end = len(paras)
    top = 0
    # One entry per open depth: (old path, new path, children numbered so far).
    stack: list[list] = []
    for i, (para, story) in enumerate(paras):
        if story != "body":
            end = min(end, i)
            break
        label, _ = numbers.label(para._p)
        text = _paragraph_text(para)
        display = f"{label} {text}" if label else text
        if checks._ANNEX_TITLE.match(display) or checks._PART_TITLE.match(display):
            end = i
            break
        if _style(para).lower().startswith("toc"):
            raise CannotRenumber(
                "The document has a table of contents, which would still show the "
                f"old numbers, {_REFUSED}.")
        if _in_table(para):
            continue
        typed = not label
        m = _TOP.match(display)
        if m and int(m.group(1)) <= checks._MAX_CLAUSE_NUMBER:
            old = (int(m.group(1)),)
        else:
            m = _SUB.match(display)
            if not m:
                continue
            old = tuple(int(x) for x in m.group(1).split("."))
        depth = len(old)
        clause = _Clause(para=i, old=old, auto=not typed,
                         span=(m.start(1), m.end(1)) if typed else (0, 0))
        if depth == 1:
            top += 1
            clause.new = (top,)
            stack = [[old, clause.new, 0]]
        else:
            parent = stack[depth - 2] if len(stack) >= depth - 1 else None
            if parent is None or parent[0] != old[:-1]:
                raise CannotRenumber(
                    f"{word.capitalize()} {_label(old)} does not sit under "
                    f"{word} {_label(old[:-1])}, so I cannot tell how it is meant "
                    f"to be numbered, {_REFUSED}.")
            if parent[2] == 0 and old[-1] != 1:
                raise CannotRenumber(
                    f"I could not tell whether “{_label(old)}” at the start of a "
                    f"paragraph is a {word} number, {_REFUSED}.")
            parent[2] += 1
            clause.new = parent[1] + (parent[2],)
            stack = stack[: depth - 1] + [[old, clause.new, 0]]
        clauses.append(clause)
    return clauses, end


# --------------------------------------------------------------------------
# Planning the edits
# --------------------------------------------------------------------------


@dataclass
class _Ref:
    start: int
    end: int
    old: str
    new: str


def _references(text: str) -> list[tuple[re.Match, list[tuple[int, int, str]]]]:
    """Each clause reference in a paragraph, with the numbers in it: the
    cross-reference check's patterns and exclusions, exactly."""
    out = []
    for m in checks._SECTION_REF.finditer(text):
        if m.group(3) or m.group(1).lower() not in checks._BODY_REF_WORDS:
            continue
        if _external(text, m):
            continue
        out.append((m, [(m.start(2), m.end(2), m.group(2))]))
    for m in checks._PLURAL_REF.finditer(text):
        if _external(text, m):
            continue
        nums = [(m.start(2) + n.start(), m.start(2) + n.end(), n.group(0))
                for n in re.finditer(r"[0-9]+(?:\.[0-9]+)*", m.group(2))]
        out.append((m, nums))
    return out


def _external(text: str, m: re.Match) -> bool:
    rest = text[m.end():]
    return bool(checks._EXTERNAL_INSTRUMENT.match(rest) or _ACRONYM_AFTER.match(rest))


def _plan(document, word: str):
    paras = _paragraphs(document)
    clauses, body_end = _structure(document, paras, word)
    if not clauses:
        raise CannotRenumber(
            f"I could not find any typed {word} numbers to renumber in this document.")
    if all(c.auto for c in clauses):
        raise CannotRenumber(
            f"Your {word} numbers are automatic, so Word keeps them in sequence "
            "itself; there is nothing for me to renumber.")

    changed = [c for c in clauses if c.old != c.new]
    if not changed:
        return paras, clauses, body_end, {}, []
    for c in changed:
        if c.auto:
            raise CannotRenumber(
                f"{word.capitalize()} {_label(c.old)} is numbered automatically by "
                f"Word, so it cannot be renumbered as a tracked change, {_REFUSED}.")

    seen: dict[tuple[int, ...], int] = {}
    for c in clauses:
        seen[c.old] = seen.get(c.old, 0) + 1
    mapping = {c.old: c.new for c in clauses if seen[c.old] == 1}
    new_paths = {c.new for c in clauses}

    edits: dict[int, list[_Ref]] = {}
    for c in changed:
        edits.setdefault(c.para, []).append(
            _Ref(c.span[0], c.span[1], _label(c.old), _label(c.new)))

    by_para = {c.para: c for c in clauses}
    changes: list[ReferenceChange] = []
    where = f"Before {word} 1"
    for i, (para, story) in enumerate(paras):
        text = _paragraph_text(para)
        if i in by_para:
            where = f"{word.capitalize()} {_label(by_para[i].new)}"
        elif i >= body_end and story == "body":
            title = checks._ANNEX_TITLE.match(text)
            if title:
                where = f"{title.group(1).capitalize()} {title.group(2)}"
        if story != "body":
            where = story
        for m, numbers in _references(text):
            rewritten = []
            for start, end, target in numbers:
                path = tuple(int(x) for x in target.split("."))
                quoted = f"“{m.group(0)}” in {where.lower() if story != 'body' else where}"
                if path in mapping:
                    new = _label(mapping[path])
                    if new != target:
                        rewritten.append(_Ref(start, end, target, new))
                elif seen.get(path, 0) > 1:
                    raise CannotRenumber(
                        f"{word.capitalize()} {target} appears {seen[path]} times, so I "
                        f"cannot tell which one {quoted} means, {_REFUSED}.")
                elif path in new_paths:
                    raise CannotRenumber(
                        f"{quoted} points at a {word} that is not there, and after "
                        f"renumbering it would point at a different one, {_REFUSED}; "
                        "tell me what it should say.")
            if rewritten:
                edits.setdefault(i, []).extend(rewritten)
                after = m.group(0)
                for r in sorted(rewritten, key=lambda r: -r.start):
                    a, b = r.start - m.start(), r.end - m.start()
                    after = after[:a] + r.new + after[b:]
                changes.append(ReferenceChange(where, m.group(0), after))
    return paras, clauses, body_end, edits, changes


# --------------------------------------------------------------------------
# Writing, then proving
# --------------------------------------------------------------------------


def renumber(content: bytes, filename: str, author: str | None = None) -> Renumbered:
    """The document with its typed numbering and references fixed as tracked
    changes. Raises CannotRenumber, with a sentence, rather than guessing."""
    original_doc = _extract(content, filename)
    word = checks._clause_word(original_doc)
    try:
        document = Document(BytesIO(content))
    except Exception as e:
        raise CannotRenumber("I could not open the document to renumber it.") from e

    paras, clauses, _, edits, changes = _plan(document, word)
    if not edits:
        if checks.numbering(original_doc):
            raise CannotRenumber(
                f"The main {word}s already run in sequence; the numbering slip is in "
                "a schedule, which numbers its paragraphs separately and which I "
                "leave for you.")
        raise CannotRenumber(
            f"The {word} numbers already run in sequence, so there is nothing to renumber.")

    expected = [_paragraph_text(p) for p, _ in paras]
    writer = RevisionWriter(document, author=author or redline.configured_author())
    for i, refs in edits.items():
        para = paras[i][0]
        text = expected[i]
        for r in sorted(refs, key=lambda r: -r.start):
            try:
                writer.replace_span(para, r.start, r.end, r.new)
            except AnchorNotSafelyEditable as e:
                raise CannotRenumber(
                    f"I could not change “{r.old}” safely because {e}, {_REFUSED}.") from e
            except Exception as e:
                raise CannotRenumber(
                    f"I could not change “{r.old}” safely, {_REFUSED}.") from e
            text = text[:r.start] + r.new + text[r.end:]
        expected[i] = text

    buffer = BytesIO()
    document.save(buffer)
    out = buffer.getvalue()

    problem = _gate(content, out, filename, expected)
    if problem:
        raise CannotRenumber(f"I renumbered it, but {problem}, {_REFUSED}.")

    renamed = [(_label(c.old), _label(c.new)) for c in clauses if c.old != c.new]
    was_wrong = [f.title for f in checks.numbering(_body_only(original_doc))]
    return Renumbered(content=out, filename=output_name(filename), clauses=renamed,
                      references=changes, word=word, was_wrong=was_wrong)


def _gate(original: bytes, output: bytes, filename: str,
          expected: list[str]) -> str | None:
    """Why the output cannot be sent, or None."""
    ok, _ = redline.verify(output, expect_revisions=True)
    if not ok:
        return "the result did not pass the checks every attachment has to pass"

    written = [_paragraph_text(p) for p, _ in _paragraphs(Document(BytesIO(output)))]
    if written != expected:
        return "the result differed from yours by more than the numbers"

    before, after = _extract(original, filename), _extract(output, filename)
    if checks.numbering(_body_only(after)):
        return "the numbering check still found a slip afterwards"
    known = {(f.title, f.anchor) for f in checks.numbering(before)}
    if any((f.title, f.anchor) not in known for f in checks.numbering(after)):
        return "the numbering check found something new afterwards"
    broken = {f.title for f in checks.cross_references(before)}
    if any(f.title not in broken for f in checks.cross_references(after)):
        return "a cross-reference no longer pointed at a clause afterwards"
    again, _ = _structure(Document(BytesIO(output)), _paragraphs(Document(BytesIO(output))))
    if any(c.old != c.new for c in again):
        return "the numbers still did not run in sequence afterwards"
    return None


def _extract(content: bytes, filename: str) -> extract.ExtractedDoc:
    return extract.extract(Attachment(
        filename=filename, content_type="application/octet-stream",
        size_bytes=len(content), content=content,
    ))


def _body_only(doc: extract.ExtractedDoc) -> extract.ExtractedDoc:
    """The clauses before the first schedule, exhibit or Part title."""
    blocks = []
    for block in doc.blocks:
        if block.kind in ("header", "footer", "note"):
            continue
        if checks._ANNEX_TITLE.match(block.display) or checks._PART_TITLE.match(block.display):
            break
        blocks.append(block)
    return extract.ExtractedDoc(blocks=blocks, docx=None, filename=doc.filename)


def _count_numbers(before: str, after: str) -> int:
    """How many numbers in one reference changed: "Sections 5 and 6" -> 2."""
    old = re.findall(r"[0-9]+(?:\.[0-9]+)*", before)
    new = re.findall(r"[0-9]+(?:\.[0-9]+)*", after)
    return sum(1 for a, b in zip(old, new, strict=False) if a != b)
