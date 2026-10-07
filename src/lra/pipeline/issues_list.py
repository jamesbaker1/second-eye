"""The issues list: the other side's draft, point by point, against ours.

On their paper the useful deliverable is not a redline of their typos, it is
the table a lawyer takes into the call: clause, what their draft says, our
position, what we propose. Rendered here, deterministically, from the review's
substantive findings; the model supplies the words, this file decides the
shape, and nothing leaves without passing the same verification as a redline.
"""

from __future__ import annotations

import logging
import re
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from lra.models import Finding
from lra.pipeline import redline

log = logging.getLogger(__name__)

COLUMNS = ("Clause", "Their draft says", "Our position", "Proposed response")

# A quoted clause longer than this is clipped: the cell is for recognising the
# clause, and the document is beside it.
QUOTE_LIMIT = 400

_PLAYBOOK_PREFIX = re.compile(r"^\s*off[ -]playbook\s*:\s*", re.IGNORECASE)
_WHERE_PREFIX = re.compile(r"^\s*(?:cl(?:ause)?\.?|schedule|sch\.?|para(?:graph)?\.?)\s*"
                           r"[\w.()]+\s*:\s*", re.IGNORECASE)


def name(filename: str) -> str:
    """"<stem> (issues list).docx", whatever the document arrived as."""
    stem, _, _ = filename.rpartition(".")
    return f"{stem or filename} (issues list).docx"


def rows(findings: list[Finding]) -> list[tuple[str, str, str, str]]:
    """One row per point, in the order given."""
    return [(_clause(f), _their_text(f), _position(f), _response(f)) for f in findings]


def build(findings: list[Finding], filename: str) -> bytes | None:
    """The issues list as a Word file, or None if there is nothing to list or
    the file did not prove readable. A failure is logged and never raised: the
    email carries the same points either way."""
    table_rows = rows(findings)
    if not table_rows:
        return None
    try:
        content = _render(table_rows, filename)
    except Exception:
        log.exception("could not build the issues list")
        return None
    ok, why = redline.verify(content)
    if not ok:
        log.error("the issues list failed verification (%s); not attaching it", why)
        return None
    return content


def _render(table_rows: list[tuple[str, str, str, str]], filename: str) -> bytes:
    document = Document()
    section = document.sections[0]
    # Four columns of prose do not fit portrait.
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2))

    props = document.core_properties
    props.author = redline.configured_author()
    props.last_modified_by = redline.configured_author()
    props.title = f"Issues list: {filename}"
    props.comments = ""
    props.keywords = ""

    document.add_heading(f"Issues list: {filename}", level=1)
    document.add_paragraph(
        "The points in their draft that depart from our positions, with a "
        "proposed response to each.")

    table = document.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    header = table.rows[0]
    _repeat_as_header(header)
    for cell, heading in zip(header.cells, COLUMNS, strict=True):
        cell.text = ""
        run = cell.paragraphs[0].add_run(heading)
        run.bold = True
    for values in table_rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    widths = (Cm(4), Cm(8), Cm(6.5), Cm(7))
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _repeat_as_header(row) -> None:
    """The heading row repeats on every page, as a long list runs to several."""
    properties = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    properties.append(flag)


def _topic(f: Finding) -> str:
    """The clause's subject from the title: "Off playbook: Indemnity - one
    way" gives "Indemnity"."""
    title = _PLAYBOOK_PREFIX.sub("", f.title or "")
    title = _WHERE_PREFIX.sub("", title)
    return re.split(r"\s+[-–—]\s+|:\s", title, maxsplit=1)[0].strip()


def _clause(f: Finding) -> str:
    where, topic = (f.where or "").strip(), _topic(f)
    if where and topic and topic.lower() not in where.lower():
        return f"{where[0].upper()}{where[1:]}: {topic}"
    return (where[:1].upper() + where[1:]) if where else topic


def _their_text(f: Finding) -> str:
    text = " ".join((f.anchor or "").split())
    if len(text) > QUOTE_LIMIT:
        text = text[: QUOTE_LIMIT - 1].rsplit(" ", 1)[0] + "…"
    return f"“{text}”" if text else ""


def _position(f: Finding) -> str:
    return " ".join((f.our_position or f.explanation or "").split())


def _response(f: Finding) -> str:
    if f.response.strip():
        return " ".join(f.response.split())
    if f.suggested_text and f.suggested_text.strip():
        return f"Replace with: “{' '.join(f.suggested_text.split())}”"
    if f.question:
        return f.question.strip()
    return "Raise with them."
