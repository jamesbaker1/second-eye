"""A signature pack: the execution copy and each party's signature page.

The last job on a deal before it signs, and a junior's job: accept every
change, take the comments out, check nothing unfinished is left in, and pull
out a signature page per party headed "Signature page to [the Agreement]
dated [date]" so each signatory signs only their own page. It is also where a
mistake is most expensive, because what goes out is what gets signed.

So the pack is all or nothing. If anything unfinished is still in the
document, or the review on this conversation still has a question open, no
execution copy is made at all, and the reply says exactly what is in the way,
clause first. An execution copy with "[NTD: check with tax]" in it is worse
than none.

Everything is reused: `clean.clean` makes the execution copy (and proves the
words did not change), `checks.placeholders` finds what was never filled in,
and `checks._signature_blocks` finds the blocks. The pages are built from the
execution copy itself, so they carry the firm's styles, and are converted to
PDF when LibreOffice is installed.
"""

from __future__ import annotations

import copy
import logging
import re
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table

from lra.models import Attachment, InboundEmail, OutboundEmail
from lra.pipeline import checks, clean, extract
from lra.pipeline.extract import Block, ExtractedDoc
from lra.pipeline.filetype import Kind, identify

log = logging.getLogger(__name__)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _w(tag: str) -> str:
    return f"{{{W}}}{tag}"


class Unsignable(Exception):
    """The file cannot be made into an execution copy at all (it would not
    clean). A sentence for the lawyer."""


class Blocked(Exception):
    """The document is not ready to sign. Carries the blockers, clause first."""

    def __init__(self, blockers: list[str]) -> None:
        super().__init__("; ".join(blockers))
        self.blockers = blockers


@dataclass
class Party:
    name: str                       # as the signature block writes it
    short: str                      # "Acme Holdings", for the reply and file name
    key: tuple[str, str] | None
    positions: list[int] = field(default_factory=list)   # original block indices


@dataclass
class SignaturePage:
    party: str
    filename: str
    content: bytes
    content_type: str


@dataclass
class Pack:
    execution: bytes
    execution_name: str
    pages: list[SignaturePage]
    removed: list[str]
    heading: str                    # "Signature page to ... dated ..."
    pdf: bool


# --------------------------------------------------------------------------
# Entry point from the handler
# --------------------------------------------------------------------------


def handle(email: InboundEmail, provider, state=None) -> str:
    """Build the pack for this message's document, or the conversation's, and
    reply. Returns a status word for the job log."""
    from lra.pipeline import intake

    note = ""
    still_open: list[str] = []
    if email.attachments:
        att = intake.pick_document(email)
        filename, content = att.filename, att.content
    elif state is not None:
        from lra import thread
        from lra.handler import _rebuilt_note

        content, landed, _ = thread.rebuild(state)
        filename = state.filename
        note = _rebuilt_note(landed)
        still_open = [q.question for q in state.open_questions]
    else:
        raise intake.Rejection(
            "Attach the Word document and I will make the execution copy and "
            "signature pages."
        )

    if identify(content, filename) is not Kind.DOCX:
        raise intake.Rejection(
            "I can only make a signature pack from a Word document. Send me the "
            ".docx and I will make the execution copy and signature pages."
        )
    try:
        pack = build(content, filename, still_open)
    except Unsignable as e:
        raise intake.Rejection(str(e)) from e
    except Blocked as e:
        provider.send(_blocked_reply(email, e.blockers, note))
        return "signature pack blocked"
    provider.send(_ready_reply(email, pack, note))
    return "signature pack"


# --------------------------------------------------------------------------
# Building it
# --------------------------------------------------------------------------


@dataclass
class Survey:
    """One document read for signing: its execution copy, who signs it, and
    what is in the way. `build` refuses on a blocker; a closing across several
    documents (closing_docs.py) reports them per document instead."""

    execution: bytes
    doc: ExtractedDoc
    parties: list[Party]
    blockers: list[str]
    title: str
    date: str
    removed: list[str]
    # Each party's signature-block lines, as the page shows them.
    lines: dict[str, list[str]] = field(default_factory=dict)

    @property
    def heading(self) -> str:
        return f"Signature page to {self.title}" + (f" dated {self.date}" if self.date else "")


def survey(content: bytes, filename: str) -> Survey:
    """Clean the document and read its signature blocks. Raises Unsignable
    only when there is no execution copy to be had at all."""
    try:
        cleaned = clean.clean(content)
    except clean.CannotClean as e:
        raise Unsignable(str(e)) from e
    execution = cleaned.content

    doc = extract.extract(Attachment(filename=_docx_name(filename), content_type=DOCX_TYPE,
                                     size_bytes=len(execution), content=execution))
    lines, origin = _lines(doc)
    blocks = checks._signature_blocks(lines)
    blockers = _blockers(lines, blocks, execution)
    parties = _parties(lines, blocks, origin)
    texts = {p.name: [line for i in p.positions for line in doc.blocks[i].text.split("\n")
                      if line.strip()] for p in parties}
    return Survey(execution=execution, doc=doc, parties=parties, blockers=blockers,
                  title=agreement_title(doc, filename), date=agreement_date(doc),
                  removed=cleaned.removed, lines=texts)


def build(content: bytes, filename: str, still_open: list[str] | None = None) -> Pack:
    found = survey(content, filename)
    blockers = found.blockers + [f"Still open from my review: {q}" for q in (still_open or [])]
    if blockers:
        raise Blocked(blockers)
    pages, pdf = _pages(found.execution, found.doc, found.parties, found.heading, filename)
    return Pack(execution=found.execution, execution_name=_execution_name(filename),
                pages=pages, removed=found.removed, heading=found.heading, pdf=pdf)


def _lines(doc: ExtractedDoc) -> tuple[ExtractedDoc, list[int]]:
    """The document with each table cell split into its lines, and for each
    line the block it came from. A signature block set in a table is one cell
    of several lines, and the signature and placeholder checks read a block
    as one line."""
    blocks: list[Block] = []
    origin: list[int] = []
    for i, b in enumerate(doc.blocks):
        parts = b.text.split("\n") if b.kind == "table-cell" else [b.text]
        for n, text in enumerate(parts):
            blocks.append(Block(len(blocks), text, b.style, b.kind, b.number if n == 0 else ""))
            origin.append(i)
    return ExtractedDoc(blocks=blocks, docx=None, filename=doc.filename), origin


# "for and on behalf of" with nobody after it.
_NOBODY_ON_BEHALF = re.compile(
    r"\bon\s+behalf\s+of\s*[:\-]?\s*(?:_{2,}|\.{3,}|…|\[[^\]]*\])?\s*$", re.IGNORECASE)


def _blockers(lines: ExtractedDoc, blocks, execution: bytes) -> list[str]:
    """What stops this being signed, clause first."""
    in_signatures = {p for b in blocks for p in b.positions}
    labels = checks._locations(lines)
    out: list[str] = []

    def where(anchor: str) -> str:
        for b in lines.blocks:
            if anchor and anchor in b.text:
                if b.index in in_signatures:
                    return "Signature page"
                if b.kind in ("header", "footer"):
                    return b.kind.capitalize()
                label = labels.get(b.index) or "Front page"
                return label[:1].upper() + label[1:]
        return "Front page"

    for f in checks.placeholders(lines):
        shown = f.anchor if len(f.anchor) <= 70 else f.anchor[:66].rstrip() + " …]"
        if f.title.startswith("Drafting note"):
            out.append(f"{where(f.anchor)}: drafting note {shown} is still in")
        else:
            out.append(f"{where(f.anchor)}: {shown} is still blank")

    for text in _highlighted_brackets(execution):
        if not any(text in line for line in out):
            out.append(f"{where(text)}: highlighted {text} is still in")

    # The wrong company on a signature page, or at the address notices go to:
    # the one party-name slip that makes a signed document the wrong one.
    # The title already says where: "The Buyer's signature block names ...".
    out.extend(f.title for f in checks.party_drift(lines))

    for b in blocks:
        for p in b.positions:
            if _NOBODY_ON_BEHALF.search(lines.blocks[p].text):
                out.append("Signature page: a block signs on behalf of nobody; "
                           "the party's name is missing")
                break

    named = [b for b in blocks if b.entity]
    dated = [b for b in named if "date" in b.fields
             and checks._field_value_present(b.fields["date"])]
    undated = [b for b in named if "date" in b.fields
               and not checks._field_value_present(b.fields["date"])]
    if dated and undated:
        out.append("Signature page: the date is filled in for "
                   + _names([_short(b.entity) for b in dated]) + " but blank for "
                   + _names([_short(b.entity) for b in undated]))

    for f in checks.leftovers(execution):
        if "watermark" in f.title.lower():
            out.append("Header: the DRAFT watermark is still on")
    return list(dict.fromkeys(out))


def _highlighted_brackets(content: bytes) -> list[str]:
    """Bracketed text left highlighted: the drafter's mark for "not agreed"."""
    doc = Document(BytesIO(content))
    found: list[str] = []
    for p in doc.element.body.iter(_w("p")):
        runs = []
        for r in p.iter(_w("r")):
            text = "".join(t.text or "" for t in r.iter(_w("t")))
            rpr = r.find(_w("rPr"))
            lit = False
            if rpr is not None:
                hl = rpr.find(_w("highlight"))
                lit = hl is not None and hl.get(_w("val"), "none") != "none"
            runs.append((text, lit))
        text = "".join(t for t, _ in runs)
        lit_at: list[bool] = []
        for t, lit in runs:
            lit_at += [lit] * len(t)
        for m in re.finditer(r"\[[^\[\]\n]{0,120}\]", text):
            if any(lit_at[m.start():m.end()]):
                found.append(m.group(0))
    return list(dict.fromkeys(found))


# --------------------------------------------------------------------------
# Title and date
# --------------------------------------------------------------------------

_INSTRUMENT = re.compile(
    r"\b(?:AGREEMENT|DEED|CONTRACT|LEASE|LICEN[CS]E|GUARANTEE|GUARANTY|UNDERTAKING|"
    r"INSTRUMENT|CHARGE|DEBENTURE|NOTE|MEMORANDUM|LETTER|TERMS|PLAN|CONSENT|"
    r"RESOLUTIONS?|AMENDMENT|SUPPLEMENT|INDENTURE|WAIVER|RELEASE)\b",
    re.IGNORECASE,
)
_THIS_INSTRUMENT = re.compile(
    r"\bTHIS\s+((?:[A-Z][\w'-]*\s+){0,6}?(?:AGREEMENT|DEED|CONTRACT|LEASE|LICEN[CS]E|"
    r"GUARANTEE|GUARANTY|INDENTURE|AMENDMENT))\b",
    re.IGNORECASE,
)
_SMALL = {"of", "and", "the", "to", "for", "in", "on", "a", "an", "by", "or"}


def _title_case(text: str) -> str:
    if not text.isupper():
        return text
    words = text.lower().split()
    return " ".join(w if (i and w in _SMALL) else w[:1].upper() + w[1:]
                    for i, w in enumerate(words))


def agreement_title(doc: ExtractedDoc, filename: str) -> str:
    """The instrument's name as its front page gives it."""
    body = [b for b in doc.blocks[:25] if b.kind not in ("header", "footer", "note")]
    for b in body:
        text = " ".join(b.text.split())
        if not text or len(text) > 90:
            continue
        titled = b.style.startswith("Title") or (text.isupper() and len(text) >= 6)
        if titled and _INSTRUMENT.search(text) and not re.search(r"\bdated\b", text, re.IGNORECASE):
            return _title_case(text.rstrip("."))
    for b in body:
        m = _THIS_INSTRUMENT.search(b.text)
        if m:
            return _title_case(" ".join(m.group(1).split()))
    return "the Agreement"


_MONTHS = (r"(?:January|February|March|April|May|June|July|August|September|October|"
           r"November|December)")
_DATE = (r"(?:\d{1,2}(?:st|nd|rd|th)?\s+(?:day\s+of\s+)?" + _MONTHS + r",?\s+\d{4}|"
         + _MONTHS + r"\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{4}|\d{4}-\d{2}-\d{2})")
_DATED = re.compile(
    r"\b(?:dated|made\s+on|entered\s+into\s+(?:on|as\s+of)|as\s+of|date\s*:)\s+"
    r"(?:the\s+)?(" + _DATE + r")", re.IGNORECASE)


def agreement_date(doc: ExtractedDoc) -> str:
    """The date the front page gives, or "" when it is not filled in."""
    for b in doc.blocks[:25]:
        m = _DATED.search(b.text)
        if m:
            return " ".join(m.group(1).split())
    return ""


# --------------------------------------------------------------------------
# Parties and their pages
# --------------------------------------------------------------------------

_INDIVIDUAL = re.compile(
    r"^\s*(?:SIGNED|EXECUTED)\b(?:\s+as\s+a\s+deed)?(?:\s+and\s+delivered)?\s+by\s+"
    r"(.+?)(?:\s+in\s+the\s+presence\b|\s*$|,)", re.IGNORECASE)


def _short(name: str) -> str:
    m = checks._ENTITY_ANY_CASE.search(name)
    base = m.group(1) if m else name
    base = re.sub(r"^(?:for\s+and\s+)?on\s+behalf\s+of\s+", "", base, flags=re.IGNORECASE)
    return _title_case(" ".join(base.split()).strip(" ,.")) or name


def _parties(lines: ExtractedDoc, blocks, origin: list[int]) -> list[Party]:
    """One party per signing entity, its blocks together. A witness's block
    and a second director's belong to the party above them."""
    parties: list[Party] = []
    for block in blocks:
        positions = [origin[p] for p in block.positions]
        texts = [lines.blocks[p].text for p in block.positions]
        if block.entity_key:
            same = next((p for p in parties if p.key == block.entity_key), None)
            if same is not None:
                same.positions += positions
                continue
            parties.append(Party(block.entity, _short(block.entity), block.entity_key,
                                 positions))
            continue
        individual = next((m.group(1) for t in texts if (m := _INDIVIDUAL.match(t))), None)
        if individual and not re.fullmatch(r"[\s_.\[\]●]*", individual):
            name = " ".join(individual.split())
            parties.append(Party(name, _title_case(name), None, positions))
        elif parties:
            parties[-1].positions += positions
        else:
            name = (block.fields.get("name") or "").strip(" _")
            name = name or f"Signatory {len(parties) + 1}"
            parties.append(Party(name, _title_case(name), None, positions))
    for p in parties:
        p.positions = sorted(set(p.positions))
    return parties


def _body_elements(document) -> list[tuple[str, object]]:
    """The body in the order `extract` reads it: ("p", paragraph) or
    ("tc", cell). Must walk exactly as `extract._extract_docx` does, so that
    block i is element i."""
    out: list[tuple[str, object]] = []

    def table(tbl) -> None:
        seen: list = []
        for row in Table(tbl, document).rows:
            for cell in row.cells:
                if any(tc is cell._tc for tc in seen):
                    continue
                seen.append(cell._tc)
                out.append(("tc", cell._tc))

    def walk(container) -> None:
        for child in container:
            if not isinstance(child.tag, str):
                continue
            if child.tag == _w("p"):
                out.append(("p", child))
            elif child.tag == _w("tbl"):
                table(child)
            elif child.tag in (_w("sdt"), _w("sdtContent"), _w("customXml")):
                walk(child)

    walk(document.element.body)
    return out


def _preamble(doc: ExtractedDoc, elements, first: int) -> list[int]:
    """The "IN WITNESS WHEREOF" line just above the first block, which a
    signature page repeats."""
    for i in range(first - 1, max(-1, first - 4), -1):
        if i < len(elements) and elements[i][0] == "p" \
                and checks._EXECUTION_CONTEXT.search(doc.blocks[i].text):
            return [i]
    return []


def _pieces(party: Party, others: set[int], elements) -> list:
    """The body elements making up one party's page, in document order."""
    wanted = [i for i in party.positions if i < len(elements)]
    if not wanted:
        return []
    lo, hi = wanted[0], wanted[-1]
    pieces: list = []
    done_tables: set[int] = set()
    for i in range(lo, hi + 1):
        kind, el = elements[i]
        if kind == "p":
            if i in wanted or i not in others:
                pieces.append(el)
            continue
        if i not in wanted:
            continue
        tr = el.getparent()
        tbl = tr.getparent()
        if id(tbl) in done_tables:
            continue
        done_tables.add(id(tbl))
        mine = [elements[j][1] for j in wanted
                if elements[j][0] == "tc" and elements[j][1].getparent().getparent() is tbl]
        rows = []
        for tc in mine:
            if tc.getparent() not in rows:
                rows.append(tc.getparent())
        theirs = {id(elements[j][1]) for j in others
                  if j < len(elements) and elements[j][0] == "tc"}
        shared = any(id(tc) in theirs for row in rows for tc in row.iter(_w("tc")))
        if not shared:
            # The rows are this party's alone: keep the table's layout.
            table_copy = copy.deepcopy(tbl)
            all_rows = tbl.findall(_w("tr"))
            positions = {all_rows.index(r) for r in rows}
            for n, tr_copy in enumerate(table_copy.findall(_w("tr"))):
                if n not in positions:
                    table_copy.remove(tr_copy)
            pieces.append(("copied", table_copy))
        else:
            # Side by side with another party: this party's cells only.
            for tc in mine:
                pieces += [p for p in tc if isinstance(p.tag, str) and p.tag == _w("p")]
    return pieces


def page_documents(execution: bytes, doc: ExtractedDoc, parties: list[Party],
                   heading) -> list[tuple[Party, bytes]]:
    """Each party's signature page as a Word file. `heading` is the page's
    header text, or a function of the party giving it. Empty when the
    document cannot be split safely."""
    source = Document(BytesIO(execution))
    elements = _body_elements(source)
    body_count = sum(1 for b in doc.blocks if b.kind not in ("header", "footer", "note"))
    if len(elements) != body_count:
        # The two readings disagree, so a page could hold the wrong text.
        log.warning("signature pack: %d elements for %d blocks", len(elements), body_count)
        return []
    first = min((p.positions[0] for p in parties if p.positions), default=0)
    preamble = _preamble(doc, elements, first)

    built: list[tuple[Party, bytes]] = []
    for party in parties:
        others = {i for p in parties if p is not party for i in p.positions}
        pieces = [elements[i][1] for i in preamble] + _pieces(party, others, elements)
        if not pieces:
            continue
        text = heading(party) if callable(heading) else heading
        built.append((party, _page_docx(execution, pieces, text)))
    return built


def _pages(execution: bytes, doc: ExtractedDoc, parties: list[Party], heading: str,
           filename: str) -> tuple[list[SignaturePage], bool]:
    built = page_documents(execution, doc, parties, heading)
    pdf = False
    from lra import convert

    if built and convert.available():
        try:
            converted = [(p, convert._convert(c, "page.docx", "pdf")) for p, c in built]
        except convert.ConversionUnavailable as e:
            log.info("signature pages stay Word files: %s", e)
        else:
            built, pdf = converted, True

    stem = _stem(filename)
    ext, kind = ("pdf", PDF_TYPE) if pdf else ("docx", DOCX_TYPE)
    pages = [SignaturePage(party.short, f"{stem} - signature page ({_safe(party.short)}).{ext}",
                           content, kind) for party, content in built]
    return pages, pdf


# Relationships a signature page has no use for unless its text points at one.
_DROPPABLE = re.compile(
    r"/(?:header|footer|image|hyperlink|oleObject|package|chart|diagram\w*|customXml|"
    r"video|audio|media)$")


def _page_docx(execution: bytes, pieces: list, heading: str) -> bytes:
    """A document holding only these pieces, in the execution copy's styles,
    with the heading as its page header and nothing else of the agreement in
    the package: no other text, notes, headers or pictures."""
    doc = Document(BytesIO(execution))
    body = doc.element.body
    sect = body.find(_w("sectPr"))
    for child in list(body):
        if child is not sect:
            body.remove(child)
    for piece in pieces:
        element = piece[1] if isinstance(piece, tuple) else copy.deepcopy(piece)
        for inner in list(element.iter(_w("sectPr"))):
            inner.getparent().remove(inner)
        for br in list(element.iter(_w("br"))):
            if br.get(_w("type")) == "page":
                br.getparent().remove(br)
        for flag in list(element.iter(_w("pageBreakBefore"))):
            flag.getparent().remove(flag)
        if sect is not None:
            sect.addprevious(element)
        else:
            body.append(element)
    if sect is not None:
        for ref in list(sect):
            if ref.tag in (_w("headerReference"), _w("footerReference"), _w("titlePg")):
                sect.remove(ref)

    used = {v for el in body.iter() for k, v in el.attrib.items() if k.startswith(f"{{{R}}}")}
    rels = doc.part.rels
    for rid in [rid for rid, rel in rels.items()
                if _DROPPABLE.search(rel.reltype) and rid not in used]:
        del rels[rid]
        rels._target_parts_by_rId.pop(rid, None)
    notes = {v for el in body.iter(_w("footnoteReference"), _w("endnoteReference"))
             for k, v in el.attrib.items() if k == _w("id")}
    for rel in rels.values():
        if rel.is_external or not rel.reltype.endswith(("/footnotes", "/endnotes")):
            continue
        _keep_notes(rel.target_part, notes)

    section = doc.sections[-1]
    header = section.header
    header.is_linked_to_previous = False
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    para.text = heading
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _keep_notes(part, keep: set[str]) -> None:
    """Empty a notes part of every note the page does not refer to. The
    separators Word draws the notes line with stay."""
    from lxml import etree

    element = getattr(part, "_element", None)
    tree = element if element is not None else etree.fromstring(part.blob)
    for note in list(tree):
        if not isinstance(note.tag, str):
            continue
        kind = note.get(_w("type"))
        if kind in ("separator", "continuationSeparator", "continuationNotice"):
            continue
        if note.get(_w("id")) not in keep:
            tree.remove(note)
    if element is None:
        part._blob = etree.tostring(tree, xml_declaration=True, encoding="UTF-8",
                                    standalone=True)


# --------------------------------------------------------------------------
# Names and replies
# --------------------------------------------------------------------------


def _stem(filename: str) -> str:
    stem = filename.rsplit(".", 1)[0] if "." in filename else filename
    return re.sub(r"\s*\((?:redline|clean)\)", "", stem).strip() or "document"


def _safe(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "", name).strip() or "party"


def _docx_name(filename: str) -> str:
    return f"{_stem(filename)}.docx"


def _execution_name(filename: str) -> str:
    return f"{_stem(filename)} (execution).docx"


def _names(items: list[str]) -> str:
    if len(items) <= 2:
        return " and ".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


def _blocked_reply(email: InboundEmail, blockers: list[str], note: str) -> OutboundEmail:
    from lra.versions import _send

    n = len(blockers)
    verdict = f"Not ready to sign. {n} thing{'s' if n != 1 else ''} to settle first."
    sections: list = [("verdict", verdict)]
    if note:
        sections.append(("summary", note))
    sections.append(("findings", "In the way", [(b, "") for b in blockers]))
    sections.append(("summary", (
        "No execution copy or signature pages yet. Sort these out, reply "
        "\"sig pages\" with the document attached, and I will make them.")))
    return _send(email, verdict, sections, [])


def _ready_reply(email: InboundEmail, pack: Pack, note: str) -> OutboundEmail:
    from lra.versions import _send

    shorts = [p.party for p in pack.pages]
    n = len(pack.pages)
    if n:
        verdict = (f"Ready to sign: execution copy plus {n} signature "
                   f"page{'s' if n != 1 else ''} ({', '.join(shorts)}).")
    else:
        verdict = "Ready to sign: execution copy attached. No signature pages."
    sections: list = [("verdict", verdict)]
    if note:
        sections.append(("summary", note))
    if not n:
        sections.append(("summary", (
            "I could not find a signature block to lift out, so there are no "
            "signature pages. The execution copy is complete.")))

    attachments = [Attachment(filename=pack.execution_name, content_type=DOCX_TYPE,
                              size_bytes=len(pack.execution), content=pack.execution)]
    attachments += [Attachment(filename=p.filename, content_type=p.content_type,
                               size_bytes=len(p.content), content=p.content)
                    for p in pack.pages]
    attachments, dropped = _fit(attachments)

    if pack.removed:
        sections.append(("findings", "Taken out of the execution copy",
                         [(r, "") for r in pack.removed]))
    said = "Not a word of the text changed."
    if n:
        said += f" Each signature page is headed \"{pack.heading}\"."
    sections.append(("summary", said))

    notes: list[str] = []
    if n and not pack.pdf:
        notes.append("The signature pages are Word files; PDF conversion isn't available here.")
    for att in dropped:
        mb = att.size_bytes / (1024 * 1024)
        notes.append(f"{att.filename} came to {mb:.0f} MB, too large to attach with "
                     "the rest, so it is not here.")
    sections.append(("notes", "", [(t, "") for t in notes]))
    return _send(email, verdict, sections, attachments)


def _fit(attachments: list[Attachment]) -> tuple[list[Attachment], list[Attachment]]:
    """Everything that fits under the send limit, dropping the largest first."""
    from lra.pipeline import reply

    kept = list(attachments)
    dropped: list[Attachment] = []
    while kept and sum(a.size_bytes for a in kept) > reply.MAX_ATTACHMENT_BYTES:
        largest = max(kept, key=lambda a: a.size_bytes)
        kept.remove(largest)
        dropped.append(largest)
    return kept, dropped

