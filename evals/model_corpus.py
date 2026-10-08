"""Documents with known answers for the model-dependent features.

evals/corpus.py scores the deterministic checks, which cannot be wrong in the
way a model can. This corpus is for the parts that are judgment: reading a
contract against the playbook, answering the other side's draft with an
issues list, and extracting key terms. Every document here is generated, and
every one carries its answer, so the first live run is scored rather than
read and remembered as an impression (evals/score_model.py).

Five sets:

  playbook     services agreements from the firm's client's side, each with
               one to three clauses planted off the starter playbook
               (skills/lra-playbook/positions). Scored on recall.
  silence      the same agreement with every clause on position. Any finding
               is a false positive, and false positives are existential
               (PRODUCT.md), so this set is scored as hard as recall.
  their_paper  the supplier's own draft, forwarded as theirs, with positions
               to push back on and ordinary typos that must not be marked up.
  key_terms    agreements whose key terms (parties, term, renewal, cap,
               governing law, notices) are known exactly.
  realistic    the three clean agreements in evals/realistic.py, drafted the
               way firms draft, for the false-positive rate on documents this
               file did not write.

One template does all of it. Every clause has an on-position text, taken from
the starter playbook's Position section, and named off-position variants taken
from its Walk-away section. A variant replaces the whole clause body, so the
numbering, the defined terms and the cross-references are identical across the
corpus and the mechanical checks stay silent on all of it (pinned by
tests/test_score_model.py): a finding on these documents is the model's.

One deliberate departure from the obvious plant. The starter playbook's
governing-law fallback accepts the law of England and Wales or the State of
New York, so a New York clause is at fallback, not off it, and a reviewer that
stays quiet about it is right. The planted governing-law clause is the
supplier's home law and courts (Texas), which is the file's walk-away.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from docx.shared import Pt

from evals import realistic

# --------------------------------------------------------------------------
# The clauses the playbook has a position on, and how a finding names them.
# --------------------------------------------------------------------------

# playbook file stem -> (clause label, words a title uses for that clause).
# A finding is matched to a planted clause by its category (`playbook`) and by
# one of these words in its title, or by its anchor falling inside the planted
# clause's text.
CLAUSES: dict[str, tuple[str, tuple[str, ...]]] = {
    "limitation-of-liability": ("Limitation of liability",
                                ("liability", "cap", "uncapped", "unlimited")),
    "indemnity": ("Indemnity", ("indemnit",)),
    "governing-law-and-disputes": ("Governing law and disputes",
                                   ("governing law", "jurisdiction", "law", "disputes",
                                    "courts", "arbitration")),
    "assignment-and-change-of-control": ("Assignment and change of control",
                                         ("assign", "novat", "subcontract",
                                          "change of control")),
    "payment-terms": ("Payment terms", ("payment", "charges", "price", "invoice",
                                        "increase")),
    "confidentiality": ("Confidentiality", ("confidential",)),
    "termination": ("Termination", ("terminat", "convenience")),
    "data-protection": ("Data protection", ("data protection", "personal data",
                                            "breach notif", "processor", "data")),
    "insurance": ("Insurance", ("insurance", "insur")),
    "intellectual-property": ("Intellectual property", ("intellectual property", "ip ",
                                                        "ip:", "ownership", "deliverables",
                                                        "licen")),
    "non-solicitation": ("Non-solicitation", ("solicit",)),
    "force-majeure": ("Force majeure", ("force majeure",)),
}


@dataclass(frozen=True)
class Parties:
    customer: str = "Kestrel Retail Limited"
    customer_short: str = "Kestrel"
    customer_number: str = "08123456"
    customer_office: str = "14 Wharf Road, London N1 7GR"
    customer_email: str = "legal@kestrelretail.co.uk"
    supplier: str = "Harbour Analytics Inc."
    supplier_office: str = "600 Congress Avenue, Austin, Texas 78701"
    supplier_email: str = "contracts@harbouranalytics.com"
    services: str = "retail demand forecasting services"
    date: str = "12 May 2026"
    effective: str = "1 June 2026"
    initial_term: str = "three years"
    renewal: str = "successive periods of 12 months"
    renewal_notice: str = "at least three months"
    floor: str = "£1,000,000"


# --------------------------------------------------------------------------
# The agreement. Each clause is (number, heading, key, paragraphs); the key
# names the playbook file whose position the on-position text follows.
# --------------------------------------------------------------------------


def _clauses(p: Parties) -> list[tuple[str, str, str, list[str]]]:
    return [
        ("1", "Definitions", "", [
            "1.1 In this Agreement:",
            "“Charges” means the charges for the Services set out in the Schedule.",
            ("“Confidential Information” means any information of a confidential nature "
            "disclosed by one party to the other in connection with this Agreement."),
            ("“Contract Year” means each period of 12 months beginning on the Effective "
            "Date or an anniversary of it."),
            f"“Effective Date” means {p.effective}.",
            ("“Deliverables” means everything the Supplier creates for the Customer in "
            "performing the Services."),
            "“Personal Data” has the meaning given in the UK GDPR.",
            f"“Services” means the {p.services} described in the Schedule.",
        ]),
        ("2", "Term", "", [
            (f"2.1 This Agreement starts on the Effective Date and continues for an "
            f"initial term of {p.initial_term}."),
            f"2.2 This Agreement then renews automatically for {p.renewal}.",
            (f"2.3 Either party may prevent a renewal by written notice given "
            f"{p.renewal_notice} before the end of the current term."),
        ]),
        ("3", "Charges and payment", "payment-terms", [
            ("3.1 The Customer shall pay each valid invoice within 45 days of receiving "
            "it."),
            ("3.2 The Charges are fixed for the initial term. After that the Supplier may "
            "increase them once in each Contract Year on 90 days' written notice, by no "
            "more than the lower of CPI and 3%."),
            ("3.3 The Customer may withhold any amount it disputes in good faith, without "
            "interest and without any suspension of the Services."),
        ]),
        ("4", "Intellectual property", "intellectual-property", [
            "4.1 The Customer owns the Deliverables and all of its own data and inputs.",
            ("4.2 The Supplier keeps its pre-existing materials and grants the Customer a "
            "perpetual, royalty-free licence to use them as part of the Deliverables."),
        ]),
        ("5", "Confidentiality", "confidentiality", [
            ("5.1 Each party shall keep the other's Confidential Information "
            "confidential and use it only to perform this Agreement."),
            ("5.2 This clause survives for five years after this Agreement ends, and "
            "indefinitely for trade secrets and Personal Data."),
        ]),
        ("6", "Data protection", "data-protection", [
            ("6.1 The Supplier processes Personal Data as processor, only on the "
            "Customer's documented instructions and on the terms in the Schedule."),
            ("6.2 The Supplier shall notify the Customer of any breach of security affecting Personal Data "
            "within 24 hours of becoming aware of it, and shall appoint a sub-processor "
            "only after giving the Customer notice and a right to object."),
        ]),
        ("7", "Indemnity", "indemnity", [
            ("7.1 The Supplier shall indemnify the Customer against any third-party "
            "claim that the Services or Deliverables infringe intellectual property "
            "rights, and against losses from its breach of clause 5 or clause 6."),
            ("7.2 The Supplier shall also indemnify the Customer against death, personal "
            "injury or property damage caused by its negligence."),
        ]),
        ("8", "Limitation of liability", "limitation-of-liability", [
            ("8.1 The Supplier's total liability in each Contract Year is limited to the "
            "greater of 200% of the Charges paid or payable in the 12 months before the "
            f"claim and {p.floor}."),
            "8.2 The Customer's total liability is limited to the Charges due.",
            ("8.3 Nothing in this Agreement limits liability under clause 7, for breach "
            "of clause 5 or clause 6, for wilful default, or for anything that cannot "
            "be limited by law."),
        ]),
        ("9", "Insurance", "insurance", [
            ("9.1 The Supplier shall maintain professional indemnity, public liability, "
            "employer's liability and cyber insurance, each for at least the amount of "
            "its liability under clause 8, for the term and six years after it."),
        ]),
        ("10", "Termination", "termination", [
            ("10.1 The Customer may terminate this Agreement for convenience on 30 "
            "days' written notice, with no termination fee."),
            ("10.2 Either party may terminate this Agreement for a material breach not "
            "remedied within 30 days of notice, and immediately if the other becomes "
            "insolvent."),
        ]),
        ("11", "Non-solicitation", "non-solicitation", [
            ("11.1 Neither party shall solicit any employee of the other who worked on "
            "the Services, for six months after they last did so. Responding to a "
            "general advertisement is not solicitation."),
        ]),
        ("12", "Force majeure", "force-majeure", [
            ("12.1 Neither party is liable for a delay caused by fire, flood, war, "
            "epidemic or government action beyond its reasonable control. This does "
            "not cover a failure of the Supplier's subcontractors or a dispute with its "
            "own staff."),
            ("12.2 If such an event continues for 30 days the Customer may terminate "
            "this Agreement and receive a pro-rata refund of Charges paid in advance."),
        ]),
        ("13", "Assignment", "assignment-and-change-of-control", [
            ("13.1 The Supplier shall not assign, novate or subcontract this Agreement "
            "without the Customer's prior written consent."),
            ("13.2 The Customer may assign this Agreement to an affiliate or to a "
            "successor to its business on notice, and may terminate it if the Supplier "
            "undergoes a change of control."),
        ]),
        ("14", "Notices", "", [
            (f"14.1 Notices shall be in writing and sent by email: to the Customer at "
            f"{p.customer_email}, and to the Supplier at {p.supplier_email}."),
        ]),
        ("15", "Governing law and disputes", "governing-law-and-disputes", [
            ("15.1 This Agreement is governed by the law of England and Wales, and the "
            "courts of England have exclusive jurisdiction."),
            ("15.2 Before starting proceedings the parties shall refer a dispute to "
            "their senior executives for 20 Business Days, without prejudice to either "
            "party's right to seek an injunction."),
        ]),
    ]


# Off-position text for a clause, keyed (clause, variant). Each is the
# playbook file's walk-away, drafted as a supplier would draft it.
OFF: dict[tuple[str, str], list[str]] = {
    ("limitation-of-liability", "customer uncapped"): [
        ("8.1 The Supplier's total liability in each Contract Year is limited to the "
        "greater of 200% of the Charges paid or payable in the 12 months before the "
        "claim and £1,000,000."),
        ("8.2 The Customer's liability under or in connection with this Agreement is "
        "unlimited."),
        ("8.3 Nothing in this Agreement limits liability under clause 7, for breach "
        "of clause 5 or clause 6, for wilful default, or for anything that cannot "
        "be limited by law."),
    ],
    ("limitation-of-liability", "low cap with indemnities inside"): [
        ("8.1 The Supplier's total liability under this Agreement, including under "
        "clause 7 and for breach of clause 5 or clause 6, is limited to 25% of the "
        "Charges paid in the three months before the claim."),
        "8.2 The Customer's total liability is limited to the Charges due.",
        ("8.3 Nothing in this Agreement limits liability for anything that cannot be "
        "limited by law."),
    ],
    ("indemnity", "one-way from the customer"): [
        ("7.1 The Customer shall indemnify the Supplier against all losses, claims and "
        "costs arising out of or in connection with this Agreement or the Services."),
        "7.2 The Supplier gives no indemnity under this Agreement.",
    ],
    ("governing-law-and-disputes", "supplier's home law and courts"): [
        ("15.1 This Agreement is governed by the law of the State of Texas, and the "
        "state and federal courts in Travis County, Texas have exclusive "
        "jurisdiction."),
        ("15.2 Before starting proceedings the parties shall refer a dispute to "
        "their senior executives for 20 Business Days, without prejudice to either "
        "party's right to seek an injunction."),
    ],
    ("assignment-and-change-of-control", "supplier free to assign"): [
        ("13.1 The Supplier may assign, novate or subcontract any of its rights and "
        "obligations under this Agreement to any person without the Customer's "
        "consent."),
        ("13.2 The Customer shall not assign this Agreement, including to a successor "
        "to its business, without the Supplier's prior written consent."),
    ],
    ("payment-terms", "increases at the supplier's discretion"): [
        ("3.1 The Customer shall pay the Charges for each Contract Year annually in "
        "advance."),
        "3.2 The Supplier may increase the Charges at any time at its discretion.",
        ("3.3 The Supplier may suspend the Services while any invoice, including a "
        "disputed invoice, remains unpaid."),
    ],
    ("confidentiality", "one-way and short"): [
        ("5.1 The Customer shall keep the Supplier's Confidential Information "
        "confidential and use it only to perform this Agreement."),
        "5.2 This clause survives for one year after this Agreement ends.",
    ],
    ("termination", "no convenience right for the customer"): [
        ("10.1 The Supplier may terminate this Agreement for convenience on seven "
        "days' written notice. The Customer has no right to terminate for "
        "convenience."),
        ("10.2 Either party may terminate this Agreement for a material breach not "
        "remedied within 90 days of notice."),
    ],
    ("data-protection", "slow breach notice and free sub-processing"): [
        ("6.1 The Supplier processes Personal Data as processor on the terms in the "
        "Schedule."),
        ("6.2 The Supplier shall notify the Customer of any breach of security affecting Personal Data "
        "within 14 days of becoming aware of it, and may appoint any sub-processor "
        "without notice to the Customer."),
    ],
    ("insurance", "cover at the supplier's discretion"): [
        ("9.1 The Supplier shall maintain such insurance as the Supplier considers "
        "appropriate."),
    ],
    ("intellectual-property", "supplier owns everything"): [
        ("4.1 The Supplier owns the Deliverables and all intellectual property "
        "derived from the Customer's data."),
        ("4.2 The Customer grants the Supplier a perpetual, irrevocable licence to "
        "use the Customer's data for any purpose."),
    ],
    ("non-solicitation", "all staff, two years"): [
        ("11.1 The Customer shall not employ or solicit any employee of the Supplier "
        "for 24 months after this Agreement ends, and shall pay the Supplier a sum "
        "equal to 18 months of that employee's salary if it does."),
    ],
    ("force-majeure", "covers subcontractor failure"): [
        ("12.1 Neither party is liable for a delay caused by any event beyond its "
        "reasonable control, including any failure of a supplier or subcontractor."),
        ("12.2 The Customer shall continue to pay the Charges in full during any "
        "such event."),
    ],
}


# --------------------------------------------------------------------------
# Specs
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Planted:
    """One clause planted off position: the finding a review owes."""

    clause: str        # playbook file stem, a key of CLAUSES
    variant: str       # key into OFF with the clause
    text: str = ""     # the planted clause's text, filled in by the spec

    @property
    def label(self) -> str:
        return CLAUSES[self.clause][0]

    @property
    def aliases(self) -> tuple[str, ...]:
        return CLAUSES[self.clause][1]


@dataclass
class ModelSpec:
    name: str
    kind: str                            # playbook | silence | their_paper | key_terms | realistic
    parties: Parties = field(default_factory=Parties)
    planted: list[Planted] = field(default_factory=list)
    # (as written in the document, the correct spelling). On their paper these
    # are typos that must NOT be marked up.
    typos: list[tuple[str, str]] = field(default_factory=list)
    # field -> fragments the answer must contain, each quoted from the document.
    key_terms: dict[str, list[str]] = field(default_factory=dict)
    builder: Callable[[], bytes] | None = None
    note: str = ""

    @property
    def filename(self) -> str:
        return f"{self.name}.docx"

    def paragraphs(self) -> list[str]:
        if self.builder is not None:
            return []
        p = self.parties
        off = {pl.clause: OFF[(pl.clause, pl.variant)] for pl in self.planted}
        out = [
            f"THIS SERVICES AGREEMENT is dated {p.date}.",
            "PARTIES",
            (f"(1) {p.customer}, a company incorporated in England and Wales with company "
            f"number {p.customer_number} whose registered office is at "
            f"{p.customer_office} (the “Customer”); and"),
            (f"(2) {p.supplier}, a Delaware corporation whose principal office is at "
            f"{p.supplier_office} (the “Supplier”)."),
            "AGREED TERMS",
        ]
        for number, heading, key, body in _clauses(p):
            out.append(f"{number}. {heading}")
            out.extend(off.get(key, body))
        out += [
            "SCHEDULE",
            (f"The Services are the {p.services}. The Charges are £240,000 a year, "
            "invoiced quarterly in arrears. The data processing terms are those in "
            "the Customer's standard processing addendum."),
            ("This Agreement has been entered into on the date stated at the beginning "
            "of it."),
            f"Signed for and on behalf of {p.customer}",
            "Name: Alison Grey",
            "Title: General Counsel",
            f"Signed for and on behalf of {p.supplier}",
            "Name: Daniel Reyes",
            "Title: Chief Executive Officer",
        ]
        for wrong, right in self.typos:
            # The first use of the word, whole-word, is the one misspelled.
            pattern = re.compile(rf"\b{re.escape(right)}\b")
            for i, text in enumerate(out):
                if pattern.search(text):
                    out[i] = pattern.sub(wrong, text, count=1)
                    break
            else:
                raise ValueError(f"{self.name}: no {right!r} to misspell")
        return out

    def planted_with_text(self) -> list[Planted]:
        return [Planted(pl.clause, pl.variant, " ".join(OFF[(pl.clause, pl.variant)]))
                for pl in self.planted]

    def build(self) -> bytes:
        if self.builder is not None:
            return self.builder()
        d = Document()
        style = d.styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        for text in self.paragraphs():
            d.add_paragraph(text)
        core = d.core_properties
        core.author = ""
        core.last_modified_by = ""
        buffer = BytesIO()
        d.save(buffer)
        return buffer.getvalue()


def _pl(clause: str, variant: str) -> Planted:
    return Planted(clause, variant)


PLAYBOOK: list[ModelSpec] = [
    ModelSpec("pb_customer_uncapped", "playbook",
              planted=[_pl("limitation-of-liability", "customer uncapped")]),
    ModelSpec("pb_one_way_indemnity", "playbook",
              planted=[_pl("indemnity", "one-way from the customer")]),
    ModelSpec("pb_texas_law", "playbook",
              planted=[_pl("governing-law-and-disputes", "supplier's home law and courts")],
              note="Texas, not New York: New York is the starter playbook's fallback."),
    ModelSpec("pb_free_assignment", "playbook",
              planted=[_pl("assignment-and-change-of-control", "supplier free to assign")]),
    ModelSpec("pb_discretionary_increases", "playbook",
              planted=[_pl("payment-terms", "increases at the supplier's discretion")]),
    ModelSpec("pb_one_way_confidentiality", "playbook",
              planted=[_pl("confidentiality", "one-way and short")]),
    ModelSpec("pb_three_operational", "playbook", planted=[
        _pl("termination", "no convenience right for the customer"),
        _pl("data-protection", "slow breach notice and free sub-processing"),
        _pl("insurance", "cover at the supplier's discretion"),
    ]),
    ModelSpec("pb_three_commercial", "playbook", planted=[
        _pl("intellectual-property", "supplier owns everything"),
        _pl("non-solicitation", "all staff, two years"),
        _pl("force-majeure", "covers subcontractor failure"),
    ]),
]

SILENCE: list[ModelSpec] = [
    ModelSpec("on_position_kestrel", "silence"),
    ModelSpec("on_position_ashdown", "silence", parties=Parties(
        customer="Ashdown Building Society", customer_short="Ashdown",
        customer_number="IP00231R", customer_office="2 Castle Street, Lewes BN7 1AA",
        customer_email="legal@ashdownbs.co.uk", supplier="Corvid Ledger Inc.",
        supplier_office="1 Market Street, San Francisco, California 94105",
        supplier_email="notices@corvidledger.com",
        services="payment reconciliation services")),
    ModelSpec("on_position_merlin", "silence", parties=Parties(
        customer="Merlin Outdoor Limited", customer_number="05512890",
        customer_office="Unit 4, Parkway, Leeds LS15 8GB",
        customer_email="contracts@merlinoutdoor.co.uk", supplier="Tessellate Labs Inc.",
        supplier_email="legal@tessellatelabs.com", services="warehouse routing services",
        date="3 August 2026", effective="1 September 2026")),
    ModelSpec("on_position_orchard", "silence", parties=Parties(
        customer="Orchard Health Partners Limited", customer_number="11002345",
        customer_office="80 Broad Street, Bristol BS1 2EP",
        customer_email="legal@orchardhealth.co.uk", supplier="Pinecrest Software Inc.",
        supplier_email="agreements@pinecrest.io",
        services="appointment scheduling services", initial_term="two years")),
    ModelSpec("on_position_fenwick", "silence", parties=Parties(
        customer="Fenwick & Rowe Limited", customer_number="04477120",
        customer_office="9 Queen Street, Edinburgh EH2 1JQ",
        customer_email="legal@fenwickrowe.co.uk", supplier="Lumen Analytics Inc.",
        supplier_email="legal@lumenanalytics.com",
        services="pricing analytics services", floor="£2,000,000")),
]

# The supplier's own draft: the playbook's walk-aways, which is what a
# supplier's form usually says, and ordinary typos a supplier leaves in. The
# typos are common words, not defined terms, so the mechanical checks do not
# see them and any mark-up of them is the model's.
THEIR_PAPER: list[ModelSpec] = [
    ModelSpec("their_draft_harbour", "their_paper",
              planted=[_pl("limitation-of-liability", "low cap with indemnities inside"),
                       _pl("governing-law-and-disputes", "supplier's home law and courts")],
              typos=[("recieving", "receiving"), ("confidencial", "confidential"),
                     ("maintian", "maintain")]),
    ModelSpec("their_draft_corvid", "their_paper",
              parties=Parties(supplier="Corvid Ledger Inc.",
                              supplier_email="notices@corvidledger.com",
                              services="payment reconciliation services"),
              planted=[_pl("indemnity", "one-way from the customer"),
                       _pl("assignment-and-change-of-control", "supplier free to assign"),
                       _pl("payment-terms", "increases at the supplier's discretion")],
              typos=[("perpetural", "perpetual"), ("remedeid", "remedied"),
                     ("contiunes", "continues")]),
    ModelSpec("their_draft_tessellate", "their_paper",
              parties=Parties(supplier="Tessellate Labs Inc.",
                              supplier_email="legal@tessellatelabs.com",
                              services="warehouse routing services"),
              planted=[_pl("termination", "no convenience right for the customer"),
                       _pl("data-protection", "slow breach notice and free sub-processing"),
                       _pl("confidentiality", "one-way and short")],
              typos=[("pre-exisitng", "pre-existing"),
                     ("negligance", "negligence"), ("proceedigns", "proceedings")]),
]


def _key_terms(p: Parties) -> dict[str, list[str]]:
    """The answer table for a key-terms document, each value quoted from it."""
    return {
        "parties": [p.customer, p.supplier.rstrip(".")],
        "term": [p.initial_term],
        "renewal": [p.renewal, p.renewal_notice],
        "cap": ["200%", "12 months", p.floor],
        "governing_law": ["England and Wales"],
        "notice": [p.customer_email, p.supplier_email],
    }


_KT_PARTIES = [
    Parties(),
    Parties(customer="Beacon Mutual Limited", customer_number="03344556",
            customer_email="notices@beaconmutual.co.uk", supplier="Quarry Data Inc.",
            supplier_email="legal@quarrydata.com", services="claims triage services",
            initial_term="four years", renewal="further periods of 24 months",
            renewal_notice="at least six months", floor="£3,500,000"),
    Parties(customer="Linden Grocers Limited", customer_number="09988776",
            customer_email="contracts@lindengrocers.co.uk", supplier="Sable Metrics Inc.",
            supplier_email="counsel@sablemetrics.com", services="shelf analytics services",
            initial_term="two years", renewal="successive periods of 6 months",
            renewal_notice="at least 60 days", floor="£750,000"),
]

KEY_TERMS: list[ModelSpec] = [
    ModelSpec(name, "key_terms", parties=p, key_terms=_key_terms(p))
    for name, p in zip(("key_terms_kestrel", "key_terms_beacon", "key_terms_linden"),
                       _KT_PARTIES, strict=True)
]

REALISTIC: list[ModelSpec] = [
    ModelSpec("realistic_uk_spa_numbered", "realistic",
              builder=lambda: realistic.uk_spa(True)),
    ModelSpec("realistic_uk_spa_typed", "realistic",
              builder=lambda: realistic.uk_spa(False)),
    ModelSpec("realistic_us_spa", "realistic", builder=realistic.us_spa),
]

ALL: list[ModelSpec] = PLAYBOOK + SILENCE + THEIR_PAPER + KEY_TERMS + REALISTIC


def by_name(name: str) -> ModelSpec:
    for spec in ALL:
        if spec.name == name:
            return spec
    raise KeyError(name)


# What the live run asks for, per set. Their paper is said in the lawyer's
# words, which is the first thing their_paper.detect reads.
INSTRUCTIONS = {
    "playbook": "",
    "silence": "",
    "realistic": "",
    "their_paper": "This is their draft. Please review it against our playbook.",
    "key_terms": ("What are the key terms? Give me the key-terms table: parties, term, "
                  "renewal, liability cap, governing law and notices."),
}


# --------------------------------------------------------------------------
# Triage: what an email asks for (src/secondeye/triage.py, docs/migration.md phase 3)
# --------------------------------------------------------------------------
#
# Emails with known plans. The phrasings come from the router, follow-up and
# intake tests, and the hard ones are the cases the rules get wrong or refuse:
# "30 thanks", "accept 3 but not 5", a document buried under two forwards, a
# reply-all with the other side on CC, a contact card beside the contract, two
# drafts and a blackline, "their draft", an undo that names nothing, and sig
# pages for one agreement against sig packets for a closing. Scored by
# evals/score_model.py (`second-eye eval --triage`), for the rules' plan
# (triage.rules_plan) and for the model's, on the same cases.
#
# Each case's `expect` names only what it is about. `intents` is the set that
# must be planned (triage's canonical names, so blackline counts as compare
# and a question as an instruction); the other keys, when present, must match
# too: document, their_paper, baseline, review_mode, answers ({question id:
# value}), undo (change numbers), undo_all, dismiss (letters), recipients.

TRIAGE_AGENT = "review@legal.firm.com"
TRIAGE_SENDER = "jim@firm.com"
TRIAGE_ADMIN = "gc@firm.com"
TRIAGE_ENV = {
    "MAIL_AGENT_ADDRESS": TRIAGE_AGENT,
    "MAIL_AGENT_ALIASES": "",
    "FIRM_DOMAINS": "firm.com",
    "PLAYBOOK_ADMINS": TRIAGE_ADMIN,
    "ALLOWED_SENDERS": "",
    "NO_AI_MATTERS": "",
    "REPLY_POLICY": "sender_only",
}
_DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass
class TriageCase:
    name: str
    email: Callable[[], object]          # -> secondeye.models.InboundEmail
    expect: dict
    thread: bool = False                 # on the standard conversation (conversation())
    closing: str | None = None           # the sender has this closing open
    on_closing: bool = False             # and the message is on its thread
    hard: bool = False


def _doc_bytes(*paragraphs: str, author: str = "", comments: tuple[str, ...] = (),
               markup_by: str = "") -> bytes:
    from secondeye.pipeline.ooxml import Revision, RevisionWriter

    d = Document()
    paras = [d.add_paragraph(p) for p in paragraphs]
    for text, para in zip(comments, paras[1:], strict=False):
        d.add_comment(para.runs, text=text, author="Pat Partner", initials="PP")
    if markup_by:
        w = RevisionWriter(d, markup_by)
        w.apply(Revision("thirty (30) days", "sixty (60) days", markup_by))
    if author:
        d.core_properties.author = author
        d.core_properties.last_modified_by = author
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


_AGREEMENT = (
    ("THIS AGREEMENT is made between Falcon Holdings Limited (the \"Buyer\") and "
     "Kestrel Supplies Limited (the \"Seller\")."),
    "1. Payment. The Buyer shall pay each invoice within thirty (30) days.",
    "2. Term. The term is thirty (13) months from the Effective Date.",
    "3. Notices. Notices shall be given under clause 7.",
)


def _att(name: str, content: bytes | None = None, ctype: str = _DOCX_TYPE):
    from secondeye.models import Attachment

    data = _doc_bytes(*_AGREEMENT) if content is None else content
    return Attachment(filename=name, content_type=ctype, size_bytes=len(data), content=data)


_VCARD = (b"BEGIN:VCARD\r\nVERSION:3.0\r\nFN:Jim Baker\r\nEMAIL:jim@firm.com\r\n"
          b"END:VCARD\r\n")
_PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
_PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"


def _mail(body: str, *attachments, subject: str = "Falcon SPA", sender: str = TRIAGE_SENDER,
          to: list[str] | None = None, cc: list[str] | None = None,
          in_reply_to: str | None = None, mid: str = "triage-1"):
    from datetime import UTC, datetime

    from secondeye.models import InboundEmail

    return InboundEmail(
        message_id=mid, thread_id=None, in_reply_to=in_reply_to, from_address=sender,
        to=[TRIAGE_AGENT] if to is None else to, cc=cc or [], subject=subject,
        text_body=body, attachments=list(attachments),
        received_at=datetime(2026, 9, 28, 9, 0, tzinfo=UTC))


def _reply(body: str, cc: list[str] | None = None):
    return _mail(body, subject="Re: Falcon SPA", in_reply_to="agent-1", cc=cc)


def conversation():
    """The conversation every thread case is on: five changes made (numbers
    1 to 5), two questions open (11: 30 or 13 months; 12: 10 or 20 business
    days), and four findings shown as A to D."""
    from secondeye import thread

    changes = [
        thread.Change(id=100 + n, anchor=a, replacement=r, title=t, category="drafting",
                      origin="review", round=1, active=True)
        for n, (a, r, t) in enumerate([
            ("Reciever", "Recipient", "Cl. 2: \"Reciever\" is misspelled"),
            ("clause 7", "clause 3", "Cl. 3: cross-reference to a missing clause"),
            ("thirty (30) days", "thirty (30) days after receipt",
             "Cl. 1: payment runs from receipt"),
            ("the Buyer", "the Purchaser", "Cl. 1: the defined term is Purchaser"),
            ("Effective Date", "Completion Date", "Cl. 2: term runs from Completion"),
        ], start=1)
    ]
    questions = [
        thread.Question(id=11, question="Should the term be 30 or 13 months?",
                        anchor="thirty (13)", options=["30", "13"], answered_with=None,
                        round=1),
        thread.Question(id=12, question="Notice period: 10 or 20 business days?",
                        anchor=None, options=["10", "20"], answered_with=None, round=1),
    ]
    flagged = [
        thread.Flagged(id=200 + i, label=label, severity="substantive", category="drafting",
                       title=title, anchor="", question=None, current=True,
                       dismissed=False, round=1)
        for i, (label, title) in enumerate([
            ("A", "Cl. 1: no interest on late payment"),
            ("B", "Cl. 4: liability cap is uncapped for the Seller"),
            ("C", "Cl. 5: no limitation period for warranty claims"),
            ("D", "Cl. 9: governing law is Texas"),
        ])
    ]
    return thread.ThreadState(
        id="triage-thread", owner=TRIAGE_SENDER, filename="Falcon SPA.docx",
        original=_doc_bytes(*_AGREEMENT), changes=changes, questions=questions,
        flagged=flagged)


_FORWARDED = """FYI - see below, can you have a look before our call?

---------- Forwarded message ---------
From: Priya Shah <priya@kestrel-legal.com>
Date: Mon, 21 Sep 2026 at 17:02
Subject: Falcon - Supply Agreement
To: Jim Baker <jim@firm.com>

Jim, our draft of the supply agreement attached. Happy to discuss.

---------- Forwarded message ---------
From: Tom Reid <tom@kestrel.com>
Date: Mon, 21 Sep 2026 at 12:40
Subject: Supply Agreement

Priya, please send this to Falcon's lawyers.
"""

_THREE = ("Falcon SPA.docx", "Disclosure Letter.docx", "Tax Deed.docx")


def _one(name: str, body: str, expect: dict, **kw) -> TriageCase:
    """A case whose email is a reply on the standard conversation."""
    return TriageCase(name, lambda: _reply(body), expect, thread=True, **kw)


TRIAGE: list[TriageCase] = [
    # --- a document ------------------------------------------------------
    TriageCase("review_plain", lambda: _mail("Quick look before I send?",
                                             _att("Falcon SPA.docx")),
               {"intents": {"review"}, "document": "Falcon SPA.docx", "their_paper": False}),
    TriageCase("review_empty_body", lambda: _mail("", _att("Falcon SPA.docx")),
               {"intents": {"review"}, "document": "Falcon SPA.docx"}),
    TriageCase("review_memo_only", lambda: _mail(
        "No redline please, just tell me what's wrong", _att("Falcon SPA.docx")),
        {"intents": {"review"}, "review_mode": "memo_only"}),
    TriageCase("review_proofread", lambda: _mail("typos only please", _att("Falcon SPA.docx")),
               {"intents": {"review"}, "review_mode": "proofread"}),
    TriageCase("vcard_and_contract", lambda: _mail(
        "Can you check this?", _att("Jim Baker.vcf", _VCARD, "text/vcard"),
        _att("Master Services Agreement.docx")),
        {"intents": {"review"}, "document": "Master Services Agreement.docx"}, hard=True),
    TriageCase("logo_and_contract", lambda: _mail(
        "Review please", _att("image001.png", _PNG, "image/png"), _att("Falcon SPA.docx")),
        {"intents": {"review"}, "document": "Falcon SPA.docx"}),
    TriageCase("forwarded_chain_doc_buried", lambda: _mail(
        _FORWARDED, _att("Supply Agreement.docx"), subject="Fwd: Fwd: Supply Agreement"),
        {"intents": {"review"}, "document": "Supply Agreement.docx", "their_paper": True},
        hard=True),
    TriageCase("their_draft_in_words", lambda: _mail(
        "Their draft attached. Usual review against our playbook please.",
        _att("Falcon SPA (Kestrel draft).docx")),
        {"intents": {"review"}, "their_paper": True}),
    TriageCase("their_draft_said_loosely", lambda: _mail(
        "This came in from the other side this morning, what do you make of it?",
        _att("Falcon SPA.docx")),
        {"intents": {"review"}, "their_paper": True}, hard=True),
    TriageCase("our_draft", lambda: _mail(
        "Our draft, going out to Kestrel tonight. Anything I've missed?",
        _att("Falcon SPA.docx")),
        {"intents": {"review"}, "their_paper": False}),
    TriageCase("their_paper_by_author", lambda: _mail(
        "Can you take a look?", _att("Falcon SPA.docx", _doc_bytes(
            *_AGREEMENT, author="priya@kestrel-legal.com"))),
        {"intents": {"review"}, "their_paper": True}),
    TriageCase("draft_and_blackline", lambda: _mail(
        "Kestrel's turn attached, with their blackline. Please review.",
        _att("Falcon SPA v4.docx"), _att("Falcon SPA v4 vs v3 blackline.docx")),
        {"intents": {"review"}, "document": "Falcon SPA v4.docx"}, hard=True),
    TriageCase("two_drafts_review_the_later", lambda: _mail(
        "v3 and v4 attached. Review v4 please, v3 is just for reference.",
        _att("Falcon SPA v3.docx"), _att("Falcon SPA v4.docx")),
        {"intents": {"review"}, "document": "Falcon SPA v4.docx"}, hard=True),
    TriageCase("compare_two", lambda: _mail(
        "What changed between these?", _att("Falcon SPA v1.docx"), _att("Falcon SPA v2.docx")),
        {"intents": {"compare"}, "baseline": "Falcon SPA v1.docx"}),
    TriageCase("blackline_with_baseline", lambda: _mail(
        "Blackline their v5 against the v3 we sent please.",
        _att("APA v3.docx"), _att("APA v5.docx")),
        {"intents": {"compare"}, "baseline": "APA v3.docx"}),
    TriageCase("blackline_baseline_by_name", lambda: _mail(
        "Run a comparison, using the Kestrel one as the base.",
        _att("SPA (Falcon markup).docx"), _att("SPA (Kestrel).docx")),
        {"intents": {"compare"}, "baseline": "SPA (Kestrel).docx"}, hard=True),
    TriageCase("compare_with_previous", lambda: _mail(
        "Compare this against the previous version you saw", _att("Falcon SPA v5.docx")),
        {"intents": {"compare"}, "baseline": "previous"}),
    TriageCase("clean_copy_attached", lambda: _mail(
        "Can you send me a clean copy of this?", _att("Falcon SPA.docx")),
        {"intents": {"clean_copy"}, "document": "Falcon SPA.docx"}),
    TriageCase("renumber", lambda: _mail("renumber please", _att("Falcon SPA.docx")),
               {"intents": {"renumber"}}),
    TriageCase("sig_pages_one_agreement", lambda: _mail(
        "Can you prepare the sig pages for this?", _att("Falcon SPA.docx")),
        {"intents": {"sig_pages"}}),
    TriageCase("sig_packets_closing", lambda: _mail(
        "Please set up the closing and send me sig packets for everyone.",
        *[_att(n) for n in _THREE]),
        {"intents": {"sig_packets"}}),
    TriageCase("sig_pages_for_the_completion", lambda: _mail(
        "I need signature pages for the completion, all three docs attached.",
        *[_att(n) for n in _THREE]),
        {"intents": {"sig_packets"}}, hard=True),
    TriageCase("signed_pages_back", lambda: _mail(
        "Kestrel's signed pages attached.", _att("Kestrel signed.pdf", _PDF, "application/pdf"),
        subject="Falcon signing"),
        {"intents": {"signed_pages_returned"}}, closing="Falcon"),
    TriageCase("closing_status", lambda: _mail(
        "where are we on signatures?", subject="Re: Falcon closing", in_reply_to="closing-1"),
        {"intents": {"closing_checklist"}}, closing="Falcon", on_closing=True),
    TriageCase("turn_comments", lambda: _mail(
        "Turn Pat's comments please", _att("Falcon SPA.docx", _doc_bytes(
            *_AGREEMENT, comments=("Use our standard terms.", "Why 13?")))),
        {"intents": {"turn_comments"}}),
    TriageCase("turn_markup", lambda: _mail(
        "Can you turn their markup? Accept what's harmless.", _att(
            "Falcon SPA (Kestrel markup).docx",
            _doc_bytes(*_AGREEMENT, markup_by="Priya Shah")),
        subject="Kestrel markup"),
        {"intents": {"turn_comments"}}),
    TriageCase("deadlines_calendar", lambda: _mail(
        "Send me the deadlines as a calendar", _att("Falcon SPA.docx")),
        {"intents": {"deadlines_calendar"}}),
    TriageCase("repair_formatting", lambda: _mail(
        "Can you fix the formatting? The fonts are all over the place.",
        _att("Falcon SPA.docx")),
        {"intents": {"repair_formatting"}}),
    # --- the firm ----------------------------------------------------------
    TriageCase("playbook_update", lambda: _mail(
        "Playbook update: our liability cap position is now 2x annual fees. Positions "
        "attached.", _att("Positions.docx"), sender=TRIAGE_ADMIN, subject="Playbook update"),
        {"intents": {"playbook_update"}}),
    TriageCase("playbook_approve", lambda: _mail("approve playbook", sender=TRIAGE_ADMIN,
                                                 subject="Re: Playbook update"),
               {"intents": {"playbook_approve"}}),
    TriageCase("playbook_undo", lambda: _mail("undo the last playbook change",
                                              sender=TRIAGE_ADMIN, subject="Playbook"),
               {"intents": {"playbook_undo"}}),
    TriageCase("audit_report", lambda: _mail("Audit report for September please",
                                             sender=TRIAGE_ADMIN, subject="Audit"),
               {"intents": {"audit_report"}}),
    TriageCase("no_ai_add", lambda: _mail("No AI for Acme", sender=TRIAGE_ADMIN,
                                          subject="Clients"),
               {"intents": {"no_ai_add"}}),
    # --- whole-message commands ----------------------------------------------
    TriageCase("help", lambda: _mail("help", subject="?"), {"intents": {"help"}}),
    TriageCase("help_in_words", lambda: _mail("What else can you do for me?", subject="?"),
               {"intents": {"help"}}, hard=True),
    TriageCase("setup", lambda: _mail("How do I set this up so it happens automatically?",
                                      subject="Setup"),
               {"intents": {"setup"}}, hard=True),
    TriageCase("revoke", lambda: _mail("disconnect", subject="DMS"), {"intents": {"revoke"}}),
    _one("stop", "stop", {"intents": {"stop"}}),
    _one("remember", "remember that", {"intents": {"remember"}}),
    _one("flag_again", "flag it again", {"intents": {"flag_again"}}),
    # --- replies on a conversation -------------------------------------------
    _one("answer", "30", {"intents": {"answer"}, "answers": {11: "30"}}),
    _one("answer_30_thanks", "30 thanks", {"intents": {"answer"}, "answers": {11: "30"}},
         hard=True),
    _one("two_answers", "30; 20", {"intents": {"answer"}, "answers": {11: "30", 12: "20"}}),
    _one("answer_in_words", "Make it 20 business days for notices.",
         {"intents": {"answer"}, "answers": {12: "20"}}, hard=True),
    _one("undo_number", "undo 2", {"intents": {"undo"}, "undo": [2]}),
    _one("accept_3_not_5", "accept 3 but not 5", {"intents": {"undo"}, "undo": [5]}, hard=True),
    _one("undo_ambiguous", "undo that", {"intents": {"undo"}, "undo": [], "undo_all": False},
         hard=True),
    _one("undo_everything", "undo everything", {"intents": {"undo"}, "undo_all": True}),
    _one("undo_by_description", "put Buyer back, don't call them the Purchaser",
         {"intents": {"undo"}, "undo": [4]}, hard=True),
    _one("dismiss", "B is fine", {"intents": {"dismiss"}, "dismiss": ["B"]}),
    _one("combined", "B is fine; 30; clean copy",
         {"intents": {"dismiss", "answer", "clean_copy"}, "dismiss": ["B"],
          "answers": {11: "30"}}),
    _one("instruction", "Make the liability cap 2x annual fees.", {"intents": {"instruction"}}),
    _one("question", "Why did you change the cross-reference in 3?",
         {"intents": {"instruction"}}),
    _one("stop_flagging", "stop flagging the Oxford comma", {"intents": {"stop_flagging"}}),
    _one("thanks", "thanks, all accepted", {"intents": {"ack"}}),
    _one("negotiation_status", "Did they accept our changes?",
         {"intents": {"negotiation_status"}}),
    _one("clean_copy_on_thread", "clean copy please", {"intents": {"clean_copy"}}),
    # --- who the reply goes to ------------------------------------------------
    TriageCase("reply_all_counterparty_cc", lambda: _reply(
        "B is fine", cc=["priya@kestrel-legal.com"]),
        {"intents": {"dismiss"}, "dismiss": ["B"], "recipients": [TRIAGE_SENDER]},
        thread=True, hard=True),
    TriageCase("bcc_cover_note", lambda: _mail(
        "Thanks Priya, see you Tuesday.", to=["priya@kestrel-legal.com"],
        subject="Re: Falcon"),
        {"intents": {"ignore"}}),
]
