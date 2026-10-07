"""Outbound cases: an email, its attachments, and what must be caught.

The corpus in corpus.py is documents. What leaves a firm is a message: a
covering note, a To line, a file name and one or more files, and the
mistakes this part of the product exists for live in the gap between them.
Each case here is one such message, built the way it would arrive when a
lawyer BCCs or forwards it, with the defects planted in it declared as
markers, `category:words in the title`, so the harness scores it without
anyone reading output.

The clean controls matter more. Every clean document in corpus.py is sent
here too, under an ordinary covering note to one of its own parties, and
the controls below are the ordinary versions of each planted case: a
blackline attached beside the clean copy, a draft marked DRAFT that is
called a draft, a disclosure letter "to follow", a note naming the other
side's law firm. Anything raised on them is a false positive.

Builders here are also what demos/outbound uses, so the demo shows exactly
what the eval measures.
"""

from __future__ import annotations

import re
import warnings
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from io import BytesIO

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from lra.models import Attachment, InboundEmail

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SENDER = "jim.baker@firm.com"
AGENT = "review@legal.example.com"


# --------------------------------------------------------------------------
# Documents
# --------------------------------------------------------------------------

NDA = [
    "Mutual Non-Disclosure Agreement",
    ("This Mutual Non-Disclosure Agreement (this “Agreement”) is made as of March 3, 2026 "
    "between Acme Holdings LLC, a Delaware limited liability company (“Acme”), and Beta "
    "Industries Inc., a New York corporation (“Beta”)."),
    ("1. Confidential Information. “Confidential Information” means any information "
    "disclosed by one party to the other that is marked confidential."),
    ("2. Obligations. Each party will hold the other party's Confidential Information in "
    "confidence and will not disclose it except as permitted by Section 3."),
    ("3. Permitted Disclosures. A party may disclose Confidential Information to its "
    "advisers who need to know it and are bound by duties of confidentiality."),
    ("4. Exclusivity. For six (6) months neither party will negotiate with any third party "
    "about the transaction this Agreement concerns."),
    ("5. Term. This Agreement continues for three (3) years from the date first written "
    "above."),
    "6. Governing Law. This Agreement is governed by the laws of the State of New York.",
]

SPA = [
    "Share Purchase Agreement",
    ("This Agreement is made as of March 3, 2026 between Acme Holdings, Inc., a Delaware "
    "corporation (“Buyer”), and Northgate Freight Ltd, a company incorporated in England "
    "and Wales (“Seller”)."),
    "1. Sale. Seller sells and Buyer purchases the Shares on the terms of this Agreement.",
    "2. Price. The price is one million dollars ($1,000,000), payable at Completion.",
    ("3. Liability. Seller's aggregate liability under this Agreement is capped at two "
    "million dollars ($2,000,000)."),
    "4. Governing Law. This Agreement is governed by the laws of the State of New York.",
    ("IN WITNESS WHEREOF, the parties have executed this Agreement as of the date first "
    "written above."),
    "ACME HOLDINGS, INC.",
    "By: ______________________",
    "Name: Jane Doe",
    "Title: Chief Executive Officer",
    "NORTHGATE FREIGHT LTD",
    "By: ______________________",
    "Name: John Smith",
    "Title: Director",
]

_W = nsdecls("w")


def docx(paragraphs: list[str], *, header: str = "", footer: str = "",
         watermark: str = "", title: str = "", custom: dict[str, str] | None = None,
         comments: list[tuple[str, str]] | None = None, inserted: str = "",
         deleted: str = "", hidden: str = "", highlighted: str = "",
         embedded: bool = False, linked: str = "", hyperlink: str = "",
         author: str = "", company: str = "") -> bytes:
    """A Word document with whatever travels in it named by keyword.

    Built with python-docx and then the package patched for the parts it
    cannot write (comments, custom properties, an embedded workbook), the
    way Word itself lays them out.
    """
    d = Document()
    first = True
    for text in paragraphs:
        if first:
            d.add_heading(text, 1)
            first = False
        else:
            d.add_paragraph(text)
    body = d.paragraphs[2]._p if len(d.paragraphs) > 2 else d.paragraphs[-1]._p
    if inserted:
        body.append(parse_xml(
            f'<w:ins {_W} w:id="901" w:author="Priya Associate" w:date="2026-03-01T00:00:00Z">'
            f'<w:r><w:t xml:space="preserve"> {inserted}</w:t></w:r></w:ins>'))
    if deleted:
        body.append(parse_xml(
            f'<w:del {_W} w:id="902" w:author="Priya Associate" w:date="2026-03-01T00:00:00Z">'
            f'<w:r><w:delText xml:space="preserve"> {deleted}</w:delText></w:r></w:del>'))
    if hidden:
        body.append(parse_xml(
            f'<w:r {_W}><w:rPr><w:vanish/></w:rPr><w:t xml:space="preserve"> {hidden}'
            "</w:t></w:r>"))
    if highlighted:
        p = d.add_paragraph()
        p._p.append(parse_xml(
            f'<w:r {_W}><w:rPr><w:highlight w:val="yellow"/></w:rPr><w:t>{highlighted}'
            "</w:t></w:r>"))
    if linked:
        p = d.add_paragraph()
        for xml in (
            '<w:r {ns}><w:fldChar w:fldCharType="begin"/></w:r>',
            ('<w:r {ns}><w:instrText xml:space="preserve"> LINK Excel.Sheet.12 "{path}" '
            '"Sheet1!R1C1:R4C3" \\a \\p </w:instrText></w:r>'),
            '<w:r {ns}><w:fldChar w:fldCharType="separate"/></w:r>',
            '<w:r {ns}><w:t>Fee table</w:t></w:r>',
            '<w:r {ns}><w:fldChar w:fldCharType="end"/></w:r>',
        ):
            p._p.append(parse_xml(xml.format(ns=_W, path=linked.replace("\\", "\\\\"))))
    if hyperlink:
        rid = d.part.relate_to(hyperlink, "http://schemas.openxmlformats.org/officeDocument/"
                               "2006/relationships/hyperlink", is_external=True)
        p = d.add_paragraph("See ")
        p._p.append(parse_xml(
            f'<w:hyperlink {_W} xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
            f'2006/relationships" r:id="{rid}"><w:r><w:t>the published terms</w:t></w:r>'
            "</w:hyperlink>"))
    section = d.sections[0]
    if header:
        section.header.paragraphs[0].text = header
    if footer:
        section.footer.paragraphs[0].text = footer
    if watermark:
        section.header.paragraphs[0]._p.append(parse_xml(
            f'<w:r {_W} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>'
            '<v:shape id="PowerPlusWaterMarkObject1" type="#_x0000_t136" '
            'style="position:absolute;width:400pt;height:100pt;rotation:315">'
            f'<v:textpath style="font-family:Calibri" string="{watermark}"/></v:shape>'
            "</w:pict></w:r>"))
    d.core_properties.author = author
    d.core_properties.last_modified_by = author
    d.core_properties.title = title
    buf = BytesIO()
    d.save(buf)
    content = buf.getvalue()
    if comments or custom or embedded or company:
        content = _patch(content, comments or [], custom or {}, embedded, company)
    return content


def _patch(content: bytes, comments: list[tuple[str, str]], custom: dict[str, str],
           embedded: bool, company: str) -> bytes:
    src = zipfile.ZipFile(BytesIO(content))
    parts = {n: src.read(n) for n in src.namelist()}
    types = parts["[Content_Types].xml"].decode()
    if comments:
        body = "".join(
            f'<w:comment w:id="{i}" w:author="{who}" w:date="2026-03-01T00:00:00Z" '
            f'w:initials="{who[:1]}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:comment>'
            for i, (who, text) in enumerate(comments))
        parts["word/comments.xml"] = f'<w:comments {_W}>{body}</w:comments>'.encode()
        rels = parts["word/_rels/document.xml.rels"].decode()
        parts["word/_rels/document.xml.rels"] = rels.replace(
            "</Relationships>",
            '<Relationship Id="rIdC1" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/comments" Target="comments.xml"/></Relationships>').encode()
        types = types.replace("</Types>", (
            '<Override PartName="/word/comments.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.wordprocessingml.comments+xml"/></Types>'))
    if custom:
        props = "".join(
            f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i + 2}" '
            f'name="{name}"><vt:lpwstr>{value}</vt:lpwstr></property>'
            for i, (name, value) in enumerate(custom.items()))
        parts["docProps/custom.xml"] = (
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
            'custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/'
            f'2006/docPropsVTypes">{props}</Properties>').encode()
        rels = parts["_rels/.rels"].decode()
        parts["_rels/.rels"] = rels.replace(
            "</Relationships>",
            '<Relationship Id="rIdP9" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/custom-properties" Target="docProps/custom.xml"/>'
            "</Relationships>").encode()
        types = types.replace("</Types>", (
            '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.custom-properties+xml"/></Types>'))
    if embedded:
        parts["word/embeddings/Microsoft_Excel_Worksheet.xlsx"] = workbook()
        rels = parts["word/_rels/document.xml.rels"].decode()
        parts["word/_rels/document.xml.rels"] = rels.replace(
            "</Relationships>",
            '<Relationship Id="rIdE1" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/package" Target="embeddings/Microsoft_Excel_Worksheet.xlsx"/>'
            "</Relationships>").encode()
        if 'Extension="xlsx"' not in types:
            types = types.replace("<Override ", (
                '<Default Extension="xlsx" ContentType="application/vnd.openxmlformats-'
                'officedocument.spreadsheetml.sheet"/><Override '), 1)
    if company:
        app = parts["docProps/app.xml"].decode()
        app = re.sub(r"<Company>[^<]*</Company>|<Company/>", "", app)
        parts["docProps/app.xml"] = app.replace(
            "</Properties>", f"<Company>{company}</Company></Properties>").encode()
    parts["[Content_Types].xml"] = types.encode()
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return out.getvalue()


def pdf(lines: list[str], *, sticky: str = "", highlight: bool = False,
        attached: bool = False, author: str = "", title: str = "",
        invariant: bool = False) -> bytes:
    """A one-page PDF export, with a reviewer's mark-up if asked. `invariant`
    leaves out the dates and random ids, so two builds are the same bytes."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import Highlight, Text
    from pypdf.generic import ArrayObject, FloatObject, NameObject, TextStringObject
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf, invariant=int(invariant))
    c.setAuthor("anonymous")
    for i, line in enumerate(lines):
        c.drawString(72, 740 - 16 * i, line[:95])
    c.showPage()
    c.save()
    w = PdfWriter(clone_from=PdfReader(BytesIO(buf.getvalue())))
    meta = {}
    if author:
        meta["/Author"] = author
    if title:
        meta["/Title"] = title
    if meta:
        w.add_metadata(meta)
    if sticky:
        note = Text(rect=(50, 550, 200, 650), text=sticky)
        note[NameObject("/T")] = TextStringObject("Bob Partner")
        w.add_annotation(page_number=0, annotation=note)
    if highlight:
        hl = Highlight(rect=(70, 720, 400, 745), quad_points=ArrayObject(
            [FloatObject(v) for v in (70, 745, 400, 745, 70, 720, 400, 720)]))
        hl[NameObject("/T")] = TextStringObject("Bob Partner")
        w.add_annotation(page_number=0, annotation=hl)
    if attached:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            w.add_attachment("fee workings.xlsx", workbook())
    out = BytesIO()
    w.write(out)
    return out.getvalue()


def deck(slides: list[str], *, notes: dict[int, str] | None = None,
         hidden: tuple[int, ...] = ()) -> bytes:
    from pptx import Presentation

    p = Presentation()
    for i, text in enumerate(slides, 1):
        slide = p.slides.add_slide(p.slide_layouts[1])
        slide.shapes.title.text = text
        if notes and i in notes:
            slide.notes_slide.notes_text_frame.text = notes[i]
        if i in hidden:
            slide._element.set("show", "0")
    p.core_properties.author = ""
    p.core_properties.last_modified_by = ""
    buf = BytesIO()
    p.save(buf)
    return buf.getvalue()


def workbook(*, hidden_sheet: str = "", comment: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    book = Workbook()
    sheet = book.active
    sheet.title = "Fees"
    for row in (["Stage", "Fee"], ["Due diligence", 40000], ["SPA", 55000], ["Total", 95000]):
        sheet.append(row)
    if comment:
        sheet["B4"].comment = Comment(comment, "Bob Partner")
    if hidden_sheet:
        other = book.create_sheet(hidden_sheet)
        other.append(["Write-off agreed with client", 15000])
        other.sheet_state = "hidden"
    book.properties.creator = ""
    buf = BytesIO()
    book.save(buf)
    return _fixed_times(buf.getvalue())


def _fixed_times(content: bytes) -> bytes:
    """The same package with fixed zip and property timestamps, so a workbook
    embedded in a document does not change the document's bytes per build."""
    source = zipfile.ZipFile(BytesIO(content))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                              rb"\g<1>2026-09-28T09:00:00Z", data)
            target.writestr(zipfile.ZipInfo(info.filename, date_time=(2026, 9, 28, 9, 0, 0)),
                            data, zipfile.ZIP_DEFLATED)
    return out.getvalue()


# --------------------------------------------------------------------------
# Messages
# --------------------------------------------------------------------------


@dataclass
class Case:
    name: str
    subject: str
    body: str
    files: Callable[[], list[tuple[str, bytes]]]  # first file is the one reviewed
    to: list[str] = field(default_factory=lambda: [AGENT])
    cc: list[str] = field(default_factory=list)
    expect: set[str] = field(default_factory=set)  # "category:words in the title"
    note: str = ""

    def email(self) -> tuple[InboundEmail, Attachment]:
        atts = [Attachment(filename=n, content_type=_type(n), size_bytes=len(c), content=c)
                for n, c in self.files()]
        email = InboundEmail(
            message_id=f"<{self.name}@eval>", from_address=SENDER, from_name="Jim Baker",
            to=self.to, cc=self.cc, subject=self.subject, text_body=self.body,
            attachments=atts, received_at=datetime(2026, 10, 4, tzinfo=UTC),
        )
        return email, atts[0]

    def external(self) -> list[str]:
        return [a for a in self.to + self.cc
                if a != AGENT and not a.endswith("@firm.com")]


def _type(name: str) -> str:
    return {"docx": DOCX, "pdf": PDF, "pptx": PPTX, "xlsx": XLSX}[name.rsplit(".", 1)[-1]]


ACME = ["legal@acme.com"]

PLANTED: list[Case] = [
    Case("out_comments", "Acme NDA", "Please find attached the NDA for signature.",
         lambda: [("Acme NDA.docx", docx(NDA, comments=[
             ("Priya Associate", "Is six months too long? Client wanted three."),
             ("Bob Partner", "Fine - they will push back anyway.")]))],
         to=ACME, expect={"leftovers:comment"},
         note="Two internal comments, one of them about the client's position."),
    Case("out_tracked_called_clean", "Acme NDA",
         "Attached is a clean copy of the NDA reflecting our call.",
         lambda: [("Acme NDA.docx", docx(NDA, inserted="and its affiliates"))],
         to=ACME, expect={"claim:clean copy", "leftovers:tracked changes"}),
    Case("out_hidden_text", "Acme NDA", "Here is the NDA.",
         lambda: [("Acme NDA.docx", docx(NDA, hidden="[Client says it will accept 12 "
                                                     "months if pushed]"))],
         to=ACME, expect={"leftovers:hidden text"}),
    Case("out_properties_other_client", "Acme NDA", "Here is the NDA.",
         lambda: [("Acme NDA.docx", docx(
             NDA, title="Bluewater Shipping LLC - Mutual NDA",
             custom={"ClientName": "Bluewater Shipping LLC", "MatterNumber": "40412-0003"}))],
         to=ACME, expect={"leftovers:bluewater shipping"},
         note="Started from the last client's NDA; its title and the document "
              "system's client property came along."),
    Case("out_embedded_workbook", "Acme NDA", "Here is the NDA.",
         lambda: [("Acme NDA.docx", docx(NDA, embedded=True))],
         to=ACME, expect={"leftovers:embedded spreadsheet"}),
    Case("out_linked_file", "Acme NDA", "Here is the NDA.",
         lambda: [("Acme NDA.docx", docx(
             NDA, linked="\\\\fileserver\\Clients\\Bluewater\\fee workings.xlsx"))],
         to=ACME, expect={"leftovers:links to"}),
    Case("out_final_says_draft", "Acme SPA",
         "Please find attached the execution version of the SPA.",
         lambda: [("Acme SPA.docx", docx(SPA, header="DRAFT 3: March 1, 2026"))],
         to=ACME, expect={"draft:marked as a draft"}),
    Case("out_final_watermark", "Acme SPA", "Here it is.",
         lambda: [("Acme SPA (execution version).docx", docx(SPA, watermark="DRAFT"))],
         to=ACME, expect={"draft:marked as a draft"},
         note="Final by its file name; a DRAFT watermark in the header."),
    Case("out_final_highlight", "Acme SPA",
         "Attached is the final version for signature.",
         lambda: [("Acme SPA.docx", docx(SPA, highlighted="Completion takes place on 31 "
                                                         "March 2026."))],
         to=ACME, expect={"draft:highlighted"}),
    Case("out_named_clean_has_markup", "Acme NDA", "Here is the NDA.",
         lambda: [("Acme NDA v3 clean.docx", docx(NDA, deleted="and its affiliates"))],
         to=ACME, expect={"claim:is called", "leftovers:tracked changes"}),
    Case("out_version_in_name", "Acme NDA", "Attached is v4 of the NDA.",
         lambda: [("Acme NDA v3.docx", docx(NDA))],
         to=ACME, expect={"attachment:version 4"}),
    Case("out_version_in_footer", "Acme NDA", "Here is version 5 of the NDA.",
         lambda: [("Acme NDA.docx", docx(NDA, footer="NY-4455667v4"))],
         to=ACME, expect={"attachment:version 5"},
         note="The document system's footer stamp says v4."),
    Case("out_note_names_other_party", "NDA",
         "Attached is the NDA with Bluewater Shipping LLC for your review.",
         lambda: [("NDA.docx", docx(NDA))],
         expect={"attachment:bluewater shipping"},
         note="The note is about Bluewater; the file is the Acme/Beta NDA."),
    Case("out_disclosure_letter_missing", "Acme SPA",
         "Attached are the SPA and the disclosure letter.",
         lambda: [("Acme SPA.docx", docx(SPA))],
         to=ACME, expect={"attachment:disclosure letter"}),
    Case("out_blackline_missing", "Acme NDA",
         "Attached are a clean version and a blackline against your draft.",
         lambda: [("Acme NDA v4.docx", docx(NDA))],
         to=ACME, expect={"attachment:blackline"}),
    Case("out_internal_marking", "Acme NDA", "Please find attached our comments.",
         lambda: [("Acme NDA.docx", docx(NDA, header="INTERNAL DRAFT - FOR INTERNAL "
                                                     "DISCUSSION ONLY"))],
         to=ACME, expect={"recipient:marked"}),
    Case("out_wrong_recipient", "NDA", "Here is the NDA.",
         lambda: [("NDA.docx", docx(NDA))],
         to=["counsel@bluewatershipping.com"], expect={"recipient:addressed to"}),
    Case("out_personal_address", "Acme NDA", "Here is the NDA.",
         lambda: [("Acme NDA.docx", docx(NDA))],
         to=ACME, cc=["jim.baker.fictional@gmail.com"], expect={"recipient:personal address"}),
    Case("out_removed_still_there", "Acme NDA",
         "I have removed the exclusivity clause as discussed.",
         lambda: [("Acme NDA.docx", docx(NDA + [("7. Exclusivity clause. The exclusivity "
                                                "clause in Section 4 survives termination.")]))],
         to=ACME, expect={"claim:removed"}),
    Case("out_wrong_kind", "Lease", "Attached is the lease for your review.",
         lambda: [("Acme NDA.docx", docx(NDA))],
         to=ACME, expect={"attachment:attachment is called"}),
    Case("out_pdf_exhibit_markup", "Acme NDA", "Here are the NDA and its exhibit.",
         lambda: [("Acme NDA.docx", docx(NDA)),
                  ("Exhibit A.pdf", pdf(["Exhibit A - Permitted Advisers",
                                         "Acme Holdings LLC: Smith & Jones LLP"],
                                        sticky="Do we have to list these?", highlight=True,
                                        attached=True))],
         to=ACME, expect={"leftovers:mark-up", "leftovers:file attached inside"},
         note="The reviewed document is clean; the exhibit beside it is not."),
    Case("out_pdf_reviewed", "Acme NDA", "Here is the NDA as a PDF.",
         lambda: [("Acme NDA.pdf", pdf(NDA, sticky="Too long?", author="Priya Associate",
                                       title="Microsoft Word - Bluewater NDA v2.docx"))],
         to=ACME, expect={"leftovers:mark-up"}),
    Case("out_deck_notes", "Acme NDA", "Here are the NDA and the board deck.",
         lambda: [("Acme NDA.docx", docx(NDA)),
                  ("Board deck.pptx", deck(["Project Acme", "Timetable", "Fees"],
                                           notes={2: "Do not mention the Bluewater bid."},
                                           hidden=(3,)))],
         to=ACME, expect={"leftovers:speaker notes", "leftovers:hidden"}),
    Case("out_workbook", "Acme NDA", "Here are the NDA and the fee estimate.",
         lambda: [("Acme NDA.docx", docx(NDA)),
                  ("Fee estimate.xlsx", workbook(hidden_sheet="Write-offs",
                                                 comment="Partner to approve discount"))],
         to=ACME, expect={"leftovers:hidden sheet", "leftovers:comment"}),
]


CONTROLS: list[Case] = [
    Case("ctl_execution_version", "Acme SPA",
         "Please find attached the execution version of the SPA for signature.",
         lambda: [("Acme SPA (execution version).docx",
                   docx(SPA, header="Execution version", footer="NY-4455667v4"))],
         to=ACME),
    Case("ctl_named_clean", "Acme NDA", "Attached is v4, clean.",
         lambda: [("Acme NDA v4 clean.docx", docx(NDA))], to=ACME),
    Case("ctl_clean_with_blackline", "Acme NDA",
         "Attached are a clean version and a blackline against your draft.",
         lambda: [("Acme NDA v4.docx", docx(NDA)),
                  ("Acme NDA v4 (blackline against v3).docx",
                   docx(NDA, inserted="and its affiliates", deleted="and advisers"))],
         to=ACME, note="A blackline beside the clean copy carries tracked changes by design."),
    Case("ctl_versions_in_one_sentence", "Acme NDA",
         "Attached is v4, which picks up your comments on v3.",
         lambda: [("Acme NDA v4.docx", docx(NDA))], to=ACME),
    Case("ctl_note_names_parties_and_lawyers", "Acme NDA",
         "Attached is the NDA between Acme Holdings LLC and Beta Industries Inc., as "
         "discussed with Smith & Jones LLP this morning.",
         lambda: [("NDA.docx", docx(NDA))], to=ACME),
    Case("ctl_note_names_a_bank", "Acme NDA",
         "Attached is the NDA. Harbour Bank plc will need to see it before signing.",
         lambda: [("Acme NDA.docx", docx(NDA))], to=ACME),
    Case("ctl_quoted_thread", "Re: Acme NDA",
         "Attached is v4 of the NDA.\n\nOn Mon, 2 Mar 2026 at 09:12, Bob Smith "
         "<bob@beta.com> wrote:\n> Here is v3, with the disclosure letter and comments from "
         "Bluewater Shipping LLC.\n> Bob",
         lambda: [("Acme NDA v4.docx", docx(NDA))], to=ACME,
         note="The other side's versions, documents and companies are quoted history."),
    Case("ctl_letter_to_follow", "Acme SPA",
         "Attached is the SPA. The disclosure letter will follow separately.",
         lambda: [("Acme SPA.docx", docx(SPA))], to=ACME),
    Case("ctl_draft_called_draft", "Acme NDA",
         "Attached is the final draft for your comments.",
         lambda: [("Acme NDA.docx", docx(NDA, header="DRAFT 3: March 1, 2026",
                                         watermark="DRAFT"))], to=ACME),
    Case("ctl_internal_marking_internally", "Acme NDA", "Can you check this before I send it?",
         lambda: [("Acme NDA.docx", docx(NDA, header="INTERNAL DRAFT"))],
         note="Marked internal, and nobody outside the firm is on the message."),
    Case("ctl_properties_name_the_parties", "Acme NDA", "Here is the NDA.",
         lambda: [("Acme NDA.docx", docx(
             NDA, title="Acme Holdings LLC / Beta Industries Inc. - NDA",
             custom={"ClientName": "Acme Holdings LLC", "MatterNumber": "40412-0007"},
             hyperlink="https://www.acme.com/terms", company="Smith & Jones LLP"))],
         to=ACME),
    Case("ctl_clean_attachments", "Acme NDA", "Here are the NDA, the exhibit, deck and fees.",
         lambda: [("Acme NDA.docx", docx(NDA)),
                  ("Exhibit A.pdf", pdf(["Exhibit A - Permitted Advisers"])),
                  ("Board deck.pptx", deck(["Project Acme", "Timetable"])),
                  ("Fee estimate.xlsx", workbook())],
         to=ACME),
    Case("ctl_clean_pdf", "Acme NDA", "Here is the NDA as a PDF.",
         lambda: [("Acme NDA.pdf", pdf(NDA, title="Mutual Non-Disclosure Agreement"))],
         to=ACME),
]


def corpus_controls() -> list[Case]:
    """Every clean document in corpus.py, sent under an ordinary note to one of
    its own parties, as a lawyer BCCing the agent would send it."""
    from evals.corpus import CLEAN
    from lra.pipeline import checks, extract

    out = []
    for spec in CLEAN:
        content = spec.build()
        doc = extract.extract(Attachment(filename=f"{spec.name}.docx", content_type=DOCX,
                                         size_bytes=len(content), content=content))
        parties = checks._parties(doc)[0]
        to = []
        if parties:
            word = re.sub(r"[^a-z]", "", parties[0][0].split()[0].lower())
            to = [f"counsel@{word}.com"]
        out.append(Case(
            f"ctl_{spec.name}", spec.name.replace("_", " "),
            "Please find attached the latest draft for your review. Happy to discuss.",
            lambda content=content, name=spec.name: [(f"{name}.docx", content)],
            to=to or [AGENT],
        ))
    return out
