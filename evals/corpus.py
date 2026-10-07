"""A synthetic corpus of legal documents, with and without planted defects.

Two jobs, and the second matters more.

**Catch rate**: does the review find defects we know are there?

**False positive rate**: does it stay quiet on documents that are correct? This
is the number that decides whether the product survives. A check that runs on
everything and is sometimes wrong gets filtered into a folder within weeks, so
the clean documents here are written to be genuinely, realistically clean:
proper defined terms, real cross-references, British and American conventions,
schedules, tables, bracketed citations, numbering that restarts per section.
Anything the checks flag in a CLEAN document is a bug in the checks.

Each document declares the defects planted in it by category, so the harness can
score without anyone hand-labelling output.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from docx.shared import Pt

from evals import realistic


@dataclass
class Spec:
    name: str
    kind: str                       # nda | spa | engagement | msa | employment | lease | minutes | brief
    clean: bool
    paragraphs: list[str]
    headings: list[tuple[int, str]] = field(default_factory=list)
    table: list[list[str]] | None = None
    expect: set[str] = field(default_factory=set)   # planted defect categories
    note: str = ""
    # A document built in Word's own terms (list numbering, tables, content
    # controls) rather than from typed paragraphs. See evals/realistic.py.
    builder: Callable[[], bytes] | None = None

    def build(self) -> bytes:
        if self.builder is not None:
            return self.builder()
        d = Document()
        style = d.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        for level, text in self.headings:
            d.add_heading(text, level)
        for text in self.paragraphs:
            d.add_paragraph(text)
        if self.table:
            t = d.add_table(rows=len(self.table), cols=len(self.table[0]))
            for r, row in enumerate(self.table):
                for c, cell in enumerate(row):
                    t.rows[r].cells[c].text = cell
        core = d.core_properties
        core.author = ""
        core.last_modified_by = ""
        buf = BytesIO()
        d.save(buf)
        return buf.getvalue()


# ==========================================================================
# CLEAN DOCUMENTS. Anything flagged here is a false positive.
# ==========================================================================

CLEAN: list[Spec] = [
    Spec(
        name="clean_nda_us",
        kind="nda",
        clean=True,
        headings=[(1, "Mutual Non-Disclosure Agreement")],
        paragraphs=[
            "This Mutual Non-Disclosure Agreement (this “Agreement”) is made as of "
            "March 3, 2026 between Acme Holdings LLC, a Delaware limited liability "
            "company (“Acme”), and Beta Industries Inc., a New York corporation "
            "(“Beta”). Acme and Beta are each a “Party” and together the “Parties”.",
            "1. Definitions. “Confidential Information” means any information "
            "disclosed by one Party to the other Party that is marked confidential or "
            "that a reasonable person would understand to be confidential. "
            "Confidential Information does not include information that becomes "
            "publicly available other than through a breach of this Agreement.",
            "2. Obligations. Each Party will hold the other Party's Confidential "
            "Information in confidence and will not disclose it to any third party "
            "except as permitted by Section 3. Each Party will protect the other "
            "Party's Confidential Information using at least the degree of care it "
            "uses to protect its own Confidential Information of like importance.",
            "3. Permitted Disclosures. A Party may disclose Confidential Information "
            "to its employees, advisers and affiliates who need to know it for the "
            "Purpose and who are bound by obligations of confidentiality no less "
            "protective than those in Section 2.",
            "4. Term. This Agreement continues for three (3) years from the date "
            "first written above. The obligations in Section 2 survive expiration for "
            "a further two (2) years.",
            "5. Return of Materials. On written request, each Party will return or "
            "destroy the other Party's Confidential Information, except for copies "
            "retained under routine backup procedures or as required by law.",
            "6. No Licence. Nothing in this Agreement grants either Party any right "
            "in the other Party's intellectual property.",
            "7. Governing Law. This Agreement is governed by the laws of the State of "
            "New York, without regard to its conflict of laws principles.",
            "8. Counterparts. This Agreement may be executed in counterparts, each of "
            "which is an original and all of which together constitute one agreement.",
        ],
        note="Plain, correct US-style NDA. Defined terms used throughout, clean numbering.",
    ),
    Spec(
        name="clean_engagement_uk",
        kind="engagement",
        clean=True,
        headings=[(1, "Engagement Letter")],
        paragraphs=[
            "Dear Ms Whitfield",
            "We are delighted that you have asked this firm to act for Whitfield "
            "Logistics Limited (the “Company”) in connection with the proposed "
            "acquisition of Northgate Freight Limited (the “Transaction”).",
            "1. Scope. We will advise the Company on the Transaction, including "
            "reviewing the share purchase agreement, conducting legal due diligence "
            "and advising on completion mechanics. Anything outside that scope will "
            "be agreed with you in writing before we begin it.",
            "2. Our Team. The engagement partner is Margaret Ellison. The day to day "
            "contact will be Thomas Reyes, a senior associate. We will tell you "
            "before we change either.",
            "3. Fees. Our fees are charged on a time basis at the rates set out in "
            "Schedule 1. We estimate our fees for the Transaction will be between "
            "£85,000 and £110,000 plus VAT and disbursements. That estimate is not "
            "a cap, and we will tell you promptly if we expect to exceed it.",
            "4. Billing. We will invoice monthly in arrears. Our invoices are payable "
            "within 30 days of receipt.",
            "5. Confidentiality. We will keep the Company's affairs confidential, "
            "except where disclosure is required by law or by our regulator.",
            "6. Termination. The Company may end this engagement at any time on "
            "written notice. We may end it on reasonable notice where we are "
            "permitted to do so by our professional rules.",
            "7. Governing Law. This engagement is governed by the law of England and "
            "Wales and the courts of England and Wales have exclusive jurisdiction.",
            "Please confirm your agreement by signing and returning the enclosed copy "
            "of this letter.",
            "Yours sincerely",
            "Margaret Ellison, Partner",
        ],
        table=[
            ["Fee earner", "Rate per hour"],
            ["Partner", "£650"],
            ["Senior associate", "£420"],
            ["Associate", "£310"],
            ["Trainee", "£180"],
        ],
        note="British conventions throughout: 'Ms' without a full stop, sterling, "
             "England and Wales, a rate table. Tests that we do not impose US style.",
    ),
    Spec(
        name="clean_msa_schedules",
        kind="msa",
        clean=True,
        headings=[(1, "Master Services Agreement")],
        paragraphs=[
            "This Master Services Agreement (this “Agreement”) is entered into as of "
            "January 12, 2026 by and between Cortland Systems Inc. (“Supplier”) and "
            "Harrow Bank plc (“Customer”).",
            "1. Services. Supplier will provide the services described in each "
            "Statement of Work executed under this Agreement (the “Services”). Each "
            "Statement of Work is governed by this Agreement and, in the event of "
            "conflict, the terms of this Agreement prevail except where the Statement "
            "of Work expressly says otherwise.",
            "2. Fees. Customer will pay the fees set out in the applicable Statement "
            "of Work. Undisputed invoices are payable within forty-five (45) days.",
            "3. Service Levels. Supplier will meet the service levels in Schedule 2. "
            "Service credits are Customer's sole financial remedy for a failure to "
            "meet a service level, without limiting its right to terminate under "
            "Section 8.",
            "4. Data Protection. Each party will comply with applicable data "
            "protection law. The data processing terms in Schedule 3 apply to any "
            "processing of personal data carried out by Supplier on Customer's behalf.",
            "5. Intellectual Property. Customer retains all right, title and interest "
            "in Customer Data. Supplier retains all right, title and interest in the "
            "Supplier Platform and in anything it develops independently of this "
            "Agreement.",
            "6. Confidentiality. Each party will protect the other's confidential "
            "information on the terms set out in Schedule 4.",
            "7. Limitation of Liability. Neither party's aggregate liability arising "
            "out of this Agreement will exceed the fees paid and payable in the "
            "twelve (12) months before the event giving rise to the claim. Nothing in "
            "this Agreement limits liability that cannot be limited by law.",
            "8. Term and Termination. This Agreement continues until terminated. "
            "Either party may terminate for material breach that remains uncured "
            "thirty (30) days after written notice.",
            "9. Governing Law. This Agreement is governed by the laws of England and "
            "Wales.",
            "Schedule 1. Form of Statement of Work.",
            "Schedule 2. Service Levels.",
            "Schedule 3. Data Processing Terms.",
            "Schedule 4. Confidentiality Terms.",
        ],
        note="Heavy cross-referencing to Schedules and Sections, all of which exist. "
             "Tests the cross-reference check against a document that is correct.",
    ),
    Spec(
        name="clean_employment",
        kind="employment",
        clean=True,
        headings=[(1, "Employment Agreement")],
        paragraphs=[
            "This Employment Agreement (this “Agreement”) is made as of April 1, 2026 "
            "between Redpoint Analytics Inc. (the “Company”) and Daniel Okonkwo (the "
            "“Executive”).",
            "1. Position. The Executive will serve as Chief Financial Officer and will "
            "report to the Chief Executive Officer.",
            "2. Compensation. The Company will pay the Executive an annual base salary "
            "of three hundred twenty thousand dollars ($320,000), payable in "
            "accordance with the Company's regular payroll practices.",
            "3. Bonus. The Executive is eligible for an annual target bonus of forty "
            "percent (40%) of base salary, subject to the achievement of objectives "
            "set by the Board.",
            "4. Benefits. The Executive is entitled to participate in the benefit "
            "plans the Company makes generally available to its senior executives.",
            "5. Restrictive Covenants. For twelve (12) months after the Executive's "
            "employment ends, the Executive will not solicit any employee of the "
            "Company whom the Executive supervised or worked with in the final year "
            "of employment.",
            "6. Termination. Either party may end the employment on sixty (60) days' "
            "written notice. The Company may end it immediately for Cause as defined "
            "in Section 7.",
            "7. Cause. “Cause” means the Executive's conviction of a felony, "
            "material breach of this Agreement that remains uncured ten (10) days "
            "after notice, or wilful misconduct materially injurious to the Company.",
            "8. Governing Law. This Agreement is governed by the laws of the State of "
            "California.",
        ],
        note="Numbers written in words and numerals correctly throughout. Tests the "
             "amount check does not fire on correct pairs.",
    ),
    Spec(
        name="clean_board_minutes",
        kind="minutes",
        clean=True,
        headings=[(1, "Minutes of a Meeting of the Board of Directors")],
        paragraphs=[
            "Minutes of a meeting of the board of directors of Ashford Partners "
            "Limited held at 10:00 a.m. on 14 February 2026 at the registered office.",
            "Present: J. Ashford (Chair), M. Ellison, P. Nakamura, S. Owusu.",
            "In attendance: R. Chen, Company Secretary.",
            "1. Quorum. The Chair noted that a quorum was present and declared the "
            "meeting open.",
            "2. Minutes of the previous meeting. The minutes of the meeting held on "
            "17 January 2026 were approved as a correct record.",
            "3. Conflicts. No director declared an interest in any matter on the "
            "agenda.",
            "4. Banking. The board considered the proposed facility with Harrow Bank "
            "plc. After discussion it was resolved that the facility be approved and "
            "that any two directors be authorised to execute the facility documents.",
            "5. Any other business. There was none.",
            "6. Close. There being no further business the Chair closed the meeting "
            "at 11:15 a.m.",
        ],
        note="British date format, initials with full stops, times. Tests we do not "
             "flag names, initials or a.m./p.m. as defects.",
    ),
    Spec(
        name="clean_brief_citations",
        kind="brief",
        clean=True,
        headings=[(1, "Memorandum in Support of Motion to Dismiss")],
        paragraphs=[
            "Defendant respectfully submits this memorandum in support of its motion "
            "to dismiss the complaint under Rule 12(b)(6).",
            "I. Standard of Review. To survive a motion to dismiss, a complaint must "
            "contain sufficient factual matter to state a claim that is plausible on "
            "its face. See Ashcroft v. Iqbal, 556 U.S. 662, 678 (2009); Bell Atlantic "
            "Corp. v. Twombly, 550 U.S. 544, 570 (2007).",
            "II. Argument. The complaint alleges no facts showing the existence of an "
            "agreement. Plaintiff pleads only that the parties “discussed terms” "
            "[Compl. ¶ 14] and that Defendant “indicated willingness” to proceed "
            "[id. ¶ 17]. Those allegations are precisely the sort of conclusory "
            "assertions Iqbal forecloses. 556 U.S. at 681.",
            "III. The Statute of Frauds Bars the Claim. Even if an agreement were "
            "adequately pleaded, it would be unenforceable. See N.Y. Gen. Oblig. Law "
            "§ 5-701(a)(1). The alleged agreement could not be performed within one "
            "year [see Compl. ¶ 22].",
            "IV. Conclusion. For the reasons stated above, the complaint should be "
            "dismissed with prejudice.",
        ],
        note="Roman numeral headings, bracketed record citations, section symbols, "
             "and years in parentheses. This document exists to break the "
             "placeholder and cross-reference checks if they are naive.",
    ),
    Spec(
        name="clean_lease_short",
        kind="lease",
        clean=True,
        headings=[(1, "Deed of Lease")],
        paragraphs=[
            "This Lease is made on 1 May 2026 between Kestrel Property Holdings "
            "Limited (the “Landlord”) and Marchmont Retail Limited (the “Tenant”).",
            "1. Demise. The Landlord lets the Premises to the Tenant for the Term at "
            "the Rent.",
            "2. Premises. The “Premises” means the ground floor retail unit at 14 "
            "Bridge Street, Manchester, shown edged red on the plan annexed to this "
            "Lease.",
            "3. Term. The “Term” means ten (10) years beginning on 1 June 2026.",
            "4. Rent. The “Rent” means fifty-two thousand pounds (£52,000) per year, "
            "payable quarterly in advance on the usual quarter days.",
            "5. Rent Review. The Rent is reviewed on the fifth anniversary of the "
            "start of the Term in accordance with Schedule 1.",
            "6. Repair. The Tenant will keep the Premises in good and substantial "
            "repair, damage by insured risks excepted.",
            "7. Alienation. The Tenant may not assign or underlet the whole or any "
            "part of the Premises without the Landlord's prior written consent, such "
            "consent not to be unreasonably withheld.",
            "8. Governing Law. This Lease is governed by the law of England and Wales.",
            "Schedule 1. Rent Review Provisions.",
        ],
        note="Addresses with numbers, a plan reference, quarter days. Tests that "
             "street numbers are not read as cross-references.",
    ),
    Spec(
        name="clean_spa_execution_us",
        kind="spa",
        clean=True,
        headings=[(1, "Share Purchase Agreement")],
        paragraphs=[
            "This Agreement is made as of March 3, 2026 between Acme Holdings, Inc., "
            "a Delaware corporation (“Buyer”), and Northgate Freight Ltd, a company "
            "incorporated in England and Wales (“Seller”).",
            "RECITALS",
            "WHEREAS, Seller is the legal and beneficial owner of the entire issued "
            "share capital of Target Logistics Ltd (the “Company”).",
            "1. Sale. Seller sells and Buyer purchases the Shares on the terms of this "
            "Agreement.",
            "2. Price. The price is one million dollars ($1,000,000), payable at "
            "Completion.",
            "3. Warranties. Seller gives the warranties in Schedule 1.",
            "4. Governing Law. This Agreement is governed by the laws of the State of "
            "New York.",
            "IN WITNESS WHEREOF, the parties have executed this Agreement as of the "
            "date first written above.",
            "ACME HOLDINGS, INC.",
            "By: ______________________",
            "Name: Jane Doe",
            "Title: Chief Executive Officer",
            "NORTHGATE FREIGHT LTD",
            "By: ______________________",
            "Name: John Smith",
            "Title: Director",
            "Schedule 1. Warranties.",
        ],
        note="An execution version with a complete American signature page. The "
             "target company is named in the recitals and does not sign; "
             "\"a Delaware corporation\" is a description, not a party.",
    ),
    Spec(
        name="clean_spa_execution_uk",
        kind="spa",
        clean=True,
        headings=[(1, "Share Purchase Agreement")],
        paragraphs=[
            "PARTIES",
            "(1) ACME LIMITED, a company incorporated in England and Wales with "
            "registered number 01234567 whose registered office is at 1 High Street, "
            "London EC1A 1AA (the “Seller”); and",
            "(2) NORTHGATE FREIGHT LIMITED, a company incorporated in England and "
            "Wales with registered number 07654321 (the “Buyer”).",
            "BACKGROUND",
            "(A) The Seller is the legal and beneficial owner of the Shares in Target "
            "Logistics Limited (the “Company”).",
            "1. Interpretation. In this Agreement the definitions in Schedule 1 apply.",
            "2. Sale. The Seller shall sell and the Buyer shall buy the Shares with "
            "effect from Completion.",
            "3. Consideration. The consideration is one million pounds (£1,000,000).",
            "EXECUTED as a DEED by ACME LIMITED acting by a director in the presence "
            "of a witness",
            "______________________",
            "Director",
            "EXECUTED as a DEED by NORTHGATE FREIGHT LIMITED acting by a director in "
            "the presence of a witness",
            "______________________",
            "Director",
            "Schedule 1. Definitions.",
        ],
        note="An English execution page. Names are written by hand at signing, so "
             "every block is blank and none of them is an omission.",
    ),
]


# Realistic drafting: Word's list numbering, definitions tables, schedules in
# parts, deliberate "[Reserved]" and "[***]", statutes cited by section. Each
# of these raised findings on a correct document before the checks learned to
# read them; see evals/realistic.py.
CLEAN += [
    Spec(name="clean_spa_uk_typed", kind="spa", clean=True, paragraphs=[],
         builder=lambda: realistic.uk_spa(auto_numbered=False),
         note="English SPA, clause numbers typed. Definitions run on into (a)/(b); "
              "Schedule 4 restarts numbering in each Part; LLP execution block."),
    Spec(name="clean_spa_uk_auto", kind="spa", clean=True, paragraphs=[],
         builder=lambda: realistic.uk_spa(auto_numbered=True),
         note="The same SPA with its clauses numbered by Word's list numbering."),
    Spec(name="clean_spa_us_auto", kind="spa", clean=True, paragraphs=[],
         builder=realistic.us_spa,
         note="US stock purchase agreement, auto-numbered, definitions table, "
              "[Reserved], [***], Title XX, Treasury Regulations, capitals on the "
              "signature page."),
    Spec(name="clean_style_separator", kind="spa", clean=True, paragraphs=[],
         builder=realistic.style_separator,
         note="A run-in heading joined by a hidden paragraph mark."),    Spec(name="clean_falcon_spa", kind="spa", clean=True, paragraphs=[],
         builder=lambda: realistic.falcon_spa(3),
         note="Project Falcon, version 3: the Buyer renamed everywhere, instalments "
              "and Schedule 2 adding up to the Purchase Price, a shareholding table "
              "totalling 100%, claims outlasting the retention."),
]



# Time limits everywhere: an initial term with automatic renewal, notice to
# stop it, a cure period, payment days, a long stop date after signing, a
# survival period and a rent-style anniversary. Every one of them agrees with
# every other, so the deadline checks must stay silent (deadlines.py).
CLEAN += [
    Spec(
        name="clean_supply_deadlines",
        kind="msa",
        clean=True,
        headings=[(1, "Supply Agreement")],
        paragraphs=[
            "This Supply Agreement (this “Agreement”) is dated 2 February 2026 and "
            "is made between Kestrel Components Limited (the “Supplier”) and "
            "Marchmont Retail Limited (the “Customer”).",
            "1. Definitions. “Business Day” means a day other than a Saturday, "
            "Sunday or public holiday in England. “Long Stop Date” means 31 March "
            "2027. “Effective Date” means 1 March 2026.",
            "2. Term. This Agreement starts on the Effective Date and continues for "
            "an initial term of twenty-four (24) months (the “Initial Term”). It then "
            "automatically renews for successive periods of twelve (12) months (each "
            "a “Renewal Period”) unless either party gives not less than ninety (90) "
            "days' written notice before the end of the Initial Term or the "
            "then-current Renewal Period.",
            "3. Conditions. If the conditions in clause 4 are not satisfied on or "
            "before the Long Stop Date, either party may terminate this Agreement by "
            "notice to the other.",
            "4. Conditions Precedent. The Customer shall deliver its first forecast "
            "within 10 Business Days after the Effective Date.",
            "5. Payment. The Customer shall pay each undisputed invoice within thirty "
            "(30) days of receipt.",
            "6. Breach. Either party may terminate this Agreement if the other "
            "commits a material breach that is not remedied within fourteen (14) "
            "days after written notice requiring it to be remedied.",
            "7. Review. The prices are reviewed on the first anniversary of the "
            "Effective Date and on each anniversary after that.",
            "8. Survival. Clause 5 survives termination for a further six (6) months.",
            "9. Governing Law. This Agreement is governed by the law of England and "
            "Wales.",
        ],
        note="Every kind of time limit the deadlines table reads, all consistent: "
             "90 days' notice against a 12-month renewal, a long stop after signing.",
    ),
]

# ==========================================================================
# DEFECTIVE DOCUMENTS. Each declares what was planted.
# ==========================================================================

def _defect(spec: Spec, **kw) -> Spec:
    return spec


DEFECTIVE: list[Spec] = [
    Spec(
        name="defect_placeholder",
        kind="nda",
        clean=False,
        headings=[(1, "Mutual Non-Disclosure Agreement")],
        paragraphs=[
            "This Agreement is made as of March 3, 2026 between Acme Holdings LLC "
            "(“Acme”) and [COUNTERPARTY NAME] (“Counterparty”).",
            "1. Definitions. “Confidential Information” means information disclosed "
            "by one party to the other. Confidential Information includes the terms "
            "of this Agreement.",
            "2. Obligations. Each party will protect the other's Confidential "
            "Information and will not disclose it without consent.",
            "3. Term. This Agreement continues for TBD years.",
            "4. Governing Law. This Agreement is governed by the laws of the State of "
            "New York.",
        ],
        expect={"placeholder"},
        note="Unfilled bracket and a TBD. Both must be caught.",
    ),
    Spec(
        name="defect_amount_mismatch",
        kind="employment",
        clean=False,
        headings=[(1, "Employment Agreement")],
        paragraphs=[
            "This Agreement is made as of April 1, 2026 between Redpoint Analytics "
            "Inc. (the “Company”) and Daniel Okonkwo (the “Executive”).",
            "1. Position. The Executive will serve as Chief Financial Officer.",
            "2. Notice. Either party may end the employment on thirty (13) days' "
            "written notice.",
            "3. Covenants. The Executive will not solicit employees for twelve (24) "
            "months after termination.",
            "4. Governing Law. This Agreement is governed by the laws of California.",
        ],
        expect={"amount"},
        note="Two word/numeral mismatches. Classic and dangerous.",
    ),
    Spec(
        name="defect_broken_crossref",
        kind="msa",
        clean=False,
        headings=[(1, "Master Services Agreement")],
        paragraphs=[
            "This Agreement is entered into as of January 12, 2026 between Cortland "
            "Systems Inc. (“Supplier”) and Harrow Bank plc (“Customer”).",
            "1. Services. Supplier will provide the Services described in each "
            "Statement of Work.",
            "2. Fees. Customer will pay the fees set out in the Statement of Work, "
            "subject to the credits described in Section 9.",
            "3. Confidentiality. Each party will protect the other's confidential "
            "information as set out in Schedule 7.",
            "4. Governing Law. This Agreement is governed by the laws of England and "
            "Wales.",
        ],
        expect={"cross-reference"},
        note="References Section 9 and Schedule 7; neither exists.",
    ),
    Spec(
        name="defect_numbering_gap",
        kind="msa",
        clean=False,
        headings=[(1, "Supply Agreement")],
        paragraphs=[
            "This Agreement is made on 2 March 2026 between Two parties.",
            "1. Supply. The Supplier will supply the Goods.",
            "2. Price. The Buyer will pay the Price.",
            "4. Delivery. The Supplier will deliver the Goods to the Site.",
            "5. Title. Title passes on payment in full.",
            "5. Risk. Risk passes on delivery.",
            "6. Governing Law. This Agreement is governed by the law of England and "
            "Wales.",
        ],
        expect={"numbering"},
        note="Skips 3 and repeats 5.",
    ),
    Spec(
        name="defect_date_contradiction",
        kind="nda",
        clean=False,
        headings=[(1, "Confidentiality Agreement")],
        paragraphs=[
            "This Agreement is entered into as of March 3, 2026 between Acme Holdings "
            "LLC and Beta Industries Inc.",
            "1. Confidentiality. Each party will protect the other's confidential "
            "information.",
            "2. Term. This Agreement continues for three (3) years from the date "
            "first written above, being March 5, 2026.",
            "3. Governing Law. New York law governs this Agreement.",
        ],
        expect={"date"},
        note="Self-referential date contradicts the opening date.",
    ),
    Spec(
        name="defect_party_name_drift",
        kind="spa",
        clean=False,
        headings=[(1, "Share Purchase Agreement")],
        paragraphs=[
            "This Agreement is made on 4 March 2026 between Northgate Freight Limited "
            "(the “Seller”) and Whitfield Logistics Limited (the “Buyer”).",
            "1. Sale. The Seller will sell and the Buyer will buy the Shares.",
            "2. Consideration. The Buyer will pay the Consideration at Completion.",
            "3. Warranties. Northgate Freight Ltd gives the warranties in Schedule 1.",
            "4. Indemnity. Whitfield Logistics LLP is indemnified against the Claims.",
            "5. Governing Law. This Agreement is governed by the law of England and "
            "Wales.",
            "Schedule 1. Warranties.",
        ],
        expect={"party-name"},
        note="Limited vs Ltd, and a Limited company that becomes an LLP.",
    ),
    Spec(
        name="defect_defined_term_misspelled",
        kind="nda",
        clean=False,
        headings=[(1, "Non-Disclosure Agreement")],
        paragraphs=[
            "This Agreement is made as of 3 March 2026 between Acme Holdings LLC and "
            "Beta Industries Inc.",
            "1. Definitions. In this Agreement, “Recipient” means the party "
            "receiving Confidential Information.",
            "2. Obligations. The Recipent will protect Confidential Information and "
            "will not disclose it.",
            "3. Return. The Recipient will return all materials on request.",
            "4. Governing Law. New York law governs this Agreement.",
        ],
        expect={"defined-term"},
        note="One-character misspelling of a defined term. Safely auto-fixable.",
    ),
    Spec(
        name="defect_term_used_before_defined",
        kind="msa",
        clean=False,
        headings=[(1, "Services Agreement")],
        paragraphs=[
            "This Agreement is made on 12 January 2026 between Cortland Systems Inc. "
            "and Harrow Bank plc.",
            "1. Charges. The Customer will pay the Charges within thirty (30) days.",
            "2. Services. The Supplier will perform the Services with reasonable care.",
            "3. Definitions. In this Agreement, “Charges” means the amounts set out "
            "in Schedule 1, and “Services” means the services described in "
            "Schedule 2.",
            "4. Term. This Agreement continues for two (2) years.",
            "5. Liability. Neither party limits liability for fraud.",
            "6. Notices. Notices must be in writing.",
            "7. Assignment. Neither party may assign without consent.",
            "8. Variation. No variation is effective unless in writing.",
            "9. Governing Law. English law governs this Agreement.",
            "10. Jurisdiction. The English courts have exclusive jurisdiction.",
            "11. Counterparts. This Agreement may be executed in counterparts.",
            "Schedule 1. Charges.",
            "Schedule 2. Services.",
        ],
        expect={"defined-term"},
        note="Charges and Services are both used in clause 1 and 2, defined in 3.",
    ),
    Spec(
        name="defect_mixed_dates",
        kind="engagement",
        clean=False,
        headings=[(1, "Engagement Letter")],
        paragraphs=[
            "This engagement begins on March 3, 2026.",
            "1. Scope. We will advise on the Transaction.",
            "2. Estimate. Our estimate is valid until 30 April 2026.",
            "3. Review. We will review our fees on 2026-09-01.",
            "4. Governing Law. English law governs this engagement.",
        ],
        expect={"date"},
        note="Three different date formats in one short letter.",
    ),
    Spec(
        name="defect_multiple",
        kind="spa",
        clean=False,
        headings=[(1, "Share Purchase Agreement")],
        paragraphs=[
            "This Agreement is made as of March 3, 2026 between Acme Holdings LLC "
            "(“Seller”) and [BUYER NAME] (“Buyer”).",
            "1. Sale. The Seller will sell the Shares to the Buyer.",
            "2. Price. The price is one million dollars ($1,500,000), payable at "
            "Completion.",
            "3. Completion. Completion takes place on the date set out in Section 8.",
            "5. Warranties. Acme Holdings Inc. gives the warranties in Schedule 1.",
            "6. Governing Law. This Agreement is governed by the laws of New York, "
            "and the estimate expires 30 April 2026.",
            "Schedule 1. Warranties.",
        ],
        expect={"placeholder", "amount", "cross-reference", "numbering",
                "party-name", "date"},
        note="Everything at once. Tests that findings do not mask each other.",
    ),
    Spec(
        name="defect_signature_page",
        kind="spa",
        clean=False,
        headings=[(1, "Share Purchase Agreement")],
        paragraphs=[
            "This Agreement is made between Acme Holdings, Inc. (“Buyer”), Northgate "
            "Freight Ltd (“Seller”) and Harbour Guarantee LLC (“Guarantor”).",
            "1. Sale. Seller sells and Buyer purchases the Shares.",
            "2. Guarantee. Guarantor guarantees the obligations of Seller.",
            "IN WITNESS WHEREOF, the parties have executed this Agreement.",
            "ACME HOLDINGS, INC.",
            "By: ______________________",
            "Name: Jane Doe",
            "Title: Chief Executive Officer",
            "BLUEWATER SHIPPING LLC",
            "By: ______________________",
            "Name:",
            "Title:",
        ],
        expect={"signature"},
        note="The previous deal's signature page: the Guarantor has no block, a "
             "stranger signs, and its block is half filled where the other is "
             "complete.",
    ),
    Spec(
        name="defect_punctuation",
        kind="msa",
        clean=False,
        headings=[(1, "Master Services Agreement")],
        paragraphs=[
            "This Agreement is made between Acme Holdings LLC (“Customer”) and Beta "
            "Services Inc. (“Supplier”).",
            "1. Services. Supplier will provide the the Services described in "
            "Schedule 1 ,and Customer will pay the Charges.",
            "2. Term. The term is two (2) years,and renews for one (1) year.",
            "3. Governing Law. This Agreement is governed by the laws of New York.",
            "Schedule 1. Services.",
        ],
        expect={"punctuation"},
        note="A repeated word, a space on the wrong side of a comma, and a missing "
             "space after one. Each is written as a tracked change.",
    ),
]


DEFECTIVE += [
    Spec(name="defect_falcon_spa", kind="spa", clean=False, paragraphs=[],
         builder=lambda: realistic.falcon_spa(3, defects=True),
         expect={"arithmetic", "party-name"},
         note="The Buyer renamed in the parties clause but not on its signature block "
              "or in the notice details; Schedule 2 totalling £12,050,000 against a "
              "£12,500,000 Purchase Price; instalments adding up to £12,750,000."),
    Spec(name="defect_uk_drafting", kind="spa", clean=False, paragraphs=[],
         builder=realistic.uk_defects,
         expect={"placeholder", "defined-term", "cross-reference", "amount", "date",
                 "party-name"},
         note="[●], [x], [date], [name of bank], drafting notes, a blank in a content "
              "control, a term defined twice differently, figures disagreeing with "
              "the words that follow them, three dollar notations, a date whose "
              "weekday is wrong, and a lower-case reference to a missing clause."),
    Spec(name="defect_moved_and_reformatted", kind="msa", clean=False, paragraphs=[],
         builder=realistic.moved_and_reformatted,
         expect={"leftovers"},
         note="Tracked changes that are a move and a formatting change."),
]


DEFECTIVE += [
    Spec(
        name="defect_deadlines",
        kind="msa",
        clean=False,
        headings=[(1, "Services Agreement")],
        paragraphs=[
            "This Services Agreement (this “Agreement”) is dated 2 February 2026 and "
            "is made between Kestrel Components Limited (the “Supplier”) and "
            "Marchmont Retail Limited (the “Customer”).",
            "1. Definitions. “Long Stop Date” means 31 December 2025.",
            "2. Term. This Agreement continues for an initial term of twelve (12) "
            "months (the “Initial Term”) and then automatically renews for "
            "successive periods of three (3) months (each a “Renewal Period”) unless "
            "either party gives not less than six (6) months' written notice before "
            "the end of the then-current Renewal Period.",
            "3. Extension. The Customer may extend the Initial Term of twenty-four "
            "(24) months once by notice.",
            "4. Payment. The Customer shall pay each invoice within thirty (30) days "
            "of receipt.",
            "5. Governing Law. This Agreement is governed by the law of England and "
            "Wales.",
        ],
        expect={"deadline"},
        note="A long stop date before the agreement is dated, six months' notice "
             "before the end of a three-month renewal, and the Initial Term given "
             "as twelve months in one clause and twenty-four in another.",
    ),
]


ALL: list[Spec] = CLEAN + DEFECTIVE


def by_name(name: str) -> Spec:
    for spec in ALL:
        if spec.name == name:
            return spec
    raise KeyError(name)
