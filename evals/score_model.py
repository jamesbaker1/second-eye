"""Score the model-dependent features against a corpus with known answers.

evals/run_checks.py scores the deterministic layer. This scores what the model
does: playbook findings, the answer to the other side's draft, the key-terms
table, and the voice of what it writes. Its job is that the first live run
produces a scorecard, not an impression.

Four ways to run it, one scorer:

  --stub perfect|bad   canned results, no network. Proves the scorer can tell
                       a right answer from a wrong one; the tests use it.
  --live               every document in evals/model_corpus.py through the real
                       review path (pipeline/review.py on Managed Agents), with
                       a dollar and a time budget. Every result is saved.
  --replay <dir>       score results saved by an earlier --live run, so the
                       scorer can change without paying for the reviews again.

  .venv/bin/python -m evals.score_model --live --budget-usd 40
  lra eval --live --budget-usd 40            (the same thing)

  lra eval --triage --stub perfect|bad       triage, canned plans, and the rules'
                                             baseline on the same emails
  lra eval --triage --live                   the first live triage run

Writes scorecard.md and scorecard.json into --out (default
work/model-eval/<time>-<how>/). Exits 0 when every target is met.

What is scored, per set (evals/model_corpus.py):

  playbook recall      each planted off-position clause is owed a finding with
                       category `playbook` whose title names the clause, or
                       whose anchor falls inside the planted text
  false positives      any finding on an on-position or realistic document;
                       any finding on a planted document that matches nothing
  their paper          any tracked change written into their text, and any
                       planted position missing from the issues list
  key terms            parties, term, renewal, cap, governing law, notices:
                       each right, wrong or missing in the table produced
  voice                agents/review_rubric.md and the prompts' reply rules:
                       summary at most one sentence, no send verdict, titles
                       lead with the clause, no system words, starter findings
                       labelled, positions and responses filled on their paper
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from evals import model_corpus
from evals.model_corpus import CLAUSES, OFF, ModelSpec
from lra.models import Attachment, Finding, Mode, ReviewResult, Severity

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ROOT = Path(__file__).resolve().parents[1]

# What a met target means. The false-positive line is PRODUCT.md's: fewer
# than one per clean document, the same bar evals/run_checks.py holds the
# deterministic layer to.
TARGETS: dict[str, tuple[str, float]] = {
    "playbook_recall": (">=", 0.9),
    "fp_per_on_position_doc": ("<", 1.0),
    "extra_findings_per_planted_doc": ("<", 1.0),
    "fp_per_realistic_doc": ("<", 1.0),
    "their_paper_edits_in_their_text": ("==", 0),
    "their_paper_positions_missing": ("==", 0),
    "key_terms_accuracy": (">=", 0.9),
    "voice_violations": ("==", 0),
    "errors": ("==", 0),
}

LABELS = {
    "playbook_recall": "Playbook recall (planted off-position clauses found)",
    "fp_per_on_position_doc": "False positives per on-position document",
    "extra_findings_per_planted_doc": "Unplanted findings per playbook document",
    "fp_per_realistic_doc": "False positives per realistic clean agreement",
    "their_paper_edits_in_their_text": "Their paper: edits written into their text",
    "their_paper_positions_missing": "Their paper: planted positions missing from the issues list",
    "key_terms_accuracy": "Key terms: fields right",
    "voice_violations": "Voice: rule breaches",
    "errors": "Reviews that failed",
}


# --------------------------------------------------------------------------
# What a review produced, and how it is saved and read back
# --------------------------------------------------------------------------


@dataclass
class Produced:
    """Everything one document's review produced that the scorer reads."""

    name: str
    result: ReviewResult | None = None
    error: str = ""
    # Why the document was not run at all (budget, time, the account).
    not_run: str = ""
    # On their paper: the redline the lawyer would get, and the issues list.
    redline: bytes | None = None
    issues_list: bytes | None = None
    # Every other file the session left (a key-terms workbook, say).
    outputs: list[tuple[str, bytes]] = field(default_factory=list)
    seconds: float = 0.0
    session_id: str = ""

    def save(self, root: Path) -> Path:
        where = root / self.name
        (where / "outputs").mkdir(parents=True, exist_ok=True)
        record = {
            "name": self.name, "error": self.error, "not_run": self.not_run,
            "seconds": round(self.seconds, 1), "session_id": self.session_id,
            "result": (self.result.model_dump(mode="json", exclude={"outputs"})
                       if self.result else None),
        }
        (where / "result.json").write_text(json.dumps(record, indent=2, ensure_ascii=False))
        for filename, data in (("redline.docx", self.redline),
                               ("issues-list.docx", self.issues_list)):
            if data:
                (where / filename).write_bytes(data)
        for filename, data in self.outputs:
            (where / "outputs" / Path(filename).name).write_bytes(data)
        return where

    @classmethod
    def load(cls, where: Path) -> Produced:
        record = json.loads((where / "result.json").read_text())
        result = ReviewResult(**record["result"]) if record.get("result") else None

        def read(name: str) -> bytes | None:
            path = where / name
            return path.read_bytes() if path.exists() else None

        outputs = sorted((p.name, p.read_bytes()) for p in (where / "outputs").glob("*")
                         if p.is_file())
        return cls(name=record["name"], result=result, error=record.get("error", ""),
                   not_run=record.get("not_run", ""), redline=read("redline.docx"),
                   issues_list=read("issues-list.docx"), outputs=outputs,
                   seconds=record.get("seconds", 0.0),
                   session_id=record.get("session_id", ""))


def load_dir(root: Path) -> dict[str, Produced]:
    return {p.name: Produced.load(p) for p in sorted(root.iterdir())
            if (p / "result.json").exists()}


# --------------------------------------------------------------------------
# Matching findings to what was planted
# --------------------------------------------------------------------------


def _norm(text: str) -> str:
    text = (text or "").replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    return " ".join(text.lower().split())


def _mentions(text: str, aliases: tuple[str, ...]) -> bool:
    text = _norm(text)
    return any(re.search(rf"\b{re.escape(a.strip())}", text) for a in aliases)


def _is_playbook(f: Finding) -> bool:
    return _norm(f.category).replace("_", "-") in ("playbook", "off-playbook")


def matches(f: Finding, planted: model_corpus.Planted) -> bool:
    """A finding answers a planted clause when it is a playbook finding and
    either its title names the clause or its anchor is inside the planted
    text. Category and clause label, as a lawyer would match them."""
    if not _is_playbook(f):
        return False
    anchor = _norm(f.anchor)
    inside = len(anchor) >= 12 and anchor in _norm(planted.text)
    return inside or _mentions(_title_topic(f.title), planted.aliases)


_OFF_PREFIX = re.compile(r"^\s*off[ -]playbook\s*:\s*", re.IGNORECASE)
_WHERE_PREFIX = re.compile(r"^\s*(?:(?:cl(?:ause)?|section|schedule|sch|para(?:graph)?)\.?\s*)?"
                           r"[0-9][\w.()]*\s*[:\-–—]?\s*", re.IGNORECASE)


def _title_topic(title: str) -> str:
    return _WHERE_PREFIX.sub("", _OFF_PREFIX.sub("", title or ""), count=1)


# --------------------------------------------------------------------------
# Voice: agents/review_rubric.md criterion 5, and the reply rules
# --------------------------------------------------------------------------

_SENTENCE_END = re.compile(r"(?<=[.!?])[\"”’)]?\s+(?=[\"“(]?[A-Z])")
_SEND_VERDICT = re.compile(
    r"\b(?:safe|ready|fine|ok|okay|good|clear)\s+to\s+(?:send|go|sign)\b"
    r"|\bcan (?:safely )?be sent\b|\bgood to go\b|\bno (?:issues|problems) (?:found|to flag)\b",
    re.IGNORECASE)
_SYSTEM_WORDS = re.compile(
    r"\b(?:checks?|checked|checker|pass|passed|passes|severity|confidence|tools?|sandbox)\b",
    re.IGNORECASE)
_QUOTED = re.compile(r"“[^”]*”|\"[^\"]*\"|‘[^’]{3,}’")
_CLAUSE_REF = re.compile(r"^\s*(?:cl(?:ause)?\.?|section|schedule|sch\.?|para(?:graph)?\.?|"
                         r"\d+(?:\.\d+)*\b)", re.IGNORECASE)
# Words that name a clause, for "the title leads with the clause". The
# playbook's own aliases, and the ordinary parts of an agreement.
_CLAUSE_WORDS = tuple(sorted({a.strip() for _, aliases in CLAUSES.values() for a in aliases}
                             | {"definition", "defined", "term", "renewal", "notice", "schedule",
                                "part", "recital", "completion", "closing", "consideration",
                                "deposit", "escrow", "warrant", "signature", "execution",
                                "dispute", "governing", "exclusiv", "audit", "service", "fee",
                                "purchase price", "interest", "set-off", "set off", "tax",
                                "condition", "long stop", "limitation", "customer", "supplier"}))


def _sentences(text: str) -> list[str]:
    return [s for s in _SENTENCE_END.split((text or "").strip()) if s.strip()]


def _unquoted(text: str) -> str:
    return _QUOTED.sub(" ", text or "")


def voice(spec: ModelSpec, result: ReviewResult) -> list[dict]:
    """Every breach of the reply rules in one result, as (rule, detail)."""
    breaches: list[dict] = []

    def breach(rule: str, detail: str) -> None:
        breaches.append({"rule": rule, "detail": detail[:160]})

    summary = result.summary or ""
    if spec.kind != "key_terms":
        # On a key-terms request the answer is the table, not one sentence.
        limit = 2 if result.mode is Mode.QUESTION else 1
        if len(_sentences(summary)) > limit:
            breach("summary longer than one sentence", summary)
    if _SEND_VERDICT.search(summary):
        breach("summary says whether to send", summary)
    word = _SYSTEM_WORDS.search(_unquoted(summary))
    if word:
        breach("system word in the summary", f"{word.group(0)!r}: {summary}")

    for f in result.findings:
        if not _leads_with_clause(f.title):
            breach("title does not lead with the clause", f.title)
        for label, text in (("title", f.title), ("explanation", f.explanation),
                            ("our_position", f.our_position), ("response", f.response),
                            ("question", f.question or "")):
            word = _SYSTEM_WORDS.search(_unquoted(text))
            if word:
                breach(f"system word in the {label}", f"{word.group(0)!r}: {text}")
        if _is_playbook(f) and not _norm(f.explanation).startswith("starter playbook"):
            breach("starter finding not labelled as the starter playbook", f.title)
        if spec.kind == "their_paper" and _is_playbook(f) and not (
                f.our_position.strip() and f.response.strip()):
            breach("position or response missing on their paper", f.title)
    return breaches


def _leads_with_clause(title: str) -> bool:
    if _CLAUSE_REF.match(_OFF_PREFIX.sub("", title or "")):
        return True
    opening = " ".join(_norm(_title_topic(title)).split()[:4])
    return any(re.search(rf"\b{re.escape(w)}", opening) for w in _CLAUSE_WORDS)


# --------------------------------------------------------------------------
# Their paper: the redline and the issues list
# --------------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def revisions(content: bytes | None) -> list[str]:
    """Every tracked change in a Word file, as "inserted '...'" or
    "deleted '...'"."""
    if not content:
        return []
    from lxml import etree

    out: list[str] = []
    with zipfile.ZipFile(BytesIO(content)) as package:
        for name in package.namelist():
            if not re.match(r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml$",
                            name):
                continue
            root = etree.fromstring(package.read(name))
            for el in root.iter(f"{_W}ins", f"{_W}del"):
                text = "".join(t.text or "" for t in el.iter(f"{_W}t", f"{_W}delText"))
                verb = "inserted" if el.tag == f"{_W}ins" else "deleted"
                out.append(f"{verb} {text!r}")
    return out


def table_rows(content: bytes | None) -> list[list[str]]:
    """The rows of every table in a Word file, header rows included."""
    if not content:
        return []
    from docx import Document

    rows: list[list[str]] = []
    for table in Document(BytesIO(content)).tables:
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
    return rows


def _issues_rows(content: bytes | None) -> list[list[str]]:
    from lra.pipeline.issues_list import COLUMNS

    return [r for r in table_rows(content) if tuple(r[:4]) != COLUMNS]


def _row_matches(row: list[str], planted: model_corpus.Planted) -> bool:
    clause, their = (row + ["", ""])[:2]
    quoted = _norm(their).strip('"').rstrip("…").strip()
    inside = len(quoted) >= 12 and quoted in _norm(planted.text)
    return inside or _mentions(_title_topic(clause), planted.aliases)


# --------------------------------------------------------------------------
# Key terms: the table, from a workbook, a Word table or the reply text
# --------------------------------------------------------------------------

KEY_TERM_FIELDS: dict[str, tuple[str, ...]] = {
    "parties": ("parties", "party"),
    "term": ("term", "initial term", "duration"),
    "renewal": ("renewal", "term"),
    "cap": ("liability cap", "cap", "limitation of liability", "liability"),
    "governing_law": ("governing law", "law"),
    "notice": ("notices", "notice"),
}


def key_term_rows(produced: Produced) -> tuple[list[tuple[str, str]], str]:
    """(term, value) pairs from whatever the review produced, and where from:
    a workbook first, then a Word table, then the reply text."""
    for filename, data in produced.outputs:
        if filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            sheet = load_workbook(BytesIO(data), read_only=True).worksheets[0]
            rows = [tuple("" if v is None else str(v) for v in r)
                    for r in sheet.iter_rows(values_only=True)]
            pairs = [(r[0], " ".join(r[1:])) for r in rows if len(r) >= 2 and r[0]]
            if pairs:
                return pairs, f"workbook {filename}"
    for filename, data in produced.outputs:
        if filename.lower().endswith(".docx"):
            pairs = [(r[0], " ".join(r[1:])) for r in table_rows(data) if len(r) >= 2]
            if pairs:
                return pairs, f"Word table {filename}"
    result = produced.result
    text = "\n".join([result.summary] + [f.explanation for f in result.findings]) \
        if result else ""
    pairs = []
    for line in text.splitlines():
        cells = [c.strip() for c in re.split(r"\s*\|\s*|\t", line.strip().strip("|")) if c.strip()]
        if len(cells) >= 2 and not set(cells[0]) <= set("-: "):
            pairs.append((cells[0].strip("*"), " ".join(cells[1:])))
            continue
        m = re.match(r"^\W*([A-Za-z][A-Za-z /&-]{1,40}?)\**\s*[:–—-]\s+(.+)$", line.strip())
        if m:
            pairs.append((m.group(1), m.group(2)))
    return pairs, "reply text" if pairs else "nothing"


def score_key_terms(spec: ModelSpec, produced: Produced) -> dict:
    pairs, source = key_term_rows(produced)
    fields: dict[str, str] = {}
    for name, fragments in spec.key_terms.items():
        aliases = KEY_TERM_FIELDS[name]
        values = [v for label, v in pairs
                  if any(re.search(rf"\b{re.escape(a)}\b", _norm(label)) for a in aliases)]
        if not values:
            fields[name] = "missing"
            continue
        joined = _norm(" ".join(values))
        fields[name] = ("right" if all(_norm(frag) in joined for frag in fragments)
                        else "wrong")
    return {"source": source, "fields": fields}


# --------------------------------------------------------------------------
# One document, then all of them
# --------------------------------------------------------------------------


def score_doc(spec: ModelSpec, produced: Produced | None) -> dict:
    card: dict = {"name": spec.name, "kind": spec.kind}
    if produced is None or produced.not_run:
        card["status"] = "not run"
        card["why"] = produced.not_run if produced else "no result"
        return card
    planted = spec.planted_with_text()
    if produced.error or produced.result is None:
        card["status"] = "error"
        card["error"] = produced.error or "no result"
        card["planted"] = [{"clause": p.label, "found": False, "by": ""} for p in planted]
        return card

    result = produced.result
    card["status"] = "scored"
    card["seconds"] = produced.seconds
    findings = list(result.findings)

    used: set[int] = set()
    rows = []
    for p in planted:
        hit = next((i for i, f in enumerate(findings) if matches(f, p)), None)
        if hit is not None:
            used.add(hit)
            # Several findings on one planted clause are one hit, not extras.
            used |= {i for i, f in enumerate(findings) if matches(f, p)}
        rows.append({"clause": p.label, "found": hit is not None,
                     "by": findings[hit].title if hit is not None else ""})
    card["planted"] = rows
    unplanted = [f for i, f in enumerate(findings) if i not in used]
    card["unplanted"] = [f"[{f.severity.value}] {f.title}" for f in unplanted]

    if spec.kind == "their_paper":
        edits = revisions(produced.redline)
        typos = [wrong for wrong, _ in spec.typos]
        issue_rows = _issues_rows(produced.issues_list)
        missing = [p.label for p in planted
                   if not any(_row_matches(r, p) for r in issue_rows)]
        card["their_paper"] = {
            "edits_in_their_text": edits,
            "typos_touched": [t for t in typos
                              if any(t in e or dict(spec.typos)[t] in e for e in edits)],
            "issues_list_rows": len(issue_rows),
            "missing_from_issues_list": missing,
            "cosmetic_reported": sum(f.severity in (Severity.STYLE, Severity.FORMATTING)
                                     for f in findings),
        }
    if spec.kind == "key_terms":
        card["key_terms"] = score_key_terms(spec, produced)
    card["voice"] = voice(spec, result)
    return card


def _ratio(n: float, d: float) -> float | None:
    return round(n / d, 3) if d else None


def score(specs: list[ModelSpec], produced: dict[str, Produced]) -> dict:
    """The scorecard as data: one entry per document, the metrics, the targets."""
    docs = [score_doc(s, produced.get(s.name)) for s in specs]
    ran = [d for d in docs if d["status"] != "not run"]
    scored = [d for d in ran if d["status"] == "scored"]

    def of(kind: str, status: str | None = None) -> list[dict]:
        return [d for d in (scored if status == "scored" else ran) if d["kind"] == kind]

    planted = [p for d in of("playbook") for p in d.get("planted", [])]
    silence = of("silence", "scored")
    realistic = of("realistic", "scored")
    playbook = of("playbook", "scored")
    theirs = of("their_paper", "scored")
    key_terms = of("key_terms", "scored")
    field_marks = [v for d in key_terms for v in d["key_terms"]["fields"].values()]
    their_planted = sum(len(d["planted"]) for d in theirs)
    their_missing = sum(len(d["their_paper"]["missing_from_issues_list"]) for d in theirs)
    voice_all = [v for d in scored for v in d["voice"]]

    metrics = {
        "playbook_recall": _ratio(sum(p["found"] for p in planted), len(planted)),
        "fp_per_on_position_doc": _ratio(sum(len(d["unplanted"]) for d in silence),
                                         len(silence)),
        "extra_findings_per_planted_doc": _ratio(sum(len(d["unplanted"]) for d in playbook),
                                                 len(playbook)),
        "fp_per_realistic_doc": _ratio(sum(len(d["unplanted"]) for d in realistic),
                                       len(realistic)),
        "their_paper_edits_in_their_text": (sum(len(d["their_paper"]["edits_in_their_text"])
                                                for d in theirs) if theirs else None),
        "their_paper_positions_missing": their_missing if theirs else None,
        "key_terms_accuracy": _ratio(field_marks.count("right"), len(field_marks)),
        "voice_violations": len(voice_all) if scored else None,
        "errors": sum(d["status"] == "error" for d in ran),
    }
    met = {k: _meets(metrics[k], *TARGETS[k]) for k in TARGETS}
    counts = {
        "documents": len(specs), "run": len(ran), "scored": len(scored),
        "not_run": len(docs) - len(ran), "planted_positions": len(planted),
        "their_paper_positions": their_planted, "key_term_fields": len(field_marks),
        "on_position_documents": len(silence), "realistic_documents": len(realistic),
    }
    return {
        "metrics": metrics, "targets": {k: list(v) for k, v in TARGETS.items()}, "met": met,
        "all_met": all(v is not False for v in met.values()) and bool(scored),
        "counts": counts,
        "voice_by_rule": dict(Counter(v["rule"] for v in voice_all).most_common()),
        "documents": docs,
    }


def _meets(value, op: str, target: float) -> bool | None:
    if value is None:
        return None
    return {">=": value >= target, "<": value < target, "==": value == target}[op]


# --------------------------------------------------------------------------
# The scorecard as markdown
# --------------------------------------------------------------------------


def _fmt(key: str, value) -> str:
    if value is None:
        return "n/a"
    if key in ("playbook_recall", "key_terms_accuracy"):
        return f"{100 * value:.0f}%"
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


def render(card: dict, title: str = "Model scorecard") -> str:
    m, c = card["metrics"], card["counts"]
    lines = [f"# {title}", ""]
    lines.append(f"{c['scored']} of {c['documents']} documents scored"
                 + (f", {c['not_run']} not run" if c["not_run"] else "")
                 + f". Result: **{'meets every target' if card['all_met'] else 'does not meet target'}**.")
    lines += ["", "| Metric | Value | Target | |", "| --- | --- | --- | --- |"]
    for key, (op, target) in TARGETS.items():
        met = card["met"][key]
        mark = "ok" if met else ("MISS" if met is False else "-")
        shown_target = (f"{op} {100 * target:.0f}%" if key in ("playbook_recall",
                                                              "key_terms_accuracy")
                        else f"{op} {target:g}")
        lines.append(f"| {LABELS[key]} | {_fmt(key, m[key])} | {shown_target} | {mark} |")
    lines += ["", (f"Planted positions: {c['planted_positions']} on the playbook set, "
                  f"{c['their_paper_positions']} on their paper. On-position documents: "
                  f"{c['on_position_documents']}. Realistic: {c['realistic_documents']}. "
                  f"Key-term fields: {c['key_term_fields']}.")]

    lines += ["", "## Documents", "", "| Document | Set | Outcome |", "| --- | --- | --- |"]
    for d in card["documents"]:
        lines.append(f"| {d['name']} | {d['kind']} | {_outcome(d)} |")

    missed = [(d["name"], p["clause"]) for d in card["documents"]
              for p in d.get("planted", []) if not p["found"] and d["kind"] == "playbook"]
    if missed:
        lines += ["", "## Planted positions missed", ""]
        lines += [f"- {name}: {clause}" for name, clause in missed]
    fps = [(d["name"], t) for d in card["documents"] for t in d.get("unplanted", [])]
    if fps:
        lines += ["", "## Findings nothing was planted for", ""]
        lines += [f"- {name}: {t}" for name, t in fps]
    tp = [d for d in card["documents"] if d.get("their_paper")]
    if any(d["their_paper"]["edits_in_their_text"] or d["their_paper"]["missing_from_issues_list"]
           for d in tp):
        lines += ["", "## Their paper", ""]
        for d in tp:
            typos = d["their_paper"]["typos_touched"]
            for e in d["their_paper"]["edits_in_their_text"]:
                planted = " (a planted typo)" if any(t in e for t in typos) else ""
                lines.append(f"- {d['name']}: {e} in their text{planted}")
            for p in d["their_paper"]["missing_from_issues_list"]:
                lines.append(f"- {d['name']}: {p} is not in the issues list")
    kt = [d for d in card["documents"] if d.get("key_terms")]
    if any(v != "right" for d in kt for v in d["key_terms"]["fields"].values()):
        lines += ["", "## Key terms", ""]
        for d in kt:
            wrong = {k: v for k, v in d["key_terms"]["fields"].items() if v != "right"}
            if wrong:
                lines.append(f"- {d['name']} (from {d['key_terms']['source']}): "
                             + ", ".join(f"{k} {v}" for k, v in wrong.items()))
    if card["voice_by_rule"]:
        lines += ["", "## Voice", ""]
        for rule, n in card["voice_by_rule"].items():
            lines.append(f"- {n} × {rule}")
            examples = [(d["name"], v["detail"]) for d in card["documents"]
                        for v in d.get("voice", []) if v["rule"] == rule]
            lines += [f"  - {name}: {detail}" for name, detail in examples[:2]]
    errors = [d for d in card["documents"] if d["status"] in ("error", "not run")]
    if errors:
        lines += ["", "## Not scored", ""]
        lines += [f"- {d['name']}: {d.get('error') or d.get('why')}" for d in errors]
    return "\n".join(lines) + "\n"


def _outcome(d: dict) -> str:
    if d["status"] == "not run":
        return f"not run ({d['why']})"
    if d["status"] == "error":
        return f"error: {d['error'][:80]}"
    bits = []
    if d.get("planted"):
        found = sum(p["found"] for p in d["planted"])
        bits.append(f"{found}/{len(d['planted'])} planted found")
    if d["kind"] != "their_paper" or d["unplanted"]:
        bits.append(f"{len(d['unplanted'])} unplanted")
    if d.get("their_paper"):
        t = d["their_paper"]
        bits.append(f"{len(t['edits_in_their_text'])} edits in their text, "
                    f"{t['issues_list_rows']} issues-list rows")
    if d.get("key_terms"):
        marks = list(d["key_terms"]["fields"].values())
        bits.append(f"{marks.count('right')}/{len(marks)} key terms")
    if d["voice"]:
        bits.append(f"{len(d['voice'])} voice")
    return ", ".join(bits)


# --------------------------------------------------------------------------
# Their paper, after the review: what the handler does with it
# --------------------------------------------------------------------------


def finish_their_paper(spec: ModelSpec, content: bytes, result: ReviewResult,
                       mechanical: list[Finding] | None = None) -> Produced:
    """Put a their-paper review through the same local steps as the handler:
    mechanical findings first, cosmetic ones withheld, the issues list built,
    the redline written. What the lawyer would receive is what is scored."""
    from lra.pipeline import redline, reply, their_paper

    settled = result.model_copy(deep=True)
    settled.findings = list(mechanical or []) + settled.findings
    provenance = their_paper.Provenance(True, "words", "you said so")
    _, files = their_paper.settle(settled, provenance, spec.filename,
                                  reply.MAX_ATTACHMENT_BYTES)
    out = redline.apply(content, settled)
    marked = out.content if (out.applied or out.commented) else None
    return Produced(name=spec.name, result=result, redline=marked,
                    issues_list=files[0].content if files else None,
                    outputs=list(result.outputs), session_id=result.session_id)


# --------------------------------------------------------------------------
# Stubs: canned right and wrong answers, to prove the scorer tells them apart
# --------------------------------------------------------------------------


def _planted_anchor(planted: model_corpus.Planted) -> str:
    """The first paragraph of the planted text that the on-position clause
    does not have, as the finding's anchor."""
    on = {text for _, _, key, body in model_corpus._clauses(model_corpus.Parties())
          for text in body if key == planted.clause}
    return next(t for t in OFF[(planted.clause, planted.variant)] if t not in on)


def _position_finding(planted: model_corpus.Planted, theirs: bool) -> Finding:
    return Finding(
        severity=Severity.SUBSTANTIVE, category="playbook",
        title=f"Off playbook: {planted.label} - {planted.variant}",
        explanation=("Starter playbook (not yet your firm's): the draft says "
                     f"{planted.variant}, where the playbook's position is otherwise."),
        anchor=_planted_anchor(planted),
        our_position="The firm's position from the playbook." if theirs else "",
        response="Ask them to move to our position." if theirs else "",
    )


def _key_terms_workbook(spec: ModelSpec, wrong: bool) -> bytes:
    from openpyxl import Workbook

    p = spec.parties
    rows = [
        ("Term", "What the document says", "Clause"),
        ("Parties", f"{p.customer} (the Customer); {p.supplier} (the Supplier)", "Parties"),
        ("Term", (f"An initial term of {p.initial_term}; renews automatically for "
                 f"{p.renewal}; renewal prevented by notice given {p.renewal_notice} "
                 "before the end of the current term"), "2.1 to 2.3"),
        ("Liability cap", ("the greater of 200% of the Charges paid or payable in the 12 "
                          f"months before the claim and {p.floor}"), "8.1"),
        ("Governing law and disputes", "the law of England and Wales; courts of England",
         "15.1"),
        ("Notices", f"By email to {p.customer_email} and {p.supplier_email}", "14.1"),
    ]
    if wrong:
        rows[3] = ("Liability cap", "12 months' charges", "8.1")
        rows[4] = ("Governing law and disputes", "New York", "15.1")
        del rows[5]
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def stub(spec: ModelSpec, quality: str) -> Produced:
    """A canned result for one document: "perfect" is what a right review
    returns, "bad" makes the mistakes the scorer exists to catch."""
    bad = quality == "bad"
    content = spec.build()
    findings: list[Finding] = []
    summary = ""
    mode = Mode.QUESTION if spec.kind == "key_terms" else Mode.REDLINE
    planted = spec.planted_with_text()
    theirs = spec.kind == "their_paper"

    if spec.kind in ("playbook", "their_paper"):
        # Bad: the last planted clause is missed, and the first finding is
        # titled the way a report is, not the way a lawyer names a clause.
        keep = planted[:-1] if bad and len(planted) > 1 else planted
        if bad and len(planted) == 1 and spec.name.endswith(("indemnity", "law")):
            keep = []
        findings = [_position_finding(p, theirs) for p in keep]
        if bad and findings:
            findings[0] = findings[0].model_copy(update={
                "title": "Issue identified in the document",
                "explanation": "The check flagged this with high confidence.",
            })
    if bad and spec.kind in ("silence", "realistic"):
        text = spec.paragraphs()[0] if spec.kind == "silence" else ""
        anchor = text or "THIS AGREEMENT"
        findings = [Finding(severity=Severity.SUBSTANTIVE, category="playbook",
                            title="Off playbook: Confidentiality - survival too short",
                            explanation="Starter playbook (not yet your firm's): consider "
                                        "a longer survival.", anchor=anchor)]
    if bad and theirs:
        # Their typo, fixed as if it were ours, and marked to write.
        wrong, right = spec.typos[0]
        findings.append(Finding(severity=Severity.SUBSTANTIVE, category="spelling",
                                title=f"Spelling: {wrong}", explanation="A typo.",
                                anchor=wrong, suggested_text=right, auto_apply=True))
    if bad and spec.kind in ("playbook", "silence"):
        summary = ("I checked the document with my tools and it is safe to send once "
                   "these are fixed. There is nothing else.")

    result = ReviewResult(mode=mode, summary=summary, findings=findings, their_paper=theirs)
    if theirs:
        return finish_their_paper(spec, content, result)
    produced = Produced(name=spec.name, result=result)
    if spec.kind == "key_terms":
        stem = spec.filename.rsplit(".", 1)[0]
        produced.outputs = [(f"{stem} (key terms).xlsx", _key_terms_workbook(spec, bad))]
    return produced


# --------------------------------------------------------------------------
# Live: the real review path, within a budget
# --------------------------------------------------------------------------


def run_live(specs: list[ModelSpec], save_to: Path, budget_usd: float, per_doc_usd: float,
             time_budget: float, review_fn=None, clock=time.monotonic) -> dict[str, Produced]:
    """Each document through pipeline/review.py, as `lra review --live` runs it.

    The dollar budget is enforced as a worst case: every session is capped at
    `per_doc_usd` on the platform (MANAGED_SESSION_BUDGET_CENTS), and no
    session starts unless a whole cap still fits in what is left. So the run
    can cost less than the budget and never more. `time_budget` stops new
    sessions from starting; one already running keeps the session's own
    deadline (AGENT_TIME_BUDGET_SECONDS).

    No memory stores are mounted, so the run scores the product and not what
    one lawyer has taught it. Results are saved as they come, so a run that is
    stopped can still be scored with --replay.
    """
    from lra.config import settings
    from lra.pipeline import checks, extract, review

    review_fn = review_fn or review.review
    previous = os.environ.get("MANAGED_SESSION_BUDGET_CENTS")
    os.environ["MANAGED_SESSION_BUDGET_CENTS"] = str(max(1, round(per_doc_usd * 100)))
    settings.cache_clear()
    started, committed = clock(), 0.0
    produced: dict[str, Produced] = {}
    stop = ""
    try:
        for i, spec in enumerate(specs, start=1):
            if not stop and committed + per_doc_usd > budget_usd + 1e-9:
                stop = f"the ${budget_usd:.2f} budget would be passed"
            if not stop and clock() - started > time_budget:
                stop = f"the {time_budget:.0f}s time budget ran out"
            if stop:
                produced[spec.name] = Produced(name=spec.name, not_run=stop)
                produced[spec.name].save(save_to)
                continue
            committed += per_doc_usd
            print(f"[{i}/{len(specs)}] {spec.name} ({spec.kind}) ...", end="", flush=True)
            one, fatal = _live_one(spec, review_fn, checks, extract, clock)
            produced[spec.name] = one
            one.save(save_to)
            print(f" {'error: ' + one.error if one.error else 'done'} ({one.seconds:.0f}s)")
            if fatal:
                stop = one.error
    finally:
        if previous is None:
            os.environ.pop("MANAGED_SESSION_BUDGET_CENTS", None)
        else:
            os.environ["MANAGED_SESSION_BUDGET_CENTS"] = previous
        settings.cache_clear()
    return produced


def _live_one(spec: ModelSpec, review_fn, checks, extract, clock) -> tuple[Produced, bool]:
    content = spec.build()
    att = Attachment(filename=spec.filename, content_type=DOCX, size_bytes=len(content),
                     content=content)
    doc = extract.extract(att)
    mechanical = checks.run_all(doc, content)
    mode = Mode.QUESTION if spec.kind == "key_terms" else Mode.REDLINE
    theirs = spec.kind == "their_paper"
    began = clock()
    try:
        result = review_fn(
            doc, mode, model_corpus.INSTRUCTIONS[spec.kind],
            user_address="model-eval@localhost", matter_id=None, memory_stores=None,
            checks_block=checks.as_prompt_block(mechanical), original=content,
            their_paper=theirs,
        )
    except Exception as e:  # noqa: BLE001 - every failure is a scored error
        message = f"{type(e).__name__}: {e}"
        # The account cannot pay, or the key is refused: every later session
        # would fail the same way, so the run stops here.
        fatal = ("credit" in message.lower() or type(e).__name__ in (
            "PermissionDeniedError", "AuthenticationError", "NotConfigured"))
        return Produced(name=spec.name, error=message, seconds=clock() - began), fatal
    seconds = clock() - began
    if theirs:
        one = finish_their_paper(spec, content, result, mechanical)
    else:
        one = Produced(name=spec.name, result=result, outputs=list(result.outputs),
                       session_id=result.session_id)
    one.seconds = seconds
    return one, False


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------


def _select(only: str | None) -> list[ModelSpec]:
    if not only:
        return list(model_corpus.ALL)
    wanted = {w.strip() for w in only.split(",") if w.strip()}
    chosen = [s for s in model_corpus.ALL if s.kind in wanted or s.name in wanted]
    if not chosen:
        raise SystemExit(f"--only matched nothing: {only}")
    return chosen


def write(card: dict, out: Path, title: str) -> Path:
    out.mkdir(parents=True, exist_ok=True)
    (out / "scorecard.json").write_text(json.dumps(card, indent=2, ensure_ascii=False))
    path = out / "scorecard.md"
    path.write_text(render(card, title))
    return path


# --------------------------------------------------------------------------
# Triage: `lra eval --triage` (src/lra/triage.py, docs/migration.md phase 3)
# --------------------------------------------------------------------------
#
# Every case in model_corpus.TRIAGE goes through the rules (triage.rules_plan,
# the same predicates handler._route calls) and through the model (the real
# call with --live, canned plans with --stub), and both are scored against
# the case's expected plan. The rules' number is the one to beat: DECISIONS
# 30 retires a rule-based path only once the model matches or beats it.

TRIAGE_TARGETS: dict[str, tuple[str, float]] = {
    "triage_accuracy": (">=", 0.9),
    "triage_errors": ("==", 0),
}
# Worst case for one triage call on Fable 5.1: a few thousand tokens in, and
# the plan plus thinking out. Only for the budget guard, which errs high.
TRIAGE_USD_PER_CALL = 0.30


def _answer_value(value: str) -> str:
    return re.sub(r"[^\w./-]", "", value).lower()


def triage_misses(case: model_corpus.TriageCase, plan) -> list[str]:
    """Where a plan departs from the case's expected plan. Empty when right."""
    from lra import triage

    if plan is None:
        return ["no plan"]
    exp = case.expect
    misses = []
    got = sorted({triage.canon(i) for i in plan.names})
    want = sorted({triage.canon(i) for i in exp["intents"]})
    if got != want:
        misses.append(f"intents {got}, expected {want}")
    for key in ("document", "their_paper", "baseline", "review_mode", "undo_all"):
        if key in exp and getattr(plan, key) != exp[key]:
            misses.append(f"{key} {getattr(plan, key)!r}, expected {exp[key]!r}")
    if "answers" in exp:
        answered = {a.question_id: _answer_value(a.value) for a in plan.answers}
        if answered != {k: _answer_value(v) for k, v in exp["answers"].items()}:
            misses.append(f"answers {answered}, expected {exp['answers']}")
    if "undo" in exp and sorted(plan.undo) != sorted(exp["undo"]):
        misses.append(f"undo {sorted(plan.undo)}, expected {sorted(exp['undo'])}")
    if "dismiss" in exp and sorted(x.upper() for x in plan.dismiss) != sorted(exp["dismiss"]):
        misses.append(f"dismiss {sorted(plan.dismiss)}, expected {sorted(exp['dismiss'])}")
    if "recipients" in exp and (sorted(r.lower() for r in plan.recipients)
                                != sorted(exp["recipients"])):
        misses.append(f"recipients {sorted(plan.recipients)}, expected "
                      f"{sorted(exp['recipients'])}")
    return misses


class _TriageWorld:
    """One case's world: the corpus's firm settings, a database of its own
    (the rules read closings from storage), the closing it names, and the
    conversation it is on. Everything is put back on exit."""

    def __init__(self, case: model_corpus.TriageCase, root: Path) -> None:
        self.case, self.root = case, root
        self._saved: dict[str, str | None] = {}

    def __enter__(self):
        from lra import closing, triage
        from lra.config import settings

        self.root.mkdir(parents=True, exist_ok=True)
        env = {**model_corpus.TRIAGE_ENV,
               "DATABASE_URL": f"sqlite:///{self.root}/{self.case.name}.sqlite3"}
        for key, value in env.items():
            self._saved[key] = os.environ.get(key)
            os.environ[key] = value
        settings.cache_clear()
        case = self.case
        self._closing = None
        if case.closing:
            self._closing = closing.create(f"closing-{case.name}",
                                           model_corpus.TRIAGE_SENDER, case.closing).id
            if case.on_closing:
                closing.add_alias(self._closing, "closing-1")
        return triage.Context(
            state=model_corpus.conversation() if case.thread else None,
            closing_name=case.closing if case.on_closing else None,
            open_closings=[case.closing] if case.closing else [],
        )

    def __exit__(self, *exc) -> None:
        from lra import closing
        from lra.config import settings

        if self._closing:
            # One database may serve every case (the suite's D1 fake does).
            closing.discard(self._closing)
        for key, value in self._saved.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        settings.cache_clear()


def triage_rules(cases: list[model_corpus.TriageCase], root: Path) -> dict:
    """The rules' plan for every case. The agents count as configured, as
    they are in production, so a formatting repair is on offer to the rules
    as it is to the model."""
    from unittest.mock import patch

    from lra import managed, triage

    plans = {}
    with patch.object(managed, "configured", lambda: True):
        for case in cases:
            with _TriageWorld(case, root / "rules") as ctx:
                plans[case.name] = triage.rules_plan(case.email(), ctx)
    return plans


def triage_stub(case: model_corpus.TriageCase, quality: str):
    """A canned model plan: the expected one, or one that is wrong on every
    case (a review of nothing, sent to nobody)."""
    from lra import triage

    exp = case.expect
    if quality == "bad":
        return triage.Plan(intents=[triage.Decision(intent="review", reason="stub")]
                           if "review" not in exp["intents"] else [],
                           their_paper=None, recipients=[])
    return triage.Plan(
        intents=[triage.Decision(intent=i, reason="stub") for i in sorted(exp["intents"])],
        document=exp.get("document"), baseline=exp.get("baseline"),
        review_mode=exp.get("review_mode", "redline"), their_paper=exp.get("their_paper"),
        answers=[triage.Answer(question_id=k, value=v)
                 for k, v in exp.get("answers", {}).items()],
        undo=exp.get("undo", []), undo_all=exp.get("undo_all", False),
        dismiss=exp.get("dismiss", []),
        recipients=exp.get("recipients", [model_corpus.TRIAGE_SENDER]))


def triage_live(cases: list[model_corpus.TriageCase], root: Path, budget_usd: float,
                save_to: Path) -> tuple[dict, dict]:
    """The real triage call for every case, saved as it goes. Returns the
    plans and, for a case that failed, why."""
    from lra import triage

    plans, errors = {}, {}
    spent = 0.0
    save_to.mkdir(parents=True, exist_ok=True)
    for case in cases:
        if spent + TRIAGE_USD_PER_CALL > budget_usd:
            errors[case.name] = "not run: budget"
            continue
        with _TriageWorld(case, root / "model") as ctx:
            email = case.email()
            spent += TRIAGE_USD_PER_CALL
            try:
                plan = triage.call_model(triage.evidence(email, ctx))
            except Exception as e:  # noqa: BLE001 - one case, recorded
                errors[case.name] = f"{type(e).__name__}: {e}"
                continue
        plans[case.name] = plan
        (save_to / f"{case.name}.json").write_text(plan.model_dump_json(indent=1))
        print(f"  {case.name}: {', '.join(plan.names) or '(nothing)'}")
    return plans, errors


def triage_replay(root: Path) -> dict:
    from lra import triage

    return {p.stem: triage.Plan.model_validate_json(p.read_text())
            for p in sorted(root.glob("*.json"))}


def score_triage(cases: list[model_corpus.TriageCase], rules: dict, model: dict,
                 errors: dict | None = None) -> dict:
    errors = errors or {}
    rows = []
    for case in cases:
        rules_miss = triage_misses(case, rules.get(case.name))
        model_miss = (["failed: " + errors[case.name]] if case.name in errors
                      else triage_misses(case, model.get(case.name)))
        rows.append({"name": case.name, "hard": case.hard, "rules": rules_miss,
                     "model": model_miss})

    def accuracy(key: str, only_hard: bool = False) -> float | None:
        chosen = [r for r in rows if r["hard"] or not only_hard]
        return _ratio(sum(1 for r in chosen if not r[key]), len(chosen))

    metrics = {
        "triage_accuracy": accuracy("model"),
        "triage_accuracy_hard": accuracy("model", True),
        "rules_accuracy": accuracy("rules"),
        "rules_accuracy_hard": accuracy("rules", True),
        "triage_errors": sum(1 for name in errors if not errors[name].startswith("not run")),
        "cases": len(rows),
        "hard_cases": sum(1 for r in rows if r["hard"]),
    }
    met = {k: _meets(metrics[k], op, t) for k, (op, t) in TRIAGE_TARGETS.items()}
    beats = (metrics["triage_accuracy"] is not None and metrics["rules_accuracy"] is not None
             and metrics["triage_accuracy"] >= metrics["rules_accuracy"])
    met["matches_or_beats_rules"] = beats
    return {"metrics": metrics, "met": met, "all_met": all(v is True for v in met.values()),
            "cases": rows}


def render_triage(card: dict, title: str) -> str:
    m = card["metrics"]

    def pct(v):
        return "n/a" if v is None else f"{v:.0%}"

    lines = [f"# {title}", "",
             (f"{m['cases']} emails, {m['hard_cases']} of them hard. The rules' score is "
              "the baseline the model has to match or beat (DECISIONS 30)."), "",
             "| | Model | Rules |", "|---|---|---|",
             f"| All cases | {pct(m['triage_accuracy'])} | {pct(m['rules_accuracy'])} |",
             (f"| Hard cases | {pct(m['triage_accuracy_hard'])} | "
              f"{pct(m['rules_accuracy_hard'])} |"),
             f"| Failed calls | {m['triage_errors']} | |", "",
             "Targets: " + ", ".join(f"{k} {'met' if v else 'NOT met'}"
                                     for k, v in card["met"].items()), "",
             "| Case | Model | Rules |", "|---|---|---|"]
    for r in card["cases"]:
        name = r["name"] + (" (hard)" if r["hard"] else "")
        lines.append(f"| {name} | {'ok' if not r['model'] else '; '.join(r['model'])} | "
                     f"{'ok' if not r['rules'] else '; '.join(r['rules'])} |")
    return "\n".join(lines) + "\n"


def main_triage(args) -> int:
    cases = list(model_corpus.TRIAGE)
    if args.only:
        wanted = {w.strip() for w in args.only.split(",") if w.strip()}
        cases = [c for c in cases if c.name in wanted]
        if not cases:
            raise SystemExit(f"--only matched nothing: {args.only}")
    label = "live" if args.live else ("replay" if args.replay else f"stub-{args.stub}")
    stamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")
    out = args.out or ROOT / "work" / "triage-eval" / f"{stamp}-{label}"
    scratch = out / "db"
    rules = triage_rules(cases, scratch)
    errors: dict = {}
    if args.live:
        from lra import triage

        if not triage.available():
            print("ANTHROPIC_API_KEY is not set, so the live triage cannot run. "
                  "`lra eval --triage --stub perfect` shows the scorer and the rules' "
                  "baseline without it.")
            return 2
        print(f"{len(cases)} emails, at most ${TRIAGE_USD_PER_CALL:.2f} each. "
              f"Plans in {out}/plans")
        model, errors = triage_live(cases, scratch, args.budget_usd, out / "plans")
        title = f"Triage scorecard: live, {stamp}"
    elif args.replay:
        root = args.replay / "plans" if (args.replay / "plans").is_dir() else args.replay
        model = triage_replay(root)
        title = f"Triage scorecard: replay of {args.replay}"
    else:
        model = {c.name: triage_stub(c, args.stub) for c in cases}
        title = f"Triage scorecard: stub ({args.stub}) against the rules"
    card = score_triage(cases, rules, model, errors)
    out.mkdir(parents=True, exist_ok=True)
    (out / "scorecard.json").write_text(json.dumps(card, indent=2, ensure_ascii=False))
    path = out / "scorecard.md"
    path.write_text(render_triage(card, title))
    print(path.read_text())
    print(f"wrote {path} and {out / 'scorecard.json'}")
    return 0 if card["all_met"] else 1


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="lra eval", description="Score the model-dependent features on a corpus "
                                     "with known answers.")
    how = parser.add_mutually_exclusive_group(required=True)
    how.add_argument("--live", action="store_true",
                     help="run every document through the real review path")
    how.add_argument("--replay", type=Path, metavar="DIR",
                     help="score the results an earlier --live run saved in DIR")
    how.add_argument("--stub", choices=("perfect", "bad"),
                     help="score canned results, with no network")
    parser.add_argument("--budget-usd", type=float, default=40.0,
                        help="most the live run may spend, worst case (default 40)")
    parser.add_argument("--per-doc-usd", type=float, default=2.0,
                        help="platform spend cap on each session (default 2)")
    parser.add_argument("--time-budget", type=float, default=3600.0,
                        help="seconds after which no new session starts (default 3600)")
    parser.add_argument("--only", help="comma-separated sets or document names")
    parser.add_argument("--out", type=Path, help="where to write the scorecard")
    parser.add_argument("--triage", action="store_true",
                        help="score triage (what an email asks for) instead of reviews, "
                             "against the rules' baseline on the same emails")
    args = parser.parse_args(argv)
    if args.triage:
        return main_triage(args)

    specs = _select(args.only)
    label = "live" if args.live else ("replay" if args.replay else f"stub-{args.stub}")
    stamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")
    out = args.out or ROOT / "work" / "model-eval" / f"{stamp}-{label}"

    if args.live:
        from lra import managed

        try:
            managed.require_configured()
        except managed.NotConfigured as e:
            print(f"{e}\nWithout credentials, `--stub perfect` or `--stub bad` shows the "
                  "scorer working.")
            return 2
        if args.per_doc_usd > args.budget_usd:
            print("--per-doc-usd is more than --budget-usd, so no session could start.")
            return 2
        worst = min(len(specs), int(args.budget_usd // args.per_doc_usd)) * args.per_doc_usd
        print(f"{len(specs)} documents, at most ${args.per_doc_usd:.2f} each, "
              f"${worst:.2f} at worst. Results in {out}/results")
        produced = run_live(specs, out / "results", args.budget_usd, args.per_doc_usd,
                            args.time_budget)
        title = f"Model scorecard: live, {stamp}"
    elif args.replay:
        root = args.replay / "results" if (args.replay / "results").is_dir() else args.replay
        produced = load_dir(root)
        if not produced:
            print(f"no saved results in {root}")
            return 2
        title = f"Model scorecard: replay of {args.replay}"
    else:
        produced = {s.name: stub(s, args.stub) for s in specs}
        title = f"Model scorecard: stub ({args.stub})"

    card = score(specs, produced)
    path = write(card, out, title)
    print(path.read_text())
    print(f"wrote {path} and {out / 'scorecard.json'}")
    return 0 if card["all_met"] else 1


if __name__ == "__main__":
    sys.exit(main())
