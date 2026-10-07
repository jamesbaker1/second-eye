"""Agreements drafted the way firms draft them, for the check corpus.

The documents in corpus.py are typed paragraphs, one per clause, with the
clause number written into the text. Real agreements are not. Their clause
numbers are drawn by Word's list numbering and are not in the text at all;
their definitions sit in tables and run on into (a) and (b); their schedules
restart numbering in each part; they say "[Reserved]" and "[***]" on purpose;
and they cite Treasury Regulations and statutes by section. A reviewer who
built three such documents found 25 findings, six of them blockers, on
agreements with nothing wrong with them, and none of that was visible to the
synthetic corpus. These builders are that reviewer's documents, kept here so
the score reflects realistic drafting.

Two changes from the originals, both to make a "clean" document clean: the US
agreement now uses the two terms it defined and never used ("Closing Payment"
and "Escrow Amount"), which the checks rightly reported.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

_LIST_ID = 90


def _add_multilevel_list(d: Document) -> int:
    """A three-level legal list: 1. / 1.1 / (a), as a firm template defines it."""
    numbering = d.part.numbering_part.element
    abstract = (
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{_LIST_ID}">'
        '<w:multiLevelType w:val="multilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        '<w:lvlText w:val="%1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        '<w:lvlText w:val="%1.%2"/></w:lvl>'
        '<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/>'
        '<w:lvlText w:val="(%3)"/></w:lvl>'
        "</w:abstractNum>"
    )
    first_num = numbering.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num")
    first_num.addprevious(parse_xml(abstract))
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{_LIST_ID}">'
        f'<w:abstractNumId w:val="{_LIST_ID}"/></w:num>'
    ))
    return _LIST_ID


def _numbered(d: Document, text: str, level: int) -> None:
    p = d.add_paragraph(text)
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/>'
        f'<w:numId w:val="{_LIST_ID}"/></w:numPr>'
    ))


def _save(d: Document) -> bytes:
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    buffer = BytesIO()
    d.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# An English share purchase agreement
# --------------------------------------------------------------------------


def uk_spa(auto_numbered: bool) -> bytes:
    d = Document()
    if auto_numbered:
        _add_multilevel_list(d)

    def cl(num: str, text: str, level: int = 0) -> None:
        if auto_numbered:
            _numbered(d, text, level)
        else:
            d.add_paragraph(f"{num}\t{text}")

    P = d.add_paragraph
    P("THIS AGREEMENT is dated 3 March 2026")
    P("PARTIES")
    P("(1)\tNORTHGATE HOLDINGS LIMITED, a company incorporated in England and Wales with "
      "company number 01234567 whose registered office is at 1 King Street, London EC2V "
      "8AU (the “Seller”); and")
    P("(2)\tBLUEWATER CAPITAL LLP, a limited liability partnership incorporated in England "
      "and Wales with registered number OC123456 whose registered office is at 20 "
      "Fenchurch Street, London EC3M 3BY (the “Buyer”).")
    P("BACKGROUND")
    P("(A)\tThe Seller is the legal and beneficial owner of the Shares.")
    P("(B)\tThe Seller has agreed to sell and the Buyer has agreed to buy the Shares on "
      "the terms of this agreement.")
    P("AGREED TERMS")
    cl("1.", "Definitions and interpretation")
    cl("1.1", "In this agreement:", 1)
    P("“Business Day” means a day other than a Saturday, Sunday or public holiday in "
      "England when banks in London are open for business.")
    P("“Company” means Northgate Freight Limited, a company incorporated in England and "
      "Wales with company number 07654321.")
    P("“Completion” means completion of the sale and purchase of the Shares in "
      "accordance with this agreement.")
    P("“Consideration Shares” means the ordinary shares in the Buyer to be issued under "
      "clause 3.2, each a “Consideration Share”.")
    P("“Disclosure Letter” means the letter from the Seller to the Buyer with any "
      "enclosures, dated the same date as this agreement.")
    P("“Material Adverse Change” means any event or circumstance which:")
    P("(a)\thas, or is reasonably likely to have, a material adverse effect on the "
      "business of the Company; or")
    P("(b)\tresults in a material adverse change in the financial position of the "
      "Company.")
    P("“Long Stop Date” means 30 June 2026 or such later date as the parties may agree "
      "in writing.")
    P("“Shares” means the 1,000 ordinary shares of £1.00 each in the capital of the "
      "Company.")
    P("“Warranties” means the warranties set out in Schedule 3.")
    cl("1.2", "References to clauses and Schedules are to the clauses of and the "
              "Schedules to this agreement, and references to paragraphs are to "
              "paragraphs of the relevant Schedule.", 1)
    cl("1.3", "Words in the singular include the plural and vice versa.", 1)
    cl("2.", "Sale and purchase")
    cl("2.1", "The Seller shall sell and the Buyer shall buy the Shares with full title "
              "guarantee, free from all Encumbrances.", 1)
    cl("3.", "Consideration")
    cl("3.1", "The consideration for the Shares is £2,500,000 (two million five hundred "
              "thousand pounds) (the “Consideration”), which shall be satisfied in "
              "accordance with clauses 3.2 and 3.3.", 1)
    cl("3.2", "£500,000 of the Consideration shall be satisfied by the allotment and "
              "issue of 50,000 Consideration Shares, credited as fully paid, at "
              "Completion.", 1)
    cl("3.3", "The balance shall be paid in cash within five (5) Business Days after "
              "Completion, together with a retention of ten per cent. (10%) held in "
              "accordance with clause 6.", 1)
    cl("4.", "Conditions")
    cl("4.1", "Completion is conditional on no Material Adverse Change having occurred "
              "before Completion. If the condition is not satisfied by the Long Stop "
              "Date, either party may terminate this agreement by notice to the other.", 1)
    cl("5.", "[Reserved]")
    cl("6.", "Retention")
    cl("6.1", "The Buyer shall hold the retention in a designated account for twelve "
              "(12) months. [The Buyer shall provide statements quarterly.]", 1)
    cl("7.", "Warranties")
    cl("7.1", "The Seller warrants to the Buyer that each Warranty is true and accurate "
              "at the date of this agreement, subject to matters Disclosed in the "
              "Disclosure Letter.", 1)
    cl("7.2", "The Seller's liability is limited as set out in Schedule 4, and no claim "
              "may be brought under paragraph 3 of Schedule 4 after the second "
              "anniversary of Completion.", 1)
    cl("8.", "Notices")
    cl("8.1", "A notice given under this agreement shall be in writing and delivered by "
              "hand or by pre-paid first-class post to the address in the Parties "
              "clause, and is deemed received at 9.00 am on the second Business Day "
              "after posting.", 1)
    cl("9.", "Governing law")
    cl("9.1", "This agreement and any dispute or claim arising out of it shall be "
              "governed by the law of England and Wales, and section 1 of the Contracts "
              "(Rights of Third Parties) Act 1999 shall not apply.", 1)
    P("This agreement has been entered into on the date stated at the beginning of it.")

    schedules = [
        ("1", "Particulars of the Company", None),
        ("2", "Completion deliverables", [
            "The Seller shall deliver the share certificates for the Shares.",
            "The Seller shall deliver resignations of the directors.",
            "The Buyer shall deliver evidence of the issue of the Consideration Shares."]),
        ("3", "Warranties", [
            "The Seller is the sole legal and beneficial owner of the Shares.",
            "The Company has complied with all applicable laws.",
            "The accounts give a true and fair view."]),
        ("4", "Limitations on the Seller's liability", [
            "The Seller shall not be liable for a Claim unless it exceeds £25,000.",
            "The aggregate liability of the Seller shall not exceed the Consideration.",
            "No Claim may be brought after the second anniversary of Completion."]),
    ]
    for number, title, paras in schedules:
        d.add_page_break()
        P(f"SCHEDULE {number}")
        P(title)
        if paras is None:
            t = d.add_table(rows=5, cols=2)
            for r, (a, b) in enumerate([
                ("Name", "Northgate Freight Limited"), ("Registered number", "07654321"),
                ("Date of incorporation", "12 May 1998"),
                ("Directors", "John Smith, Jane Doe"), ("Secretary", "N/A")]):
                t.rows[r].cells[0].text = a
                t.rows[r].cells[1].text = b
        elif number == "4":
            P("Part 1 – Financial limits")
            for i, p in enumerate(paras[:2], 1):
                P(f"{i}.\t{p}")
            P("Part 2 – Time limits")
            for i, p in enumerate(paras[2:], 1):
                P(f"{i}.\t{p}")
        else:
            for i, p in enumerate(paras, 1):
                P(f"{i}.\t{p}")
    P("EXECUTED as a deed by NORTHGATE HOLDINGS LIMITED acting by a director in the "
      "presence of a witness:")
    P("Director: ______________________")
    P("Witness: ______________________")
    P("SIGNED by John Brown for and on behalf of BLUEWATER CAPITAL LLP:")
    P("Member: ______________________")
    section = d.sections[0]
    section.header.paragraphs[0].text = "Execution version"
    section.footer.paragraphs[0].text = "UK-ACTIVE-123456789.1"
    return _save(d)


# --------------------------------------------------------------------------
# An American stock purchase agreement
# --------------------------------------------------------------------------


def us_spa() -> bytes:
    d = Document()
    _add_multilevel_list(d)

    def cl(text: str, level: int = 0) -> None:
        _numbered(d, text, level)

    P = d.add_paragraph
    d.add_heading("STOCK PURCHASE AGREEMENT", 0)
    P("This STOCK PURCHASE AGREEMENT (this “Agreement”) is entered into as of March 3, "
      "2026, by and among Acme Parent Holdings, Inc., a Delaware corporation "
      "(“Parent”), Acme Merger Sub LLC, a Delaware limited liability company "
      "(“Buyer”), and Bluewater Capital LLC, a Delaware limited liability company "
      "(“Seller”).")
    P("RECITALS")
    P("WHEREAS, Seller owns all of the issued and outstanding shares of common stock "
      "(the “Shares”) of Target Robotics Inc., a Delaware corporation (the "
      "“Company”);")
    P("NOW, THEREFORE, in consideration of the mutual covenants herein, the parties "
      "agree as follows:")
    cl("DEFINITIONS")
    cl("Definitions. Capitalized terms used herein have the meanings set forth in the "
       "table below.", 1)
    t = d.add_table(rows=4, cols=2)
    table = [
        ("“Business Day”", ("any day other than a Saturday, Sunday or a day on which "
                            "banks in New York City are authorized to close.")),
        ("“Company Intellectual Property”", ("all intellectual property owned by the "
                                             "Company, including all company "
                                             "intellectual property licensed to third "
                                             "parties.")),
        ("“Material Adverse Effect”", ("any event that has a material adverse effect "
                                       "on the Company.")),
        ("“Closing Payment”", "the amount set forth in Section 2.2."),
    ]
    for r, (a, b) in enumerate(table):
        t.rows[r].cells[0].text = a
        t.rows[r].cells[1].text = b
    cl("PURCHASE AND SALE")
    cl("Purchase and Sale. At the Closing, Seller shall sell to Buyer, and Buyer shall "
       "purchase from Seller, the Shares.", 1)
    cl("Purchase Price. The aggregate purchase price shall be $25,000,000 (twenty-five "
       "million dollars) (the “Purchase Price”), of which $2,500,000 (the “Escrow "
       "Amount”) shall be deposited with the Escrow Agent.", 1)
    cl("Closing. The closing (the “Closing”) shall take place on the third (3rd) "
       "Business Day after the conditions in Article VI are satisfied. At the Closing, "
       "Buyer shall pay the Closing Payment to Seller and deposit the Escrow Amount.", 1)
    cl("REPRESENTATIONS AND WARRANTIES OF SELLER")
    cl("Organization. As of the date hereof and since December 31, 2024, the Company has "
       "been duly organized and in good standing.", 1)
    cl("Intellectual Property. Schedule 3.2 lists each Patent owned by the Company. No "
       "Parent entity of the Company holds any Patent.", 1)
    cl("[Reserved].", 1)
    cl("[RESERVED]", 1)
    cl("Taxes. The Company has filed all returns required under Title XX of the Social "
       "Security Act and Section 1.1502-6 of the Treasury Regulations.", 1)
    cl("COVENANTS")
    cl("Conduct of Business. From the date hereof until the Closing, except as set forth "
       "in Section 4.1 of the Disclosure Schedule, Seller shall cause the Company to "
       "operate in the ordinary course.", 1)
    cl("Confidentiality. Each party shall keep confidential all information marked "
       "“Confidential”, as provided in Sections 7.2 and 7.3.", 1)
    cl("INDEMNIFICATION")
    cl("Seller shall indemnify Buyer for losses arising from breach of Article III, "
       "subject to Section 5.2, up to [***].", 1)
    cl("The cap is ten percent (10%) of the Purchase Price.", 1)
    cl("CONDITIONS")
    cl("No Material Adverse Effect shall have occurred.", 1)
    cl("MISCELLANEOUS")
    cl("Notices. All notices shall be delivered to the addresses on the signature "
       "pages.", 1)
    cl("Governing Law. This Agreement is governed by the laws of the State of "
       "Delaware.", 1)
    cl("Counterparts. This Agreement may be executed in counterparts.", 1)
    P("IN WITNESS WHEREOF, the parties have executed this Agreement as of the date "
      "first written above.")
    for name, who in [("ACME PARENT HOLDINGS, INC.", "Jane Doe"),
                      ("ACME MERGER SUB LLC", "Jane Doe"),
                      ("BLUEWATER CAPITAL LLC", "John Smith")]:
        P(name)
        P("By: ______________________")
        P(f"Name: {who}")
        P("Title: President")
    d.add_page_break()
    P("EXHIBIT A")
    P("Form of Escrow Agreement")
    P("[To be attached.]")
    return _save(d)


# --------------------------------------------------------------------------
# A run-in heading, made with a style separator
# --------------------------------------------------------------------------


def style_separator() -> bytes:
    """A hidden paragraph mark, which is how Word joins a run-in heading to its
    paragraph. It hides nothing a reader could want to see."""
    d = Document()
    p = d.add_paragraph("1. Definitions")
    p._p.get_or_add_pPr().append(
        parse_xml(f'<w:rPr {nsdecls("w")}><w:vanish/><w:specVanish/></w:rPr>'))
    d.add_paragraph("In this agreement the following words have the following meanings.")
    return _save(d)


# --------------------------------------------------------------------------
# Defects a UK drafter actually leaves in
# --------------------------------------------------------------------------


def uk_defects() -> bytes:
    d = Document()
    P = d.add_paragraph
    P("THIS AGREEMENT is dated Monday, 3 March 2026 between Northgate Holdings Limited "
      "(the “Seller”) and Bluewater Capital LLP (the “Buyer”).")
    P("1. Definitions. “Completion Date” means 10 April 2026. “Long Stop Date” means "
      "[●]. “Buyer’s Solicitors” means [x] LLP.")
    P("“Completion Date” means the fifth Business Day after the Conditions are "
      "satisfied.")
    P("2. Price. The price is £2,500,000 (two million three hundred thousand pounds) "
      "payable within 10 (fifteen) Business Days.")
    P("3. Payment. The Seller shall pay US$50,000 to the Agent and $ 20,000 to the Buyer "
      "and USD 10,000 to the Registrar, being EUR5,000 in aggregate.")
    P("4. Deposit. The Buyer shall pay a deposit of £[●] on or before [date] into the "
      "account of [name of bank].")
    P("5. Notices. [Note to draft: Buyer to confirm notice details.] Notices shall be "
      "served as set out in clause 12.")
    P("6. Completion shall occur on 03/04/2026 or such other date as the parties agree, "
      "and not later than 30 April 2026, being a Friday.")
    P("7. The deposit shall be released [within [x] days] in accordance with clause 6. "
      "[NTD: check with tax.]")
    P("8. The Warranties are given as at the date of this agreement and repeated at "
      "Completion. The Seller shall procure that Northgate Holdings Ltd does not "
      "dispose of any asset.")
    P("9. The quoted passage reads “[...] the court shall have regard to all the "
      "circumstances”.")
    # A template's form field holding a blank, in a block-level content control.
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr/><w:sdtContent><w:p><w:r>'
        "<w:t>Purchase price: [INSERT AMOUNT]</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    )
    d.element.body.insert(len(d.element.body) - 1, sdt)
    return _save(d)


def moved_and_reformatted() -> bytes:
    """Tracked changes that are a move and a formatting change, not an
    insertion or a deletion. They are in the recipient's Review pane all the
    same."""
    d = Document()
    p = d.add_paragraph()
    p._p.append(parse_xml(
        f'<w:moveTo {nsdecls("w")} w:id="1" w:author="Opposing Counsel" '
        'w:date="2026-03-01T00:00:00Z"><w:r><w:t>Moved sentence.</w:t></w:r></w:moveTo>'))
    q = d.add_paragraph()
    q._p.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:rPr><w:b/><w:rPrChange w:id="2" '
        'w:author="Opposing Counsel" w:date="2026-03-01T00:00:00Z"><w:rPr/></w:rPrChange>'
        "</w:rPr><w:t>Bolded by the other side.</w:t></w:r>"))
    return _save(d)


# --------------------------------------------------------------------------
# Project Falcon: the deal the demo runs on
# --------------------------------------------------------------------------


def falcon_spa(version: int = 3, defects: bool = False) -> bytes:
    """An English SPA whose arithmetic, timings and party names are the point.

    Version 2 names Falcon Holdings Limited as the Buyer throughout. Version 3
    renames the Buyer to Falcon Topco Limited in the parties clause. Clean,
    version 3 makes the change everywhere, and every figure adds up: the
    instalments to the Purchase Price, Schedule 2 to the Purchase Price, the
    shareholding to 100%. With `defects`, version 3 leaves the old name on
    the Buyer's signature block and notice details, Schedule 2 and the
    instalments no longer add up, the claims window closes before the
    retention is released, and a breach can be cured after the notice that
    terminates for it has run out.
    """
    d = Document()
    P = d.add_paragraph
    buyer = "FALCON HOLDINGS LIMITED" if version < 3 else "FALCON TOPCO LIMITED"
    stale = "FALCON HOLDINGS LIMITED" if defects else buyer
    title = lambda s: " ".join(w.capitalize() for w in s.split())  # noqa: E731

    P("THIS AGREEMENT is dated 1 September 2026")
    P("PARTIES")
    P("(1)\tBETA INDUSTRIES LIMITED, a company incorporated in England and Wales with "
      "company number 04567123 whose registered office is at 1 King Street, London EC2V "
      "8AU (the “Seller”); and")
    P(f"(2)\t{buyer}, a company incorporated in England and Wales with company number "
      "11223344 whose registered office is at 20 Fenchurch Street, London EC3M 3BY (the "
      "“Buyer”).")
    P("BACKGROUND")
    P("(A)\tThe Seller is the legal and beneficial owner of the Shares.")
    P("(B)\tThe Seller has agreed to sell and the Buyer has agreed to buy the Shares on "
      "the terms of this agreement.")
    P("AGREED TERMS")
    P("1.\tDefinitions")
    P("1.1\tIn this agreement:")
    P("“Business Day” means a day other than a Saturday, Sunday or public holiday in "
      "England when banks in London are open for business.")
    P("“Company” means Falcon Target Limited, a company incorporated in England and "
      "Wales with company number 07654321.")
    P("“Completion” means completion of the sale and purchase of the Shares in "
      "accordance with this agreement.")
    P("“Completion Date” means 30 October 2026.")
    P("“Long Stop Date” means 31 December 2026.")
    P("“Retention Amount” means £1,250,000.")
    P("“Shares” means the 1,000 ordinary shares of £1.00 each in the capital of the "
      "Company.")
    P("2.\tSale and purchase")
    P("2.1\tThe Seller shall sell and the Buyer shall buy the Shares with full title "
      "guarantee.")
    P("3.\tConsideration")
    P("3.1\tThe consideration for the Shares is £12,500,000 (twelve million five "
      "hundred thousand pounds) (the “Purchase Price”).")
    P("3.2\tThe Purchase Price shall be paid in instalments as follows:")
    P("(a)\t£10,000,000 on Completion;")
    P("(b)\t£1,250,000 on the first anniversary of Completion; and")
    P(f"(c)\t£{'1,500,000' if defects else '1,250,000'} on the second anniversary of "
      "Completion.")
    P("3.3\tThe Purchase Price shall be allocated as set out in Schedule 2.")
    P("3.4\tCompletion shall take place on the Completion Date.")
    P("4.\tConditions")
    P("4.1\tCompletion is conditional on the Buyer obtaining clearance from the "
      "Competition and Markets Authority. If the condition is not satisfied by the Long "
      "Stop Date, either party may terminate this agreement by notice to the other.")
    P("5.\tRetention")
    P("5.1\tThe Buyer shall hold the Retention Amount in escrow and release it to the "
      "Seller 18 months after Completion.")
    P("6.\tClaims")
    P(f"6.1\tNo claim may be brought under the Warranties unless notice of it is given "
      f"within {'12' if defects else '24'} months after Completion.")
    P("7.\tTermination")
    P("7.1\tEither party may terminate this agreement by giving not less than 30 days' "
      "written notice to the other if the other commits a material breach which is not "
      f"remedied within {'60 days' if defects else '10 Business Days'} of being "
      "notified of it.")
    P("8.\tNotices")
    P("8.1\tA notice given under this agreement shall be in writing and delivered by "
      "hand or by pre-paid first-class post to the following addresses, and is deemed "
      "received on the second Business Day after posting:")
    P("Seller: Beta Industries Limited, 1 King Street, London EC2V 8AU")
    P(f"Buyer: {title(stale)}, 20 Fenchurch Street, London EC3M 3BY")
    P("9.\tGoverning law")
    P("9.1\tThis agreement is governed by the law of England and Wales.")
    P("This agreement has been entered into on the date stated at the beginning of it.")

    d.add_page_break()
    P("SCHEDULE 1")
    P("The Company and its shareholders")
    t = d.add_table(rows=4, cols=2)
    for r, (a, b) in enumerate([("Shareholder", "Percentage"), ("Beta Industries Limited",
                                "85%"), ("Management shareholders", "15%"),
                                ("Total", "100%")]):
        t.rows[r].cells[0].text = a
        t.rows[r].cells[1].text = b
    d.add_page_break()
    P("SCHEDULE 2")
    P("Allocation of the Purchase Price")
    shares = "£10,550,000" if defects else "£11,000,000"
    total = "£12,050,000" if defects else "£12,500,000"
    t = d.add_table(rows=5, cols=2)
    for r, (a, b) in enumerate([("Asset", "Amount"), ("Shares", shares),
                                ("Intellectual property", "£1,000,000"),
                                ("Goodwill", "£500,000"), ("Total", total)]):
        t.rows[r].cells[0].text = a
        t.rows[r].cells[1].text = b

    P("EXECUTED as a deed by BETA INDUSTRIES LIMITED acting by a director in the "
      "presence of a witness:")
    P("Director: ______________________")
    P("Witness: ______________________")
    P(f"EXECUTED as a deed by {stale} acting by a director in the presence of a "
      "witness:")
    P("Director: ______________________")
    P("Witness: ______________________")
    return _save(d)
