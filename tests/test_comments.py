"""Turning comments and markup (skills/lra-comments, src/secondeye/comments.py).

Three layers, each tested offline. The Word mechanics: a reply Word shows in
the comment's thread and a "Resolved" it shows as resolved, which means every
comments part agreeing (pipeline/wordcomments.py), checked part by part and by
pipeline/validate.py. The skill: its scripts run from the built bundle, in a
fresh interpreter, as the sandbox would run them, with and without pydantic.
And the host: the request is recognised, the detached session is faked
(tests/fake_sessions.py), what comes back is examined on our side, a problem
goes back to the agent once, and the reply is built from the evidence.
"""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import zipfile
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from lxml import etree

from secondeye import comments, handler, managed, skillsync
from secondeye.mail.console import ConsoleProvider
from secondeye.models import Attachment, Finding, InboundEmail, Mode, ReviewResult, Severity
from secondeye.pipeline import ooxml, redline, turning, validate, wordcomments
from secondeye.pipeline.ooxml import Revision, RevisionWriter
from tests import fake_sessions as fs
from tests.conftest import documents

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"
W16CID = "{http://schemas.microsoft.com/office/word/2016/wordml/cid}"
ME = "Reviewer"

CLAUSES = [
    "1. Definitions",
    '1.1 "fees" means the fees set out in Schedule 1.',
    "4. Payment",
    "4.2 Invoices are payable within 45 days of receipt.",
    "7. Confidentiality",
    "7.1 Each party shall keep the other's informaton confidential.",
    "9. Liability",
    "9.4 The Supplier's total liability is capped at our standard cap.",
    "12. Governing law",
    "12.1 This Agreement is governed by the laws of England.",
]
PARTNER_COMMENTS = [
    ('1.1 "fees" means', "Capitalise the defined term."),
    ("4.2 Invoices are payable", "Change to 30 days."),
    ("7.1 Each party shall", "Typo: informaton."),
    ("9.4 The Supplier's total", "Use our standard cap."),
    ("12.1 This Agreement", "Fine as is, I think?"),
]


def save(d) -> bytes:
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def commented(markup: bool = False) -> bytes:
    """An agreement a partner has commented on, and, with `markup`, that the
    other side has marked up."""
    d = Document()
    paragraphs = {text.split(" ", 1)[0]: d.add_paragraph(text) for text in CLAUSES}
    for start, text in PARTNER_COMMENTS:
        para = paragraphs[start.split(" ", 1)[0]]
        d.add_comment(para.runs, text=text, author="Pat Partner", initials="PP")
    if markup:
        w = RevisionWriter(d, "Opposing Counsel")
        w.apply(Revision("informaton", "information", "Opposing Counsel"))
        w.apply(Revision("capped at our standard cap", "capped at the Fees",
                         "Opposing Counsel"))
        w.insert_paragraph_after("governed by the laws of England.",
                                 "12.2 Disputes go to arbitration in Geneva.")
    return save(d)


def marked_up() -> bytes:
    """The other side's markup and nothing else."""
    d = Document()
    for text in CLAUSES:
        d.add_paragraph(text)
    w = RevisionWriter(d, "Opposing Counsel")
    w.apply(Revision("informaton", "information", "Opposing Counsel"))
    w.apply(Revision("payable within 45 days", "payable within 60 days", "Opposing Counsel"))
    w.apply(Revision("capped at our standard cap", "capped at the Fees", "Opposing Counsel"))
    w.insert_paragraph_after("governed by the laws of England.",
                             "12.2 Disputes go to arbitration in Geneva.")
    return save(d)


def part(content: bytes, name: str) -> etree._Element:
    return etree.fromstring(zipfile.ZipFile(BytesIO(content)).read(name))


def edit(content: bytes, *changes: tuple[str, str]) -> bytes:
    findings = [Finding(severity=Severity.SUBSTANTIVE, category="instruction", title=a,
                        explanation="", anchor=a, suggested_text=t, auto_apply=True)
                for a, t in changes]
    out = redline.apply(content, ReviewResult(mode=Mode.REDLINE, summary="",
                                              findings=findings),
                        author=ME, comment_questions=False)
    assert len(out.applied) == len(changes), out.notes
    return out.content


def by_text(content: bytes) -> dict[str, wordcomments.Comment]:
    return {c.text: c for c in wordcomments.read(content)}


def turned_as_asked(original: bytes) -> bytes:
    """What the agent does for "turn the comments": four done and resolved,
    one left open with a question."""
    ids = {c.text: c.id for c in wordcomments.read(original)}
    content = edit(original, ('"fees" means', '"Fees" means'),
                   ("within 45 days", "within 30 days"), ("informaton", "information"))
    for text, reply, resolve in [
        ("Capitalise the defined term.", "Done.", True),
        ("Change to 30 days.", "Done: now 30 days.", True),
        ("Typo: informaton.", "Fixed.", True),
        ("Fine as is, I think?", "Agreed, no change needed.", True),
        ("Use our standard cap.", "They asked for 'our standard cap' — 1x fees or 2x?",
         False),
    ]:
        content, _ = wordcomments.reply(content, ids[text], reply, ME, resolve=resolve)
    return content


# --------------------------------------------------------------------------
# Reading comments
# --------------------------------------------------------------------------


def test_every_comment_is_read_with_its_anchor_clause_and_author():
    found = wordcomments.read(commented())
    assert [c.text for c in found] == [t for _, t in PARTNER_COMMENTS]
    cap = by_text(commented())["Use our standard cap."]
    assert cap.author == "Pat Partner" and cap.clause == "9.4"
    assert cap.anchored == "9.4 The Supplier's total liability is capped at our standard cap."
    assert cap.parent == "" and cap.replies == [] and cap.done is False
    assert re.match(r"\d{4}-\d{2}-\d{2}T", cap.date)


def test_the_anchor_is_the_text_as_the_reader_sees_it():
    """Their deletion is not part of what a comment is about; their insertion is."""
    cap = by_text(commented(markup=True))["Use our standard cap."]
    assert cap.anchored == "9.4 The Supplier's total liability is capped at the Fees."


def test_the_clause_is_the_number_word_draws():
    """A numbered list, not typed numbers: the label comes from numbering.xml,
    the same reader the review uses."""
    d = Document()
    d.add_paragraph("Liability", style="List Number")
    d.add_paragraph("Nothing limits liability for fraud.", style="List Number")
    para = d.add_paragraph("It is capped.", style="List Number")
    d.add_comment(para.runs, text="How much?", author="P")
    assert wordcomments.read(save(d))[0].clause == "3"


def test_a_document_with_no_comments_has_none():
    d = Document()
    d.add_paragraph("Nothing to say.")
    assert wordcomments.read(save(d)) == []
    with pytest.raises(wordcomments.CommentError):
        wordcomments.reply(save(d), "0", "Hello", ME)


# --------------------------------------------------------------------------
# Replies and "Resolved", as Word writes them
# --------------------------------------------------------------------------


def test_a_reply_is_threaded_in_every_part_word_reads():
    original = commented()
    cap = by_text(original)["Use our standard cap."]
    out, reply_id = wordcomments.reply(original, cap.id, "1x fees or 2x?", ME)

    # comments.xml: the reply, signed, with a paraId on its paragraph.
    reply = next(c for c in part(out, "word/comments.xml").iter(W + "comment")
                 if c.get(W + "id") == reply_id)
    assert reply.get(W + "author") == ME and reply.get(W + "initials") == "R"
    reply_para = reply.findall(W + "p")[-1].get(
        "{http://schemas.microsoft.com/office/word/2010/wordml}paraId")
    root = next(c for c in part(out, "word/comments.xml").iter(W + "comment")
                if c.get(W + "id") == cap.id)
    root_para = root.findall(W + "p")[-1].get(
        "{http://schemas.microsoft.com/office/word/2010/wordml}paraId")
    assert reply_para and root_para and reply_para != root_para

    # commentsExtended.xml: the reply points at its parent; both open.
    extended = {e.get(W15 + "paraId"): e for e in
                part(out, "word/commentsExtended.xml").iter(W15 + "commentEx")}
    assert extended[reply_para].get(W15 + "paraIdParent") == root_para
    assert extended[root_para].get(W15 + "paraIdParent") is None
    assert extended[reply_para].get(W15 + "done") == "0"

    # commentsIds.xml: a durable id for each; people.xml: the author.
    durable = {e.get(W16CID + "paraId"): e.get(W16CID + "durableId")
               for e in part(out, "word/commentsIds.xml").iter(W16CID + "commentId")}
    assert reply_para in durable and root_para in durable
    assert all(int(v, 16) < 0x80000000 for v in durable.values())
    people = [p.get(W15 + "author") for p in part(out, "word/people.xml").iter(W15 + "person")]
    assert people == [ME]

    # document.xml: the reply's range and reference beside its parent's.
    body = part(out, "word/document.xml")
    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
        assert reply_id in [m.get(W + "id") for m in body.iter(W + tag)], tag

    # Wired in: content types and relationships.
    types = zipfile.ZipFile(BytesIO(out)).read("[Content_Types].xml").decode()
    rels = zipfile.ZipFile(BytesIO(out)).read("word/_rels/document.xml.rels").decode()
    for name, kind in (("commentsExtended", "commentsExtended+xml"),
                       ("commentsIds", "commentsIds+xml"), ("people", "people+xml")):
        assert f'PartName="/word/{name}.xml"' in types and kind in types
        assert f'Target="{name}.xml"' in rels

    # The comments root declares what it uses as ignorable, as Word does.
    comments_root = part(out, "word/comments.xml")
    assert "w14" in comments_root.get(
        "{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable").split()

    report = validate.validate(out)
    assert report.ok, report
    assert redline.verify(out)[0]
    read_back = by_text(out)
    assert read_back["1x fees or 2x?"].parent == cap.id
    assert read_back["Use our standard cap."].replies == [reply_id]


def test_resolving_marks_the_thread_done_and_reopening_undoes_it():
    original = commented()
    typo = by_text(original)["Typo: informaton."]
    out, reply_id = wordcomments.reply(original, typo.id, "Fixed.", ME, resolve=True)
    assert by_text(out)["Typo: informaton."].done and by_text(out)["Fixed."].done
    extended = list(part(out, "word/commentsExtended.xml").iter(W15 + "commentEx"))
    assert {e.get(W15 + "done") for e in extended} == {"1"}

    reopened = wordcomments.resolve(out, reply_id, done=False)
    assert not by_text(reopened)["Typo: informaton."].done
    again = wordcomments.resolve(reopened, typo.id)
    assert by_text(again)["Typo: informaton."].done
    assert validate.validate(again).ok
    # Resolving touches nothing but the done flags.
    others = {c.id: (c.text, c.done) for c in wordcomments.read(again)
              if c.id not in (typo.id, reply_id)}
    assert others == {c.id: (c.text, False) for c in wordcomments.read(original)
                      if c.id != typo.id}


def test_a_reply_to_a_reply_joins_the_first_comments_thread():
    """Word threads are one level deep."""
    original = commented()
    cap = by_text(original)["Use our standard cap."]
    once, first = wordcomments.reply(original, cap.id, "1x or 2x?", ME)
    twice, second = wordcomments.reply(once, first, "Or 1.5x?", ME)
    found = by_text(twice)
    assert found["Or 1.5x?"].parent == cap.id
    assert found["Use our standard cap."].replies == [first, second]
    assert validate.validate(twice).ok


def test_a_reply_uses_the_parts_the_file_already_has():
    """A file Word saved already has commentsExtended, commentsIds and
    people; a second reply adds to them rather than creating second copies."""
    once, _ = wordcomments.reply(commented(), "0", "Done.", ME, resolve=True)
    twice, _ = wordcomments.reply(once, "1", "Done.", "Jane Associate", resolve=True)
    names = zipfile.ZipFile(BytesIO(twice)).namelist()
    assert sum(1 for n in names if "commentsExtended" in n) == 1
    assert sum(1 for n in names if "commentsIds" in n) == 1
    rels = zipfile.ZipFile(BytesIO(twice)).read("word/_rels/document.xml.rels").decode()
    assert rels.count('Target="commentsExtended.xml"') == 1
    people = [p.get(W15 + "author") for p in part(twice, "word/people.xml").iter(W15 + "person")]
    assert people == [ME, "Jane Associate"]
    assert validate.validate(twice).ok


def test_an_extensible_part_already_in_the_file_is_kept_up_to_date():
    once, _ = wordcomments.reply(commented(), "0", "Done.", ME)
    pkg = wordcomments.Package(once)
    pkg.ensure_part("extensible")
    with_cex = pkg.save()
    twice, _ = wordcomments.reply(with_cex, "1", "Done.", ME)
    cex = part(twice, "word/commentsExtensible.xml")
    ns = "{http://schemas.microsoft.com/office/word/2018/wordml/cex}"
    entries = list(cex.iter(ns + "commentExtensible"))
    assert len(entries) == 1 and entries[0].get(ns + "dateUtc")
    assert validate.validate(twice).ok


def test_a_reply_to_a_comment_that_is_not_there_raises_and_writes_nothing():
    with pytest.raises(wordcomments.CommentError, match="no comment with id 99"):
        wordcomments.reply(commented(), "99", "Hello", ME)


def test_ids_of_new_comments_are_above_every_revision_id():
    out, reply_id = wordcomments.reply(commented(markup=True), "3", "Which?", ME)
    ids = [int(i) for i in re.findall(r'w:id="(\d+)"', zipfile.ZipFile(BytesIO(out))
                                      .read("word/document.xml").decode())]
    assert int(reply_id) == max(ids)


# --------------------------------------------------------------------------
# The validator knows what a sound thread is
# --------------------------------------------------------------------------


def _rewrite(content: bytes, name: str, change) -> bytes:
    source = zipfile.ZipFile(BytesIO(content))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == name:
                data = change(data.decode()).encode()
            z.writestr(info.filename, data)
    return out.getvalue()


def test_the_validator_refuses_a_dangling_comment_reference():
    bad = _rewrite(commented(), "word/document.xml",
                   lambda x: x.replace('w:commentReference w:id="4"',
                                       'w:commentReference w:id="44"'))
    report = validate.validate(bad)
    assert not report.ok and "comment 44" in str(report)


def test_the_validator_refuses_a_reply_to_a_reply_and_unknown_paragraphs():
    once, _ = wordcomments.reply(commented(), "3", "1x or 2x?", ME)
    ext = part(once, "word/commentsExtended.xml")
    entries = list(ext.iter(W15 + "commentEx"))
    root_para, reply_para = entries[0].get(W15 + "paraId"), entries[1].get(W15 + "paraId")
    nested = _rewrite(once, "word/commentsExtended.xml", lambda x: x.replace(
        "</w15:commentsEx>",
        f'<w15:commentEx w15:paraId="{root_para}" w15:paraIdParent="{reply_para}" '
        'w15:done="0"/></w15:commentsEx>'))
    assert "one level deep" in str(validate.validate(nested))
    unknown = _rewrite(once, "word/commentsExtended.xml",
                       lambda x: x.replace(reply_para, "0ABCDEF0", 1))
    assert "which no comment has" in str(validate.validate(unknown))


def test_the_validator_refuses_a_comments_part_word_cannot_find():
    once, _ = wordcomments.reply(commented(), "3", "1x or 2x?", ME)
    unwired = _rewrite(once, "[Content_Types].xml",
                       lambda x: x.replace('PartName="/word/commentsIds.xml"',
                                           'PartName="/word/elsewhere.xml"'))
    assert "commentsIds.xml has no content type" in str(validate.validate(unwired))


# --------------------------------------------------------------------------
# The examination: what the returned file did
# --------------------------------------------------------------------------


def test_a_well_turned_file_passes_and_says_what_happened():
    original = commented()
    exam = turning.examine(original, turned_as_asked(original), ME)
    assert exam.ok, exam.report()
    statuses = {c.text: (c.status, c.clause) for c in exam.comments}
    assert statuses["Use our standard cap."] == ("open", "9.4")
    assert [s for s, _ in statuses.values()].count("resolved") == 4
    assert [c["kind"] for c in exam.ours] == ["replaced"] * 3
    assert "4 resolved, 1 open" in exam.report()


def _drop_comment(content: bytes, cid: str) -> bytes:
    """A file that has lost a comment and its thread, and is otherwise
    sound: every part that named them no longer does."""
    thread = [cid] + next(c.replies for c in wordcomments.read(content) if c.id == cid)
    paras: list[str] = []
    for tid in thread:
        match = re.search(rf'<w:comment w:id="{tid}".*?</w:comment>', zipfile.ZipFile(
            BytesIO(content)).read("word/comments.xml").decode(), re.DOTALL)
        paras += re.findall(r'w14:paraId="([0-9A-F]+)"', match.group(0))
        content = _rewrite(content, "word/comments.xml", lambda x, tid=tid: re.sub(
            rf'<w:comment w:id="{tid}".*?</w:comment>', "", x, flags=re.DOTALL))
        content = _rewrite(content, "word/document.xml", lambda x, tid=tid: re.sub(
            rf'<w:(commentRangeStart|commentRangeEnd) w:id="{tid}"/>|<w:r><w:rPr><w:rStyle '
            rf'w:val="CommentReference"/></w:rPr><w:commentReference w:id="{tid}"/></w:r>',
            "", x))
    for name, tag in (("word/commentsExtended.xml", "w15:commentEx"),
                      ("word/commentsIds.xml", "w16cid:commentId")):
        for para in paras:
            content = _rewrite(content, name, lambda x, para=para, tag=tag: re.sub(
                rf'<{tag} [^>]*paraId="{para}"[^>]*/>', "", x))
    return content


def test_a_deleted_comment_is_a_problem():
    original = commented()
    bad = _drop_comment(turned_as_asked(original), "4")
    assert validate.validate(bad).ok, "the file itself is sound"
    exam = turning.examine(original, bad, ME)
    assert not exam.ok
    assert any("was deleted" in p and "cl. 12.1" in p for p in exam.problems)


def test_their_change_removed_without_a_decision_is_a_problem():
    original = marked_up()
    d = Document(BytesIO(original))
    change = next(c for c in ooxml.tracked_changes(d) if c.kind == "new paragraph")
    ooxml.decide(d, reject=[change.id])
    exam = turning.examine(original, save(d), ME)
    assert any("neither accepted nor rejected" in p for p in exam.problems)
    assert turning.examine(original, save(d), ME, rejected=[change.id]).ok


def test_a_decision_recorded_the_wrong_way_round_is_caught_by_the_replay():
    original = marked_up()
    d = Document(BytesIO(original))
    change = next(c for c in ooxml.tracked_changes(d) if c.kind == "new paragraph")
    ooxml.decide(d, reject=[change.id])
    exam = turning.examine(original, save(d), ME, accepted=[change.id])
    assert any("recorded decisions" in p or "text differs" in p for p in exam.problems)


def test_a_change_signed_by_someone_else_is_a_problem():
    original = commented()
    d = Document(BytesIO(original))
    RevisionWriter(d, "Pat Partner").apply(Revision("45 days", "30 days", "Pat Partner"))
    exam = turning.examine(original, save(d), ME)
    assert any('signed "Pat Partner"' in p for p in exam.problems)


# --------------------------------------------------------------------------
# The skill, run from the built bundle as the sandbox runs it
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def bundle(tmp_path_factory) -> Path:
    return skillsync.build(tmp_path_factory.mktemp("skill"), "lra-comments")


def run(bundle: Path, script: str, *args: str, shim: bool = False, ok: bool = True) -> dict:
    env = {k: v for k, v in os.environ.items() if k not in ("PYTHONPATH", "SECOND_EYE_AUTHOR")}
    if shim:
        env["SECOND_EYE_FORCE_SHIM"] = "1"
    done = subprocess.run([sys.executable, str(bundle / "scripts" / script), *args],
                          capture_output=True, text=True, env=env, cwd=bundle,
                          timeout=120, check=False)
    assert done.stdout, done.stderr
    assert (done.returncode == 0) == ok, done.stdout
    return json.loads(done.stdout)


def test_the_bundle_has_the_shape_the_skills_api_requires(bundle):
    head = (bundle / "SKILL.md").read_text().split("---")[1]
    name = next(line for line in head.splitlines() if line.startswith("name:"))
    description = next(line for line in head.splitlines() if line.startswith("description:"))
    assert name.split(":", 1)[1].strip() == "lra-comments"
    assert 0 < len(description.split(":", 1)[1].strip()) <= 1024
    assert (bundle / "shims" / "pydantic" / "__init__.py").exists()
    for relative in skillsync.BUNDLED:
        assert (bundle / "lib" / "secondeye" / relative).read_bytes() == (
            skillsync.PACKAGE / relative).read_bytes()
    for script in ("list_comments.py", "reply_comment.py", "resolve_comment.py",
                   "accept_reject.py", "edit.py", "finish.py"):
        assert (bundle / "scripts" / script).exists()


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_turning_comments_runs_from_the_bundle(bundle, tmp_path, shim):
    source = commented()
    original = tmp_path / "workspace" / "MSA.docx"
    original.parent.mkdir()
    original.write_bytes(source)
    work = tmp_path / "turn" / "work.docx"
    work.parent.mkdir()
    work.write_bytes(original.read_bytes())
    outputs = tmp_path / "outputs"

    listed = run(bundle, "list_comments.py", str(work), shim=shim)
    ids = {c["text"]: c["id"] for c in listed["comments"]}
    assert listed["open_threads"] == 5
    assert {c["clause"] for c in listed["comments"]} == {"1.1", "4.2", "7.1", "9.4", "12.1"}

    changes = tmp_path / "turn" / "changes.json"
    changes.write_text(json.dumps([
        {"anchor": '"fees" means', "text": '"Fees" means', "kind": "replace",
         "title": "Defined term"},
        {"anchor": "within 45 days", "text": "within 30 days", "kind": "replace",
         "title": "30 days"},
        {"anchor": "informaton", "text": "information", "kind": "replace", "title": "Typo"},
    ]))
    edited = run(bundle, "edit.py", str(work), "--changes", str(changes), "--author", ME,
                 shim=shim)
    assert edited["applied"] == ["Defined term", "30 days", "Typo"] and edited["verified"]

    for text, reply in [("Capitalise the defined term.", "Done."),
                        ("Change to 30 days.", "Done: now 30 days."),
                        ("Typo: informaton.", "Fixed.")]:
        answered = run(bundle, "reply_comment.py", str(work), "--id", ids[text], "--text",
                       reply, "--resolve", "--author", ME, shim=shim)
        assert answered["resolved"] is True and answered["thread"] == ids[text]
    run(bundle, "reply_comment.py", str(work), "--id", ids["Fine as is, I think?"],
        "--text", "Agreed, no change needed.", "--author", ME, shim=shim)
    run(bundle, "resolve_comment.py", str(work), "--id", ids["Fine as is, I think?"],
        shim=shim)
    run(bundle, "reply_comment.py", str(work), "--id", ids["Use our standard cap."],
        "--text", "They asked for 'our standard cap' — 1x fees or 2x?", "--open",
        "--author", ME, shim=shim)

    out = outputs / "MSA (comments turned).docx"
    done = run(bundle, "finish.py", str(original), str(work), "--out", str(out),
               "--author", ME, shim=shim)
    assert done["file"] == str(out) and out.exists()
    assert "4 resolved, 1 open" in done["report"]
    record = json.loads((outputs / "turned.json").read_text())
    assert record["ok"] and record["author"] == ME
    assert validate.validate(out.read_bytes()).ok
    assert original.read_bytes() == source, "the lawyer's file is untouched"


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_turning_markup_runs_from_the_bundle_with_a_blackline(bundle, tmp_path, shim):
    original = tmp_path / "MSA.docx"
    original.write_bytes(marked_up())
    work = tmp_path / "turn" / "work.docx"
    work.parent.mkdir()
    work.write_bytes(original.read_bytes())

    listed = run(bundle, "accept_reject.py", str(work), "--list", shim=shim)
    changes = {(c["kind"], c["text"]): c["id"] for c in listed["tracked_changes"]}
    typo = changes[("replaced", "informaton → information")]
    cap = changes[("replaced", "capped at our standard cap → capped at the Fees")]
    decided = run(bundle, "accept_reject.py", str(work), "--accept", typo,
                  "--reject", cap, shim=shim)
    assert decided["accepted"] == [typo] and decided["rejected"] == [cap]
    assert len(decided["still_tracked"]) == 2

    counter = tmp_path / "turn" / "counter.json"
    counter.write_text(json.dumps([{"anchor": "capped at our standard cap",
                                    "text": "capped at 1.5x the Fees", "kind": "replace",
                                    "title": "Our fallback"}]))
    run(bundle, "edit.py", str(work), "--changes", str(counter), "--author", ME, shim=shim)

    out = tmp_path / "outputs" / "MSA (our turn).docx"
    line = tmp_path / "outputs" / "MSA (blackline against theirs).docx"
    done = run(bundle, "finish.py", str(original), str(work), "--out", str(out),
               "--blackline", str(line), "--author", ME, shim=shim)
    assert done["blackline"] == str(line)
    record = json.loads((tmp_path / "outputs" / "turned.json").read_text())
    assert [c["id"] for c in record["accepted"]] == [typo]
    assert [c["id"] for c in record["rejected"]] == [cap]
    assert {c["kind"] for c in record["still_tracked"]} == {"replaced", "new paragraph"}
    assert [c["text"] for c in record["ours"]] == [
        "capped at our standard cap → capped at 1.5x the Fees"]

    ours = Document(BytesIO(out.read_bytes()))
    authors = {c.author for c in ooxml.tracked_changes(ours)}
    assert authors == {"Opposing Counsel", ME}, "their undecided changes are still theirs"
    blackline = line.read_bytes()
    assert redline.verify(blackline, expect_revisions=True)[0]
    from secondeye.pipeline.extract import visible_text
    text = " ".join(visible_text(p._p) for p in Document(BytesIO(blackline)).paragraphs)
    assert "1.5x the Fees" in text


def test_finish_refuses_a_file_that_lost_a_comment(bundle, tmp_path):
    original = tmp_path / "MSA.docx"
    original.write_bytes(commented())
    work = tmp_path / "work.docx"
    work.write_bytes(_drop_comment(turned_as_asked(commented()), "4"))
    out = tmp_path / "outputs" / "MSA (comments turned).docx"
    refused = run(bundle, "finish.py", str(original), str(work), "--out", str(out),
                  "--author", ME, ok=False)
    assert refused["error"].startswith("NOTHING WAS WRITTEN")
    assert "was deleted" in refused["error"]
    assert not out.exists()
    assert json.loads((out.parent / "turned.json").read_text())["ok"] is False


def test_a_decision_without_the_script_is_refused(bundle, tmp_path):
    """Their change removed by hand, not by accept_reject.py: finish says so."""
    original = tmp_path / "MSA.docx"
    original.write_bytes(marked_up())
    d = Document(BytesIO(marked_up()))
    ooxml.decide(d, accept=[ooxml.tracked_changes(d)[0].id])
    work = tmp_path / "work.docx"
    work.write_bytes(save(d))
    refused = run(bundle, "finish.py", str(original), str(work), "--out",
                  str(tmp_path / "o.docx"), "--author", ME, ok=False)
    assert "neither accepted nor rejected" in refused["error"]


def test_script_errors_are_json(bundle, tmp_path):
    work = tmp_path / "w.docx"
    work.write_bytes(commented())
    assert "no comment with id 77" in run(bundle, "reply_comment.py", str(work), "--id",
                                          "77", "--text", "x", ok=False)["error"]
    assert "no tracked change with id 5" in run(
        bundle, "accept_reject.py", str(work), "--accept", "5", ok=False)["error"]


# --------------------------------------------------------------------------
# The host: the request, the session, the reply
# --------------------------------------------------------------------------


class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return f"agent-msg-{len(self.sent)}"


@pytest.fixture
def env(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/c.sqlite3")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    monkeypatch.setenv("ALLOWLIST_CONTACT", "Jim Baker <jim@firm.com>")
    monkeypatch.setenv("REDLINE_AUTHOR", ME)
    fs.configure(monkeypatch, MANAGED_COMMENTS_AGENT_ID="agent_comments",
                 SANDBOX_COMMENTS_SKILL_ID="skill_comments", MANAGED_PLAYBOOK_AGENT_ID="")
    now = [0.0]
    monkeypatch.setattr(managed.time, "monotonic", lambda: now[0])
    monkeypatch.setattr(managed.time, "sleep", lambda s: now.__setitem__(0, now[0] + s))
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    yield provider
    from secondeye.config import settings
    settings.cache_clear()


def email(body: str, content: bytes | None, subject: str = "MSA",
          name: str = "MSA v3.docx") -> InboundEmail:
    attachments = [] if content is None else [
        Attachment(filename=name, content_type=DOCX, size_bytes=len(content),
                   content=content)]
    return InboundEmail(message_id=f"<m-{abs(hash(body))}@firm.com>", thread_id="t1",
                        from_address="jim@firm.com", to=["review@legal.firm.com"],
                        subject=subject, text_body=body, attachments=attachments,
                        received_at=datetime.now(UTC))


@pytest.mark.parametrize("body", [
    "Turn the comments please.", "Can you turn these?", "turn these", "Turn this round by 5pm.",
    "Please deal with the partner's comments.", "Address the client comments on this.",
    "Turn Pat's comments", "Work through the comments and send it back.",
])
def test_the_request_is_recognised_when_a_commented_document_is_attached(env, body):
    assert comments.intent(email(body, commented())) == comments.TURN


@pytest.mark.parametrize("body", [
    ("Turn their markup: accept their typo fixes, reject the cap change, counter 9.4 "
     "with our fallback."),
    "Accept their changes to clause 7 and reject the rest.",
])
def test_the_markup_request_is_recognised_when_their_changes_are_attached(env, body):
    assert comments.intent(email(body, marked_up())) == comments.TURN


@pytest.mark.parametrize("body,content", [
    ("Please review this.", "commented"),
    ("Turn this into a clean copy.", "commented"),
    ("Turn the comments please.", "plain"),
    ("Turn the comments please.", None),
    ("Accept their changes.", "commented"),
])
def test_other_requests_are_not_turning(env, body, content):
    d = Document()
    d.add_paragraph("No comments here.")
    source = {"commented": commented(), "plain": save(d), None: None}[content]
    assert comments.intent(email(body, source)) is None


def agent(outputs: dict[str, bytes], text: str = "Turned.") -> fs.Turn:
    return fs.Turn(fs.message(text), fs.idle(), outputs=outputs)


def record_of(original: bytes, turned: bytes, **decided) -> bytes:
    return json.dumps(turning.examine(original, turned, ME, **decided).as_dict()).encode()


def test_turn_the_comments_end_to_end(env):
    original = commented()
    turned = turned_as_asked(original)
    fake = fs.DetachedSessions([agent({"MSA v3 (comments turned).docx": turned,
                                       "turned.json": record_of(original, turned)})])
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Turn the comments please.", original))

    out = env.sent[-1]
    assert out.text_body.splitlines()[0] == (
        "Turned 5 comments: 4 done and resolved, 1 left open "
        "(cl. 9.4: They asked for 'our standard cap' — 1x fees or 2x?).")
    attached = documents(out)
    assert [a.filename for a in attached] == ["MSA v3 (comments turned).docx"]
    assert attached[0].content == turned
    assert "3 tracked changes of ours, for you to accept or reject." in out.text_body

    created = fake.created[0]
    assert created["agent"] == "agent_comments"
    opening = created["initial_events"][0]["content"][0]["text"]
    assert "Use our standard cap." in opening and '--author "Reviewer"' in opening
    assert "/mnt/session/outputs/MSA v3 (comments turned).docx" in opening
    assert opening.rstrip().endswith("short summary.")
    assert fake.deleted_sessions == ["sesn_fake"]


def test_turn_their_markup_end_to_end_attaches_our_turn_and_a_blackline(env):
    original = marked_up()
    d = Document(BytesIO(original))
    changes = {(c.kind, c.text): c.id for c in ooxml.tracked_changes(d)}
    accept = [changes[("replaced", "informaton → information")]]
    reject = [changes[("replaced", "capped at our standard cap → capped at the Fees")]]
    ooxml.decide(d, accept=accept, reject=reject)
    ours = edit(save(d), ("capped at our standard cap", "capped at 1.5x the Fees"))
    line, _ = turning.blackline(original, ours, ME)
    record = record_of(original, ours, accepted=accept, rejected=reject)
    fake = fs.DetachedSessions([agent({"MSA v3 (our turn).docx": ours,
                                       "MSA v3 (blackline against theirs).docx": line,
                                       "turned.json": record})])
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Turn their markup: accept their typo fixes, reject the cap "
                             "change, counter 9.4 with our fallback.", original))

    out = env.sent[-1]
    assert [a.filename for a in documents(out)] == [
        "MSA v3 (our turn).docx", "MSA v3 (blackline against theirs).docx"]
    assert "Turned their markup: accepted 1 change, rejected 1, left 2 tracked." \
        in out.text_body
    assert "1 tracked change of ours, for you to accept or reject." in out.text_body
    opening = fake.created[0]["initial_events"][0]["content"][0]["text"]
    assert "--blackline" in opening and "counter 9.4 with our fallback" in opening


def test_a_problem_goes_back_to_the_agent_once_and_the_fix_is_sent(env):
    original = commented()
    good = turned_as_asked(original)
    bad = _drop_comment(good, "4")
    fake = fs.DetachedSessions([
        agent({"MSA v3 (comments turned).docx": bad, "turned.json": b"{}"}),
        agent({"MSA v3 (comments turned).docx": good}),
    ])
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Turn the comments please.", original))

    correction = fake.sent[0][0]
    assert correction["type"] == "user.message"
    text = correction["content"][0]["text"]
    assert "was deleted" in text and "finish.py" in text and "File check: passed" in text
    out = env.sent[-1]
    assert documents(out)[0].content == good
    assert out.text_body.startswith("Turned 5 comments")


def test_a_file_that_still_fails_is_not_attached(env):
    original = commented()
    bad = _drop_comment(turned_as_asked(original), "4")
    fake = fs.DetachedSessions([agent({"MSA v3 (comments turned).docx": bad}),
                                agent({"MSA v3 (comments turned).docx": bad})])
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Turn the comments please.", original))
    out = env.sent[-1]
    assert documents(out) == []
    assert out.text_body.startswith("I couldn't return a turned copy of MSA v3.docx")
    assert "was deleted" in out.text_body
    assert len(fake.sent) == 1, "one correction, not a loop"


def test_no_file_at_all_is_asked_for_once_then_reported(env):
    fake = fs.DetachedSessions([agent({}), agent({})])
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Turn the comments please.", commented()))
    assert "finish.py" in fake.sent[0][0]["content"][0]["text"]
    assert documents(env.sent[-1]) == []


@pytest.mark.parametrize("missing", ["MANAGED_COMMENTS_AGENT_ID", "SANDBOX_COMMENTS_SKILL_ID"])
def test_not_set_up_says_so_and_touches_nothing(env, monkeypatch, missing):
    settings_now = {"MANAGED_COMMENTS_AGENT_ID": "agent_comments",
                    "SANDBOX_COMMENTS_SKILL_ID": "skill_comments"}
    fs.configure(monkeypatch, **{**settings_now, missing: ""})
    fake = fs.DetachedSessions([])
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Turn the comments please.", commented()))
    assert "isn't set up yet" in env.sent[-1].text_body
    assert documents(env.sent[-1]) == []
    assert fake.created == []


def test_the_comments_agent_has_its_skill_and_no_custom_tools(monkeypatch):
    """Detached, so a custom tool would wait for ever; and lra-comments is
    offered to this agent only."""
    from secondeye.config import settings

    monkeypatch.setenv("SANDBOX_COMMENTS_SKILL_ID", "skill_comments")
    settings.cache_clear()
    try:
        body = managed.agent_body(
            managed.load_manifest(managed.AGENTS / "comments.agent.yaml"), [])
        assert not [t for t in body["tools"] if t.get("type") == "custom"]
        assert {"type": "custom", "skill_id": "skill_comments", "version": "latest"} \
            in body["skills"]
        assert "skill_comments" not in [s["skill_id"] for s in managed.skill_refs()]
        assert "lra-comments" in body["system"]
    finally:
        settings.cache_clear()


def test_the_summary_lists_each_open_question_when_there_are_several():
    exam = turning.Examination(verified=True, verification="ok", comments=[
        turning.CommentOutcome("1", "P", "9.4", "", "", "open", "1x or 2x?"),
        turning.CommentOutcome("2", "P", "", "the Supplier shall", "", "open", "Which?"),
        turning.CommentOutcome("3", "P", "4.2", "", "", "resolved", "Done."),
        turning.CommentOutcome("4", "P", "7", "", "", "untouched"),
    ])
    assert comments.summary(exam) == [
        "Turned 3 of 4 comments: 1 done and resolved, 2 left open:",
        "- cl. 9.4: 1x or 2x?",
        '- "the Supplier shall": Which?',
        "Not answered: cl. 7.",
    ]
