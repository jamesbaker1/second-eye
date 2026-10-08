"""The checks on documents drafted the way firms draft them.

Every check here runs on every document, so a finding on a correct agreement
is the failure that matters most (PRODUCT.md). Each fix below is pinned from
both sides: the realistic drafting that used to raise a false finding now
stays silent, and the defect the check exists for is still caught.
"""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from evals import realistic
from secondeye.models import Attachment, Severity
from secondeye.pipeline import checks, extract

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def text_doc(*paragraphs):
    raw = "\n".join(paragraphs).encode()
    return extract.extract(Attachment(filename="t.txt", content_type="text/plain",
                                      size_bytes=len(raw), content=raw))


def docx_doc(content: bytes, name: str = "t.docx"):
    return extract.extract(Attachment(filename=name, content_type=DOCX,
                                      size_bytes=len(content), content=content))


def save(d) -> bytes:
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    buffer = BytesIO()
    d.save(buffer)
    return buffer.getvalue()


def auto_numbered(*items: tuple[int, str], before: tuple[str, ...] = ()) -> bytes:
    """A document whose clauses are numbered by Word: (level, text) pairs."""
    d = Document()
    realistic._add_multilevel_list(d)
    for text in before:
        d.add_paragraph(text)
    for level, text in items:
        if level < 0:
            d.add_paragraph(text)
        else:
            realistic._numbered(d, text, level)
    return save(d)


def titles(findings):
    return [f.title for f in findings]


# --- Word's automatic numbering ---------------------------------------------

def test_the_number_word_draws_is_read_from_the_list_definition():
    doc = docx_doc(auto_numbered(
        (0, "Definitions"), (1, "In this agreement:"), (2, "first"), (2, "second"),
        (1, "Interpretation"), (0, "Sale"), (1, "The Seller shall sell."), (2, "restarts"),
    ))
    numbered = [(b.number, b.text) for b in doc.blocks if b.number]
    assert numbered == [
        ("1.", "Definitions"), ("1.1", "In this agreement:"), ("(a)", "first"),
        ("(b)", "second"), ("1.2", "Interpretation"), ("2.", "Sale"),
        ("2.1", "The Seller shall sell."), ("(a)", "restarts"),
    ]
    # The number is not in the text, because it is not in the runs: an anchor
    # that quoted it could never be found again.
    assert doc.blocks[1].text == "In this agreement:"
    assert doc.blocks[1].display == "1.1 In this agreement:"
    assert "(paragraph, numbered 1.1) In this agreement:" in doc.as_prompt()


def test_a_start_override_on_the_list_instance_restarts_the_count():
    d = Document()
    realistic._add_multilevel_list(d)
    numbering = d.part.numbering_part.element
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="91"><w:abstractNumId w:val="90"/>'
        '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="5"/></w:lvlOverride></w:num>'))
    realistic._numbered(d, "One", 0)
    p = d.add_paragraph("Five")
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="91"/></w:numPr>'))
    doc = docx_doc(save(d))
    assert [b.number for b in doc.blocks if b.number] == ["1.", "5."]


def test_numbering_carried_by_the_paragraph_style_is_read_too():
    """Most firm templates number clauses through the heading styles."""
    d = Document()
    realistic._add_multilevel_list(d)
    style = d.styles["Heading 2"]
    style.element.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="1"/><w:numId w:val="90"/></w:numPr>'))
    realistic._numbered(d, "Definitions", 0)
    d.add_paragraph("Meaning", style="Heading 2")
    doc = docx_doc(save(d))
    assert [b.number for b in doc.blocks if b.number] == ["1.", "1.1"]


def test_a_list_this_reader_cannot_follow_is_admitted_not_guessed():
    d = Document()
    p = d.add_paragraph("Numbered by a list that is not defined")
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="77"/></w:numPr>'))
    assert docx_doc(save(d)).numbering_unresolved is True


def test_the_body_is_read_in_order_and_into_content_controls():
    content = realistic.uk_defects()
    doc = docx_doc(content)
    assert any("[INSERT AMOUNT]" in b.text for b in doc.blocks)
    assert any("INSERT AMOUNT" in f.anchor for f in checks.placeholders(doc))


def test_a_definitions_table_is_read_where_it_sits():
    doc = docx_doc(realistic.us_spa())
    texts = [b.text for b in doc.blocks]
    assert texts.index("“Business Day”") < texts.index(
        next(t for t in texts if t.startswith("Purchase and Sale.")))


# --- cross-references --------------------------------------------------------

def test_references_to_auto_numbered_clauses_are_not_broken():
    doc = docx_doc(auto_numbered(
        (0, "Price"), (1, "The price is payable under Section 2.1."),
        (0, "Payment"), (1, "As provided in Sections 1.1 and 2.1, and clause 1."),
        before=("EXHIBIT A",),
    ))
    assert checks.cross_references(doc) == []


def test_a_broken_reference_in_an_auto_numbered_document_is_still_caught():
    doc = docx_doc(auto_numbered((0, "Price"), (1, "Subject to Section 9.1, the price.")))
    assert titles(checks.cross_references(doc)) == [
        "Reference to Section 9.1, which does not exist"]


def test_no_clause_numbers_at_all_means_no_clause_reference_findings():
    """Only an "EXHIBIT A" title is visible, so every clause number is unknown."""
    doc = text_doc("EXHIBIT A", "Subject to Section 2.2 and Section 5.2.")
    assert checks.cross_references(doc) == []


def test_unresolved_numbering_silences_clause_references():
    d = Document()
    d.add_paragraph("1. Price")
    p = d.add_paragraph("Subject to Section 7.2.")
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="77"/></w:numPr>'))
    assert checks.cross_references(docx_doc(save(d))) == []


@pytest.mark.parametrize("sentence", [
    "Tax is due under Section 1.1502-6 of the Treasury Regulations.",
    "Except as set out in Section 4.1 of the Disclosure Schedule.",
    "Subject to Section 3 of the Escrow Agreement.",
    "As required by section 9 of the Companies Act 2006.",
    "Processing under article 6(1)(f) of the UK GDPR.",
])
def test_a_section_of_another_instrument_is_not_a_broken_reference(sentence):
    assert checks.cross_references(text_doc("1. Scope", sentence)) == []


def test_lower_case_and_plural_references_are_checked():
    doc = text_doc("1. Scope", "2. Price", "3. Payment",
                   "As set out in clause 12.", "Subject to Sections 2 and 4.",
                   "Under clause 3 and clauses 1 to 2.")
    assert sorted(titles(checks.cross_references(doc))) == [
        "Reference to Section 4, which does not exist",
        "Reference to clause 12, which does not exist",
    ]


# --- numbering ---------------------------------------------------------------

def test_each_part_of_a_schedule_restarts_its_numbering():
    doc = docx_doc(realistic.uk_spa(auto_numbered=False))
    assert checks.numbering(doc) == []


def test_a_duplicate_is_named_in_the_documents_own_word():
    doc = text_doc("1. Scope", "2. Price", "2. Payment", "See clause 1.")
    assert titles(checks.numbering(doc)) == ["Clause 2 appears 2 times"]


# --- placeholders -------------------------------------------------------------

@pytest.mark.parametrize("text", [
    "3.3 [Reserved].", "[RESERVED]", "[Intentionally Omitted]", "[Redacted]",
    "up to [***].", "reads “[...] the court shall”", "[Intentionally left blank]",
    "[The Buyer shall provide statements quarterly.]",
    "under Title XX of the Social Security Act",
])
def test_deliberate_brackets_and_numerals_are_not_placeholders(text):
    assert checks.placeholders(text_doc(text)) == []


@pytest.mark.parametrize("text, shown", [
    ("“Long Stop Date” means [●].", "[●]"),
    ("a deposit of £[●] on completion", "[●]"),
    ("[x] LLP", "[x]"),
    ("released [within [x] days]", "[x]"),
    ("on or before [date]", "[date]"),
    ("into the account of [name of bank]", "[name of bank]"),
    ("a fee of $XX,XXX", "$XX,XXX"),
    ("dated 20XX", "20XX"),
])
def test_the_blanks_a_drafter_leaves_are_caught(text, shown):
    found = checks.placeholders(text_doc(text))
    assert [f.anchor for f in found] == [shown]
    assert found[0].severity is Severity.BLOCKER


def test_a_drafting_note_is_caught_once_as_a_note():
    found = checks.placeholders(text_doc(
        "5. Notices. [Note to draft: Buyer to confirm notice details.] As set out.",
        "7. Release. [NTD: check with tax.]",
    ))
    assert titles(found) == [
        "Drafting note left in: [Note to draft: Buyer to confirm notice details.]",
        "Drafting note left in: [NTD: check with tax.]",
    ]


def test_a_placeholder_title_does_not_repeat_itself():
    assert titles(checks.placeholders(text_doc("Price: TBD"))) == ["Unfilled TBD"]


def test_labelled_rules_under_an_execution_line_are_places_to_sign():
    doc = text_doc(
        "SIGNED by John Brown for and on behalf of BLUEWATER CAPITAL LLP:",
        "Member: ______________________",
        "Authorised Representative ______________________",
    )
    assert checks.placeholders(doc) == []


def test_a_labelled_blank_in_the_body_is_still_a_blank():
    doc = text_doc("12. Notices. Notices to the Buyer shall be sent to:",
                   "Address: ______________________")
    assert [f.anchor for f in checks.placeholders(doc)] == ["______________________"]


# --- defined terms -------------------------------------------------------------

def test_a_real_word_near_a_defined_term_is_not_a_misspelling():
    doc = text_doc(
        'Acme Parent Holdings, Inc. ("Parent") and the Seller.',
        "3.2 Schedule 3.2 lists each Patent owned by the Company. No Parent entity "
        "holds any Patent.",
    )
    assert checks.term_misspellings(doc) == []


def test_a_legend_in_quotes_is_not_a_definition():
    doc = text_doc(*([f"{i}. Clause {i}." for i in range(1, 12)] + [
        'Each party shall keep confidential all information marked "Confidential".']))
    assert not any("Confidential" in t for t in titles(checks.defined_terms(doc)))


def test_a_term_used_only_in_its_plural_is_used():
    doc = text_doc(*([f"{i}. Clause {i}." for i in range(1, 12)] + [
        ('"Consideration Shares" means the shares to be issued, each a '
         '"Consideration Share".'),
        "The Buyer shall issue 50,000 Consideration Shares.",
    ]))
    assert checks.defined_terms(doc) == []


def test_a_shorter_term_inside_a_longer_one_is_not_a_use():
    doc = text_doc(
        '1. Definitions. "Consideration Shares" means the shares issued under clause 3.',
        "2. The Consideration Shares are issued at Completion.",
        '3. Price. The price is £1 (the "Consideration").',
    )
    assert not any("used in" in t for t in titles(checks.defined_terms(doc)))


def test_a_term_used_before_its_definition_is_still_caught():
    doc = text_doc("1. The Consideration is payable.", "2. Payment.",
                   '3. Price. The price is £1 (the "Consideration").')
    assert '"Consideration" is used in clause 1 but defined in clause 3' in titles(
        checks.defined_terms(doc))


def test_lower_case_inside_a_definition_table_or_its_limbs_is_the_definition():
    assert checks.term_capitalization(docx_doc(realistic.us_spa())) == []
    assert checks.term_capitalization(
        docx_doc(realistic.uk_spa(auto_numbered=True))) == []


def test_lower_case_use_elsewhere_is_still_reported():
    doc = text_doc(
        '"Material Adverse Change" means any event which:',
        "(a) has a material adverse effect; or",
        "(b) results in a material adverse change.",
        "4. Conditions. There has been no material adverse change.",
    )
    found = checks.term_capitalization(doc)
    assert len(found) == 1 and found[0].anchor == "material adverse change"


def test_a_term_defined_twice_differently_is_caught():
    doc = text_doc('1. "Completion Date" means 10 April 2026.',
                   '6. "Completion Date" means the fifth Business Day after the Conditions.')
    found = checks.duplicate_definitions(doc)
    assert titles(found) == ['"Completion Date" is defined twice, differently']
    assert found[0].severity is Severity.SUBSTANTIVE


def test_a_pointer_and_an_exhibits_own_definitions_are_not_duplicates():
    doc = text_doc(
        '1. "Completion" has the meaning given in clause 5.',
        '5. Completion (“Completion”) takes place at the offices of the Buyer.',
        '"Agreement" means this agreement.',
        "EXHIBIT A",
        '"Agreement" means this escrow agreement.',
    )
    assert checks.duplicate_definitions(doc) == []


def test_an_undefined_capitalised_term_used_twice_is_reported():
    doc = text_doc('1. The Buyer ("Buyer") shall pay the Escrow Agent.',
                   "2. The Escrow Agent shall hold the funds.")
    found = checks.undefined_terms(doc)
    assert titles(found) == [
        '"Escrow Agent" is capitalised like a defined term but is never defined']
    assert found[0].severity is Severity.STYLE


def test_names_statutes_places_and_sentence_starts_are_not_undefined_terms():
    doc = text_doc(
        "1. Governing Law. The laws of the State of Delaware apply, and the "
        "Companies Act applies. Jane Doe signs. Jane Doe dates it.",
        "2. Governing Law. The Supreme Court and the High Court and the "
        "Financial Conduct Authority. Northgate Freight Limited and the "
        "Northgate Freight Limited board.",
        "3. The Purchase Price is paid. The Purchase Price is final.",
        '4. Price. the "Purchase Price" means £1.',
        "5. Board. The Board and the Board decide.",
    )
    assert checks.undefined_terms(doc) == []


# --- amounts, currency, dates ---------------------------------------------------

def test_figures_then_words_that_disagree_are_caught():
    found = checks.amounts(text_doc(
        "The price is £2,500,000 (two million three hundred thousand pounds) payable "
        "within 10 (fifteen) Business Days."))
    assert sorted(titles(found)) == [
        "'10 (fifteen)' disagrees with itself",
        "'£2,500,000 (two million three hundred thousand pounds)' disagrees with itself",
    ]


def test_figures_then_words_that_agree_are_silent():
    assert checks.amounts(text_doc(
        "The price is £2,500,000 (two million five hundred thousand pounds) and "
        "$25,000,000 (twenty-five million US dollars), within 10 (ten) days, under "
        "clause 3 (Definitions), £1.50 (one pound fifty pence).")) == []


def test_an_amount_anchor_is_whole_words():
    found = checks.amounts(text_doc(
        "Either party may end the employment on thirty (13) days' written notice."))
    anchor = found[0].anchor
    assert anchor.startswith("employment") and "thirty (13)" in anchor


def test_three_notations_for_one_currency_are_noted_once():
    found = checks.currency_notation(text_doc(
        "Pay US$50,000 to the Agent, $ 20,000 to the Buyer and USD 10,000 to the "
        "Registrar."))
    assert len(found) == 1
    assert found[0].severity is Severity.FORMATTING
    assert '"US$50,000"' in found[0].title


def test_one_notation_throughout_is_silent():
    assert checks.currency_notation(text_doc(
        "Pay $50,000 and $20,000, and £1,000 and HK$5,000.")) == []


def test_a_weekday_that_does_not_match_its_date_is_caught():
    found = checks.weekday_mismatch(text_doc(
        "THIS AGREEMENT is dated Monday, 3 March 2026.",
        "Not later than 30 April 2026, being a Friday."))
    assert len(found) == 2
    assert "3 March 2026 is a Tuesday" in found[0].title


def test_a_weekday_that_matches_is_silent():
    assert checks.weekday_mismatch(text_doc(
        "Dated Tuesday, 3 March 2026, and by Thursday, April 30, 2026, being a "
        "Thursday.")) == []


def test_an_ambiguous_slashed_date_is_read_the_documents_way_and_not_applied():
    found = checks.date_consistency(text_doc(
        "Dated 3 March 2026.", "Completion on 03/04/2026.", "Long stop 30 April 2026."))
    assert len(found) == 1
    assert found[0].suggested_text == "3 April 2026"      # day first, as the document
    assert found[0].auto_apply is False
    assert found[0].options == ["3 April 2026", "4 March 2026"]


def test_an_unambiguous_slashed_date_is_still_converted():
    found = checks.date_consistency(text_doc(
        "Dated 3 March 2026.", "Completion on 25/04/2026.", "Long stop 30 April 2026."))
    assert found[0].suggested_text == "25 April 2026" and found[0].auto_apply


def test_a_later_date_in_a_representation_is_not_the_opening_date():
    doc = text_doc("Entered into as of March 3, 2026.",
                   "3.1 As of the date hereof and since December 31, 2024, the Company "
                   "has been in good standing.")
    assert checks.contradictory_dates(doc) == []


def test_the_opening_date_restated_in_another_format_is_not_a_contradiction():
    doc = text_doc("Entered into as of March 3, 2026.",
                   "The term runs from the date first written above, being 3 March 2026.")
    assert checks.contradictory_dates(doc) == []


# --- party names ------------------------------------------------------------------

def test_a_name_in_capitals_on_the_signature_page_is_the_same_name():
    doc = text_doc("Between Acme Merger Sub LLC (“Buyer”) and Bluewater Capital LLC.",
                   "ACME MERGER SUB LLC", "BLUEWATER CAPITAL LLC")
    assert checks.party_names(doc) == []


# --- leftovers ----------------------------------------------------------------------

def test_a_style_separator_is_not_hidden_text():
    assert checks.leftovers(realistic.style_separator()) == []


def test_hidden_text_in_a_run_is_still_reported():
    d = Document()
    run = d.add_paragraph("Visible. ").add_run("Our fallback is 5%.")
    run.font.hidden = True
    assert "The document contains hidden text" in titles(checks.leftovers(save(d)))


def test_a_moved_or_reformatted_passage_is_a_tracked_change():
    assert "The document still contains tracked changes" in titles(
        checks.leftovers(realistic.moved_and_reformatted()))


# --- where each finding is -------------------------------------------------------------

def test_a_finding_says_which_clause_it_is_in():
    found = checks.run_all(text_doc(
        "1. Scope", "3.2 The fee is thirty (13) days.",
        "SCHEDULE 4", "1. Limits", "2. The cap is TBD."))
    where = {f.category: f.where for f in found}
    assert where["amount"] == "clause 3.2"
    assert where["placeholder"] == "Schedule 4, paragraph 2"


def test_an_auto_numbered_clause_is_named_by_its_drawn_number():
    doc = docx_doc(auto_numbered((0, "Price"), (1, "The price is TBD.")))
    found = [f for f in checks.run_all(doc) if f.category == "placeholder"]
    assert found[0].where == "clause 1.1"


# --- the realistic corpus ----------------------------------------------------------------

@pytest.mark.parametrize("build", [
    lambda: realistic.uk_spa(auto_numbered=False),
    lambda: realistic.uk_spa(auto_numbered=True),
    realistic.us_spa,
    realistic.style_separator,
], ids=["uk-typed", "uk-auto", "us-auto", "style-separator"])
def test_a_realistic_clean_agreement_raises_nothing(build):
    content = build()
    found = checks.run_all(docx_doc(content), content)
    assert titles(found) == []


def test_the_realistic_defects_are_all_caught():
    content = realistic.uk_defects()
    found = checks.run_all(docx_doc(content), content)
    joined = "\n".join(titles(found))
    for expected in (
        "[●]", "[x]", "[date]", "[name of bank]", "[INSERT AMOUNT]",
        "Drafting note left in", '"Completion Date" is defined twice, differently',
        "Reference to clause 12", "'10 (fifteen)' disagrees",
        "Dollar amounts are written 3 ways", "3 March 2026 is a Tuesday",
        "30 April 2026 is a Thursday", '"Northgate Holdings Ltd" should probably be',
    ):
        assert expected in joined, expected
    slashed = next(f for f in found if f.anchor == "03/04/2026")
    assert slashed.auto_apply is False and slashed.suggested_text == "3 April 2026"
