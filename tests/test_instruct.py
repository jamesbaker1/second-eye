"""Carrying out a free-form instruction.

The model call is faked so these run offline; what is exercised is everything
around it, which is where the product promises live: that a change lands as a
tracked change, that drafting is marked as new language, that a concern is
stated and the change made anyway, and that an instruction-driven edit can be
undone like any other.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document

from secondeye import handler
from secondeye.mail.console import ConsoleProvider
from secondeye.models import Attachment, InboundEmail, ReviewResult
from secondeye.pipeline import instruct

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return f"agent-msg-{len(self.sent)}"


@pytest.fixture
def captured(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/i.sqlite3")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    from secondeye.config import settings

    settings.cache_clear()
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    yield provider
    settings.cache_clear()


def contract() -> bytes:
    d = Document()
    d.add_paragraph("1. Term. This Agreement continues for twelve (12) months.")
    d.add_paragraph("2. Indemnity. The Supplier indemnifies the Customer without limit.")
    d.add_paragraph("7. Notices. Any notice must be in writing.")
    d.add_paragraph("8. Governing Law. New York law governs.")
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def first_email() -> InboundEmail:
    raw = contract()
    return InboundEmail(
        message_id="first-1", thread_id="thread-1",
        from_address="jim@firm.com", to=["review@legal.firm.com"],
        subject="Supply agreement", text_body="Have a look?",
        attachments=[Attachment(filename="Supply.docx", content_type=DOCX,
                                size_bytes=len(raw), content=raw)],
        received_at=datetime.now(UTC),
    )


def reply_email(body: str, mid: str) -> InboundEmail:
    return InboundEmail(
        message_id=mid, thread_id="thread-1", in_reply_to="agent-msg-1",
        from_address="jim@firm.com", to=["review@legal.firm.com"],
        subject="Re: Supply agreement", text_body=body,
        received_at=datetime.now(UTC),
    )


def start(captured, monkeypatch):
    monkeypatch.setattr(
        handler.review, "review",
        lambda doc, mode, instructions, **k: ReviewResult(
            mode=mode, summary="A short supply agreement.", findings=[]),
    )
    handler.handle(first_email())


def fake_outcome(understood="Understood.", changes=(), concerns=(), questions=(),
                 declined=()):
    o = instruct.Outcome()
    o.understood = understood
    o.changes = list(changes)
    o.concerns = list(concerns)
    o.questions = list(questions)
    o.declined = list(declined)
    return o


def with_outcome(monkeypatch, outcome):
    monkeypatch.setattr(instruct, "carry_out", lambda *a, **k: outcome)
    monkeypatch.setattr(handler.instruct, "carry_out", lambda *a, **k: outcome)


# --- a change lands as a tracked change ----------------------------------

def test_an_instruction_produces_a_tracked_change(captured, monkeypatch):
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        understood="Changing the term from twelve months to twenty-four.",
        changes=[{"kind": "replace", "anchor": "twelve (12) months",
                  "text": "twenty-four (24) months",
                  "title": "Term extended to 24 months", "drafting": False}],
    ))
    handler.handle(reply_email("make the term 24 months", "r2"))

    out = captured.sent[1]
    assert out.attachments, "no updated document came back"
    from secondeye.pipeline import redline

    ok, why = redline.verify(out.attachments[0].content, expect_revisions=True)
    assert ok, why
    assert "twenty-four (24) months" in _xml(out.attachments[0].content)


def _xml(content: bytes) -> str:
    import zipfile

    return zipfile.ZipFile(BytesIO(content)).read("word/document.xml").decode()


def test_the_reply_restates_what_it_understood(captured, monkeypatch):
    """So the lawyer can see whether it read the instruction correctly, before
    they open anything."""
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        understood="Changing the term from twelve months to twenty-four.",
        changes=[{"kind": "replace", "anchor": "twelve (12) months",
                  "text": "twenty-four (24) months", "title": "Term", "drafting": False}],
    ))
    handler.handle(reply_email("make the term 24 months", "r2"))
    # The verdict leads, as in every reply; what it understood comes straight after.
    verdict, understood = captured.sent[1].text_body.split("\n\n")[:2]
    assert verdict == "Done."
    assert understood == "Changing the term from twelve months to twenty-four."


# --- drafting -------------------------------------------------------------

def test_a_drafted_clause_is_inserted_as_a_new_paragraph(captured, monkeypatch):
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        understood="Adding a force majeure clause after the notices clause.",
        changes=[{"kind": "insert_after",
                  "anchor": "7. Notices. Any notice must be in writing.",
                  "text": "7A. Force Majeure. Neither party is liable for a failure "
                          "to perform caused by an event beyond its reasonable control.",
                  "title": "Force majeure clause added", "drafting": True}],
    ))
    handler.handle(reply_email("add a force majeure clause after 7", "r2"))

    out = captured.sent[1]
    xml = _xml(out.attachments[0].content)
    assert "Force Majeure" in xml
    assert "7. Notices. Any notice must be in writing." in xml, "the anchor was consumed"


def test_new_language_is_marked_as_new_language(captured, monkeypatch):
    """Text the agent wrote deserves a closer read than a correction to text
    that was already there."""
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        changes=[{"kind": "insert_after",
                  "anchor": "7. Notices. Any notice must be in writing.",
                  "text": "7A. Force Majeure. Neither party is liable.",
                  "title": "Force majeure clause added", "drafting": True}],
    ))
    handler.handle(reply_email("add a force majeure clause", "r2"))
    body = captured.sent[1].text_body
    assert "new wording:" in body
    assert "  1. new wording: Force majeure clause added." in body


# --- flag, then comply ----------------------------------------------------

def test_a_concern_is_stated_and_the_change_made_anyway(captured, monkeypatch):
    """Decision 22. It says what it thinks once; the lawyer decides."""
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        understood="Removing the liability cap.",
        changes=[{"kind": "replace",
                  "anchor": "The Supplier indemnifies the Customer without limit.",
                  "text": "The Supplier indemnifies the Customer without limit or cap.",
                  "title": "Cap removed", "drafting": False}],
        concerns=[{"about": "This leaves you with an uncapped indemnity",
                   "why": "Most counterparties will not sign it."}],
    ))
    handler.handle(reply_email("take the cap off the indemnity", "r2"))

    out = captured.sent[1]
    body = out.text_body
    assert "uncapped indemnity" in body
    assert body.startswith("Done, but check this first.")
    assert out.attachments, "it flagged but did not comply"


def test_the_concern_comes_before_the_change_list(captured, monkeypatch):
    """Burying it under a list of edits would be a way of technically
    disclosing it."""
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        changes=[{"kind": "replace", "anchor": "twelve (12) months",
                  "text": "twenty-four (24) months", "title": "Term", "drafting": False}],
        concerns=[{"about": "That doubles the commitment", "why": "Worth a check."}],
    ))
    handler.handle(reply_email("make it 24 months", "r2"))
    body = captured.sent[1].text_body
    assert body.index("doubles the commitment") < body.index("Tracked in the attached copy")


# --- asking, and declining ------------------------------------------------

def test_an_ambiguous_instruction_asks_and_changes_nothing(captured, monkeypatch):
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        understood="I need to know which party you mean.",
        questions=["Which party's liability did you want capped?"],
    ))
    handler.handle(reply_email("cap the liability", "r2"))
    out = captured.sent[1]
    assert "Which party's liability" in out.text_body
    assert out.attachments == []


def test_something_out_of_scope_is_declined_in_the_reply(captured, monkeypatch):
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        understood="That is outside what I do.",
        declined=[{"what": "advise on whether to sign",
                   "why": "I review the document, not the transaction."}],
    ))
    handler.handle(reply_email("should we sign this?", "r2"))
    assert "I review the document, not the transaction" in captured.sent[1].text_body


# --- it joins the ledger --------------------------------------------------

def test_an_instruction_driven_change_can_be_undone(captured, monkeypatch):
    """The whole point of routing it through the ledger."""
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        changes=[{"kind": "replace", "anchor": "twelve (12) months",
                  "text": "twenty-four (24) months",
                  "title": "Term extended to 24 months", "drafting": False}],
    ))
    handler.handle(reply_email("make the term 24 months", "r2"))
    assert captured.sent[1].attachments

    handler.handle(reply_email("undo the term change", "r3"))
    out = captured.sent[2]
    assert "reversed" in out.text_body
    assert out.attachments == [], "the change was not actually removed"


def test_a_later_instruction_builds_on_the_earlier_one(captured, monkeypatch):
    """"tighten that further" must mean further than the change already made."""
    start(captured, monkeypatch)
    with_outcome(monkeypatch, fake_outcome(
        changes=[{"kind": "replace", "anchor": "twelve (12) months",
                  "text": "twenty-four (24) months", "title": "Term", "drafting": False}],
    ))
    handler.handle(reply_email("make it 24 months", "r2"))

    seen = {}

    def capture_doc(doc, **kwargs):
        seen["text"] = doc.as_prompt()
        return fake_outcome(understood="Noted.")

    monkeypatch.setattr(handler.instruct, "carry_out", capture_doc)
    handler.handle(reply_email("actually make it longer still", "r3"))
    assert "twenty-four" in seen["text"], (
        "the second instruction saw the original, not the current document"
    )


# --- failure --------------------------------------------------------------

def test_a_failed_instruction_says_nothing_was_altered(captured, monkeypatch):
    start(captured, monkeypatch)
    monkeypatch.setattr(
        handler.instruct, "carry_out",
        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("model down")),
    )
    handler.handle(reply_email("tighten the indemnity", "r2"))
    body = captured.sent[1].text_body
    assert "you have not lost anything" in body


# --- the unit under it ----------------------------------------------------

def test_outcome_converts_changes_into_findings():
    o = fake_outcome(changes=[
        {"kind": "replace", "anchor": "a", "text": "b", "title": "t", "drafting": False},
        {"kind": "insert_after", "anchor": "c", "text": "d", "title": "u", "drafting": True},
    ])
    findings = o.as_findings()
    assert [f.edit_kind for f in findings] == ["replace", "insert_after"]
    assert [f.category for f in findings] == ["instruction", "drafting"]
    assert all(f.auto_apply for f in findings)


def test_a_change_missing_its_anchor_or_text_is_dropped():
    o = fake_outcome(changes=[
        {"kind": "replace", "anchor": "", "text": "b", "title": "t"},
        {"kind": "replace", "anchor": "a", "text": "", "title": "t"},
    ])
    assert o.as_findings() == []


# --- derived files, produced in the container ----------------------------

def with_artefacts(monkeypatch, outcome, files):
    outcome.artefacts = files
    with_outcome(monkeypatch, outcome)


def test_a_produced_file_comes_back_as_an_attachment(captured, monkeypatch):
    """"Give me a table of every payment date" is not an edit. It is a new
    document, which is exactly what the container is for."""
    start(captured, monkeypatch)
    outcome = fake_outcome(understood="Built a schedule of the payment dates.")
    with_artefacts(monkeypatch, outcome, [("Payment schedule.xlsx", b"PK\x03\x04fake")])
    handler.handle(reply_email("give me a table of every payment date", "r2"))

    out = captured.sent[1]
    assert [a.filename for a in out.attachments] == ["Payment schedule.xlsx"]
    assert out.attachments[0].content_type.endswith("spreadsheetml.sheet")
    assert "Payment schedule.xlsx" in out.text_body


def test_a_produced_file_is_named_in_the_reply(captured, monkeypatch):
    start(captured, monkeypatch)
    outcome = fake_outcome(understood="Pulled the obligations into a schedule.")
    with_artefacts(monkeypatch, outcome, [("Obligations.docx", b"PK\x03\x04fake")])
    handler.handle(reply_email("pull the obligations into a schedule", "r2"))
    assert "I have also attached 1 file" in captured.sent[1].text_body


def test_a_file_and_a_redline_can_both_come_back(captured, monkeypatch):
    """"Cap the liability and give me a table of the caps" is both at once."""
    start(captured, monkeypatch)
    outcome = fake_outcome(
        understood="Capped the liability and built a table of the caps.",
        changes=[{"kind": "replace",
                  "anchor": "The Supplier indemnifies the Customer without limit.",
                  "text": "The Supplier indemnifies the Customer up to the Cap.",
                  "title": "Liability capped", "drafting": False}],
    )
    with_artefacts(monkeypatch, outcome, [("Caps.xlsx", b"PK\x03\x04fake")])
    handler.handle(reply_email("cap the liability and table the caps", "r2"))

    names = [a.filename for a in captured.sent[1].attachments]
    assert "Supply (redline).docx" in names
    assert "Caps.xlsx" in names


def test_the_reply_opens_by_saying_a_file_came_back_when_nothing_changed(captured, monkeypatch):
    start(captured, monkeypatch)
    outcome = fake_outcome(understood="Counted them.")
    with_artefacts(monkeypatch, outcome, [("Counts.xlsx", b"PK\x03\x04fake")])
    handler.handle(reply_email("how many times does this say reasonable efforts", "r2"))
    out = captured.sent[1]
    assert out.text_body.startswith("Done. 1 file attached.")
    # And the subject is left alone, so the reply threads under the request.
    assert " - " not in out.subject.removeprefix("Re: ").replace("Acme / Beta", "")


# --- the boundary that must not move -------------------------------------

def test_a_redline_the_associate_wrote_itself_is_never_attached(captured, monkeypatch):
    """Edits go through make_changes and the local writer, so an "(redline)"
    file left in the session outputs is not a deliverable, whatever it holds."""
    from secondeye import managed

    start(captured, monkeypatch)
    from io import BytesIO

    from openpyxl import Workbook

    book = BytesIO()
    Workbook().save(book)
    run = managed.SessionRun(session_id="s", outputs=[
        ("Supply (redline).docx", b"PK-improvised"), ("Caps.xlsx", book.getvalue())])

    def fake_session(**kw):
        kw["tools"]["make_changes"]({"understood": "Tabled the caps.", "changes": []})
        return run

    monkeypatch.setattr(managed, "run_session", fake_session)
    handler.handle(reply_email("table the caps", "r2"))
    assert [a.filename for a in captured.sent[1].attachments] == ["Caps.xlsx"]


def test_the_writer_is_still_unreachable_from_the_session():
    """Edits go through the tested writer and the ledger. Code in the session
    produces new files only, and the associate's prompt says so in terms."""
    import inspect

    from secondeye.pipeline import instruct, ooxml, redline

    for module in (ooxml, redline):
        assert "managed" not in inspect.getsource(module).lower()

    system = instruct.SYSTEM
    assert "NEW files only" in system
    assert "never edits the document under review" in system


def test_a_question_is_answered_in_the_first_line(captured, monkeypatch):
    """ "What's the notice period?" got "Nothing to change": the agent could
    only report edits."""
    start(captured, monkeypatch)
    o = fake_outcome(understood="Read clause 9.")
    o.answer = "30 days' written notice (cl. 9.2), either party."
    with_outcome(monkeypatch, o)
    handler.handle(reply_email("what's the notice period?", "r2"))
    out = captured.sent[1]
    assert out.text_body.startswith("30 days' written notice (cl. 9.2), either party.")
    assert out.attachments == []
