"""The reason for a substantive change sits in the Word margin beside it.

A lawyer decides on a tracked change in Word. If the reason for it is only in
the covering email, every substantive change means switching windows. A typo
fix explains itself, and a comment beside each one would bury the ones that
matter.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO

import pytest
from docx import Document

from secondeye import convert
from secondeye.models import Finding, Mode, ReviewResult, Severity
from secondeye.pipeline import redline
from secondeye.pipeline.ooxml import RevisionWriter, author_initials

WHY = ("A cap of fees paid in the last month leaves the client with almost no "
       "recourse for a serious breach. Twelve months is the market position. "
       "This third sentence is detail the margin does not need.")


def save(d) -> bytes:
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def doc(*paragraphs: str) -> bytes:
    d = Document()
    for text in paragraphs:
        d.add_paragraph(text)
    return save(d)


def part(content: bytes, name: str) -> str:
    package = zipfile.ZipFile(BytesIO(content))
    return package.read(name).decode() if name in package.namelist() else ""


def run(content: bytes, *findings: Finding, **kw) -> redline.RedlineOutput:
    return redline.apply(content, ReviewResult(mode=Mode.REDLINE, summary="",
                                               findings=list(findings)), **kw)


def substantive(anchor="one (1) month", new="twelve (12) months",
                why=WHY, severity=Severity.SUBSTANTIVE, category="liability",
                **kw) -> Finding:
    return Finding(severity=severity, category=category, title="Raise the cap",
                   explanation=why, anchor=anchor, suggested_text=new,
                   auto_apply=True, **kw)


def typo() -> Finding:
    return Finding(severity=Severity.FORMATTING, category="typo", title="Fix typo",
                   explanation="Reciever is misspelt.", anchor="Reciever",
                   suggested_text="Recipient", auto_apply=True)


TEXT = "The Reciever's liability is capped at fees paid in one (1) month."


def test_a_substantive_change_carries_its_reason_in_the_margin():
    out = run(doc(TEXT), substantive())
    assert out.applied and out.explained == out.applied
    assert out.commented == [], "a reason is not a question left for the lawyer"
    comments = part(out.content, "word/comments.xml")
    assert "leaves the client with almost no recourse" in comments
    assert "Twelve months is the market position." in comments
    assert "third sentence" not in comments
    ok, why = redline.verify(out.content, expect_revisions=True)
    assert ok, why


def test_the_comment_spans_the_change_itself():
    out = run(doc(TEXT), substantive())
    body = part(out.content, "word/document.xml")
    span = re.search(r"<w:commentRangeStart [^>]*/>(.*?)<w:commentRangeEnd", body, re.DOTALL)
    assert span, "no comment range in the body"
    inside = span.group(1)
    assert "<w:ins " in inside and "<w:del " in inside
    assert "twelve (12) months" in inside and "one (1) month" in inside
    assert "<w:commentReference " in body


def test_a_typo_fix_gets_no_comment():
    out = run(doc(TEXT), typo())
    assert out.applied and out.explained == []
    assert "<w:commentRangeStart" not in part(out.content, "word/document.xml")


def test_only_the_substantive_change_is_explained_when_both_land():
    out = run(doc(TEXT), typo(), substantive())
    assert len(out.applied) == 2
    assert [f.title for f in out.explained] == ["Raise the cap"]
    assert part(out.content, "word/document.xml").count("<w:commentRangeStart") == 1


@pytest.mark.parametrize("finding", [
    substantive(why=""),
    substantive(severity=Severity.STYLE),
    substantive(category="punctuation"),
    # A defined term's capitalisation, even marked substantive.
    substantive(anchor="Reciever's", new="RECIEVER'S", category="defined-term"),
], ids=["no-explanation", "style", "mechanical-category", "case-only"])
def test_changes_that_explain_themselves_get_no_comment(finding):
    out = run(doc(TEXT), finding)
    assert out.applied and out.explained == []


def test_a_blocker_is_explained():
    out = run(doc(TEXT), substantive(severity=Severity.BLOCKER))
    assert len(out.explained) == 1


def test_a_drafted_clause_is_explained_on_the_new_paragraph():
    out = run(doc("7. Termination.", "8. Notices."), Finding(
        severity=Severity.STYLE, category="drafting", title="Add force majeure",
        explanation="The draft has no force majeure clause. The client asked for one.",
        anchor="7. Termination.", suggested_text="7A. Force majeure.",
        edit_kind="insert_after", auto_apply=True))
    assert len(out.explained) == 1
    body = part(out.content, "word/document.xml")
    new = next(p for p in re.findall(r"<w:p[ >].*?</w:p>", body, re.DOTALL)
               if "7A. Force majeure." in p)
    assert "<w:commentRangeStart" in new and "<w:commentReference" in new
    assert "no force majeure clause" in part(out.content, "word/comments.xml")
    ok, why = redline.verify(out.content, expect_revisions=True)
    assert ok, why


def test_questions_still_become_comments_beside_explained_changes():
    out = run(doc(TEXT, "Payment is due in thirty (13) days."), substantive(),
              Finding(severity=Severity.BLOCKER, category="amount", title="mismatch",
                      explanation="", anchor="thirty (13)", question="30 or 13?"))
    assert len(out.explained) == 1 and len(out.commented) == 1
    comments = part(out.content, "word/comments.xml")
    assert "30 or 13?" in comments and "market position" in comments
    ids = re.findall(r'<w:comment [^>]*w:id="(\d+)"', comments)
    assert len(ids) == len(set(ids)) == 2
    ok, why = redline.verify(out.content, expect_revisions=True)
    assert ok, why


def test_margin_note_is_short_and_ends_on_a_word():
    long = "This clause " + "really " * 60 + "matters. Second."
    note = redline.margin_note(long)
    assert len(note) <= 200 and note.endswith("…")
    assert redline.margin_note("One. Two. Three.") == "One. Two."
    assert redline.margin_note("  Spread\n over   lines. ") == "Spread over lines."


# --- initials follow the configured author --------------------------------

@pytest.mark.parametrize("author, initials", [
    ("Reviewer", "R"), ("Jane Associate", "JA"), ("jane q. associate", "JQA"),
    ("", ""),
])
def test_initials_come_from_the_author(author, initials):
    assert author_initials(author) == initials


def test_comment_initials_follow_the_author_not_the_tool():
    out = run(doc(TEXT, "Payment is due in thirty (13) days."), substantive(),
              Finding(severity=Severity.BLOCKER, category="amount", title="q",
                      explanation="", anchor="thirty (13)", question="30 or 13?"),
              author="Jane Associate")
    comments = part(out.content, "word/comments.xml")
    assert re.findall(r'w:initials="([^"]*)"', comments) == ["JA", "JA"]
    assert "LRA" not in comments
    assert 'w:author="Jane Associate"' in comments


def test_the_default_author_signs_r():
    out = run(doc(TEXT), substantive())
    assert 'w:initials="R"' in part(out.content, "word/comments.xml")


def test_writer_comment_takes_explicit_initials():
    d = Document(BytesIO(doc(TEXT)))
    RevisionWriter(d, "Jane Associate").comment("one (1) month", "Why?", initials="J")
    assert 'w:initials="J"' in part(save(d), "word/comments.xml")


def test_a_header_change_is_made_without_a_comment():
    """Comments live in the body's story; a header edit still lands."""
    d = Document()
    d.add_paragraph("Body text.")
    d.sections[0].header.paragraphs[0].text = "DRAFT liability one (1) month"
    out = run(save(d), substantive())
    assert out.applied and out.explained == []
    ok, why = redline.verify(out.content, expect_revisions=True)
    assert ok, why


@pytest.mark.skipif(not convert.available(), reason="LibreOffice is not installed")
def test_an_explained_redline_opens_in_libreoffice():
    out = run(doc(TEXT), typo(), substantive())
    ok, why = convert.opens_in_libreoffice(out.content)
    assert ok, why
