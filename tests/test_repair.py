"""Formatting repair: the session is not trusted, the gate is.

Nothing here runs a session. Each test plays the part of one, leaving the kind
of file a model might plausibly hand back in the outputs, and checks what
reaches the lawyer.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.shared import Pt

from lra import managed
from lra.pipeline import repair
from lra.pipeline.ooxml import Revision, RevisionWriter

PARAGRAPHS = [
    "1. Term. This Agreement continues for twelve months.",
    "2. Payment. The Buyer shall pay as set out in Section 1.",
]


def document(paragraphs=PARAGRAPHS, size=None) -> bytes:
    d = Document()
    for text in paragraphs:
        run = d.add_paragraph().add_run(text)
        if size:
            run.font.size = Pt(size)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def container_returns(monkeypatch, content: bytes | None, text="- Made body text 11pt"):
    files = [("Agreement (repaired).docx", content)] if content else []
    monkeypatch.setattr(managed, "run_job",
                        lambda *a, **k: managed.SessionRun(
                            session_id="sesn", messages=["Let me unzip it first.", text],
                            outputs=files))


def test_a_faithful_repair_is_returned_with_what_changed(monkeypatch):
    container_returns(monkeypatch, document(size=11))
    out = repair.repair(document(size=9), "Agreement.docx")
    assert out.filename == "Agreement (repaired).docx"
    assert out.changes == ["Made body text 11pt"]


def test_one_changed_word_stops_it(monkeypatch):
    tampered = [PARAGRAPHS[0].replace("twelve", "twenty"), PARAGRAPHS[1]]
    container_returns(monkeypatch, document(tampered))
    with pytest.raises(repair.RepairRejected, match="wording"):
        repair.repair(document(), "Agreement.docx")


def test_reordered_paragraphs_stop_it(monkeypatch):
    container_returns(monkeypatch, document(list(reversed(PARAGRAPHS))))
    with pytest.raises(repair.RepairRejected, match="wording"):
        repair.repair(document(), "Agreement.docx")


def test_quietly_accepting_a_tracked_change_stops_it(monkeypatch):
    d = Document(BytesIO(document()))
    RevisionWriter(d, "Opposing Counsel").apply(
        Revision("twelve months", "six months", "Opposing Counsel"))
    buf = BytesIO()
    d.save(buf)
    accepted = [PARAGRAPHS[0].replace("twelve months", "six months"), PARAGRAPHS[1]]
    container_returns(monkeypatch, document(accepted))
    with pytest.raises(repair.RepairRejected, match="tracked changes"):
        repair.repair(buf.getvalue(), "Agreement.docx")


def test_a_file_that_is_not_word_stops_it(monkeypatch):
    container_returns(monkeypatch, b"PK\x03\x04 not really")
    with pytest.raises(repair.RepairRejected, match="checks every attachment"):
        repair.repair(document(), "Agreement.docx")


def test_nothing_produced_is_a_sentence(monkeypatch):
    container_returns(monkeypatch, None)
    with pytest.raises(repair.RepairRejected, match="did not produce"):
        repair.repair(document(), "Agreement.docx")


def test_the_container_being_down_is_a_sentence(monkeypatch):
    def down(*a, **k):
        raise managed.JobFailed("off")

    monkeypatch.setattr(managed, "run_job", down)
    with pytest.raises(repair.RepairRejected, match="Nothing was changed"):
        repair.repair(document(), "Agreement.docx")
