"""The Cloudflare deployment: D1, R2, encryption, the Workflow's entry, the mail path.

`conftest.py` can run the whole suite in D1 mode. These tests force it for
themselves, because they are about things that only exist in that mode: where a
document physically ends up, what an operator with the bucket but not the key
can read, and what happens when a container dies half way through a review.
"""

from __future__ import annotations

import re
import sqlite3
from datetime import UTC, datetime, timedelta
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from secondeye import archive, crypto, d1, edge, memory, oauth, store, thread
from secondeye.models import Attachment, InboundEmail, OutboundEmail
from tests.fake_edge import FakeEdge

ROOT = Path(__file__).resolve().parents[1]
KEY = "MDEyMzQ1Njc4OWFiY2RlZjAxMjM0NTY3ODlhYmNkZWY="        # 32 bytes, base64
OTHER_KEY = "ZmVkY2JhOTg3NjU0MzIxMGZlZGNiYTk4NzY1NDMyMTA="


@pytest.fixture
def cloud(monkeypatch, tmp_path):
    """D1 and R2, played by the fake Worker, with a data key set."""
    from secondeye.config import settings

    monkeypatch.setenv("DATA_KEY", KEY)
    monkeypatch.setenv("EDGE_URL", "https://edge.test")
    monkeypatch.setenv("EDGE_SECRET", "s3cret")
    monkeypatch.setenv("MAIL_PROVIDER", "cloudflare")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    settings.cache_clear()
    fake = FakeEdge(str(tmp_path / "d1.sqlite3"))
    monkeypatch.setattr(store, "is_d1", lambda: True)
    monkeypatch.setattr(edge, "call", fake.call)
    d1._applied_scripts.clear()
    thread._migrated.clear()
    yield fake
    d1._applied_scripts.clear()
    thread._migrated.clear()
    settings.cache_clear()


def docx_bytes(text="1. The term is thirty days.") -> bytes:
    d = Document()
    d.add_paragraph(text)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


# --- encryption ------------------------------------------------------------


def test_what_is_sealed_opens_and_is_not_readable_without_the_key(cloud, monkeypatch):
    from secondeye.config import settings

    sealed = crypto.seal(b"privileged and confidential")
    assert sealed.startswith(b"enc:v1:") and b"privileged" not in sealed
    assert crypto.unseal(sealed) == b"privileged and confidential"
    # Same plaintext, different ciphertext: a nonce per message.
    assert crypto.seal(b"privileged and confidential") != sealed

    monkeypatch.setenv("DATA_KEY", OTHER_KEY)
    settings.cache_clear()
    with pytest.raises(crypto.KeyMissing):
        crypto.unseal(sealed)


def test_a_rotated_key_still_reads_what_the_old_one_sealed(cloud, monkeypatch):
    from secondeye.config import settings

    old = crypto.seal(b"sealed last quarter")
    monkeypatch.setenv("DATA_KEY", OTHER_KEY)
    monkeypatch.setenv("DATA_KEY_PREVIOUS", KEY)
    settings.cache_clear()
    assert crypto.unseal(old) == b"sealed last quarter"
    # New writes use the new key only.
    monkeypatch.setenv("DATA_KEY_PREVIOUS", "")
    settings.cache_clear()
    assert crypto.unseal(crypto.seal(b"new")) == b"new"


def test_with_a_key_set_plaintext_is_refused_not_read(cloud):
    """Anyone who can write the bucket or the table, but does not hold the
    key, could otherwise plant a document or a token and have it read as ours."""
    with pytest.raises(crypto.NotSealed):
        crypto.unseal(b"PK\x03\x04 a planted docx")
    with pytest.raises(crypto.NotSealed):
        crypto.unseal_text("planted-token")


def test_without_a_key_plaintext_is_read_as_it_is(monkeypatch):
    from secondeye.config import settings

    monkeypatch.setenv("DATA_KEY", "")
    monkeypatch.setenv("DATA_KEY_PREVIOUS", "")
    settings.cache_clear()
    try:
        assert crypto.unseal(b"PK\x03\x04 laptop docx") == b"PK\x03\x04 laptop docx"
        assert crypto.unseal_text("plain-token") == "plain-token"
    finally:
        settings.cache_clear()


def test_a_malformed_key_is_refused_loudly(monkeypatch):
    from secondeye.config import settings

    monkeypatch.setenv("DATA_KEY", "dG9vLXNob3J0")  # valid base64, nine bytes
    settings.cache_clear()
    try:
        with pytest.raises(crypto.KeyMissing, match="32 bytes"):
            crypto.seal(b"x")
    finally:
        settings.cache_clear()


def test_the_service_will_not_start_on_d1_without_a_key(cloud, monkeypatch):
    from secondeye import main
    from secondeye.config import settings

    monkeypatch.setenv("DATA_KEY", "")
    settings.cache_clear()
    with pytest.raises(crypto.KeyMissing, match="DATA_KEY"):
        main._refuse_to_store_unencrypted()


# --- where documents end up ------------------------------------------------


def test_a_document_goes_to_object_storage_sealed_and_the_row_holds_a_reference(cloud):
    original = docx_bytes()
    thread.start("t1", "jim@firm.com", "NDA.docx", original, None)

    [(key, stored)] = cloud.blobs.items()
    assert key.startswith("thread/")
    assert stored.startswith(b"enc:v1:") and b"word/document.xml" not in stored

    raw = sqlite3.connect(cloud.path).execute("SELECT original FROM threads").fetchone()[0]
    assert bytes(raw).startswith(b"blobref:v1:thread/")
    assert thread.load("t1").original == original


def test_a_newer_version_replaces_the_stored_document_rather_than_orphaning_it(cloud):
    thread.start("t1", "jim@firm.com", "NDA.docx", docx_bytes("one"), None)
    thread.start("t1", "jim@firm.com", "NDA.docx", docx_bytes("two"), None)
    # The first is kept, once, as the conversation's earlier version (for
    # "blackline against v1"), and goes with the conversation.
    assert len(cloud.blobs) == 2
    thread.purge(owner="jim@firm.com")
    assert len(cloud.blobs) == 0


def test_retention_deletes_the_document_not_just_the_row(cloud, monkeypatch):
    from secondeye.config import settings

    thread.start("old", "jim@firm.com", "Old.docx", docx_bytes(), None)
    thread.start("new", "jim@firm.com", "New.docx", docx_bytes(), None)
    long_ago = (datetime.now(UTC) - timedelta(days=45)).isoformat()
    with store.connect() as c:
        c.execute("UPDATE threads SET updated_at = ? WHERE id = 'old'", (long_ago,))

    monkeypatch.setenv("THREAD_RETENTION_DAYS", "30")
    settings.cache_clear()
    assert thread.purge() == 1
    assert thread.load("old") is None and thread.load("new") is not None
    assert len(cloud.blobs) == 1


def test_a_lawyers_conversations_can_all_be_deleted_on_request(cloud):
    thread.start("a", "jim@firm.com", "A.docx", docx_bytes(), None)
    thread.start("b", "jane@firm.com", "B.docx", docx_bytes(), None)
    assert thread.purge(owner="Jim@Firm.com") == 1
    assert thread.load("a") is None and thread.load("b") is not None
    assert len(cloud.blobs) == 1


def test_archived_attachments_follow_the_same_rules(cloud, monkeypatch):
    from secondeye.config import settings

    monkeypatch.setenv("ARCHIVE_ENABLED", "true")
    settings.cache_clear()
    content = docx_bytes()
    email = InboundEmail(
        message_id="<a1@firm.com>", from_address="jim@firm.com",
        to=["review@legal.firm.com"], subject="Supply agreement",
        attachments=[Attachment(filename="Supply.docx", content_type="x",
                                size_bytes=len(content), content=content)],
        received_at=datetime.now(UTC),
    )
    archive.store(email, owner="jim@firm.com", direction="received",
                  attachment_text={"Supply.docx": "the term is thirty days"})

    assert [k.split("/")[0] for k in cloud.blobs] == ["archive"]
    found, _ = archive.latest_version("jim@firm.com", "Supply.docx")
    assert found == content
    # Full-text search works on D1: it is SQLite, FTS5 included.
    assert archive.search("jim@firm.com", "thirty")

    archive.purge(owner="jim@firm.com")
    assert cloud.blobs == {}


def test_document_system_tokens_are_not_stored_in_the_clear(cloud):
    oauth.save(oauth.Token(user_address="jim@firm.com", system="imanage",
                           access_token="access-abc", refresh_token="refresh-xyz",
                           scopes=["read"], expires_at=None))
    row = sqlite3.connect(cloud.path).execute(
        "SELECT access_token, refresh_token FROM oauth_tokens").fetchone()
    assert "access-abc" not in row[0] and "refresh-xyz" not in row[1]
    token = oauth.get("jim@firm.com", "imanage")
    assert (token.access_token, token.refresh_token) == ("access-abc", "refresh-xyz")


# --- a review that dies half way -------------------------------------------


def test_a_second_claim_on_a_message_never_takes_it_over_however_old(cloud):
    """No lease: a Workflow retries its step under the same job id
    (claim_for), and nothing else redelivers a message."""
    assert store.claim("job-1", "<m1@firm.com>", "jim@firm.com")
    assert not store.claim("job-2", "<m1@firm.com>", "jim@firm.com")
    ancient = (datetime.now(UTC) - timedelta(days=2)).isoformat()
    with store.connect() as c:
        c.execute("UPDATE jobs SET status = 'reviewing', updated_at = ?", (ancient,))
    assert not store.claim("job-3", "<m1@firm.com>", "jim@firm.com")
    assert store.claim_for("job-1", "<m1@firm.com>", "jim@firm.com")
    assert not hasattr(store, "LEASE_MINUTES") and not hasattr(store, "touch")


def test_a_finished_job_is_never_taken_over_however_old(cloud):
    store.claim("job-1", "<m1@firm.com>", "jim@firm.com")
    store.record("job-1", "<m1@firm.com>", "replied", "jim@firm.com", {})
    ancient = (datetime.now(UTC) - timedelta(days=2)).isoformat()
    with store.connect() as c:
        c.execute("UPDATE jobs SET updated_at = ?", (ancient,))
    assert not store.claim("job-2", "<m1@firm.com>", "jim@firm.com")


# --- threading without knowing our own Message-ID --------------------------


def test_a_reply_finds_its_thread_through_the_references_chain(cloud):
    state = thread.start("t1", "jim@firm.com", "NDA.docx", docx_bytes(), None)
    thread.add_alias(state.id, "<first@firm.com>")
    found = thread.resolve(
        "jim@firm.com", None, "<third@firm.com>",
        "<an-id-cloudflare-chose-that-we-never-saw@cf>",
        ["<first@firm.com>", "<an-id-cloudflare-chose-that-we-never-saw@cf>"],
    )
    assert found == "t1"
    # The owner check still holds: someone else quoting the chain gets nothing.
    assert thread.resolve("eve@firm.com", None, "<x@firm.com>", None,
                          ["<first@firm.com>"]) != "t1"


# --- the mail path ----------------------------------------------------------


def test_the_provider_only_accepts_our_own_worker(cloud):
    from secondeye.mail import get_provider

    provider = get_provider()
    assert provider.name == "cloudflare"
    assert provider.verify({"authorization": "Bearer s3cret"}, b"")
    assert not provider.verify({"authorization": "Bearer wrong"}, b"")
    assert not provider.verify({}, b"")


def test_a_reply_is_sent_in_the_shape_the_email_service_takes(cloud):
    from secondeye.mail import get_provider

    sent_id = get_provider().send(OutboundEmail(
        to=["jim@firm.com"], subject="Re: NDA - Looks good", text_body="Fine.",
        html_body="<p>Fine.</p>", in_reply_to="<first@firm.com>",
        attachments=[Attachment(filename="NDA (redline).docx", content_type="application/x",
                                size_bytes=3, content=b"abc")],
    ))
    assert sent_id == "cf-1"
    [out] = cloud.sent
    assert out["from"] == {"email": "review@legal.firm.com", "name": "Second Eye"}
    assert out["headers"] == {"In-Reply-To": "<first@firm.com>", "References": "<first@firm.com>"}
    assert out["attachments"] == [
        {"filename": "NDA (redline).docx", "type": "application/x", "content": "YWJj"}]


def test_a_reply_too_big_for_cloudflare_fails_before_it_is_sent(cloud):
    from secondeye.mail import get_provider
    from secondeye.mail.cloudflare import MessageTooLarge

    big = Attachment(filename="big.docx", content_type="x", size_bytes=1,
                     content=b"0" * (5 * 1024 * 1024))
    with pytest.raises(MessageTooLarge) as refused:
        get_provider().send(OutboundEmail(to=["jim@firm.com"], subject="s",
                                          text_body="t", attachments=[big]))
    assert cloud.sent == []
    # What a caller needs to resend without the file, and know it will fit.
    assert refused.value.fits_without_attachments
    assert refused.value.size > refused.value.limit


def test_a_body_too_big_on_its_own_says_so(cloud):
    from secondeye.mail import get_provider
    from secondeye.mail.cloudflare import MessageTooLarge

    with pytest.raises(MessageTooLarge) as refused:
        get_provider().send(OutboundEmail(to=["jim@firm.com"], subject="s",
                                          text_body="\u00e9" * (3 * 1024 * 1024)))
    # Measured in bytes: three million two-byte characters is over five MiB.
    assert not refused.value.fits_without_attachments


# --- a message from the Worker, end to end ---------------------------------


def raw_message(body="Could you send me a clean copy of this?") -> bytes:
    msg = EmailMessage()
    msg["From"] = "Jim Baker <jim@firm.com>"
    msg["To"] = "review@legal.firm.com"
    msg["Subject"] = "NDA"
    msg["Message-ID"] = "<queued-1@firm.com>"
    msg.set_content(body)
    d = Document()
    d.add_paragraph("1. The term is thirty days.")
    d.core_properties.author = "Jane Associate"
    buf = BytesIO()
    d.save(buf)
    msg.add_attachment(buf.getvalue(), maintype="application", subtype="octet-stream",
                       filename="NDA.docx")
    return msg.as_bytes()


def prepare(client, sealed: bytes, job: str = "rf-cloudflare-1", **headers):
    return client.post("/prepare", content=sealed,
                       headers={"x-job-id": job, **headers})


def test_a_sealed_message_from_the_worker_becomes_a_sent_reply(cloud):
    from fastapi.testclient import TestClient

    from secondeye import main

    client = TestClient(main.app)
    sealed = crypto.seal(raw_message())          # what the Worker leaves in R2
    auth = {"authorization": "Bearer s3cret"}

    assert prepare(client, sealed).status_code == 401
    assert prepare(client, sealed, **auth).json() == {"status": "done"}

    [reply] = cloud.sent
    assert reply["to"] == ["jim@firm.com"]
    assert reply["subject"] == "Re: NDA"
    assert reply["text"].startswith("Clean copy attached")
    assert reply["attachments"][0]["filename"] == "NDA (clean).docx"
    # The job is recorded, in D1, so another instance for the same message
    # (a second delivery the Worker did not catch) sends nothing.
    assert prepare(client, sealed, job="rf-cloudflare-2", **auth).json() == {
        "status": "duplicate"}
    assert len(cloud.sent) == 1


def test_a_message_that_cannot_be_parsed_is_acknowledged_not_retried(cloud):
    from fastapi.testclient import TestClient

    from secondeye import main

    response = prepare(TestClient(main.app), crypto.seal(b"Subject: no sender\r\n\r\nhello"),
                       authorization="Bearer s3cret")
    assert response.status_code == 200 and response.json() == {"status": "unparseable"}


def test_a_message_no_key_opens_is_refused_loudly_so_the_step_retries(cloud, monkeypatch,
                                                                       caplog):
    """A key problem is not a bad message. Acknowledging it dropped a lawyer's
    document without a word; failing it lets the Workflow retry the step, then
    take its failure path, where the sender is told."""
    from fastapi.testclient import TestClient

    from secondeye import main
    from secondeye.config import settings

    sealed = crypto.seal(raw_message())
    monkeypatch.setenv("DATA_KEY", OTHER_KEY)
    settings.cache_clear()
    client = TestClient(main.app, raise_server_exceptions=False)
    with caplog.at_level("ERROR"):
        response = prepare(client, sealed, authorization="Bearer s3cret")
    assert response.status_code == 500
    assert "could not unseal" in caplog.text
    assert cloud.sent == []

    # And a plaintext message where only sealed ones belong, the same.
    response = prepare(client, raw_message(), job="rf-cloudflare-3",
                       authorization="Bearer s3cret")
    assert response.status_code == 500


# --- the adapter's own edges ------------------------------------------------


def test_every_schema_in_the_codebase_splits_into_statements_sqlite_accepts():
    scripts = [store.SCHEMA, thread.THREAD_SCHEMA, memory.MEMORY_SCHEMA,
               archive.ARCHIVE_SCHEMA, oauth.TOKEN_SCHEMA]
    conn = sqlite3.connect(":memory:")
    for script in scripts:
        statements = d1._split(script)
        assert statements and all(";" not in s for s in statements)
        for statement in statements:
            conn.execute(statement)


def test_schema_setup_costs_one_round_trip_per_process_not_one_per_call(cloud):
    for _ in range(5):
        thread.init()
    first = cloud.db_requests
    for _ in range(5):
        thread.init()
    assert cloud.db_requests == first


def test_a_database_failure_is_the_error_type_existing_handlers_catch(cloud, monkeypatch):
    def down(*a, **k):
        raise edge.EdgeError("worker unreachable")

    monkeypatch.setattr(edge, "call", down)
    with pytest.raises(sqlite3.OperationalError):
        store.seen("<m@firm.com>")


def worker_source() -> str:
    """Every module of the Worker, as one text."""
    return "\n".join(f.read_text() for f in sorted((ROOT / "cloudflare" / "src").glob("*.ts")))


def test_the_worker_and_the_application_agree_on_the_wire():
    """A static contract. The fake Worker proves the Python half; this pins the
    names the TypeScript half has to keep for that proof to mean anything."""
    worker = worker_source()
    for name in ('"/internal/db"', '"/internal/send"', '"/internal/blob/"', "statements",
                 "last_row_id", "changes", "columns", "rows", "$bytes", "messageId",
                 "new Request(`http://container${path}`", "enc:v1:"):
        assert name in worker, name
    # Same key shape on both sides, or the Worker rejects what blobs.stash and
    # jobstate.put write.
    assert r"^(thread|archive|doc|job)\/[0-9a-f]{32}$" in worker


# --- the spending guards -----------------------------------------------------


def test_the_deployed_config_cannot_run_up_a_bill_by_accident():
    """Cloudflare has no spending cap, so the caps are configuration, and
    configuration drifts. Loosening any of these should be a decision somebody
    makes in a diff, with this test to change alongside it."""
    import json

    # The shared config, and every deployment's rendered from it (the
    # primary one, Jim's, is private and only in the development repository).
    for path in [ROOT / "cloudflare" / "wrangler.jsonc",
                 *sorted((ROOT / "deployments").glob("*/wrangler.jsonc"))]:
        raw = path.read_text()
        config = json.loads(re.sub(r"^\s*//.*$", "", raw, flags=re.MULTILINE))

        [container] = config["containers"]
        assert container["max_instances"] == 1, path
        assert 0 < int(config["vars"]["MAX_EMAILS_PER_DAY"]) <= 100, path
        # An empty allowlist admits nobody at the edge, but say who it is for.
        assert config["vars"]["ALLOWED_SENDERS"].strip(), path
        # Nothing redelivers a message behind the Workflow's back.
        assert "queues" not in config, path

    worker = (ROOT / "cloudflare" / "src" / "index.ts").read_text()
    # Both checks run before anything is stored for review or woken.
    # (While paused, by the var or the D1 switch, an admitted message is held before the cap is
    # counted: holding wakes nothing and costs nothing, and the cap's refusal
    # is a bounce, which the switch promises not to send.)
    handler = worker[worker.index("async email("):worker.index("async fetch(")]
    assert handler.index("isAllowed(") < handler.index("countToday(") < handler.rindex("env.DOCS.put(")
    assert handler.index("unauthenticated(") < handler.index("loadPause(env)")
    assert handler.index("unauthenticated(") < handler.index("countToday(")


