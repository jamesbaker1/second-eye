"""The negotiation ledger and the covering-note lie detector (skills/lra-negotiation).

The Project Falcon demo: we sent seven points on their SPA; their v2 comes
back with "We've accepted all your points except the liability cap." It
accepts five, keeps the cap, keeps New York law without a word, deletes the
fraud carve-out in 9.4 and moves the long-stop date. The reply leads with
"Not what they said."

The session is scripted (tests/fake_sessions.py), but what it does in the
sandbox is real: it runs the skill's own by_clause.py and
validate_negotiation.py, as the agent would, on the evidence file the host
mounted. The documents, the comparison and the ledger are real throughout.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

from lra import handler, managed, memory, negotiation, skillsync
from lra.config import settings
from lra.mail.console import ConsoleProvider
from lra.models import Attachment, Finding, InboundEmail, Mode, ReviewResult, Severity
from lra.pipeline import compare, redline
from lra.pipeline.identity import EntryMode
from tests import fake_sessions as fs
from tests.conftest import documents

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SKILL = skillsync.SKILLS / "lra-negotiation"

# --- Project Falcon -----------------------------------------------------------

V1 = [
    "Share Purchase Agreement: Project Falcon",
    "3.2 Deposit. The Buyer shall pay a deposit of 5% of the Price, which is non-refundable.",
    "4.1 Price adjustment. The Price is not subject to any adjustment.",
    "7.1 Warranty claims. The Buyer may bring a Warranty claim within 12 months of Completion.",
    "8.2 Tax indemnity. The Seller gives no tax indemnity.",
    "9.3 De minimis. No claim may be brought for less than £50,000.",
    "9.4 Fraud. Nothing in this clause 9 limits the liability of any party for fraud.",
    "10.1 Liability cap. The Seller's total liability is capped at 100% of the Price.",
    "11.3 Non-compete. The Seller shall not compete with the Business for 6 months.",
    ("12.2 Long-stop date. If Completion has not occurred by 31 March 2027, either party "
     "may terminate this Agreement."),
    "14 Governing law. This Agreement is governed by the laws of the State of New York.",
]

V2 = [
    "Share Purchase Agreement: Project Falcon",
    ("3.2 Deposit. The Buyer shall pay a deposit of 5% of the Price, which is refundable "
     "if Completion does not occur."),
    "4.1 Price adjustment. The Price is subject to a completion accounts adjustment.",
    "7.1 Warranty claims. The Buyer may bring a Warranty claim within 24 months of Completion.",
    "8.2 Tax indemnity. The Seller indemnifies the Buyer for pre-Completion tax.",
    "9.3 De minimis. No claim may be brought for less than £50,000.",
    "10.1 Liability cap. The Seller's total liability is capped at 100% of the Price.",
    "11.3 Non-compete. The Seller shall not compete with the Business for 24 months.",
    ("12.2 Long-stop date. If Completion has not occurred by 31 January 2027, either party "
     "may terminate this Agreement."),
    "14 Governing law. This Agreement is governed by the laws of the State of New York.",
]

# Our seven points on v1, as the round-one review put them on the issues list.
POINTS = [
    ("3.2", "Deposit", "The deposit is refundable if Completion does not occur.",
     "non-refundable"),
    ("4.1", "Price adjustment", "A completion accounts adjustment.",
     "not subject to any adjustment"),
    ("7.1", "Warranty claims", "24 months to bring a Warranty claim.", "within 12 months"),
    ("8.2", "Tax indemnity", "A full pre-Completion tax indemnity.", "gives no tax indemnity"),
    ("10.1", "Liability cap", "A cap at 25% of the Price.", "capped at 100% of the Price"),
    ("11.3", "Non-compete", "A 24-month non-compete.", "for 6 months"),
    ("14", "Governing law", "English law and the English courts.", "State of New York"),
]


def point_findings() -> list[Finding]:
    return [Finding(severity=Severity.SUBSTANTIVE, category="playbook",
                    title=f"Off playbook: {topic} - as drafted", where=f"clause {number}",
                    explanation="Starter playbook (not yet your firm's): off position.",
                    anchor=anchor, our_position=position,
                    response=f"Ask for: {position}")
            for number, topic, position, anchor in POINTS]


COVER = "We've accepted all your points except the liability cap."


def forward(body_above: str, cover: str, filename: str, content: bytes,
            message_id: str, subject: str = "Fwd: Project Falcon SPA") -> InboundEmail:
    body = f"""{body_above}

---------- Forwarded message ---------
From: Dana Other <dana@otherside.com>
Date: Mon, 28 Sep 2026 at 10:02
Subject: Project Falcon SPA
To: Jim <jim@firm.com>

Jim,

{cover}

Best,
Dana

On Fri, 25 Sep 2026 at 17:40, Jim Baker <jim@firm.com> wrote:
> Dana, our issues list is attached. We have also accepted your cap.
"""
    return InboundEmail(
        message_id=message_id, from_address="jim@firm.com", to=["review@firm.com"],
        subject=subject, text_body=body,
        attachments=[Attachment(filename=filename, content_type=DOCX,
                                size_bytes=len(content), content=content)],
        received_at=datetime.now(UTC))


def docx(paragraphs: list[str]) -> bytes:
    d = Document()
    for text in paragraphs:
        d.add_paragraph(text)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return f"<sent-{len(self.sent)}@firm.com>"


@pytest.fixture(autouse=True)
def firm(monkeypatch, tmp_path):
    fs.configure(monkeypatch, DATABASE_URL=f"sqlite:///{tmp_path}/n.sqlite3",
                 FIRM_DOMAINS="firm.com", MAIL_AGENT_ADDRESS="review@firm.com",
                 SANDBOX_NEGOTIATION_SKILL_ID="skill_negotiation")
    monkeypatch.setattr(memory, "stores_for", lambda *a, **k: {})
    # The comparison's risk notes are a model call; nothing here reaches the API.
    monkeypatch.setattr(compare, "_assess", lambda *a, **k: None)
    monkeypatch.setattr(handler, "SLOW_NOTICE_AFTER", 0)
    now = [0.0]
    monkeypatch.setattr(managed.time, "monotonic", lambda: now[0])
    monkeypatch.setattr(managed.time, "sleep", lambda s: now.__setitem__(0, now[0] + s))
    negotiation._migrated.clear()
    yield
    settings.cache_clear()


@pytest.fixture
def captured(monkeypatch):
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    return provider


def round_one(monkeypatch, captured):
    """Their v1 arrives; the review puts seven points on the issues list."""
    real = handler.review.review

    def _review(doc, mode, instructions, **kwargs):
        assert not kwargs.get("negotiation_block"), "nothing to read against yet"
        return ReviewResult(mode=mode, summary="", findings=point_findings(),
                            their_paper=kwargs.get("their_paper", False))

    monkeypatch.setattr(handler.review, "review", _review)
    handler.handle(forward("Their draft. Issues list please.", "Please find our draft attached.",
                           "Project Falcon SPA v1.docx", docx(V1), "<falcon-1@firm.com>"))
    monkeypatch.setattr(handler.review, "review", real)
    return captured.sent[-1]


# --- the agent, in the sandbox ------------------------------------------------

LABELS = {"10.1": "the cap", "14": "governing law"}
UNASKED = {"9.4": "deletes the fraud carve-out in 9.4",
           "12.2": "moves the long-stop date from 31 March to 31 January"}


def run_script(script: str, *args: str) -> subprocess.CompletedProcess:
    env = {k: v for k, v in os.environ.items() if k != "PYTHONPATH"}
    return subprocess.run([sys.executable, str(SKILL / "scripts" / script), *args],
                          capture_output=True, text=True, env=env, timeout=60, check=False)


def agent_in_the_sandbox(evidence_path: Path, out: Path) -> dict:
    """What the agent does with the skill: group the evidence by clause, decide,
    and record through the validator. Returns the validator's answer."""
    grouped = json.loads(run_script("by_clause.py", "--evidence", str(evidence_path)).stdout)
    evidence = json.loads(evidence_path.read_text())
    points, cited = [], set()
    for clause in grouped["clauses"]:
        for p in clause["points"]:
            at = [c["index"] for c in clause["changes"]]
            cited.update(at)
            points.append({
                "id": p["id"], "status": "accepted" if at else "rejected",
                "label": LABELS.get(clause["clause"], p["clause"].split(": ")[-1].lower()),
                "mentioned": clause["clause"] == "10.1",
                "changes": at,
                "evidence": (f"Clause {clause['clause']} now reads as we asked." if at else
                             f"Clause {clause['clause']} is unchanged from v1."),
            })
    unrequested = [{"clause": c["where"], "summary": UNASKED[c["where"]], "risk": "high",
                    "mentioned": False, "changes": [c["index"]]}
                   for c in grouped["changes_at_clauses_we_did_not_raise"]]
    law = next(p["id"] for p in evidence["points"] if p["clause"].startswith("Clause 14"))
    draft = {"points": points, "unrequested": unrequested, "claims": [{
        "quote": COVER, "verdict": "false",
        "evidence": "Clause 14 still says New York law, which we asked to change, and v2 "
                    "also deletes 9.4 and moves the long-stop date.",
        "points": [law], "changes": [u["changes"][0] for u in unrequested]}]}
    draft_path = evidence_path.parent / "draft.json"
    draft_path.write_text(json.dumps(draft))
    done = run_script("validate_negotiation.py", str(draft_path), "--evidence",
                      str(evidence_path), "--out", str(out))
    assert done.returncode == 0, done.stdout + done.stderr
    return json.loads(done.stdout)


REVIEW_OF_V2 = {"summary": "", "findings": [{
    "severity": "substantive", "category": "playbook",
    "title": "Off playbook: Fraud - carve-out removed", "where": "clause 9",
    "explanation": "Starter playbook (not yet your firm's): fraud is never capped.",
    "anchor": "No claim may be brought for less than £50,000.",
    "our_position": "Liability for fraud is never limited.",
    "response": "Reinstate 9.4."}]}


class Negotiator:
    """Runs the sandbox half when the session is created: by then the host has
    uploaded the evidence, exactly as the agent would find it mounted."""

    def __init__(self, tmp_path: Path):
        self.tmp = tmp_path
        self.answer: dict = {}

    def outputs(self, fake) -> dict[str, bytes]:
        [(_, evidence)] = [(n, d) for n, d in fake.uploaded if n == negotiation.EVIDENCE_FILE]
        path = self.tmp / "workspace" / negotiation.EVIDENCE_FILE
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(evidence)
        out = self.tmp / "outputs" / negotiation.OUTPUT_FILE
        self.answer = agent_in_the_sandbox(path, out)
        return {negotiation.OUTPUT_FILE: out.read_bytes()}


def detached_session(tmp_path) -> tuple[fs.DetachedSessions, Negotiator]:
    agent = Negotiator(tmp_path)
    turn = fs.Turn(fs.idle("end_turn"),
                   outputs={"findings.json": json.dumps(REVIEW_OF_V2).encode()})

    class Fake(fs.DetachedSessions):
        def create(self, **kwargs):
            made = super().create(**kwargs)
            self._turn.outputs.update(agent.outputs(self))
            return made

    return Fake([turn]), agent


FALCON_REPLY = (
    "Not what they said. Besides the cap, v2 deletes the fraud carve-out in 9.4 and moves "
    "the long-stop date from 31 March to 31 January.\n"
    "5 of your 7 points were accepted; 2 were not — and one of those isn't mentioned in "
    "their email.\n"
    "\n"
    "Not accepted:\n"
    "  - Clause 10.1: Liability cap. Rejected: Clause 10.1 is unchanged from v1.\n"
    "  - Clause 14: Governing law. Rejected, not mentioned in their email: Clause 14 is "
    "unchanged from v1.\n"
    "\n"
    "Changes you didn't ask for:\n"
    "  - Clause 9.4: deletes the fraud carve-out in 9.4. High risk, not mentioned in their "
    "email.\n"
    "  - Clause 12.2: moves the long-stop date from 31 March to 31 January. High risk, not "
    "mentioned in their email.\n"
    "\n"
    "What their email says:\n"
    "  - “We've accepted all your points except the liability cap.” Not true: Clause 14 "
    "still says New York law, which we asked to change, and v2 also deletes 9.4 and moves "
    "the long-stop date.\n"
    "\n"
    "Accepted:\n"
    "  - Clause 3.2: Deposit; Clause 4.1: Price adjustment; Clause 7.1: Warranty claims; "
    "Clause 8.2: Tax indemnity; Clause 11.3: Non-compete.\n"
    "\n"
    "Project Falcon SPA v2 (issues list, updated).docx has every point with its status.\n"
)


def test_falcon_the_reply_says_their_email_is_not_what_v2_does(monkeypatch, captured,
                                                               tmp_path):
    round_one(monkeypatch, captured)
    assert len(negotiation.points(next(iter(_ledger_threads())))) == 7

    fake, agent = detached_session(tmp_path)
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(forward("Their v2 is in. Can you check it?", COVER,
                               "Project Falcon SPA v2.docx", docx(V2), "<falcon-2@firm.com>"))

    out = captured.sent[-1]
    print(out.text_body)
    assert out.text_body.startswith(FALCON_REPLY)
    assert agent.answer["points"] == 7 and agent.answer["accepted"] == 5

    # The session was told, fenced, and given the evidence as a file.
    [create] = fake.created
    opening = create["initial_events"][0]["description"]
    assert "<negotiation-" in opening and "validate_negotiation.py" in opening
    assert COVER in opening and "Dana, our issues list is attached" not in opening
    assert any(r["mount_path"] == f"/workspace/{negotiation.EVIDENCE_FILE}"
               for r in create["resources"] if r["type"] == "file")
    evidence = json.loads(dict(fake.uploaded)[negotiation.EVIDENCE_FILE])
    assert evidence["cover"] == {"from": "dana@otherside.com",
                                 "text": f"Jim,\n\n{COVER}"}  # the sign-off is not the note
    assert len(evidence["points"]) == 7
    assert {c["where"] for c in evidence["changes"]} == {
        "3.2", "4.1", "7.1", "8.2", "9.4", "11.3", "12.2"}

    # The HTML leads the same way, and the regular review follows it.
    assert out.html_body.index("Not what they said.") < out.html_body.index("Their draft")
    assert "Their draft: 1 point to push back on." in out.text_body
    assert "Since the version I reviewed" not in out.text_body, "the list above replaces it"

    names = [a.filename for a in documents(out)]
    assert "Project Falcon SPA v2 (issues list, updated).docx" in names
    assert "Project Falcon SPA v2 (changes since last time).docx" in names
    updated = next(a for a in documents(out) if "updated" in a.filename).content
    ok, why = redline.verify(updated)
    assert ok, why
    rows = table_rows(updated, 0)
    assert rows[0] == ["Clause", "What we asked", "Status", "Their version"]
    status = {r[0]: r[2] for r in rows[1:]}
    assert status["Clause 14: Governing law"] == "Rejected (not mentioned in their email)"
    assert status["Clause 10.1: Liability cap"] == "Rejected"
    assert status["Clause 3.2: Deposit"] == "Accepted"
    unasked = table_rows(updated, 1)
    assert [r[0] for r in unasked[1:]] == ["Clause 9.4", "Clause 12.2"]

    # The ledger moved with the document, carries the statuses, and holds the
    # new round's point without repeating the seven.
    [ledger] = _ledger_threads()
    held = negotiation.points(ledger)
    assert [p.status for p in held[:7]].count("accepted") == 5
    assert {p.clause: p.status for p in held}["Clause 14: Governing law"] == "rejected"
    assert len(held) == 8 and held[-1].round == 2 and held[-1].status == "open"


def test_falcon_did_they_accept_our_changes_later_on_the_thread(monkeypatch, captured,
                                                                tmp_path):
    round_one(monkeypatch, captured)

    # Asked before their v2 has come in.
    handler.handle(question("<q-1@firm.com>", in_reply_to="<falcon-1@firm.com>"))
    assert captured.sent[-1].text_body.startswith(
        "Nothing from them yet: I haven't seen a version of theirs since we raised our 7 points.")

    fake, _ = detached_session(tmp_path)
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(forward("Their v2 is in. Can you check it?", COVER,
                               "Project Falcon SPA v2.docx", docx(V2), "<falcon-2@firm.com>"))

    handler.handle(question("<q-2@firm.com>", in_reply_to="<falcon-2@firm.com>"))
    out = captured.sent[-1]
    assert out.text_body.startswith(FALCON_REPLY.split("\n\nNot accepted")[0])
    assert [a.filename for a in out.attachments] == [
        "Project Falcon SPA v2 (issues list, updated).docx"]
    assert out.in_reply_to == "<q-2@firm.com>"


def question(message_id: str, in_reply_to: str) -> InboundEmail:
    return InboundEmail(message_id=message_id, from_address="jim@firm.com",
                        to=["review@firm.com"], subject="Re: Project Falcon SPA",
                        text_body="Did they accept our changes?", in_reply_to=in_reply_to,
                        received_at=datetime.now(UTC))


def _ledger_threads() -> list[str]:
    from lra.store import connect

    with connect() as c:
        return [r[0] for r in c.execute(
            "SELECT DISTINCT thread_id FROM negotiation_points ORDER BY thread_id")]


def table_rows(content: bytes, which: int) -> list[list[str]]:
    table = Document(BytesIO(content)).tables[which]
    return [[cell.text for cell in row.cells] for row in table.rows]


# --- when the session's account does not hold together -------------------------


def test_a_file_the_host_refuses_is_left_out_and_said_so(monkeypatch, captured, tmp_path):
    round_one(monkeypatch, captured)
    bad = json.dumps({"points": [], "unrequested": [], "claims": []}).encode()
    fake = fs.DetachedSessions([fs.Turn(fs.idle("end_turn"), outputs={
        "findings.json": json.dumps(REVIEW_OF_V2).encode(), negotiation.OUTPUT_FILE: bad})])
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(forward("Their v2.", COVER, "Project Falcon SPA v2.docx", docx(V2),
                               "<falcon-2@firm.com>"))
    body = captured.sent[-1].text_body
    assert body.startswith("Their draft")
    assert "did not hold together, so I have left it out rather than guess" in body
    assert all(p.status == "open" for p in negotiation.points(_ledger_threads()[0])[:7])


def test_no_file_at_all_is_said_too(monkeypatch, captured):
    round_one(monkeypatch, captured)
    fake = fs.DetachedSessions([fs.Turn(fs.idle("end_turn"), outputs={
        "findings.json": json.dumps(REVIEW_OF_V2).encode()})])
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(forward("Their v2.", COVER, "Project Falcon SPA v2.docx", docx(V2),
                               "<falcon-2@firm.com>"))
    assert "I couldn't settle this time which of the 7 points we raised" in \
        captured.sent[-1].text_body


def test_without_the_skill_synced_their_version_gets_an_ordinary_review(monkeypatch, captured):
    monkeypatch.setenv("SANDBOX_NEGOTIATION_SKILL_ID", "")
    settings.cache_clear()
    round_one(monkeypatch, captured)
    seen: dict = {}

    def _review(doc, mode, instructions, **kwargs):
        seen.update(kwargs)
        return ReviewResult(mode=mode, summary="", findings=[], their_paper=True)

    monkeypatch.setattr(handler.review, "review", _review)
    handler.handle(forward("Their v2.", COVER, "Project Falcon SPA v2.docx", docx(V2),
                           "<falcon-2@firm.com>"))
    assert seen["negotiation_block"] == "" and seen["negotiation_evidence"] is None
    assert "what they said" not in captured.sent[-1].text_body


def test_our_own_edit_of_the_document_is_not_read_as_theirs(monkeypatch, captured):
    round_one(monkeypatch, captured)
    seen: dict = {}

    def _review(doc, mode, instructions, **kwargs):
        seen.update(kwargs)
        return ReviewResult(mode=mode, summary="", findings=[])

    monkeypatch.setattr(handler.review, "review", _review)
    mine = docx(V1[:-1] + ["14 Governing law. This Agreement is governed by English law."])
    handler.handle(InboundEmail(
        message_id="<mine@firm.com>", from_address="jim@firm.com", to=["review@firm.com"],
        subject="Falcon SPA", text_body="I tidied clause 14, quick look?",
        attachments=[Attachment(filename="Project Falcon SPA v1.docx", content_type=DOCX,
                                size_bytes=len(mine), content=mine)],
        received_at=datetime.now(UTC)))
    assert seen["negotiation_block"] == ""


# --- the skill's scripts, run as the container runs them ---------------------------


def evidence(tmp_path, cover: str = COVER) -> Path:
    path = tmp_path / "evidence.json"
    path.write_text(json.dumps({
        "ours": {"filename": "v1.docx", "how": "as we returned it"},
        "theirs": {"filename": "v2.docx", "label": "v2"},
        "points": [
            {"id": "P1", "clause": "Clause 10.1: Liability cap", "asked": "25%",
             "source": "issue", "round": 1},
            {"id": "P2", "clause": "Clause 14: Governing law", "asked": "English law",
             "source": "issue", "round": 1},
            {"id": "P3", "clause": "", "asked": "Please also extend the long-stop and drop "
             "the non-compete.", "source": "email", "round": 1},
        ],
        "changes": [
            {"index": 0, "kind": "deleted", "where": "9.4", "before": "9.4 Fraud.", "after": ""},
            {"index": 1, "kind": "replaced", "where": "12.2", "before": "31 March",
             "after": "31 January"},
        ],
        "cover": {"from": "dana@otherside.com", "text": cover},
    }))
    return path


def point(pid, status="rejected", changes=(), mentioned=False, **extra) -> dict:
    return {"id": pid, "status": status, "label": pid, "mentioned": mentioned,
            "changes": list(changes), "evidence": f"{pid} is as it was.", **extra}


def unasked(*changes) -> dict:
    return {"clause": "9.4", "summary": "deletes 9.4", "risk": "high", "mentioned": False,
            "changes": list(changes)}


def validate(tmp_path, draft: dict, cover: str = COVER) -> tuple[int, dict, Path]:
    d = tmp_path / "draft.json"
    d.write_text(json.dumps(draft))
    out = tmp_path / "outputs" / "negotiation.json"
    done = run_script("validate_negotiation.py", str(d), "--evidence",
                      str(evidence(tmp_path, cover)), "--out", str(out))
    return done.returncode, json.loads(done.stdout), out


def test_the_validator_records_a_file_that_holds(tmp_path):
    code, said, out = validate(tmp_path, {
        "points": [point("P1", mentioned=True), point("P2"),
                   point("P3.1", "changed", [1], ask="extend the long-stop"),
                   point("P3.2", ask="drop the non-compete")],
        "unrequested": [unasked(0)],
        "claims": [{"quote": "we’ve accepted all your points", "verdict": "false",
                    "evidence": "14 is unchanged.", "points": ["P2"], "changes": []}]})
    assert code == 0, said
    assert said["points"] == 4 and json.loads(out.read_text())["points"][2]["id"] == "P3.1"


def test_the_validator_refuses_what_does_not_hold_and_writes_nothing(tmp_path):
    code, said, out = validate(tmp_path, {
        "points": [point("P1", "accepted"), point("P1"), point("P9"),
                   point("P2.1", changes=[7], ask="x")],
        "unrequested": [],
        "claims": [{"quote": "We accepted everything.", "verdict": "false",
                    "evidence": "No.", "points": [], "changes": []}]})
    assert code == 1 and not out.exists()
    error = said["error"]
    assert error.startswith("NOTHING WAS RECORDED")
    for words in ("points[0] (P1) is accepted: cite the change that shows it",
                  "point P1 is answered more than once",
                  "there is no point P9 in the evidence",
                  "only a point from our covering email can be split",
                  "point P3 (no clause) is not answered",
                  "change(s) 0, 1 are not accounted for",
                  "is not word for word in their covering email",
                  "cites change 7",
                  "is false: cite the points or changes"):
        assert words in error, words
    assert "validate_negotiation.py" in error


def test_a_claim_needs_a_covering_email_and_a_shape_error_is_named(tmp_path):
    code, said, _ = validate(tmp_path, {
        "points": [point("P1"), point("P2"), point("P3", mentioned=True)],
        "unrequested": [unasked(0, 1)],
        "claims": [{"quote": "x", "verdict": "maybe", "evidence": "y", "points": [],
                    "changes": []}]}, cover="")
    assert code == 1
    assert "negotiation.claims[0].verdict is 'maybe'" in said["error"]


def test_by_clause_puts_each_point_beside_the_changes_at_its_clause(tmp_path):
    path = evidence(tmp_path)
    data = json.loads(path.read_text())
    data["changes"].append({"index": 2, "kind": "replaced", "where": "under 10.1",
                            "before": "100%", "after": "90%"})
    path.write_text(json.dumps(data))
    grouped = json.loads(run_script("by_clause.py", "--evidence", str(path)).stdout)
    by = {c["clause"]: c for c in grouped["clauses"]}
    assert [c["index"] for c in by["10.1"]["changes"]] == [2]
    assert by["14"]["changes"] == []
    assert [c["where"] for c in grouped["changes_at_clauses_we_did_not_raise"]] == ["9.4", "12.2"]


def test_the_skill_builds_as_instructions_and_scripts_with_nothing_bundled(tmp_path):
    bundle = skillsync.build(tmp_path, "lra-negotiation")
    head = (bundle / "SKILL.md").read_text().split("---")[1]
    assert "name: lra-negotiation" in head
    description = next(line for line in head.splitlines() if line.startswith("description:"))
    assert 0 < len(description.split(":", 1)[1].strip()) <= 1024
    assert not (bundle / "lib").exists()
    for name in ("negotiation.schema.json", "scripts/rules.py",
                 "scripts/validate_negotiation.py", "scripts/by_clause.py"):
        assert (bundle / name).exists(), name


@pytest.mark.parametrize("manifest, attached", [
    ("review.agent.yaml", True),
    ("associate.agent.yaml", False), ("playbook.agent.yaml", False),
])
def test_the_skill_is_attached_to_the_reviewers_once_synced(monkeypatch, manifest, attached):
    monkeypatch.setenv("SANDBOX_NEGOTIATION_SKILL_ID", "skill_neg")
    settings.cache_clear()
    body = managed.agent_body(managed.load_manifest(managed.AGENTS / manifest), [])
    ref = {"type": "custom", "skill_id": "skill_neg", "version": "latest"}
    assert (ref in body["skills"]) is attached
    assert "extra_skills" not in body


# --- the host's parts ----------------------------------------------------------------


def test_the_covering_email_is_the_forwarded_message_without_its_history():
    email = forward("Look?", COVER, "a.docx", docx(["1. A."]), "<c@firm.com>")
    assert negotiation.cover_note(email) == ("dana@otherside.com",
                                             f"Jim,\n\n{COVER}")


def test_an_outlook_forward_and_a_colleagues_forward():
    outlook = InboundEmail(
        message_id="<o@firm.com>", from_address="jim@firm.com", to=["review@firm.com"],
        subject="FW: SPA", received_at=datetime.now(UTC), text_body=(
            "See below.\n\n________________________________\nFrom: Dana Other "
            "<dana@otherside.com>\nSent: 28 September 2026 10:02\nTo: Jim Baker\n"
            "Subject: SPA\n\nAll points agreed bar the cap.\n"))
    assert negotiation.cover_note(outlook) == ("dana@otherside.com",
                                               "All points agreed bar the cap.")
    colleague = outlook.model_copy(update={"text_body": outlook.text_body.replace(
        "dana@otherside.com", "amy@firm.com")})
    assert negotiation.cover_note(colleague) == ("", "")


@pytest.fixture
def gmail_lawyer(monkeypatch):
    """Production's shape: the firm's domain is not the lawyer's mailbox's."""
    monkeypatch.setenv("FIRM_DOMAINS", "legal.firm.com")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("ALLOWED_SENDERS", "fictional.lawyer@gmail.com")
    settings.cache_clear()
    yield
    settings.cache_clear()


def lawyer_writes(body: str, sender: str = "fictional.lawyer@gmail.com") -> InboundEmail:
    content = docx(["1. A."])
    return InboundEmail(
        message_id="<g@gmail.com>", from_address=sender, to=["review@legal.firm.com"],
        subject="Falcon SPA - agreed form", received_at=datetime.now(UTC), text_body=body,
        attachments=[Attachment(filename="SPA.docx", content_type=DOCX,
                                size_bytes=len(content), content=content)])


def test_the_allowlisted_lawyer_outside_the_firms_domains_is_not_the_other_side(gmail_lawyer):
    """Falcon rehearsal, step 7a with production's FIRM_DOMAINS: the lawyer on
    Gmail asked for a final check, and their own words were read as the
    other side's covering email ("I couldn't settle this time which of the 9
    points we raised they accepted")."""
    email = lawyer_writes("Harrowgate have agreed the form overnight. Final check, please.")
    assert negotiation.cover_note(email) == ("", "")


def test_the_thread_owner_is_not_the_other_side(gmail_lawyer):
    email = lawyer_writes("Final check, please.", sender="fictional.lawyer@outlook.com")
    assert negotiation.cover_note(email, owner="fictional.lawyer@outlook.com") == ("", "")
    assert negotiation.cover_note(email)[0] == "fictional.lawyer@outlook.com"


def test_their_email_is_the_counterpartys_in_the_forwarded_chain(gmail_lawyer):
    """The lawyer forwards their own reply to Dana: the first quoted From is
    the lawyer, and the covering email is Dana's, below it."""
    email = lawyer_writes(
        "Check this against our points.\n\n---------- Forwarded message ---------\n"
        "From: Jim Baker <fictional.lawyer@gmail.com>\nDate: Mon, 28 Sep 2026 at 11:00\n"
        "Subject: Re: SPA\nTo: Dana <dana@otherside.com>\n\nThanks Dana, will revert.\n\n"
        "From: Dana Other <dana@otherside.com>\nSent: 28 September 2026 10:02\n"
        "To: Jim Baker\nSubject: SPA\n\nAll points agreed bar the cap.\n")
    assert negotiation.cover_note(email) == ("dana@otherside.com",
                                             "All points agreed bar the cap.")


@pytest.mark.parametrize("said, asks", [
    ("Did they accept our changes?", True),
    ("did the other side accept all of our points", True),
    ("Have they taken our mark-up?", True),
    ("What did they accept?", True),
    ("Which of our changes made it in?", True),
    ("Can you accept the changes in clause 4?", False),
    ("Undo 2", False),
])
def test_the_question_in_the_lawyers_own_words(said, asks):
    email = InboundEmail(message_id="<a@b>", from_address="jim@firm.com", to=[],
                         subject="Re: SPA", text_body=said, received_at=datetime.now(UTC))
    assert negotiation.asks_status(email) is asks


def test_a_point_still_open_is_not_raised_twice():
    negotiation.record_points("t1", 1, [{"source": "issue", "clause": "Clause 10.1: Cap",
                                         "asked": "25%", "anchor": "100% of the Price"}])
    assert negotiation.record_points("t1", 2, [
        {"source": "issue", "clause": "Clause 10.1: Cap", "asked": "25%"},
        {"source": "issue", "clause": "Clause 10: Cap", "asked": "Lower it",
         "anchor": "100% of the Price"},
        {"source": "issue", "clause": "Clause 9.4: Fraud", "asked": "Reinstate it"},
    ]) == 1
    assert [p.round for p in negotiation.points("t1")] == [1, 2]


def test_the_copy_we_were_bccd_on_is_the_version_we_sent(monkeypatch, captured):
    """Our email to them makes asks, and the document on it is what they will
    change. Both go on the ledger."""
    monkeypatch.setattr(handler.review, "review", lambda doc, mode, instructions, **kw:
                        ReviewResult(mode=mode, summary="", findings=[]))
    sent = docx(V1)
    email = InboundEmail(
        message_id="<bcc@firm.com>", from_address="jim@firm.com",
        to=["dana@otherside.com"], subject="Falcon SPA",
        text_body="Dana, we need the cap at 25% and English law. Please confirm by Friday.",
        attachments=[Attachment(filename="Project Falcon SPA v1.docx", content_type=DOCX,
                                size_bytes=len(sent), content=sent)],
        received_at=datetime.now(UTC))
    from lra.pipeline import identity

    assert identity.entry_mode(email) is EntryMode.BCC_SENT
    handler.handle(email)
    [ledger] = _ledger_threads()
    [asks] = negotiation.points(ledger)
    assert asks.source == "email" and "cap at 25%" in asks.asked
    assert negotiation.latest_sent(ledger)[1] == sent


def test_the_words_when_everything_they_said_is_true():
    s = negotiation.Settled(
        outcome={"points": [point("P1", "accepted", [0], mentioned=True)],
                 "unrequested": [], "claims": [{"quote": "q", "verdict": "true",
                                                "evidence": "e", "points": ["P1"],
                                                "changes": [0]}]},
        ledger={"P1": {"clause": "Clause 4: Term"}}, version="v3", cover=True)
    assert s.verdict() == "Matches what they said."
    assert s.tally() == "Your one point was accepted."
    s.outcome["unrequested"] = [{"clause": "9", "summary": "adds a break fee", "risk": "medium",
                                 "mentioned": False, "changes": [1]}]
    assert s.verdict() == "Matches what they said, but v3 also adds a break fee."


def test_the_words_when_there_was_no_covering_email():
    s = negotiation.Settled(
        outcome={"points": [point("P1"), point("P2")], "unrequested": [unasked(0)],
                 "claims": []},
        ledger={}, version="their version", cover=False)
    assert s.verdict() == "They sent no covering note. Their version deletes 9.4."
    assert s.tally() == "None of your 2 points were accepted."


def test_the_version_label_comes_from_the_filename():
    assert negotiation.version_label("Project Falcon SPA v2.docx") == "v2"
    assert negotiation.version_label("SPA (Draft 3).docx") == "draft3"
    assert negotiation.version_label("SPA.docx") == "their version"


def test_the_first_message_fences_the_evidence():
    from lra.pipeline import extract, review

    raw = docx(["1. A."])
    doc = extract.extract(Attachment(filename="a.docx", content_type=DOCX,
                                     size_bytes=len(raw), content=raw))
    text = review.first_message(doc, Mode.REDLINE, "", None, "", "", "", ["a.docx"],
                                negotiation_block="P1 | Clause 1 | asked: x")
    fence = text.split("<negotiation-", 1)[1].split(">", 1)[0]
    assert f"</negotiation-{fence}>" in text
    assert text.index(negotiation.TASK) > text.index(f"</negotiation-{fence}>")
    # The writer still works on the document, not the evidence.
    assert "`/workspace/a.docx`" not in text or "negotiation-evidence" not in text.split(
        "run the writer", 1)[-1]
