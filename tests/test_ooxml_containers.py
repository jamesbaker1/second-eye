"""The places text hides, and the places an edit must refuse to go.

Word puts visible text inside containers the naive reader never opens:
hyperlinks, fields, content controls, other people's tracked changes. Two
things go wrong when the writer cannot see them. The text on either side of the
container becomes adjacent in its model, so a phrase that is not visibly
contiguous matches and the edit splices itself around the container. And the
ambiguity guard counts occurrences in the same blind model, so a phrase the
lawyer can see twice is counted once and silently changed in one place.

Every test here is a document-damaging case the adversarial audit produced and
a real document can reproduce. The posture they enforce is the one in
docs/redlining.md: refusing an edit is always acceptable, a wrong edit never is.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from lra.models import Finding, Mode, ReviewResult, Severity
from lra.pipeline import redline
from lra.pipeline.ooxml import (
    AnchorAmbiguous,
    AnchorNotSafelyEditable,
    Revision,
    RevisionWriter,
    _all_paragraphs,
    _paragraph_text,
)
from lra.pipeline.validate import validate

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


# --- builders -------------------------------------------------------------

def save(doc) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def reopen(doc):
    """Round-trip through the package, as redline.apply does."""
    return Document(BytesIO(save(doc)))


def xml_of(content: bytes, part: str = "word/document.xml") -> str:
    return zipfile.ZipFile(BytesIO(content)).read(part).decode()


def edit(anchor, replacement, auto_apply=True) -> ReviewResult:
    return ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.FORMATTING, category="typo", title="t",
                explanation="", anchor=anchor, suggested_text=replacement,
                auto_apply=auto_apply)
    ])


def sub(parent, tag, text=None, **attrs):
    element = etree.SubElement(parent, qn(tag))
    for name, value in attrs.items():
        element.set(qn(name.replace("_", ":")), value)
    if text is not None:
        element.text = text
    return element


def run_with_text(parent, text):
    run = sub(parent, "w:r")
    sub(run, "w:t", text)
    return run


def apply_to(content: bytes, anchor: str, replacement: str | None):
    doc = Document(BytesIO(content))
    RevisionWriter(doc, author="A").apply(Revision(anchor, replacement, "A"))
    return save(doc)


# --- text Word displays that the writer used to be blind to ---------------

def test_text_inside_a_content_control_is_visible_to_the_writer():
    """[run "Governed by the laws of "][sdt "Delaware"][run ", without regard"].
    Skipping the content control made "laws of , without regard" look like a
    real phrase in the document."""
    d = Document()
    p = d.add_paragraph()
    p.add_run("Governed by the laws of ")
    sdt = sub(p._p, "w:sdt")
    sub(sdt, "w:sdtPr")
    run_with_text(sub(sdt, "w:sdtContent"), "Delaware")
    p.add_run(", without regard to conflicts.")

    text = _paragraph_text(reopen(d).paragraphs[0])
    assert text == "Governed by the laws of Delaware, without regard to conflicts."


def test_text_inside_a_simple_field_is_visible_to_the_writer():
    d = Document()
    p = d.add_paragraph()
    p.add_run("See Section ")
    field = sub(p._p, "w:fldSimple", w_instr=" REF _Ref12345 \\r \\h ")
    run_with_text(field, "4.2")
    p.add_run(" for indemnity.")

    assert _paragraph_text(reopen(d).paragraphs[0]) == "See Section 4.2 for indemnity."


def test_an_anchor_that_reads_across_a_content_control_is_refused():
    """The anchor exists in the old text model and nowhere on the page. Editing
    it stranded the control's text mid-sentence."""
    d = Document()
    p = d.add_paragraph()
    p.add_run("Governed by the laws of ")
    sdt = sub(p._p, "w:sdt")
    sub(sdt, "w:sdtPr")
    run_with_text(sub(sdt, "w:sdtContent"), "Delaware")
    p.add_run(", without regard to conflicts.")
    content = save(d)

    out = redline.apply(content, edit("laws of , without regard", "laws of New York"))
    assert not out.applied
    assert out.content == content


def test_an_edit_reaching_into_a_content_control_is_refused():
    d = Document()
    p = d.add_paragraph()
    p.add_run("Governed by the laws of ")
    sdt = sub(p._p, "w:sdt")
    sub(sdt, "w:sdtPr")
    run_with_text(sub(sdt, "w:sdtContent"), "Delaware")
    p.add_run(", without regard to conflicts.")

    doc = reopen(d)
    with pytest.raises(AnchorNotSafelyEditable):
        RevisionWriter(doc, "A").apply(Revision("of Delaware, without", "of New York,", "A"))


# --- cross-reference fields ----------------------------------------------

def complex_field_paragraph():
    """[begin][ REF _Ref12345 ][separate][result "4.2"][end] -- what Word emits
    for a cross-reference inserted from the References ribbon."""
    d = Document()
    p = d.add_paragraph()
    p.add_run("See Section ")
    sub(sub(p._p, "w:r"), "w:fldChar", w_fldCharType="begin")
    sub(sub(p._p, "w:r"), "w:instrText", " REF _Ref12345 \\r \\h ")
    sub(sub(p._p, "w:r"), "w:fldChar", w_fldCharType="separate")
    run_with_text(p._p, "4.2")
    sub(sub(p._p, "w:r"), "w:fldChar", w_fldCharType="end")
    p.add_run(" for indemnity.")
    return save(d)


def test_an_edit_over_a_cross_reference_result_is_refused():
    """The result run is an ordinary w:r, so the writer used to move "4.2" into
    a w:del and leave ` REF _Ref12345 ` behind. The next field update -- F9,
    printing, simply opening the file -- regenerates the old number and the
    tracked change the lawyer accepted evaporates."""
    content = complex_field_paragraph()
    out = redline.apply(content, edit("Section 4.2 for indemnity",
                                      "Section 4.3 for indemnity"))
    assert not out.applied
    assert out.content == content
    assert out.notes and "field" in out.notes[0]


def test_the_field_instruction_and_its_result_stay_together():
    content = complex_field_paragraph()
    out = redline.apply(content, edit("Section 4.2 for indemnity",
                                      "Section 4.3 for indemnity"))
    body = xml_of(out.content)
    assert body.count("<w:fldChar") == 3
    assert ">4.2<" in body


def test_an_edit_beside_a_field_still_lands():
    """Refusing the field result must not refuse the rest of the paragraph."""
    content = complex_field_paragraph()
    out = redline.apply(content, edit("for indemnity", "for indemnification"))
    assert len(out.applied) == 1
    assert redline.verify(out.content, expect_revisions=True)[0]


# --- another author's tracked insertion -----------------------------------

def mixed_authorship_document():
    d = Document()
    p = d.add_paragraph()
    p.add_run("The Recipient shall pay ")
    ins = sub(p._p, "w:ins", w_id="9000", w_author="A. Associate",
              w_date="2026-01-01T00:00:00Z")
    run_with_text(ins, "thirty (13)")
    p.add_run(" days after closing.")
    return save(d)


def test_an_edit_straddling_another_authors_insertion_is_refused():
    """The covered runs had different parents, so parent.remove() raised part
    way through the loop -- after the w:ins and w:del were spliced in. The
    lawyer was told the edit was not made while the file read "thirty (30) days
    days after closing"."""
    content = mixed_authorship_document()
    doc = Document(BytesIO(content))
    with pytest.raises(AnchorNotSafelyEditable):
        RevisionWriter(doc, "LRA").apply(
            Revision("thirty (13) days", "thirty (30) days", "LRA"))


def test_a_refusal_across_authorship_leaves_the_text_exactly_as_it_was():
    content = mixed_authorship_document()
    doc = Document(BytesIO(content))
    with pytest.raises(AnchorNotSafelyEditable):
        RevisionWriter(doc, "LRA").apply(
            Revision("thirty (13) days", "thirty (30) days", "LRA"))
    assert _paragraph_text(doc.paragraphs[0]) == (
        "The Recipient shall pay thirty (13) days after closing."
    )
    assert "<w:del " not in xml_of(save(doc))


def test_one_findings_refusal_does_not_stop_another_from_landing():
    d = Document(BytesIO(mixed_authorship_document()))
    d.add_paragraph("This agreement is governed by the laws of Delaware.")
    result = ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.BLOCKER, category="amount", title="a",
                explanation="", anchor="thirty (13) days",
                suggested_text="thirty (30) days", auto_apply=True),
        Finding(severity=Severity.SUBSTANTIVE, category="law", title="b",
                explanation="", anchor="Delaware", suggested_text="New York",
                auto_apply=True),
    ])
    out = redline.apply(save(d), result)
    assert [f.title for f in out.applied] == ["b"]
    assert "days days" not in xml_of(out.content)


def test_an_unexpected_failure_abandons_the_redline_rather_than_attaching_it(monkeypatch):
    """Every refusal the writer knows how to make is typed. An untyped one means
    the document may be in a state nobody predicted, and attaching it while
    reporting the finding as merely skipped is the worst of both."""
    def boom(self, revision):
        raise RuntimeError("something nobody predicted")

    monkeypatch.setattr(redline.RevisionWriter, "apply", boom)
    content = Path("samples/simple.docx").read_bytes()
    out = redline.apply(content, edit("Mutual Non-Disclosure Agreement",
                                      "Mutual Nondisclosure Agreement"))
    assert out.content == content
    assert not out.applied
    assert out.notes and "not attached" in out.notes[0]


# --- merged cells and linked headers are one paragraph, not several -------

def test_a_cell_merged_across_columns_is_not_counted_as_several_occurrences():
    """row.cells returns one entry per grid position, so a cell spanning three
    columns was yielded three times -- the same lxml element -- and _locate
    called a unique phrase ambiguous."""
    d = Document()
    table = d.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Schedule of Fees payable on Closing"
    content = save(d)

    paragraphs = list(_all_paragraphs(Document(BytesIO(content))))
    assert len(paragraphs) == len({id(p._p) for p in paragraphs})

    out = apply_to(content, "Schedule of Fees payable on Closing", "Fee Schedule")
    assert redline.verify(out, expect_revisions=True)[0]


def test_a_cell_merged_across_rows_is_not_counted_as_several_occurrences():
    d = Document()
    table = d.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(2, 0)).text = "Fees payable on Closing"
    out = apply_to(save(d), "Fees payable on Closing", "Closing fees")
    assert redline.verify(out, expect_revisions=True)[0]


def test_a_header_linked_to_the_previous_section_is_not_counted_twice():
    """section.header on a linked section returns the previous section's header
    definition, so the same paragraph was yielded once per section."""
    from docx.enum.section import WD_SECTION

    d = Document()
    d.sections[0].header.paragraphs[0].text = "ACME CONFIDENTIAL"
    d.add_paragraph("Body text.")
    d.add_section(WD_SECTION.NEW_PAGE)
    assert d.sections[1].header.is_linked_to_previous

    out = apply_to(save(d), "ACME CONFIDENTIAL", "ACME - CONFIDENTIAL")
    assert redline.verify(out)[0]


def test_text_in_a_nested_table_is_visible_to_the_writer():
    """Text the writer cannot see is worse than text it cannot edit: an anchor
    that occurs twice looks unique, and gets edited."""
    d = Document()
    outer = d.add_table(rows=1, cols=1)
    outer.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Payable on Closing"
    d.add_paragraph("Payable on Closing")

    doc = Document(BytesIO(save(d)))
    with pytest.raises(AnchorAmbiguous):
        RevisionWriter(doc, "A").apply(Revision("Payable on Closing", "Payable at Closing", "A"))


# --- looking at a document must not restructure it ------------------------

def test_searching_for_an_anchor_adds_no_header_or_footer_parts():
    """python-docx's header accessors are get-or-create, so merely locating an
    anchor materialised six parts and six sectPr references in a document that
    had none. "Preserve headers, footers, numbering, styles and embedded
    objects" is the first non-negotiable in docs/redlining.md."""
    original = Path("samples/simple.docx").read_bytes()
    before = set(zipfile.ZipFile(BytesIO(original)).namelist())

    out = redline.apply(original, edit("text that is nowhere in this file", "x"))
    assert not out.applied

    doc = Document(BytesIO(original))
    list(_all_paragraphs(doc))
    after = set(zipfile.ZipFile(BytesIO(save(doc))).namelist())
    assert after == before
    assert "headerReference" not in xml_of(save(doc))


# --- revision ids are unique across every part, not just the body ---------

def test_revision_ids_do_not_collide_with_ids_in_a_header():
    """Headers carry other authors' revisions too, and the writer edits them.
    Allocating from the body's high-water mark alone restarted at 1."""
    d = Document()
    header = d.sections[0].header.paragraphs[0]
    header.text = "Header"
    sub(header._p, "w:ins", w_id="7000", w_author="A. Associate",
        w_date="2026-01-01T00:00:00Z")
    d.add_paragraph("The term is thirty days.")

    doc = Document(BytesIO(save(d)))
    writer = RevisionWriter(doc, "A")
    writer.apply(Revision("thirty", "sixty", "A"))
    ids = [int(i) for i in re.findall(r'w:id="(\d+)"', xml_of(save(doc)))]
    assert ids and min(ids) > 7000


# --- a header redline is a redline ----------------------------------------

def test_a_header_only_redline_is_attached_not_discarded():
    """The markup lands in word/header1.xml, which the verifier never read, so
    a correct redline was thrown away with "I could not produce a safe edited
    copy". Removing a DRAFT marker from a header is exactly the leftover the
    product advertises catching."""
    d = Document()
    d.sections[0].header.paragraphs[0].text = "DRAFT - Acme Holdings"
    d.add_paragraph("Body text about the fee.")
    content = save(d)

    out = redline.apply(content, edit("DRAFT - Acme Holdings", "Acme Holdings"))
    assert len(out.applied) == 1
    assert out.content != content
    assert "<w:ins " in xml_of(out.content, "word/header1.xml")
    assert redline.verify(out.content, expect_revisions=True)[0]


def test_a_malformed_header_revision_is_caught_by_the_validator():
    """Nothing is sent without passing, and a header is part of "nothing"."""
    d = Document()
    d.sections[0].header.paragraphs[0].text = "DRAFT - Acme Holdings"
    d.add_paragraph("Body.")
    content = save(d)

    source = zipfile.ZipFile(BytesIO(content))
    items = {n: source.read(n) for n in source.namelist()}
    items["word/header1.xml"] = items["word/header1.xml"].decode().replace(
        "<w:p>",
        '<w:p><w:del w:id="3" w:author="A" w:date="2026-01-01T00:00:00Z">'
        "<w:r><w:t>oops</w:t></w:r></w:del>",
        1,
    ).encode()
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in items.items():
            out.writestr(name, data)

    report = validate(buf.getvalue())
    assert not report.ok
    assert any("delText" in e for e in report.errors)


# --- a single word is a word, not a substring -----------------------------

def test_a_single_word_anchor_is_not_ambiguous_against_a_longer_word():
    """The one deterministic check that sets auto_apply anchors on a bare word,
    so this path is on by default in conservative mode."""
    d = Document()
    d.add_paragraph("The Indemnitee shall notify the Indemnitees of any claim.")
    out = apply_to(save(d), "Indemnitee", "Indemnified Party")
    assert ">Indemnified Party<" in xml_of(out)
    assert redline.verify(out, expect_revisions=True)[0]


def test_a_single_word_anchor_never_lands_inside_a_longer_word():
    d = Document()
    d.add_paragraph("The Indemnitees listed in Schedule 2 shall be notified.")
    out = redline.apply(save(d), edit("Indemnitee", "Indemnified Party"))
    assert not out.applied
    assert "Indemnified Partys" not in xml_of(out.content)


def test_a_padded_phrase_anchor_still_matches_mid_word():
    """checks.py pads a phrase with the characters around it to make it unique,
    and the padding starts and ends mid-word on purpose."""
    d = Document()
    d.add_paragraph("1. Term. The term is thirty (13) months from the Effective Date.")
    out = apply_to(save(d), "he term is thirty (13) months from the E",
                   "he term is thirty (30) months from the E")
    assert redline.verify(out, expect_revisions=True)[0]


# --- tabs, breaks, and text boxes ----------------------------------------

def test_tabs_and_line_breaks_read_the_same_way_the_agent_sees_them():
    """extract.py feeds the agent para.text, which renders w:tab as a tab. The
    writer rendered it as nothing, so "2.<TAB>Fees.<TAB>[AMOUNT] is payable"
    -- the default formatting of every numbered agreement -- was unlocatable."""
    d = Document()
    p = d.add_paragraph()
    p.add_run("2.")
    p.add_run().add_tab()
    p.add_run("Fees.")
    p.add_run().add_break()
    p.add_run("[AMOUNT] is payable on Closing.")

    paragraph = reopen(d).paragraphs[0]
    assert _paragraph_text(paragraph) == paragraph.text
    assert "\t" in _paragraph_text(paragraph)


def numbered(text: str, bold: bool = False):
    """One paragraph typed as one run, the way a document that is not
    auto-numbered holds its clauses: "9.3<TAB>The liability ..." with the tab
    a w:tab child of the same run as the number and the words."""
    d = Document()
    p = d.add_paragraph()
    run = p.add_run()
    run.bold = bold
    first, _, rest = text.partition("\t")
    sub(run._r, "w:t", first)
    sub(run._r, "w:tab")
    sub(run._r, "w:t", rest).set(
        "{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return d


def decided_text(content: bytes, accepting: bool) -> str:
    """The paragraph as it reads once every change is accepted, or rejected."""
    from lra.pipeline.ooxml import decide, tracked_changes

    doc = Document(BytesIO(content))
    ids = [c.id for c in tracked_changes(doc)]
    decide(doc, accept=ids if accepting else (), reject=() if accepting else ids)
    return _paragraph_text(reopen(doc).paragraphs[0])


def assert_sound(content: bytes) -> None:
    report = validate(content)
    assert report.ok, str(report)
    ok, reason = redline.verify(content, expect_revisions=True)
    assert ok, reason


ORIGINAL = "9.3\tThe liability shall not exceed 100 per cent. of the Price."


def test_an_edit_after_a_tab_in_the_same_run_lands_and_keeps_the_tab():
    """Rehearsal of Project Falcon: "I left "shall not exceed 100 per cent..."
    alone because the text to change shares a run with ... a tab". A clause
    typed "9.3<TAB>text" as one run could never be edited, answered or turned,
    though the tab was nowhere near the words being changed."""
    out = redline.apply(save(numbered(ORIGINAL, bold=True)),
                        edit("100 per cent.", "50 per cent."))
    assert len(out.applied) == 1, out.notes
    assert_sound(out.content)
    assert decided_text(out.content, True) == ORIGINAL.replace("100", "50")
    assert decided_text(out.content, False) == ORIGINAL
    body = etree.fromstring(xml_of(out.content).encode())
    tabs = body.findall(f".//{{{W}}}tab")
    assert len(tabs) == 1
    tab_run = tabs[0].getparent()
    assert tab_run.getparent().tag == f"{{{W}}}p", "the tab is not part of the change"
    assert tab_run.find(f"{{{W}}}rPr/{{{W}}}b") is not None, "formatting kept"
    for run in body.iter(f"{{{W}}}r"):
        assert run.find(f"{{{W}}}rPr/{{{W}}}b") is not None


def test_an_edit_before_a_tab_in_the_same_run_lands_and_keeps_the_tab():
    d = numbered("Fees\t[AMOUNT] is payable on Closing.")
    out = redline.apply(save(d), edit("Fees", "Charges"))
    assert len(out.applied) == 1, out.notes
    assert_sound(out.content)
    assert decided_text(out.content, True) == "Charges\t[AMOUNT] is payable on Closing."
    assert decided_text(out.content, False) == "Fees\t[AMOUNT] is payable on Closing."


def test_the_clause_number_beside_a_tab_can_be_renumbered():
    out = redline.apply(save(numbered(ORIGINAL)), edit("9.3", "9.4"))
    assert len(out.applied) == 1, out.notes
    assert_sound(out.content)
    assert decided_text(out.content, True) == ORIGINAL.replace("9.3", "9.4")


def test_an_anchor_that_crosses_a_tab_is_changed_with_real_tabs():
    """The tab inside the span goes into the w:del as a deleted w:tab, and a
    tab in the replacement arrives as a w:tab inside the w:ins, never as the
    character "\\t" inside a w:t, which Word does not show as a tab."""
    d = Document()
    p = d.add_paragraph()
    p.add_run("2.")
    p.add_run().add_tab()
    p.add_run("Fees. [AMOUNT] is payable on Closing.")
    out = redline.apply(save(d), edit("2.\tFees. [AMOUNT] is payable",
                                      "2.\tFees. $50,000 is payable"))
    assert len(out.applied) == 1, out.notes
    assert_sound(out.content)
    assert decided_text(out.content, True) == "2.\tFees. $50,000 is payable on Closing."
    assert decided_text(out.content, False) == "2.\tFees. [AMOUNT] is payable on Closing."
    body = xml_of(out.content)
    assert "\t" not in "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", body))
    assert re.search(r"<w:del [^>]*>.*<w:tab/>.*</w:del>", body)
    assert re.search(r"<w:ins [^>]*>.*<w:tab/>.*</w:ins>", body)


def test_a_change_across_a_tab_inside_one_run_is_changed_and_undone_exactly():
    out = redline.apply(save(numbered(ORIGINAL)),
                        edit("9.3\tThe liability", "9.3\tThe total liability"))
    assert len(out.applied) == 1, out.notes
    assert_sound(out.content)
    assert decided_text(out.content, True) == ORIGINAL.replace("The", "The total")
    assert decided_text(out.content, False) == ORIGINAL


def test_an_insertion_right_after_a_tab_lands():
    doc = Document(BytesIO(save(numbered(ORIGINAL))))
    para = doc.paragraphs[0]
    RevisionWriter(doc, author="A").insert_at(para, len("9.3\t"), "Subject to clause 9.4, ")
    content = save(doc)
    assert_sound(content)
    assert decided_text(content, True) == ORIGINAL.replace(
        "\tThe", "\tSubject to clause 9.4, The")
    assert decided_text(content, False) == ORIGINAL


def test_a_run_holding_a_break_as_well_as_the_words_is_still_refused():
    """Only tabs are split out. A line break could be a page or column break,
    which is layout the lawyer did not ask to have moved."""
    d = Document()
    run = d.add_paragraph().add_run("The term is thirty days.")
    run.add_break()
    content = save(d)
    out = redline.apply(content, edit("thirty days", "sixty days"))
    assert not out.applied
    assert out.content == content


def test_an_edit_after_a_tab_still_lands():
    d = Document()
    p = d.add_paragraph()
    p.add_run("2.")
    p.add_run().add_tab()
    p.add_run("Fees. The Reciever pays on Closing.")
    out = apply_to(save(d), "Reciever", "Recipient")
    assert ">Recipient<" in xml_of(out)


def textbox_run(paragraph, text):
    """A floating text box: mc:AlternateContent > ... > w:txbxContent > w:p."""
    run = sub(paragraph._p, "w:r")
    alternate = etree.SubElement(run, f"{{{MC}}}AlternateContent")
    fallback = etree.SubElement(alternate, f"{{{MC}}}Fallback")
    picture = sub(fallback, "w:pict")
    box = sub(picture, "w:txbxContent")
    run_with_text(sub(box, "w:p"), text)
    return run


def test_a_text_box_does_not_leak_into_the_paragraphs_text():
    """_run_text read descendants while _set_run_text wrote direct children, so
    the two disagreed about a run holding a text box, the offsets drifted, and
    the split wrote the wrong slice of a paragraph it had already scrambled."""
    d = Document()
    p = d.add_paragraph()
    p.add_run("The term is thirty days.")
    textbox_run(p, "SIDEBAR NOTE")

    paragraph = reopen(d).paragraphs[0]
    assert _paragraph_text(paragraph) == "The term is thirty days."


def test_an_edit_beside_a_text_box_neither_duplicates_nor_scrambles_it():
    d = Document()
    p = d.add_paragraph()
    p.add_run("The term is thirty days.")
    textbox_run(p, "SIDEBAR NOTE")
    content = save(d)

    out = redline.apply(content, edit("thirty days", "sixty days"))
    assert len(out.applied) == 1
    body = xml_of(out.content)
    assert body.count("SIDEBAR NOTE") == 1
    assert redline.verify(out.content, expect_revisions=True)[0]


def test_a_run_sharing_itself_with_a_text_box_is_refused():
    d = Document()
    p = d.add_paragraph()
    run = p.add_run("The term is thirty days.")
    alternate = etree.SubElement(run._r, f"{{{MC}}}AlternateContent")
    fallback = etree.SubElement(alternate, f"{{{MC}}}Fallback")
    box = sub(sub(fallback, "w:pict"), "w:txbxContent")
    run_with_text(sub(box, "w:p"), "SIDEBAR NOTE")
    content = save(d)

    out = redline.apply(content, edit("thirty days", "sixty days"))
    assert not out.applied
    assert out.content == content


# --- privileged text must not reach the logs ------------------------------

def test_neither_the_anchor_nor_the_title_reaches_the_log(caplog):
    """An anchor is text quoted verbatim from a client's contract and a title
    routinely names the parties. Anchor misses are not rare -- this module
    exists to refuse them -- so this runs on ordinary documents at the default
    INFO level."""
    import logging

    result = ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.BLOCKER, category="amount",
                title="Acme Holdings LLC is capped at $4,000,000 in clause 9.2",
                explanation="", anchor="the liability of Acme Holdings LLC is "
                "capped at $4,000,000", suggested_text="the liability of Acme "
                "Holdings LLC is capped at $2,000,000", auto_apply=True),
    ])
    with caplog.at_level(logging.INFO):
        redline.apply(Path("samples/simple.docx").read_bytes(), result)

    logged = "\n".join(r.getMessage() for r in caplog.records)
    assert "Acme" not in logged
    assert "4,000,000" not in logged
    assert "amount" in logged, "the log still has to be useful"


def test_paragraph_dedup_survives_garbage_collection():
    """The dedup used id() on lxml proxies, which are collected once nothing
    holds them, and CPython reuses the freed address. An id already in the set
    could therefore belong to a different element, so a header paragraph looked
    already-seen and was skipped. It passed on one machine and failed in CI on
    the same commit, which is the worst shape a bug can have."""
    import gc

    from docx.enum.section import WD_SECTION

    from lra.pipeline.ooxml import _all_paragraphs, _paragraph_text

    d = Document()
    d.sections[0].header.paragraphs[0].text = "ACME CONFIDENTIAL"
    for i in range(40):
        d.add_paragraph(f"{i}. A clause with text in it.")
    d.add_section(WD_SECTION.NEW_PAGE)
    d.add_paragraph("Text in the second section.")

    doc = Document(BytesIO(save(d)))
    # Collect aggressively while walking, which is what makes the address reuse
    # observable rather than a once-in-a-while flake.
    texts = []
    for paragraph in _all_paragraphs(doc):
        gc.collect()
        texts.append(_paragraph_text(paragraph))

    assert "ACME CONFIDENTIAL" in texts, "the header paragraph was skipped"
    assert texts.count("ACME CONFIDENTIAL") == 1, "the header was yielded twice"
    assert "Text in the second section." in texts
