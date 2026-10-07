"""A closing by email: packets out, signed pages in, the executed set, the
checklist.

The platform is faked (tests/fake_sessions.py), but the sandbox is not: the
"agent" in these tests is a script of the model's decisions (who signs,
which page a scan is, whether it is signed), and every tool it uses is the
real lra-closing skill, built as `lra skills sync` builds it and run from its
own folder by a fresh interpreter, exactly as the container would. So what
the host receives is what the skill really writes, and what is exercised is
everything around the model: the state the session is given, the validator
on both sides, the files kept, who may file pages, and the reply.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace as NS
from unittest.mock import patch

import pytest
from docx import Document
from pypdf import PdfReader

import lra.config
from lra import closing, handler, managed, skillsync
from lra.mail.console import ConsoleProvider
from lra.models import Attachment, InboundEmail
from lra.pipeline import closing_docs, closing_state
from tests import fake_sessions as fs

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OWNER = "jim@firm.com"

SPA_FRONT = [
    "SHARE PURCHASE AGREEMENT",
    "This Agreement is dated 3 March 2026 and is made between:",
    '(1) ACME HOLDINGS LIMITED, a company incorporated in England (the "Seller"); and',
    '(2) BETA TRADING LLC, a Delaware limited liability company (the "Buyer").',
    "1. Sale. The Seller shall sell the Shares to the Buyer.",
    ("4. Conditions. 4.1 Completion is conditional on: (a) merger clearance; and "
     "(b) the Buyer's board approving the purchase."),
    ("6. Completion deliverables. At Completion the Seller shall deliver the share "
     "certificates and a stock transfer form."),
    "IN WITNESS WHEREOF the parties have signed this Agreement on the date above.",
]
LETTER_FRONT = [
    "DISCLOSURE LETTER",
    "This letter is dated 3 March 2026.",
    "1. Disclosures. The Seller discloses the matters in the schedule.",
    "IN WITNESS WHEREOF this letter has been signed.",
]
SIGNERS = (("ACME HOLDINGS LIMITED", "Jane Smith"), ("BETA TRADING LLC", "Bob Jones"))


def agreement(front=SPA_FRONT, signers=SIGNERS) -> bytes:
    d = Document()
    for text in front:
        d.add_paragraph(text)
    for entity, who in signers:
        for line in (f"Signed for and on behalf of {entity}", "By: ______________",
                     f"Name: {who}", "Title: Director"):
            d.add_paragraph(line)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def attach(name: str, content: bytes, kind: str = DOCX) -> Attachment:
    return Attachment(filename=name, content_type=kind, size_bytes=len(content), content=content)


# --- LibreOffice, present or not ------------------------------------------------------


@pytest.fixture(scope="session")
def soffice_dir(tmp_path_factory) -> Path:
    """A folder holding a `soffice` that runs tests/fake_soffice.py."""
    folder = tmp_path_factory.mktemp("bin")
    script = Path(__file__).with_name("fake_soffice.py")
    launcher = folder / "soffice"
    launcher.write_text(f"#!{sys.executable}\nimport runpy\n"
                        f"runpy.run_path({str(script)!r}, run_name='__main__')\n")
    launcher.chmod(0o755)
    return folder


@pytest.fixture(autouse=True, params=["no-office", "office"])
def office(request, monkeypatch, soffice_dir) -> bool:
    """Every test twice: once where conversion fails as it does with no
    LibreOffice, once with the stand-in laying documents out as real
    multi-page PDFs. Either way the stand-in is first on PATH, so the
    scripts in their subprocesses never reach a real LibreOffice this machine
    may have, and the results are the same on a laptop and in CI."""
    from lra import convert

    present = request.param == "office"
    monkeypatch.setenv("PATH", f"{soffice_dir}{os.pathsep}{os.environ.get('PATH', '')}")
    monkeypatch.setenv("LRA_FAKE_SOFFICE_LAYOUT", "split" if present else "fail")
    monkeypatch.setattr(convert, "binary", lambda: str(soffice_dir / "soffice"))
    return present


def at(office: bool, page: int) -> str:
    """How the tally names a page: its number when the document could be laid
    out, nothing when it could not. Never a guess."""
    return f" p.{page}" if office else ""


# --- the sandbox, as the container runs it -----------------------------------------


@pytest.fixture(scope="module")
def bundle(tmp_path_factory) -> Path:
    return skillsync.build(tmp_path_factory.mktemp("skill"), "lra-closing")


def run_script(bundle: Path, script: str, *args: str, shim: bool = False) -> dict:
    env = {k: v for k, v in os.environ.items() if k != "PYTHONPATH"}
    if shim:
        env["LRA_FORCE_SHIM"] = "1"
    done = subprocess.run([sys.executable, str(bundle / "scripts" / script), *args],
                          capture_output=True, text=True, env=env, cwd=bundle, timeout=120,
                          check=False)
    assert done.stdout, done.stderr
    return json.loads(done.stdout)


class Sandbox(fs.FakeSessions):
    """The platform, with the agent's work done when the outputs are listed:
    the mounted files are laid out as /workspace, `agent` runs the skill's
    scripts on them, and what it leaves in the outputs folder comes back."""

    def __init__(self, bundle: Path, root: Path, agent):
        super().__init__([fs.idle()])
        self.bundle, self.root, self.agent = bundle, root, agent
        self._done = False

    def list(self, **kw):
        if "scope_id" in kw and not self._done:
            self._done = True
            ws, out = self.root / "workspace", self.root / "outputs"
            ws.mkdir(parents=True)
            out.mkdir(parents=True)
            for name, data in self.uploaded:
                (ws / managed.mount_name(name)).write_bytes(data)
            self.agent(NS(ws=ws, out=out, root=self.root,
                          run=lambda script, *a: run_script(self.bundle, script, *a)))
            self.outputs = [(p.name, p.read_bytes()) for p in sorted(out.iterdir())]
        return super().list(**kw)


class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return f"<agent-{len(self.sent)}@legal.firm.com>"


@pytest.fixture
def env(monkeypatch, tmp_path, bundle):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/c.sqlite3")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    monkeypatch.setenv("ALLOWED_SENDERS", "")
    fs.configure(monkeypatch, MANAGED_CLOSING_AGENT_ID="agent_closing")
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    sessions: list[Sandbox] = []

    def send(inbound: InboundEmail, agent=None) -> Sandbox | None:
        fake = Sandbox(bundle, tmp_path / f"session{len(sessions)}", agent or (lambda s: None))
        sessions.append(fake)
        with patch.object(managed, "anthropic_client", lambda: fake.client):
            handler.handle(inbound)
        return fake

    yield NS(provider=provider, send=send, sessions=sessions, tmp=tmp_path)
    lra.config.settings.cache_clear()


_n = iter(range(1, 10_000))


def email(body: str, attachments=(), sender: str = OWNER, subject: str = "Project Falcon",
          reply_to: str | None = None) -> InboundEmail:
    n = next(_n)
    return InboundEmail(
        message_id=f"<m{n}@firm.com>", thread_id=None, from_address=sender,
        to=["review@legal.firm.com"], subject=subject, text_body=body,
        in_reply_to=reply_to, references=[reply_to] if reply_to else [],
        attachments=list(attachments), received_at=datetime.now(UTC))


def last(env):
    return env.provider.sent[-1]


def sent_id(env) -> str:
    return f"<agent-{len(env.provider.sent)}@legal.firm.com>"


# --- the model's decisions, scripted ----------------------------------------------

PLAN = {
    "closing": "Project Falcon",
    "return": "Scan or photograph each signed page and email it to Jim Baker (jim@firm.com).",
    "documents": [{"id": "spa", "file": "Falcon SPA.docx", "short": "SPA"},
                  {"id": "dl", "file": "Disclosure Letter.docx", "short": "Disclosure Letter"}],
    "signatories": [
        {"id": "acme-director", "name": "Acme's director", "person": "Jane Smith",
         "capacity": "Director",
         "sign": [{"document": "spa", "party": "ACME HOLDINGS LIMITED"},
                  {"document": "dl", "party": "Acme Holdings"}]},
        {"id": "beta-director", "name": "Beta Trading's director", "person": "Bob Jones",
         "capacity": "Director",
         "sign": [{"document": "spa", "party": "Beta Trading LLC"},
                  {"document": "dl", "party": "BETA TRADING LLC"}]},
    ],
}


def write_update(s, state: dict, message: list[str], attach: list[str]) -> dict:
    update = s.root / "update.json"
    update.write_text(json.dumps({"schema": closing_state.UPDATE_SCHEMA, "state": state,
                                  "message": message, "attach": attach}))
    return s.run("validate_state.py", "--update", str(update),
                 "--previous", str(s.ws / closing.STATE_FILE), "--outputs", str(s.out),
                 "--held", str(s.ws), "--write", str(s.out / closing.UPDATE_FILE))


def packets_agent(s):
    surveyed = s.run("survey.py", str(s.ws / "Falcon SPA.docx"),
                     str(s.ws / "Disclosure Letter.docx"))
    assert [d["title"] for d in surveyed["documents"]] == [
        "Share Purchase Agreement", "Disclosure Letter"]
    (s.root / "plan.json").write_text(json.dumps(PLAN))
    built = s.run("packets.py", "--plan", str(s.root / "plan.json"), "--workspace", str(s.ws),
                  "--out", str(s.out), "--draft", str(s.root / "draft.json"))
    assert "error" not in built, built
    state = json.loads((s.root / "draft.json").read_text())
    names = [p["file"] for p in built["packets"]] + [d["execution_file"]
                                                      for d in built["documents"]]
    done = write_update(s, state, [("Packets for 2 signatories across 2 documents; "
                                    "execution copies attached.")], names)
    assert "error" not in done, done


def sign(page_pdf: bytes, signed: bool = True) -> bytes:
    """The page as it comes back: a wet-ink squiggle over the signature line."""
    from pypdf import PdfWriter
    from reportlab.pdfgen import canvas

    page = PdfReader(BytesIO(page_pdf)).pages[0]
    writer = PdfWriter()
    added = writer.add_page(page)
    if signed:
        overlay = BytesIO()
        c = canvas.Canvas(overlay, pagesize=(float(page.mediabox.width),
                                             float(page.mediabox.height)))
        c.setLineWidth(1.5)
        c.bezier(120, 520, 160, 600, 200, 440, 260, 540)
        c.save()
        added.merge_page(PdfReader(BytesIO(overlay.getvalue())).pages[0])
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def packet_page(env, signatory: str, n: int) -> bytes:
    """Page n (1 is the cover) of a packet the lawyer was sent."""
    out = env.provider.sent[0]
    packet = next(a for a in out.attachments
                  if a.filename == f"Signature packet - {signatory}.pdf")
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_page(PdfReader(BytesIO(packet.content)).pages[n - 1])
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def scan(*pages: bytes) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for p in pages:
        writer.add_page(PdfReader(BytesIO(p)).pages[0])
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def photo(page_pdf: bytes) -> bytes:
    import pypdfium2

    image = pypdfium2.PdfDocument(page_pdf)[0].render(scale=1.2).to_pil()
    buf = BytesIO()
    image.convert("RGB").save(buf, format="JPEG", quality=70)
    return buf.getvalue()


def filing_agent(accept: dict[tuple[str, int], str], reject: dict[tuple[str, int], str],
                 message: list[str], compile_when_done: bool = True):
    """The model having looked at each page: `accept` maps (attachment, page)
    to the page id it is; `reject` maps it to the reason it is not filed."""
    def agent(s):
        state = json.loads((s.ws / closing.STATE_FILE).read_text())
        pages = {p["id"]: p for p in state["pages"]}
        docs = {d["id"]: d for d in state["documents"]}
        for (name, n), pid in accept.items():
            target = f"signed - {pid}.pdf"
            args = [str(s.ws / name), "--out", str(s.out / target)]
            if name.lower().endswith(".pdf"):
                args += ["--page", str(n)]
            got = s.run("receive.py", *args)
            assert "error" not in got, got
            if got["stamp"]:
                assert got["stamp"]["page"] == pid
            pages[pid]["status"] = "received"
            pages[pid]["received"] = {
                "file": target, "source": f"{name}, page {n}", "from": OWNER,
                "at": "2026-09-28T10:00:00Z", "signed": True,
                "version_ref": docs[pages[pid]["document"]]["ref"], "note": ""}
        for (name, n), reason in reject.items():
            state["rejected"].append({"source": f"{name}, page {n}", "reason": reason})
        attach: list[str] = []
        if compile_when_done and all(p["status"] == "received" for p in pages.values()):
            (s.root / "draft.json").write_text(json.dumps(state))
            built = s.run("compile.py", "--state", str(s.root / "draft.json"),
                          "--dir", str(s.out), "--dir", str(s.ws), "--out", str(s.out))
            assert "error" not in built, built
            state["executed"] = {"files": built["files"], "index": built["index"],
                                 "at": "2026-09-28T10:00:00Z"}
            attach = built["files"] + [built["index"]]
        done = write_update(s, state, message, attach)
        assert "error" not in done, done
    return agent


def open_closing(env):
    env.send(email("Sig packets for the closing please.",
                   [attach("Falcon SPA.docx", agreement()),
                    attach("Disclosure Letter.docx", agreement(LETTER_FRONT))]),
             packets_agent)
    return sent_id(env)


# --- 1. packets ------------------------------------------------------------------


def test_sig_packets_for_the_closing_are_one_per_signatory_with_a_cover_sheet(env, office):
    open_closing(env)
    out = last(env)
    # Laid out, each document's blocks sit alone on its page 2.
    n = at(office, 2)
    assert out.text_body.startswith(
        f"0 of 4 in. Waiting on: Acme's director — SPA{n}, Disclosure Letter{n}; "
        f"Beta Trading's director — SPA{n}, Disclosure Letter{n}.")
    names = sorted(a.filename for a in out.attachments)
    assert names == ["Disclosure Letter (execution).docx", "Falcon SPA (execution).docx",
                     "Signature packet - Acme's director.pdf",
                     "Signature packet - Beta Trading's director.pdf"]
    packet = next(a for a in out.attachments if "Beta" in a.filename).content
    pages = [p.extract_text() for p in PdfReader(BytesIO(packet)).pages]
    assert len(pages) == 3
    assert "signature packet" in pages[0] and "How to return" in pages[0]
    assert "jim@firm.com" in pages[0]
    assert pages[1].startswith("Share Purchase Agreement — signature page")
    assert "BETA TRADING LLC" in pages[1] and "ACME" not in pages[1]
    assert "page spa--beta-director" in pages[1]
    assert pages[2].startswith("Disclosure Letter — signature page")

    stored = closing.open_for(OWNER)[0]
    assert stored.name == "Project Falcon"
    assert {n for n, k in closing.held(stored.id) if k == "execution"} == {
        "Falcon SPA (execution).docx", "Disclosure Letter (execution).docx"}
    session = env.sessions[-1].created[-1]
    assert session["agent"] == "agent_closing"


def test_a_document_not_ready_to_sign_builds_nothing_and_says_why_clause_first(env, bundle,
                                                                               tmp_path):
    front = list(SPA_FRONT)
    front[4] = "1. Sale. The Seller shall sell the Shares. [NTD: Seller to confirm]"
    (tmp_path / "SPA.docx").write_bytes(agreement(front))
    (tmp_path / "plan.json").write_text(json.dumps({**PLAN, "documents": [
        {"id": "spa", "file": "SPA.docx", "short": "SPA"}], "signatories": [
        {"id": "a", "name": "A", "sign": [{"document": "spa", "party": "Acme Holdings"}]}]}))
    got = run_script(bundle, "packets.py", "--plan", str(tmp_path / "plan.json"),
                     "--workspace", str(tmp_path), "--out", str(tmp_path / "out"),
                     "--draft", str(tmp_path / "d.json"))
    assert got["error"] == ("Not ready for signature: SPA: Clause 1: drafting note "
                            "[NTD: Seller to confirm] is still in")
    assert not (tmp_path / "out").exists()


def test_a_party_the_document_does_not_have_is_refused(tmp_path):
    (tmp_path / "SPA.docx").write_bytes(agreement())
    plan = {**PLAN, "documents": [{"id": "spa", "file": "SPA.docx"}], "signatories": [
        {"id": "g", "name": "Gamma", "sign": [{"document": "spa", "party": "Gamma Ltd"}]}]}
    with pytest.raises(closing_docs.ClosingError, match="has no party 'Gamma Ltd'"):
        closing_docs.packets(plan, tmp_path, tmp_path / "out")


def _fake_libreoffice(monkeypatch):
    """LibreOffice as the container has it: the signature blocks land on the
    document's last page, after 41 pages of text."""
    from reportlab.platypus import PageBreak, Paragraph

    def to_pdf(docx: bytes) -> bytes:
        paragraphs = closing_docs._docx_paragraphs(docx)

        def flow(styles):
            items = []
            if any(t.startswith("IN WITNESS") for t in paragraphs) and len(paragraphs) > 12:
                for n in range(40):
                    items += [Paragraph(f"Schedule page {n + 1}.", styles["body"]),
                              PageBreak()]
            for t in paragraphs:
                if t.startswith("IN WITNESS"):
                    items.append(PageBreak())
                items.append(Paragraph(closing_docs._esc(t) or "&nbsp;", styles["body"]))
            return items
        return closing_docs._render(flow)

    monkeypatch.setattr(closing_docs, "_libreoffice", lambda: True)
    monkeypatch.setattr(closing_docs, "to_pdf", to_pdf)


def test_with_a_word_processor_each_page_is_numbered_and_the_tally_names_it(tmp_path,
                                                                            monkeypatch):
    _fake_libreoffice(monkeypatch)
    (tmp_path / "Falcon SPA.docx").write_bytes(agreement())
    (tmp_path / "Disclosure Letter.docx").write_bytes(agreement(LETTER_FRONT))
    built = closing_docs.packets(PLAN, tmp_path, tmp_path / "out")
    spa = next(d for d in built["documents"] if d["id"] == "spa")
    assert spa["pages_total"] == 42
    numbers = {p["id"]: p["page"] for p in built["pages"]}
    assert numbers["spa--beta-director"] == 42
    cover = PdfReader(BytesIO((tmp_path / "out" / "Signature packet - Beta Trading's "
                                                 "director.pdf").read_bytes())).pages[0]
    assert "p.42" in cover.extract_text()


def test_the_tally_is_counted_from_the_state_and_names_what_is_outstanding():
    state = closing_state.empty("Project Falcon")
    state["documents"] = [{"id": "spa", "title": "Share Purchase Agreement", "short": "SPA"},
                          {"id": "dl", "title": "Disclosure Letter"},
                          {"id": "tsa", "title": "Transitional Services Agreement",
                           "short": "TSA"}]
    state["signatories"] = [{"id": "acme", "name": "Acme's director"},
                            {"id": "beta", "name": "Beta Trading's director"}]
    state["pages"] = [
        {"id": "spa--acme", "document": "spa", "signatory": "acme", "status": "received"},
        {"id": "dl--acme", "document": "dl", "signatory": "acme", "status": "received"},
        {"id": "tsa--acme", "document": "tsa", "signatory": "acme", "status": "received"},
        {"id": "spa--beta", "document": "spa", "signatory": "beta", "status": "awaiting",
         "page": 42},
        {"id": "dl--beta", "document": "dl", "signatory": "beta", "status": "awaiting",
         "page": 9},
    ]
    assert closing_state.tally(state) == (
        "3 of 5 in. Waiting on: Beta Trading's director — SPA p.42, "
        "Disclosure Letter p.9.")


# --- 2. signed pages ---------------------------------------------------------------


def test_signed_pages_are_read_filed_and_counted_and_an_unsigned_one_is_not(env, office):
    thread = open_closing(env)
    returned = scan(sign(packet_page(env, "Acme's director", 2)),
                    sign(packet_page(env, "Acme's director", 3)),
                    sign(packet_page(env, "Beta Trading's director", 3), signed=False))
    env.send(email("Acme's pages, and Beta's letter page.",
                   [attach("Signed pages.pdf", returned, "application/pdf")],
                   reply_to=thread),
             filing_agent({("Signed pages.pdf", 1): "spa--acme-director",
                           ("Signed pages.pdf", 2): "dl--acme-director"},
                          {("Signed pages.pdf", 3): "Disclosure Letter for Beta Trading is "
                                                    "unsigned."},
                          ["Filed 2: Acme's director — SPA, Disclosure Letter.",
                           "Not filed: Disclosure Letter for Beta Trading is unsigned."]))
    body = last(env).text_body
    n = at(office, 2)
    assert body.startswith(f"2 of 4 in. Waiting on: Beta Trading's director — SPA{n}, "
                           f"Disclosure Letter{n}.")
    assert "Not filed: Disclosure Letter for Beta Trading is unsigned." in body
    stored = closing.open_for(OWNER)[0]
    signed = {n for n, k in closing.held(stored.id) if k == "signed"}
    assert signed == {"signed - spa--acme-director.pdf", "signed - dl--acme-director.pdf"}
    page = closing.file_of(stored.id, "signed - spa--acme-director.pdf")
    assert len(PdfReader(BytesIO(page)).pages) == 1
    # The next session is given the state and what is held, not the packets.
    env.send(email("status", reply_to=sent_id(env)))
    assert last(env).text_body.startswith("2 of 4 in.")
    assert len(env.sessions) == 3 and env.sessions[-1].uploaded == []


def test_a_forwarded_photo_from_the_lawyer_finds_the_open_closing(env, office):
    open_closing(env)
    picture = photo(sign(packet_page(env, "Beta Trading's director", 2)))
    env.send(email("Fwd: signed signature pages from Beta's counsel",
                   [attach("IMG_0412.jpg", picture, "image/jpeg"),
                    attach("image001.png", b"\x89PNG\r\n\x1a\n" + b"0" * 900, "image/png")],
                   subject="Fwd: Falcon - Beta signature"),
             filing_agent({("IMG_0412.jpg", 1): "spa--beta-director"}, {},
                          ["Filed 1: Beta Trading's director — SPA (photo)."]))
    n = at(office, 2)
    assert last(env).text_body.startswith(
        f"1 of 4 in. Waiting on: Acme's director — SPA{n}, Disclosure Letter{n}; "
        f"Beta Trading's director — Disclosure Letter{n}.")
    mounted = {n for n, _ in env.sessions[-1].uploaded}
    assert "IMG_0412.jpg" in mounted and "image001.png" not in mounted
    assert closing.STATE_FILE in mounted


def test_a_colleague_on_the_thread_is_told_to_go_through_the_owner(env):
    thread = open_closing(env)
    page = sign(packet_page(env, "Beta Trading's director", 2))
    env.send(email("Beta's page attached.", [attach("p.pdf", page, "application/pdf")],
                   sender="sam@firm.com", reply_to=thread))
    assert last(env).to == ["sam@firm.com"]
    assert last(env).text_body.startswith("This closing is run by jim@firm.com")
    assert len(env.sessions) == 2 and env.sessions[-1].created == []
    assert closing_state.tally(closing.open_for(OWNER)[0].state).startswith("0 of 4 in.")


def test_the_other_side_writing_to_the_agent_directly_gets_nothing(env):
    thread = open_closing(env)
    before = len(env.provider.sent)
    page = sign(packet_page(env, "Beta Trading's director", 2))
    env.send(email("Signed pages attached", [attach("p.pdf", page, "application/pdf")],
                   sender="counsel@betalaw.com", reply_to=thread))
    assert len(env.provider.sent) == before
    assert env.sessions[-1].created == []


def test_an_update_that_files_an_unsigned_page_is_not_recorded(env):
    thread = open_closing(env)
    page = sign(packet_page(env, "Beta Trading's director", 2), signed=False)

    def careless(s):
        state = json.loads((s.ws / closing.STATE_FILE).read_text())
        s.run("receive.py", str(s.ws / "p.pdf"), "--page", "1",
              "--out", str(s.out / "signed - spa--beta-director.pdf"))
        p = next(p for p in state["pages"] if p["id"] == "spa--beta-director")
        p["status"] = "received"
        p["received"] = {"file": "signed - spa--beta-director.pdf", "source": "p.pdf, page 1",
                         "signed": False, "version_ref": p["ref"]}
        refused = write_update(s, state, ["Filed."], [])
        assert "NOTHING WAS WRITTEN" in refused["error"]
        assert "not marked signed" in refused["error"]
        # The model ignores the refusal and writes the file itself.
        (s.out / closing.UPDATE_FILE).write_text(json.dumps(
            {"schema": closing_state.UPDATE_SCHEMA, "state": state, "message": ["Filed."],
             "attach": []}))

    env.send(email("Beta's page", [attach("p.pdf", page, "application/pdf")],
                   reply_to=thread), careless)
    assert last(env).text_body.startswith("Nothing changed on the closing")
    stored = closing.open_for(OWNER)[0]
    assert closing_state.tally(stored.state).startswith("0 of 4 in.")
    assert not [n for n, k in closing.held(stored.id) if k == "signed"]


def test_a_page_signed_on_an_old_version_does_not_count(env, bundle, tmp_path):
    state = closing_state.empty("X")
    state["documents"] = [{"id": "spa", "title": "SPA", "file": "a.docx", "ref": "BBBBBB"}]
    state["signatories"] = [{"id": "b", "name": "B", "party": "B"}]
    state["pages"] = [{"id": "spa--b", "document": "spa", "signatory": "b", "ref": "BBBBBB",
                       "status": "received",
                       "received": {"file": "s.pdf", "source": "x", "signed": True,
                                    "version_ref": "AAAAAA"}}]
    errors = closing_state.validate_state(state, None, {"s.pdf"})
    assert errors == [("page spa--b: signed on version 'AAAAAA', not the current BBBBBB; "
                       "a stale page goes in rejected")]


def test_a_received_page_cannot_quietly_go_back_to_waiting():
    before = closing_state.empty("X")
    before["documents"] = [{"id": "spa", "title": "SPA", "file": "a.docx", "ref": "AAAAAA"}]
    before["signatories"] = [{"id": "b", "name": "B", "party": "B"}]
    before["pages"] = [{"id": "spa--b", "document": "spa", "signatory": "b", "ref": "AAAAAA",
                        "status": "received",
                        "received": {"file": "s.pdf", "source": "x", "signed": True,
                                     "version_ref": "AAAAAA"}}]
    after = json.loads(json.dumps(before))
    after["pages"][0].update(status="awaiting", received=None)
    assert any("was received and is now awaiting" in e
               for e in closing_state.validate_state(after, before, {"s.pdf"}))
    # A new version of the document is the one reason it may.
    after["documents"][0]["ref"] = after["pages"][0]["ref"] = "CCCCCC"
    assert closing_state.validate_state(after, before, set()) == []


# --- 3. the executed set -------------------------------------------------------------


def test_the_last_page_in_compiles_the_executed_set_and_the_closing_index(env):
    thread = open_closing(env)
    pages = {pid: sign(packet_page(env, who, n)) for pid, who, n in (
        ("spa--acme-director", "Acme's director", 2), ("dl--acme-director", "Acme's director", 3),
        ("spa--beta-director", "Beta Trading's director", 2),
        ("dl--beta-director", "Beta Trading's director", 3))}
    names = list(pages)
    env.send(email("All pages", [attach("all.pdf", scan(*pages.values()), "application/pdf")],
                   reply_to=thread),
             filing_agent({("all.pdf", i + 1): pid for i, pid in enumerate(names)}, {},
                          ["Filed 4. Executed set and closing index attached."]))
    out = last(env)
    assert out.text_body.startswith("All 4 in.\nExecuted set compiled: 2 documents and the "
                                    "closing index.")
    got = {a.filename: a.content for a in out.attachments}
    assert set(got) == {"Falcon SPA (executed).pdf", "Disclosure Letter (executed).pdf",
                        "Project Falcon - closing index.pdf"}
    spa = [p.extract_text() for p in PdfReader(BytesIO(got["Falcon SPA (executed).pdf"])).pages]
    assert "Seller shall sell the Shares" in spa[0]
    assert "page spa--acme-director" in spa[-2] and "page spa--beta-director" in spa[-1]
    # The execution copy's own signature page is gone, not kept beside the signed ones.
    assert len(spa) == 3 and "LRA ref" not in spa[0]
    index = " ".join(" ".join(p.extract_text().split()) for p in
                      PdfReader(BytesIO(got["Project Falcon - closing index.pdf"])).pages)
    for words in ("Project Falcon: closing index", "Share Purchase Agreement",
                  "Disclosure Letter", "3 March 2026", "ACME HOLDINGS LIMITED",
                  "Beta Trading's director", "dl--beta-director"):
        assert words in index
    stored = closing.load(closing.find(email("x", reply_to=thread)).id)
    assert stored.status == "complete"


def test_substitution_puts_the_signed_pages_where_the_execution_copys_were(tmp_path,
                                                                          monkeypatch):
    _fake_libreoffice(monkeypatch)
    (tmp_path / "Falcon SPA.docx").write_bytes(agreement())
    (tmp_path / "Disclosure Letter.docx").write_bytes(agreement(LETTER_FRONT))
    out = tmp_path / "out"
    built = closing_docs.packets(PLAN, tmp_path, out)
    state = closing_state.empty("Project Falcon")
    state.update(documents=built["documents"], pages=built["pages"],
                 signatories=[{"id": s["id"], "name": s["name"], "party": ""}
                              for s in PLAN["signatories"]])
    for p in state["pages"]:
        page, _ = closing_docs.receive(sign(_packet_page(out, p["signatory"], p["document"])),
                                       "scan.pdf")
        (out / f"signed - {p['id']}.pdf").write_bytes(page)
        p["status"] = "received"
        p["received"] = {"file": f"signed - {p['id']}.pdf", "source": "scan", "signed": True,
                         "version_ref": p["ref"]}
    result = closing_docs.compile_set(state, [out], tmp_path / "set")
    executed = PdfReader(BytesIO((tmp_path / "set" / "Falcon SPA (executed).pdf").read_bytes()))
    texts = [p.extract_text() for p in executed.pages]
    assert len(texts) == 43, "41 pages of text, then the two signed pages for page 42"
    assert "LRA ref" in texts[41] and "LRA ref" in texts[42]
    assert result["method"] == "libreoffice"


def _packet_page(out: Path, signatory: str, document: str) -> bytes:
    name = next(s["name"] for s in PLAN["signatories"] if s["id"] == signatory)
    reader = PdfReader(BytesIO((out / f"Signature packet - {name}.pdf").read_bytes()))
    n = 1 if document == "spa" else 2
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_page(reader.pages[n])
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_a_signature_block_sharing_a_page_with_the_agreement_keeps_every_word(
        tmp_path, monkeypatch, office):
    """What a short document does under the real LibreOffice: the blocks sit
    on page 1 with the text. The page number is the true one, p.1, and the
    executed copy keeps that page, with the signed pages after it, rather than
    replacing it and losing the agreement's words."""
    if not office:
        pytest.skip("a layout question; without a word processor there is no page")
    monkeypatch.setenv("LRA_FAKE_SOFFICE_LAYOUT", "one")
    (tmp_path / "Falcon SPA.docx").write_bytes(agreement())
    (tmp_path / "Disclosure Letter.docx").write_bytes(agreement(LETTER_FRONT))
    out = tmp_path / "out"
    built = closing_docs.packets(PLAN, tmp_path, out)
    assert built["method"] == "libreoffice"
    assert {p["page"] for p in built["pages"]} == {1}
    state = closing_state.empty("Project Falcon")
    state.update(documents=built["documents"], pages=built["pages"],
                 signatories=[{"id": s["id"], "name": s["name"], "party": ""}
                              for s in PLAN["signatories"]])
    for p in state["pages"]:
        page, _ = closing_docs.receive(sign(_packet_page(out, p["signatory"], p["document"])),
                                       "scan.pdf")
        (out / f"signed - {p['id']}.pdf").write_bytes(page)
        p["status"] = "received"
        p["received"] = {"file": f"signed - {p['id']}.pdf", "source": "scan", "signed": True,
                         "version_ref": p["ref"]}
    closing_docs.compile_set(state, [out], tmp_path / "set")
    spa = [p.extract_text() for p in PdfReader(BytesIO(
        (tmp_path / "set" / "Falcon SPA (executed).pdf").read_bytes())).pages]
    assert len(spa) == 3
    assert "Seller shall sell the Shares" in spa[0] and "LRA ref" not in spa[0]
    assert "page spa--acme-director" in spa[1] and "page spa--beta-director" in spa[2]


def test_a_document_that_cannot_be_converted_gets_no_page_number(tmp_path, monkeypatch,
                                                                office):
    monkeypatch.setenv("LRA_FAKE_SOFFICE_LAYOUT", "fail")
    (tmp_path / "Falcon SPA.docx").write_bytes(agreement())
    (tmp_path / "Disclosure Letter.docx").write_bytes(agreement(LETTER_FRONT))
    built = closing_docs.packets(PLAN, tmp_path, tmp_path / "out")
    assert built["method"] == "text"
    assert {p["page"] for p in built["pages"]} == {None}


def test_compile_refuses_while_a_page_is_outstanding(tmp_path):
    state = closing_state.empty("X")
    state["pages"] = [{"id": "a--b", "document": "a", "signatory": "b", "status": "awaiting",
                       "ref": "A"}]
    with pytest.raises(closing_docs.ClosingError, match="Not all pages are in: a--b"):
        closing_docs.compile_set(state, [tmp_path], tmp_path)


# --- 4. the checklist ----------------------------------------------------------------


def test_the_closing_checklist_is_kept_on_the_thread_and_follows_the_pages(env):
    thread = open_closing(env)

    def checklist_agent(s):
        state = json.loads((s.ws / closing.STATE_FILE).read_text())
        state["checklist"] = [
            {"id": "cp-clearance", "kind": "condition", "item": "Merger clearance",
             "source": "SPA cl. 4.1(a)", "responsible": "Buyer", "status": "open"},
            {"id": "cp-board", "kind": "condition", "item": "Buyer board approval",
             "source": "SPA cl. 4.1(b)", "responsible": "Buyer", "status": "open"},
            {"id": "del-certs", "kind": "deliverable",
             "item": "Share certificates and stock transfer form", "source": "SPA cl. 6",
             "responsible": "Seller", "status": "open"},
            {"id": "sign-spa", "kind": "signing", "item": "SPA signed by both parties",
             "source": "SPA", "responsible": "All", "status": "done",
             "pages": ["spa--acme-director", "spa--beta-director"]},
        ]
        (s.root / "draft.json").write_text(json.dumps(state))
        s.run("checklist.py", "--state", str(s.root / "draft.json"),
              "--out", str(s.out / "Closing checklist.pdf"))
        done = write_update(s, state, ["SPA: 2 conditions (cl. 4.1), 1 deliverable (cl. 6)."],
                            ["Closing checklist.pdf"])
        assert "error" not in done, done

    env.send(email("closing checklist please", reply_to=thread), checklist_agent)
    out = last(env)
    assert out.text_body.startswith("0 of 4 in.")
    assert "Checklist: 0 of 4 done." in out.text_body, "the signing item follows its pages"
    text = PdfReader(BytesIO(out.attachments[0].content)).pages[0].extract_text()
    assert "SPA cl. 4.1(a)" in text and "Merger clearance" in text
    stored = closing.open_for(OWNER)[0]
    assert [i["id"] for i in stored.state["checklist"]][:1] == ["cp-clearance"]


def test_a_signing_item_is_done_exactly_when_its_pages_are_in():
    state = closing_state.empty("X")
    state["pages"] = [{"id": "p1", "status": "received"}, {"id": "p2", "status": "awaiting"}]
    state["checklist"] = [{"id": "s", "kind": "signing", "item": "i", "status": "done",
                           "pages": ["p1", "p2"]}]
    assert closing_state.normalise(state)["checklist"][0]["status"] == "open"
    state["pages"][1]["status"] = "received"
    assert closing_state.normalise(state)["checklist"][0]["status"] == "done"


# --- routing ------------------------------------------------------------------------


@pytest.mark.parametrize("body,expected", [
    ("Sig packets for the closing please.", closing.START),
    ("Can you prepare the signing packets", closing.START),
    ("Closing checklist please", closing.START),
    ("sig pages please", ""),
    ("Please review this SPA.", ""),
])
def test_what_opens_a_closing(env, body, expected):
    got = closing.route(email(body, [attach("SPA.docx", agreement())]))
    assert got[0] == expected


def test_off_the_thread_a_word_file_is_a_review_even_with_a_closing_open(env):
    open_closing(env)
    msg = email("Please review the signed NDA", [attach("NDA.docx", agreement())])
    assert closing.route(msg)[0] == ""


def test_without_the_closing_agent_nothing_changes_and_it_says_why(env, monkeypatch):
    monkeypatch.setenv("MANAGED_CLOSING_AGENT_ID", "")
    lra.config.settings.cache_clear()
    env.send(email("Sig packets for the closing please.", [attach("SPA.docx", agreement())]))
    assert last(env).text_body.startswith("Closings by email aren't set up yet")
    assert env.sessions[-1].created == []


def test_a_first_session_that_records_nothing_leaves_no_closing_behind(env):
    env.send(email("Sig packets for the closing please.", [attach("SPA.docx", agreement())]))
    assert last(env).text_body.startswith("Nothing changed on the closing")
    assert closing.open_for(OWNER) == []


def test_purging_a_lawyer_takes_their_closings_and_signed_pages(env):
    open_closing(env)
    cid = closing.open_for(OWNER)[0].id
    assert closing.purge(OWNER) == 1
    assert closing.load(cid) is None and closing.held(cid) == []


def test_two_open_closings_and_a_forward_that_names_neither_is_asked_about(env):
    closing.create("c1", OWNER, "Project Falcon")
    closing.create("c2", OWNER, "Project Osprey")
    page = sign(scan(closing_docs._text_page("X — signature page", ["Signed"], "")))
    env.send(email("Signed pages attached", [attach("p.pdf", page, "application/pdf")],
                   subject="Pages"))
    assert last(env).text_body.startswith("Which closing are these for: Project")


# --- the skill and the agent ----------------------------------------------------------


def test_the_closing_skill_bundle_carries_the_real_modules_and_a_valid_frontmatter(bundle):
    head = (bundle / "SKILL.md").read_text().split("---")[1]
    fields = dict(line.split(":", 1) for line in head.strip().splitlines())
    assert fields["name"].strip() == "lra-closing"
    assert 0 < len(fields["description"].strip()) <= 1024
    for relative in skillsync.CLOSING_BUNDLED:
        assert (bundle / "lib" / "lra" / relative).read_bytes() == (
            skillsync.PACKAGE / relative).read_bytes()
    assert (bundle / "shims" / "pydantic" / "__init__.py").exists()
    assert (bundle / "reference" / "state.md").exists()
    for relative in skillsync.CLOSING_BUNDLED:
        top = [line for line in (bundle / "lib" / "lra" / relative).read_text().splitlines()
               if line.startswith(("import ", "from "))]
        assert not [line for line in top if "lra.config" in line or "anthropic" in line
                    or "lra.pipeline import intake" in line], relative


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_the_scripts_run_from_the_bundle(bundle, tmp_path, shim):
    (tmp_path / "Falcon SPA.docx").write_bytes(agreement())
    (tmp_path / "Disclosure Letter.docx").write_bytes(agreement(LETTER_FRONT))
    got = run_script(bundle, "survey.py", str(tmp_path / "Falcon SPA.docx"), shim=shim)
    assert got["documents"][0]["parties"][1]["name"] == "BETA TRADING LLC"
    (tmp_path / "plan.json").write_text(json.dumps(PLAN))
    built = run_script(bundle, "packets.py", "--plan", str(tmp_path / "plan.json"),
                       "--workspace", str(tmp_path), "--out", str(tmp_path / "out"),
                       "--draft", str(tmp_path / "draft.json"), shim=shim)
    assert built["tally"].startswith("0 of 4 in.")
    packet = tmp_path / "out" / "Signature packet - Acme's director.pdf"
    got = run_script(bundle, "receive.py", str(packet), "--page", "2",
                     "--out", str(tmp_path / "one.pdf"), shim=shim)
    assert got["stamp"]["page"] == "spa--acme-director" and got["pages_in_file"] == 3
    bad = run_script(bundle, "receive.py", str(tmp_path / "plan.json"),
                     "--out", str(tmp_path / "x.pdf"), shim=shim)
    assert bad["error"] == "plan.json is not a PDF or an image"


def test_the_closing_agent_has_no_custom_tools_and_its_own_skill(env, monkeypatch):
    monkeypatch.setenv("SANDBOX_CLOSING_SKILL_ID", "skill_closing")
    lra.config.settings.cache_clear()
    body = managed.agent_body(managed.load_manifest(managed.AGENTS / "closing.agent.yaml"),
                              closing.CUSTOM_TOOLS)
    assert body["system"] == closing.SYSTEM
    assert [t["type"] for t in body["tools"]] == ["agent_toolset_20260401"]
    assert body["skills"][-1] == {"type": "custom", "skill_id": "skill_closing",
                                  "version": "latest"}
    assert "extra_skills" not in body
    review = managed.agent_body(managed.load_manifest(managed.AGENTS / "review.agent.yaml"), [])
    assert "skill_closing" not in [s["skill_id"] for s in review["skills"]]
