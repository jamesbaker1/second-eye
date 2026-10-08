"""The other side's draft: detected conservatively, answered not polished.

Open question O17. Detection is the lawyer's own words first, then a forward
from outside the firm, then an email address outside the firm in the
document's metadata; a name is never judged. On their paper nothing cosmetic
becomes a tracked change in their text, the substantive points become an
issues list, and the verdict counts what to push back on. The model is
stubbed throughout: what is tested is everything around it.
"""

from __future__ import annotations

import zipfile
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from secondeye import handler
from secondeye.mail.console import ConsoleProvider
from secondeye.models import Attachment, Finding, InboundEmail, Mode, ReviewResult, Severity
from secondeye.pipeline import issues_list, redline, reply, review, their_paper
from secondeye.pipeline.identity import EntryMode
from tests.conftest import documents

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SAMPLE = Path("samples/simple.docx")


@pytest.fixture(autouse=True)
def firm(monkeypatch, tmp_path):
    """jim@firm.com is inside the firm; everyone else is not."""
    from secondeye.config import settings

    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/t.sqlite3")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@firm.com")
    settings.cache_clear()
    yield
    settings.cache_clear()


def an_email(body="Can you look at this?", subject="MSA", content=None,
             to=("review@firm.com",)) -> InboundEmail:
    raw = content if content is not None else SAMPLE.read_bytes()
    return InboundEmail(
        message_id=f"m-{subject}-{abs(hash(body))}", from_address="jim@firm.com",
        to=list(to), subject=subject, text_body=body,
        attachments=[Attachment(filename="Acme MSA.docx", content_type=DOCX,
                                size_bytes=len(raw), content=raw)],
        received_at=datetime.now(UTC),
    )


GMAIL_FORWARD = """Can you look at this before my call?

---------- Forwarded message ---------
From: Dana Other <dana@otherside.com>
Date: Mon, 21 Sep 2026 at 10:02
Subject: MSA
To: Jim <jim@firm.com>

Please find our draft attached. Please delete the indemnity.
"""

OUTLOOK_FORWARD = """Quick look please.

________________________________
From: Dana Other <dana@otherside.com>
Sent: 21 September 2026 10:02
To: Jim Baker
Subject: MSA
"""


def detect(email, content=None, entry=EntryMode.FORWARD, mode=Mode.REDLINE):
    return their_paper.detect(email, content, entry, mode)


# --- detection: the lawyer's words -------------------------------------------

@pytest.mark.parametrize("said", [
    "Their draft, can you check it?",
    "This came in on their paper this morning.",
    "Opposing counsel's latest - please review.",
    "Just in from the other side.",
    "The other side’s draft of the MSA.",
    "From opposing counsel, can you check the cap?",
])
def test_the_lawyer_saying_it_is_theirs_is_enough(said):
    got = detect(an_email(said))
    assert got.theirs and got.source == "words"


@pytest.mark.parametrize("said", [
    "Can you check this before it goes to opposing counsel?",
    "I want to send this to the other side today.",
    "Quick look?",
])
def test_naming_the_other_side_is_not_calling_it_theirs(said):
    assert not detect(an_email(said)).theirs


def test_our_draft_overrides_a_forward_from_outside():
    """The lawyer's words win, whichever way they go."""
    e = an_email(GMAIL_FORWARD.replace("Can you look at this before my call?",
                                       "This is our draft, they sent it back unchanged."),
                 subject="Fwd: MSA")
    got = detect(e)
    assert not got.theirs and got.source == "words"


def test_not_their_paper_is_read_as_ours():
    assert not detect(an_email("Not their paper - it is ours, check it")).theirs


def test_their_words_in_the_forwarded_message_are_not_the_lawyers():
    """Opposing counsel writing "our draft" must not claim it for the firm,
    and the counterparty's covering note is never the lawyer speaking."""
    e = an_email(GMAIL_FORWARD, subject="Fwd: MSA")
    got = detect(e)
    assert got.theirs and got.source == "forward"


def test_their_draft_and_our_markup_together_is_left_to_the_evidence():
    assert their_paper.from_words("Our mark-up of their draft") is None


# --- detection: forwards -----------------------------------------------------

def test_a_gmail_forward_from_outside_the_firm_is_theirs():
    got = detect(an_email(GMAIL_FORWARD, subject="Fwd: MSA"))
    assert got.theirs and "dana@otherside.com" in got.reason


def test_an_outlook_forward_from_outside_the_firm_is_theirs():
    got = detect(an_email(OUTLOOK_FORWARD, subject="FW: MSA"))
    assert got.theirs and got.source == "forward"


def test_a_reply_chain_is_not_a_forward():
    """Outlook quotes a reply with the same From/Sent block as a forward."""
    assert not detect(an_email(OUTLOOK_FORWARD, subject="RE: MSA")).theirs


def test_a_forward_from_a_colleague_is_ours():
    body = GMAIL_FORWARD.replace("dana@otherside.com", "partner@firm.com")
    assert not detect(an_email(body, subject="Fwd: MSA")).theirs


def test_a_forward_of_the_lawyers_own_mail_is_ours(monkeypatch):
    """Production's lawyer writes from a personal address outside the firm's
    domains; forwarding their own message must not make it the other side's."""
    body = GMAIL_FORWARD.replace("dana@otherside.com", "jim@firm.com")
    assert not detect(an_email(body, subject="Fwd: MSA")).theirs
    monkeypatch.setenv("ALLOWED_SENDERS", "fictional.lawyer@gmail.com")
    from secondeye.config import settings

    settings.cache_clear()
    body = GMAIL_FORWARD.replace("dana@otherside.com", "fictional.lawyer@gmail.com")
    assert not detect(an_email(body, subject="Fwd: MSA")).theirs


def test_a_forward_header_with_no_address_says_nothing():
    body = OUTLOOK_FORWARD.replace("Dana Other <dana@otherside.com>", "Dana Other")
    assert not detect(an_email(body, subject="FW: MSA")).theirs


def test_a_document_already_sent_is_never_theirs():
    e = an_email("Attached is our response to their draft.", to=("client@acme.com",))
    assert not detect(e, entry=EntryMode.BCC_SENT).theirs


def test_a_proofread_or_a_question_is_left_as_asked():
    e = an_email(GMAIL_FORWARD, subject="Fwd: MSA")
    assert not detect(e, mode=Mode.PROOFREAD).theirs
    assert not detect(e, mode=Mode.QUESTION).theirs
    assert detect(e, mode=Mode.MEMO_ONLY).theirs


# --- detection: the document's own metadata ----------------------------------

def with_metadata(author="", last="", revision_author="") -> bytes:
    d = Document()
    d.add_paragraph("1. Liability. The Supplier's liability is capped at one month of fees.")
    d.core_properties.author = author
    d.core_properties.last_modified_by = last
    buf = BytesIO()
    d.save(buf)
    content = buf.getvalue()
    if not revision_author:
        return content
    src = zipfile.ZipFile(BytesIO(content))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(
                    b"<w:body>",
                    b'<w:body><w:p><w:ins w:id="90" w:author="' + revision_author.encode()
                    + b'" w:date="2026-09-20T00:00:00Z"><w:r><w:t>New</w:t></w:r></w:ins></w:p>',
                )
            dst.writestr(item, data)
    return out.getvalue()


def test_an_outside_address_as_last_saved_by_is_theirs():
    content = with_metadata(last="dana@otherside.com")
    got = detect(an_email(content=content), content)
    assert got.theirs and got.reason == "it was last saved by dana@otherside.com"


def test_an_outside_tracked_change_author_is_theirs():
    content = with_metadata(revision_author="Dana Other (dana@otherside.com)")
    got = detect(an_email(content=content), content)
    assert got.theirs and "tracked changes are by dana@otherside.com" in got.reason


def test_a_name_is_never_judged():
    """"Dana Other" could be a colleague. Guessing is how the firm's own
    draft loses its typo fixes."""
    content = with_metadata(author="Dana Other", last="J Smith", revision_author="Dana Other")
    assert not detect(an_email(content=content), content).theirs


def test_an_inside_address_in_the_metadata_is_ours():
    content = with_metadata(author="jim@firm.com", last="partner@legal.firm.com")
    assert not detect(an_email(content=content), content).theirs


def test_a_file_that_is_not_word_says_nothing():
    assert their_paper.metadata_addresses(b"%PDF-1.7 not a zip") == []


# --- what changes on their paper ---------------------------------------------

def finding(severity, category="playbook", **kw) -> Finding:
    base = {"severity": severity, "category": category, "title": f"{category} point",
                "explanation": "Why it matters.", "anchor": "capped at one month of fees"}
    base.update(kw)
    return Finding(**base)


def test_cosmetic_findings_are_withheld_and_counted_blockers_never():
    fs = [finding(Severity.STYLE, "style"), finding(Severity.FORMATTING, "numbering"),
          finding(Severity.BLOCKER, "placeholder"), finding(Severity.SUBSTANTIVE)]
    kept, withheld = their_paper.withhold_cosmetic(fs)
    assert withheld == 2
    assert [f.severity for f in kept] == [Severity.BLOCKER, Severity.SUBSTANTIVE]


def test_the_note_says_why_and_how_to_say_otherwise():
    p = their_paper.Provenance(True, "forward", "you forwarded it from dana@otherside.com")
    line = their_paper.note(p, 7)
    assert "dana@otherside.com" in line and "7 typo and style points" in line
    assert '"our draft"' in line
    said = their_paper.note(their_paper.Provenance(True, "words", "you said so"), 0)
    assert "our draft" not in said and "typo" not in said


# --- the issues list ---------------------------------------------------------

PLAYBOOK = Finding(
    severity=Severity.SUBSTANTIVE, category="playbook",
    title="Off playbook: Limitation of liability - cap at one month's fees",
    explanation="Starter playbook (not yet your firm's): their cap is one month.",
    anchor="The Supplier's liability is capped at one month of fees.",
    where="clause 1", our_position="A cap of 12 months' fees, with the usual carve-outs.",
    response="Ask for 12 months' fees; accept 6 at most.",
)


def table_of(content: bytes) -> list[list[str]]:
    table = Document(BytesIO(content)).tables[0]
    return [[c.text for c in row.cells] for row in table.rows]


def test_the_issues_list_is_a_verified_four_column_table():
    content = issues_list.build([PLAYBOOK], "Acme MSA.docx")
    assert content is not None
    assert redline.verify(content) == (True, "ok")
    rows = table_of(content)
    assert rows[0] == list(issues_list.COLUMNS)
    assert rows[1] == [
        "Clause 1: Limitation of liability",
        "“The Supplier's liability is capped at one month of fees.”",
        "A cap of 12 months' fees, with the usual carve-outs.",
        "Ask for 12 months' fees; accept 6 at most.",
    ]


def test_the_issues_list_says_nothing_about_how_it_was_made():
    content = issues_list.build([PLAYBOOK], "Acme MSA.docx")
    core = zipfile.ZipFile(BytesIO(content)).read("docProps/core.xml").decode()
    assert "python-docx" not in core


def test_a_point_without_a_position_falls_back_to_what_the_finding_says():
    bare = Finding(severity=Severity.SUBSTANTIVE, category="indemnity",
                   title="Cl. 5: indemnity is one-way", explanation="Only we indemnify.",
                   anchor="The Customer shall indemnify the Supplier",
                   suggested_text="Each party shall indemnify the other")
    (clause, _said, position, response), = issues_list.rows([bare])
    assert clause == "indemnity is one-way"
    assert position == "Only we indemnify."
    assert response == "Replace with: “Each party shall indemnify the other”"
    asked = bare.model_copy(update={"suggested_text": None, "question": "Mutual?"})
    assert issues_list.rows([asked])[0][3] == "Mutual?"
    plain = bare.model_copy(update={"suggested_text": None})
    assert issues_list.rows([plain])[0][3] == "Raise with them."


def test_a_long_clause_is_clipped_in_the_table():
    long = PLAYBOOK.model_copy(update={"anchor": "word " * 200})
    said = issues_list.rows([long])[0][1]
    assert len(said) <= issues_list.QUOTE_LIMIT + 2 and said.endswith("…”")


def test_nothing_to_list_is_no_file():
    assert issues_list.build([], "Acme MSA.docx") is None


def test_the_name_follows_the_document():
    assert issues_list.name("Acme MSA.pdf") == "Acme MSA (issues list).docx"


def test_settle_does_nothing_on_our_paper():
    result = ReviewResult(mode=Mode.REDLINE, summary="",
                          findings=[finding(Severity.STYLE, "style")])
    assert their_paper.settle(result, their_paper.OURS, "a.docx", 10**7) == ([], [])
    assert len(result.findings) == 1 and not result.their_paper


# --- the model is told ----------------------------------------------------------

def test_the_first_message_carries_the_rule_only_on_their_paper():
    from secondeye.pipeline import extract

    doc = extract.extract(Attachment(filename="a.docx", content_type=DOCX,
                                     size_bytes=0, content=SAMPLE.read_bytes()))
    args = (doc, Mode.REDLINE, "", None, "", "", "", [])
    assert review.THEIR_PAPER_RULES not in review.first_message(*args)
    theirs = review.first_message(*args, their_paper=True)
    assert review.THEIR_PAPER_RULES in theirs
    assert review.MODE_RULES[Mode.REDLINE] in theirs
    # Outside the document's fence: it is our instruction, not their text.
    assert theirs.index(review.THEIR_PAPER_RULES) < theirs.index("<document-")


def test_the_report_schema_takes_a_position_and_a_response():
    props = review.findings_schema()["properties"]["findings"]["items"]["properties"]
    assert "our_position" in props and "response" in props
    parsed = review._Report()
    parsed.record("", [{"severity": "substantive", "category": "playbook",
                        "title": "t", "explanation": "e", "anchor": "a",
                        "our_position": "p", "response": "r"}])
    assert parsed.findings[0].our_position == "p" and parsed.findings[0].response == "r"


# --- the verdict -----------------------------------------------------------------

def verdict(findings):
    result = ReviewResult(mode=Mode.REDLINE, summary="", findings=findings,
                          their_paper=True)
    out = reply.compose(an_email(), result, None, "Acme MSA.docx", [], findings,
                        their_paper=True)
    return out.text_body.splitlines()[0], out.text_body


def test_the_verdict_counts_points_to_push_back_on():
    points = [finding(Severity.SUBSTANTIVE) for _ in range(5)]
    first, body = verdict(points)
    assert first == "Their draft: 5 points to push back on."
    assert "To push back on:" in body and "Your call" not in body


def test_the_verdict_names_errors_in_their_draft_too():
    first, body = verdict([finding(Severity.SUBSTANTIVE),
                           finding(Severity.BLOCKER, "placeholder")])
    assert first == "Their draft: 1 point to push back on, and 1 error in it."
    assert "Errors in their draft:" in body and "Fix before sending" not in body
    assert verdict([finding(Severity.BLOCKER, "placeholder")])[0] == (
        "Their draft: 1 error in it, nothing to push back on.")
    assert verdict([])[0] == "Their draft: nothing to push back on."


# --- end to end through the handler ---------------------------------------------

class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return "captured"


TYPO = Finding(severity=Severity.FORMATTING, category="typo", title="Cl. 1: typo",
               explanation="", anchor="disclosed by the Discloser to the Reciever",
               suggested_text="disclosed by the Discloser to the Receiver", auto_apply=True)
POINT = Finding(
    severity=Severity.SUBSTANTIVE, category="playbook",
    title="Off playbook: Governing law - New York", where="clause 4",
    explanation="Starter playbook (not yet your firm's): we ask for English law.",
    anchor="This Agreement is governed by the laws of the State of New York",
    our_position="English law and the courts of England.",
    response="Ask for English law; New York only if they take the cap.",
)


@pytest.fixture
def captured(monkeypatch):
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    return provider


def stubbed(monkeypatch, seen: dict, outputs=None):
    def _review(doc, mode, instructions, **kwargs):
        seen.update(kwargs)
        result = ReviewResult(mode=mode, summary="", findings=[TYPO, POINT],
                              their_paper=kwargs.get("their_paper", False))
        result.outputs = outputs or []
        return result

    monkeypatch.setattr(handler.review, "review", _review)
    monkeypatch.setattr(handler.checks, "run_all", lambda *a, **k: [])
    monkeypatch.setattr(handler.email_checks, "run_all", lambda *a, **k: [])


def test_a_forwarded_draft_from_the_other_side_gets_an_issues_list(captured, monkeypatch):
    seen: dict = {}
    stubbed(monkeypatch, seen)
    handler.handle(an_email(GMAIL_FORWARD, subject="Fwd: MSA"))

    assert seen["their_paper"] is True
    out = captured.sent[0]
    body = out.text_body
    assert body.startswith("Their draft: 1 point to push back on.")
    assert "Acme MSA (issues list).docx has the point as a table" in body
    assert "1 typo and style point left alone" in body
    names = [a.filename for a in documents(out)]
    # No redline: the only change on offer was their typo.
    assert names == ["Acme MSA (issues list).docx"]
    rows = table_of(documents(out)[0].content)
    assert rows[1][0] == "Clause 4: Governing law"
    assert "Reciever" not in body


def test_the_sessions_redline_of_their_typos_is_not_attached(captured, monkeypatch):
    """The model may still write their typo into the session's redline. The
    local writer runs without it, the two differ, and the session's is
    thrown away."""
    marked = redline.apply(SAMPLE.read_bytes(), ReviewResult(
        mode=Mode.REDLINE, summary="", findings=[TYPO])).content
    stubbed(monkeypatch, {}, outputs=[("Acme MSA (redline).docx", marked)])
    handler.handle(an_email(GMAIL_FORWARD, subject="Fwd: MSA"))
    assert "(redline)" not in " ".join(a.filename for a in documents(captured.sent[0]))


def test_our_own_draft_is_reviewed_as_before(captured, monkeypatch):
    seen: dict = {}
    stubbed(monkeypatch, seen)
    handler.handle(an_email("Quick look before I send?"))

    assert seen["their_paper"] is False
    out = captured.sent[0]
    assert not out.text_body.startswith("Their draft")
    names = [a.filename for a in documents(out)]
    assert names == ["Acme MSA (redline).docx"]
