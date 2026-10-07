"""Comparison: what changed between two versions, as real tracked changes.

The claim the email makes is that accepting every change yields the later
version and rejecting every change yields the earlier one. Most of these tests
hold the output to that claim through a reader that shares no code with the
writer: python-docx for structure, and a regex over the raw XML for the two
views.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO

from docx import Document
from docx.shared import Pt

from lra.pipeline import compare, redline
from lra.pipeline.validate import validate

EARLIER = [
    "1. Definitions",
    '1.1 "Term" means thirty (30) days from the Effective Date.',
    "2. Payment. The Buyer shall pay the Price within ten days of invoice.",
    ("3. Liability. The Supplier's liability is capped at the fees paid in the "
    "prior twelve months."),
    "4. Notices shall be sent by post to the registered office of the recipient.",
    "5. Governing law. This Agreement is governed by the laws of England and Wales.",
]


def build(paragraphs) -> bytes:
    d = Document()
    for text in paragraphs:
        d.add_paragraph(text)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def body_xml(content: bytes) -> str:
    return zipfile.ZipFile(BytesIO(content)).read("word/document.xml").decode()


def _paragraphs(xml: str) -> list[str]:
    return re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.DOTALL)


def accepted(content: bytes) -> list[str]:
    """Accept-all, by regex: drop deletions, keep everything else."""
    out = []
    for p in _paragraphs(body_xml(content)):
        p = re.sub(r"<w:del [^>]*>.*?</w:del>", "", p, flags=re.DOTALL)
        text = "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", p))
        if text.strip():
            out.append(_unescape(text))
    return out


def rejected(content: bytes) -> list[str]:
    """Reject-all, by regex: drop insertions, read deleted text as text."""
    out = []
    for p in _paragraphs(body_xml(content)):
        p = re.sub(r"<w:ins [^>]*>.*?</w:ins>", "", p, flags=re.DOTALL)
        text = "".join(re.findall(r"<w:(?:t|delText)(?: [^>]*)?>([^<]*)</w:(?:t|delText)>", p))
        if text.strip():
            out.append(_unescape(text))
    return out


def _unescape(text: str) -> str:
    return (text.replace("&quot;", '"').replace("&apos;", "'")
            .replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&"))


def test_identical_versions_produce_nothing():
    raw = build(EARLIER)
    result = compare.compare(raw, build(EARLIER))
    assert result.identical
    assert result.content is None


def test_a_changed_number_is_one_tracked_replacement():
    later = list(EARLIER)
    later[1] = '1.1 "Term" means sixty (60) days from the Effective Date.'
    result = compare.compare(build(EARLIER), build(later))

    assert [c.kind for c in result.changes] == ["replaced"]
    assert result.changes[0].where == "1.1"
    assert result.proven
    xml = body_xml(result.content)
    # One phrase, not two edits with a space between them.
    assert "sixty (60<" in xml
    assert xml.count("<w:ins ") == 1 and xml.count("<w:del ") == 1
    assert accepted(result.content) == later
    assert rejected(result.content) == EARLIER


def test_inserted_and_deleted_paragraphs_round_trip():
    later = list(EARLIER)
    later.insert(3, "2A. Interest accrues on late payments at 8% per annum.")
    del later[5]                                   # clause 4, the notices
    result = compare.compare(build(EARLIER), build(later))

    assert sorted(c.kind for c in result.changes) == ["deleted", "inserted"]
    assert result.proven
    assert accepted(result.content) == later
    assert rejected(result.content) == EARLIER


def test_a_deleted_paragraph_has_its_mark_deleted_too():
    """Otherwise accepting the change leaves an empty numbered clause behind."""
    later = EARLIER[:3] + EARLIER[4:]
    result = compare.compare(build(EARLIER), build(later))
    xml = body_xml(result.content)
    assert re.search(r"<w:pPr>.*?<w:rPr><w:del [^>]*/></w:rPr>", xml, flags=re.DOTALL)


def test_a_paragraph_inserted_at_the_very_top():
    later = ["DRAFT 2 - SUBJECT TO CONTRACT"] + EARLIER
    result = compare.compare(build(EARLIER), build(later))
    assert result.proven
    assert accepted(result.content) == later


def test_a_clause_that_moves_and_loses_its_cap_is_one_change():
    later = list(EARLIER)
    moved = later.pop(3)
    later.insert(4, moved.replace("is capped at the fees paid in the prior twelve months",
                                  "is not capped"))
    result = compare.compare(build(EARLIER), build(later))
    kinds = [c.kind for c in result.changes]
    assert kinds == ["moved"]
    assert "not capped" in result.changes[0].after
    assert result.proven


def test_a_rewritten_paragraph_is_swapped_whole_not_shredded():
    later = list(EARLIER)
    later[4] = ("4. Notices may be given by email to the address each party has "
                "most recently notified, and take effect on the next business day.")
    result = compare.compare(build(EARLIER), build(later))
    xml = body_xml(result.content)
    assert xml.count("<w:ins ") <= 2
    assert accepted(result.content) == later


def test_formatting_in_the_earlier_version_survives():
    d = Document()
    p = d.add_paragraph("The term is ")
    bold = p.add_run("thirty")
    bold.bold = True
    bold.font.size = Pt(14)
    p.add_run(" days from signing.")
    buf = BytesIO()
    d.save(buf)

    result = compare.compare(buf.getvalue(), build(["The term is sixty days from signing."]))
    assert result.proven
    xml = body_xml(result.content)
    inserted = re.search(r"<w:ins [^>]*>(.*?)</w:ins>", xml, flags=re.DOTALL).group(1)
    assert "<w:b/>" in inserted, "the replacement should carry the bold it replaced"


def test_output_passes_the_same_gate_as_every_redline():
    later = list(EARLIER)
    later[5] = "5. Governing law. This Agreement is governed by the laws of New York."
    later.append("6. Counterparts. This Agreement may be signed in counterparts.")
    result = compare.compare(build(EARLIER), build(later))
    ok, reason = redline.verify(result.content, expect_revisions=True)
    assert ok, reason
    report = validate(result.content)
    assert report.ok and report.authors == {compare.AUTHOR}


def test_another_authors_revisions_are_left_alone():
    d = Document(BytesIO(build(EARLIER)))
    from lra.pipeline.ooxml import Revision, RevisionWriter

    RevisionWriter(d, "Opposing Counsel").apply(
        Revision("ten days", "fourteen days", "Opposing Counsel"))
    buf = BytesIO()
    d.save(buf)
    earlier = buf.getvalue()

    later = list(EARLIER)
    later[2] = later[2].replace("ten days", "fourteen days")
    later[5] = later[5].replace("England and Wales", "New York")
    result = compare.compare(earlier, build(later))

    assert redline.revision_authors(result.content) == {"Opposing Counsel", compare.AUTHOR}
    # Their insertion already reads "fourteen days", so that is not a change.
    assert [c.where for c in result.changes] == ["5"]
    assert any("tracked changes" in n for n in result.notes)


def test_a_change_inside_a_hyperlink_is_reported_not_forced():
    d = Document(BytesIO(build(EARLIER)))
    p = d.add_paragraph("See ")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "top")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "the old policy"
    run.append(t)
    link.append(run)
    p._p.append(link)
    buf = BytesIO()
    d.save(buf)

    later = EARLIER + ["See the new policy"]
    result = compare.compare(buf.getvalue(), build(later))
    change = result.changes[-1]
    assert not change.landed and "hyperlink" in change.why_not
    # Still true whatever landed: rejecting everything gives back the earlier text.
    assert result.content is None or rejected(result.content)[: len(EARLIER)] == EARLIER


def test_text_comparison_lists_changes_without_markup():
    from lra.pipeline.extract import Block, ExtractedDoc

    def doc(lines):
        return ExtractedDoc([Block(i, t, "Normal", "paragraph") for i, t in enumerate(lines)],
                            None, "x.pdf")

    later = list(EARLIER)
    later[1] = later[1].replace("thirty (30)", "sixty (60)")
    result = compare.compare_text(doc(EARLIER), doc(later))
    assert result.content is None
    assert [c.kind for c in result.changes] == ["replaced"]
    assert not result.changes[0].landed


def test_a_failed_assessment_keeps_the_change_list(monkeypatch):
    later = list(EARLIER)
    later[1] = later[1].replace("thirty (30)", "sixty (60)")
    result = compare.compare(build(EARLIER), build(later))

    def boom(*args, **kwargs):
        raise RuntimeError("no network")

    monkeypatch.setattr(compare, "_assess", boom)
    from lra.pipeline.extract import ExtractedDoc

    out = compare.explain(result, ExtractedDoc([], None, "later.docx"))
    assert len(out.changes) == 1 and not out.changes[0].risk
    assert any("could not assess" in n for n in out.notes)


def _doc(paragraphs: list[str]) -> bytes:
    d = Document()
    for text in paragraphs:
        d.add_paragraph(text)
    buffer = BytesIO()
    d.save(buffer)
    return buffer.getvalue()


def test_consecutive_new_paragraphs_compare_and_prove():
    """A paragraph inserted beside one just inserted copied its neighbour's
    paragraph-mark revision, id and all, and every comparison that added two
    paragraphs in a row failed verification. Blacklines across several turns
    always do."""
    result = compare.compare(_doc(["a one two", "b three four"]),
                             _doc(["a one two", "b three four", "c five", "d six", "e seven"]))
    assert result.content is not None and result.proven is True
    after_deleted = compare.compare(_doc(["a one two", "b three four gone", "c five"]),
                                    _doc(["a one two", "x new one", "y new two", "c five"]))
    assert after_deleted.content is not None and after_deleted.proven is True
