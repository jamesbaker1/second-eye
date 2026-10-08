"""Documents are read here, not in a container, and the writer never leaves.

The associate's sandbox converts formats we cannot read locally and nothing
else. What is tested is everything that decides whether that is reached at
all, and the guarantee that matters most: tracked changes are written by
tested local code, whether it runs here or as the skill in the session.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from secondeye import managed
from secondeye.models import Attachment
from secondeye.pipeline import extract
from secondeye.pipeline.filetype import Kind, identify

OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def att(content: bytes, name: str) -> Attachment:
    return Attachment(filename=name, content_type="application/octet-stream",
                      size_bytes=len(content), content=content)


def pptx_bytes() -> bytes:
    import zipfile

    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/presentation.xml", "<p/>")
    return buf.getvalue()


def xlsx_bytes() -> bytes:
    import zipfile

    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/workbook.xml", "<w/>")
    return buf.getvalue()


# --- nothing model-backed runs half-configured -----------------------------

def test_conversion_in_a_session_needs_the_agent_ids(monkeypatch):
    """A legacy .doc reaches a session only when the agents are set up;
    otherwise it is refused with the sentence that asks for a .docx."""
    from secondeye.config import settings

    monkeypatch.setenv("MANAGED_REVIEW_AGENT_ID", "")
    settings.cache_clear()
    try:
        assert managed.configured() is False
        with pytest.raises(managed.JobFailed):
            managed.convert_to_docx(b"x", "old.doc")
    finally:
        settings.cache_clear()


# --- format recognition ---------------------------------------------------

def test_a_powerpoint_deck_is_recognised():
    assert identify(pptx_bytes(), "deck.pptx") is Kind.PPTX


def test_an_excel_workbook_is_recognised():
    assert identify(xlsx_bytes(), "model.xlsx") is Kind.XLSX


def test_a_deck_and_a_word_file_are_told_apart():
    d = Document()
    d.add_paragraph("x")
    buf = BytesIO()
    d.save(buf)
    assert identify(buf.getvalue(), "a.docx") is Kind.DOCX
    assert identify(pptx_bytes(), "a.docx") is Kind.PPTX


# --- behaviour when it is off --------------------------------------------

def test_a_legacy_doc_is_still_refused_helpfully_without_a_converter():
    from secondeye.config import settings

    settings.cache_clear()
    with pytest.raises(extract.Unreadable) as e:
        extract.extract(att(OLE + b"\x00" * 400, "old.doc"))
    assert "save it as .docx" in str(e.value)


def test_a_deck_is_refused_with_its_own_sentence_without_a_converter():
    from secondeye.config import settings

    settings.cache_clear()
    with pytest.raises(extract.Unreadable) as e:
        extract.extract(att(pptx_bytes(), "deck.pptx"))
    assert "PowerPoint" in str(e.value)


# --- the guarantee that matters ------------------------------------------

def test_the_writer_never_reaches_for_a_session():
    """The writer is the load-bearing component and is tested against cases a
    model improvising in a container has no guarantees about. It runs in the
    session too, as a skill script, and that script calls this same code."""
    import inspect

    from secondeye.pipeline import ooxml, redline

    for module in (ooxml, redline):
        source = inspect.getsource(module).lower()
        assert "managed" not in source and "anthropic" not in source, (
            f"{module.__name__} reaches for the platform; tracked changes must "
            "be written by tested local code"
        )


def test_a_docx_is_never_converted_in_a_session():
    """It is read precisely in-process; only formats we cannot open are sent."""
    from secondeye.pipeline.filetype import SANDBOX_CONVERTIBLE, SANDBOX_READABLE

    assert Kind.DOCX not in SANDBOX_CONVERTIBLE | SANDBOX_READABLE
    assert Kind.PDF not in SANDBOX_CONVERTIBLE | SANDBOX_READABLE


# --- decks and workbooks are read in-process, not in a container ----------

def test_a_deck_is_read_locally():
    """The container has python-pptx; so do we. It was never buying capability,
    only convenience, and it charges thirty days of retention for it."""
    from io import BytesIO as _BytesIO

    from pptx import Presentation
    from pptx.util import Inches

    d = Presentation()
    slide = d.slides.add_slide(d.slide_layouts[5])
    slide.shapes.title.text = "Project Falcon"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
    box.text_frame.text = "Indemnity capped at $1,500,000."
    buf = _BytesIO()
    d.save(buf)

    doc = extract.extract(att(buf.getvalue(), "deck.pptx"))
    text = " ".join(b.text for b in doc.blocks)
    assert "Project Falcon" in text
    assert "1,500,000" in text
    assert any(b.kind == "page-break" for b in doc.blocks)


def test_a_workbook_is_read_locally_with_values_not_formulas():
    from io import BytesIO as _BytesIO

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Payments"
    ws.append(["Milestone", "Amount"])
    ws.append(["Signing", 500000])
    ws["B3"] = "=B2*2"
    buf = _BytesIO()
    wb.save(buf)

    doc = extract.extract(att(buf.getvalue(), "payments.xlsx"))
    text = " ".join(b.text for b in doc.blocks)
    assert "[sheet Payments]" in text
    assert "Signing | 500000" in text
    assert "=B2*2" not in text, "a formula leaked instead of its value"


def test_checks_still_run_over_a_deck():
    """A deck cannot carry a tracked change, but the mechanical checks are
    still worth running over it."""
    from io import BytesIO as _BytesIO

    from pptx import Presentation
    from pptx.util import Inches

    d = Presentation()
    slide = d.slides.add_slide(d.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
    box.text_frame.text = "Term: thirty (13) months from signing."
    buf = _BytesIO()
    d.save(buf)

    from secondeye.pipeline import checks

    content = buf.getvalue()
    found = checks.run_all(extract.extract(att(content, "deck.pptx")), content)
    assert any(f.category == "amount" for f in found)


def test_nothing_is_read_in_a_session_any_more():
    from secondeye.pipeline.filetype import SANDBOX_READABLE

    assert SANDBOX_READABLE == set()


def test_a_damaged_deck_is_reported_not_crashed():
    import zipfile
    from io import BytesIO as _BytesIO

    buf = _BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/presentation.xml", "not xml at all")
    with pytest.raises(extract.Unreadable) as e:
        extract.extract(att(buf.getvalue(), "broken.pptx"))
    assert "would not open" in str(e.value)


def test_every_import_in_src_is_a_declared_dependency():
    """CI installs from pyproject, not from whatever happens to be in the venv.
    python-pptx and openpyxl were added to the venv with uv and the declaration
    was lost when another change rewrote that block, so the suite passed here
    and failed on a clean checkout."""
    import ast
    import pathlib
    import sys
    import tomllib

    root = pathlib.Path(__file__).resolve().parent.parent
    declared = tomllib.loads((root / "pyproject.toml").read_text())
    names = declared["project"]["dependencies"]
    names += declared["project"]["optional-dependencies"]["dev"]
    # "python-pptx>=1.0" -> "python-pptx" -> the module it provides
    packages = {n.split(">")[0].split("=")[0].split("[")[0].strip().lower()
                for n in names}
    provides = {
        "python-docx": "docx", "python-pptx": "pptx", "pydantic-settings":
        "pydantic_settings", "pypdfium2": "pypdfium2", "openpyxl": "openpyxl",
        "pyyaml": "yaml", "pillow": "PIL",
    }
    available = {provides.get(p, p.replace("-", "_")) for p in packages}
    # evals/ and demos/ sit beside the package in a checkout; `second-eye eval` and
    # `second-eye live-check` put the checkout on the path and say so when it is not.
    available |= set(sys.stdlib_module_names) | {"secondeye", "evals", "demos"}

    missing: set[str] = set()
    for path in (root / "src").rglob("*.py"):
        tree = ast.parse(path.read_text())
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                roots = [a.name.split(".")[0] for a in node.names]
            elif isinstance(node, ast.ImportFrom) and node.level == 0 and node.module:
                roots = [node.module.split(".")[0]]
            else:
                continue
            missing |= {r for r in roots if r not in available}
    assert not missing, f"imported but not declared in pyproject.toml: {sorted(missing)}"
