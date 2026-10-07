"""Comparison and clean copy, from inbound email to sent reply.

The model is stubbed, as everywhere else. Everything a lawyer would open is
real: the routing, the comparison, the proof, the scrubber, the reply.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document

from lra import handler
from lra.mail.console import ConsoleProvider
from lra.models import Attachment, InboundEmail
from lra.pipeline import checks, compare, intake
from lra.pipeline.ooxml import Revision, RevisionWriter

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

V1 = [
    "1. Term. This Agreement continues for twelve (12) months.",
    "2. Liability. The Supplier's liability is capped at the fees paid.",
    "3. Governing law. New York law governs this Agreement.",
]
V2 = [
    "1. Term. This Agreement continues for twenty-four (24) months.",
    "2. Liability. The Supplier's liability is capped at the fees paid.",
    "3. Governing law. New York law governs this Agreement.",
    "4. Exclusivity. The Buyer shall purchase only from the Supplier.",
]


class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return "captured"


@pytest.fixture
def captured(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/versions.sqlite3")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    from lra.config import settings

    settings.cache_clear()
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    yield provider
    settings.cache_clear()


def docx(paragraphs) -> bytes:
    d = Document()
    for text in paragraphs:
        d.add_paragraph(text)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def attach(name: str, content: bytes) -> Attachment:
    return Attachment(filename=name, content_type=DOCX,
                      size_bytes=len(content), content=content)


def email(body: str, *attachments: Attachment, message_id="v-1") -> InboundEmail:
    return InboundEmail(
        message_id=message_id, from_address="jim@firm.com",
        to=["review@legal.firm.com"], subject="Supply agreement",
        text_body=body, attachments=list(attachments),
        received_at=datetime.now(UTC),
    )


def assessed(monkeypatch):
    def _assess(comparison, later, instructions, house_style):
        comparison.summary = "They doubled the term and added exclusivity."
        for change in comparison.changes:
            if "Exclusivity" in change.after:
                change.risk, change.impact = "high", "Locks the Buyer to one supplier."
                change.response = "Delete clause 4."
            else:
                change.risk, change.impact = "medium", "Doubles the commitment."

    monkeypatch.setattr(compare, "_assess", _assess)


# --- comparison -----------------------------------------------------------


def test_two_versions_and_the_word_compare(captured, monkeypatch):
    assessed(monkeypatch)
    handler.handle(email("Can you compare these? They sent v2 back last night.",
                         attach("Supply v2.docx", docx(V2)),
                         attach("Supply v1.docx", docx(V1))))

    assert len(captured.sent) == 1
    out = captured.sent[0]
    assert out.to == ["jim@firm.com"]
    assert out.subject == "Re: Supply agreement"
    assert out.text_body.startswith("2 changes. 2 need your attention.")
    # Ordered by the filenames, not by the order they were attached.
    assert "Supply v1.docx as the earlier version" in out.text_body
    assert "version numbers in the filenames" in out.text_body
    # The risky one leads, and its words are quoted, not just described.
    assert out.text_body.index("Key: ") < out.text_body.index("shall purchase only")
    assert "shall purchase only from the Supplier" in out.text_body
    assert "Delete clause 4" in out.text_body

    assert [a.filename for a in out.attachments] == ["Supply v2 (comparison).docx"]
    ok, reason = handler.redline.verify(out.attachments[0].content, expect_revisions=True)
    assert ok, reason
    assert "I checked both" in out.text_body


def test_the_model_failing_still_sends_the_complete_list(captured, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("no network")

    monkeypatch.setattr(compare, "_assess", boom)
    handler.handle(email("what changed between these two?",
                         attach("a (1).docx", docx(V1)), attach("a (2).docx", docx(V2))))
    out = captured.sent[0]
    assert "2 changes between the two versions" in out.text_body
    assert "twenty-four (24" in out.text_body
    assert out.attachments


def test_identical_versions_say_so_and_attach_nothing(captured):
    handler.handle(email("compare please",
                         attach("a v1.docx", docx(V1)), attach("a v2.docx", docx(V1))))
    out = captured.sent[0]
    assert out.text_body.startswith("No changes")
    assert out.attachments == []


def test_compare_against_the_last_version_uses_the_archive(captured, monkeypatch):
    assessed(monkeypatch)
    monkeypatch.setattr(
        handler.versions.archive, "latest_version",
        lambda owner, filename: (docx(V1), "2026-09-15T10:00:00+00:00"),
    )
    handler.handle(email("Compare this to the last version you saw.",
                         attach("Supply agreement.docx", docx(V2))))
    out = captured.sent[0]
    assert "copy I last saw on 2026-09-15" in out.text_body
    assert out.attachments


def test_no_earlier_version_is_a_sentence_not_an_error(captured, monkeypatch):
    monkeypatch.setattr(handler.versions.archive, "latest_version", lambda o, f: None)
    handler.handle(email("compare to the previous draft",
                         attach("Supply agreement.docx", docx(V2))))
    assert "do not have an earlier version" in captured.sent[0].text_body


def test_two_documents_without_the_word_compare_are_not_compared(captured):
    handler.handle(email("Please review.",
                         attach("a.docx", docx(V1)), attach("b.docx", docx(V2))))
    out = captured.sent[0]
    assert "more than one document" in out.text_body
    assert '"compare"' in out.text_body


def test_a_forwarded_request_to_compare_is_not_ours(captured, monkeypatch):
    """Quoted text is the other side's, and must not route anything."""
    monkeypatch.setattr(handler.review, "review", lambda *a, **k: pytest.fail(
        "should have been rejected for two attachments, not reviewed"))
    body = ("See below.\n\nOn Mon, Sep 14, 2026, counsel@other.com wrote:\n"
            "> please compare against our last turn")
    handler.handle(email(body, attach("a.docx", docx(V1)), attach("b.docx", docx(V2))))
    assert "more than one document" in captured.sent[0].text_body


def test_order_falls_back_to_when_each_was_saved():
    older, newer = Document(), Document()
    older.add_paragraph("x")
    newer.add_paragraph("y")
    older.core_properties.modified = datetime(2026, 9, 1, tzinfo=UTC)
    newer.core_properties.modified = datetime(2026, 9, 9, tzinfo=UTC)

    def raw(d):
        buf = BytesIO()
        d.save(buf)
        return buf.getvalue()

    first, second, how = intake.pick_pair(
        email("compare", attach("Final.docx", raw(newer)), attach("Final final.docx", raw(older))))
    assert (first.filename, second.filename) == ("Final final.docx", "Final.docx")
    assert "last saved" in how


# --- across formats: the draft went out as Word and came back as a PDF --------


def pdf(paragraphs) -> bytes:
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf)
    y = 800
    for text in paragraphs:
        c.drawString(40, y, text)
        y -= 20
    c.save()
    return buf.getvalue()


def as_pdf(name: str, paragraphs) -> Attachment:
    content = pdf(paragraphs)
    return Attachment(filename=name, content_type="application/pdf",
                      size_bytes=len(content), content=content)


def as_text(name: str, paragraphs) -> Attachment:
    content = "\n\n".join(paragraphs).encode()
    return Attachment(filename=name, content_type="text/plain",
                      size_bytes=len(content), content=content)


def test_word_against_the_pdf_the_other_side_returned_is_a_word_comparison(captured,
                                                                            monkeypatch):
    """Litera Compare's headline case. This used to fall to compare_text with
    no attachment; now the PDF is rebuilt as Word and the comparison is real
    tracked changes in the earlier version, proved like every other one."""
    assessed(monkeypatch)
    handler.handle(email("Compare these; they sent it back as a PDF.",
                         attach("Supply v1.docx", docx(V1)),
                         as_pdf("Supply v2.pdf", V2)))
    [out] = captured.sent
    assert out.text_body.startswith("2 changes. 2 need your attention.")
    assert [a.filename for a in out.attachments] == ["Supply v2 (comparison).docx"]
    ok, reason = handler.redline.verify(out.attachments[0].content, expect_revisions=True)
    assert ok, reason
    # The orientation says which side was rebuilt and from what.
    assert "Supply v2.pdf is a PDF, so I rebuilt it as a Word document" in out.text_body
    assert "twenty-four (24" in out.text_body
    assert "there is no marked-up copy" not in out.text_body


def test_a_pdf_earlier_version_says_the_markup_lacks_its_layout(captured, monkeypatch):
    assessed(monkeypatch)
    handler.handle(email("compare", as_pdf("Supply v1.pdf", V1), attach("Supply v2.docx", docx(V2))))
    [out] = captured.sent
    assert out.attachments
    assert "does not have the original's layout" in out.text_body
    assert "Supply v1.pdf is a PDF, so I rebuilt it" in out.text_body


def test_word_against_a_text_file(captured, monkeypatch):
    assessed(monkeypatch)
    handler.handle(email("compare", attach("Supply v1.docx", docx(V1)),
                         as_text("Supply v2.txt", V2)))
    [out] = captured.sent
    assert [a.filename for a in out.attachments] == ["Supply v2 (comparison).docx"]
    ok, reason = handler.redline.verify(out.attachments[0].content, expect_revisions=True)
    assert ok, reason
    assert "Supply v2.txt is a text file, so I rebuilt it" in out.text_body


def test_a_deck_is_still_a_list_only_comparison(captured):
    from pptx import Presentation

    deck_file = Presentation()
    slide = deck_file.slides.add_slide(deck_file.slide_layouts[1])
    slide.shapes.title.text = "Supply agreement"
    slide.placeholders[1].text = V2[0]
    buf = BytesIO()
    deck_file.save(buf)
    deck = buf.getvalue()
    handler.handle(email("compare", attach("Supply v1.docx", docx(V1)),
                         Attachment(filename="Supply v2.pptx",
                                    content_type="application/octet-stream",
                                    size_bytes=len(deck), content=deck)))
    [out] = captured.sent
    assert out.attachments == []
    assert "there is no marked-up copy" in out.text_body


def test_compare_to_the_one_i_sent_resolves_through_the_thread_with_the_archive_off(
        captured, monkeypatch):
    """The archive is off by default and the help text promises this anyway.
    The lawyer's conversations hold every document reviewed, so the earlier
    version is found there: by name first."""
    from tests.test_handler import stub_review

    monkeypatch.setattr(handler.review, "review", stub_review())
    monkeypatch.setattr(handler.versions.archive, "latest_version", lambda o, f: None)
    handler.handle(email("Quick look please.", attach("Supply agreement.docx", docx(V1)),
                         message_id="v-review"))
    assert len(captured.sent) == 1

    handler.handle(email("Compare this to the version I sent you last week.",
                         attach("Supply agreement.docx", docx(V2)), message_id="v-compare"))
    out = captured.sent[1]
    assert out.text_body.startswith("2 changes")
    assert "copy I last saw on" in out.text_body
    assert out.attachments
    ok, reason = handler.redline.verify(out.attachments[0].content, expect_revisions=True)
    assert ok, reason


def test_compare_to_the_one_i_sent_with_nothing_held_is_still_a_sentence(captured, monkeypatch):
    monkeypatch.setattr(handler.versions.archive, "latest_version", lambda o, f: None)
    handler.handle(email("compare to the previous draft", attach("Fresh.docx", docx(V2))))
    assert "do not have an earlier version" in captured.sent[0].text_body


# --- clean copy -----------------------------------------------------------


def dirty() -> bytes:
    d = Document()
    d.add_paragraph("1. The term is thirty days.")
    d.core_properties.author = "Jane Associate"
    w = RevisionWriter(d, "Jane Associate")
    w.apply(Revision("thirty days", "sixty days", "Jane Associate"))
    w.comment("The term", "check with client")
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_asking_for_a_clean_copy_returns_one(captured):
    handler.handle(email("Can you send me a clean copy of this for the client?",
                         attach("NDA.docx", dirty())))
    out = captured.sent[0]
    assert out.text_body.startswith("Clean copy attached")
    assert [a.filename for a in out.attachments] == ["NDA (clean).docx"]
    assert checks.leftovers(out.attachments[0].content) == []
    assert "Jane Associate" in out.text_body and "1 comment" in out.text_body


def test_telling_a_client_a_clean_copy_is_attached_is_not_a_request():
    claim = email("Dear Bob, attaching a clean copy as discussed.",
                  attach("NDA.docx", dirty()))
    assert not intake.wants_clean_copy(claim)
    assert intake.wants_clean_copy(email("please scrub this", attach("NDA.docx", dirty())))


def commented_pdf(**encrypt) -> Attachment:
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import Text
    from pypdf.generic import NameObject, TextStringObject

    w = PdfWriter(clone_from=PdfReader(BytesIO(pdf(V1))))
    w.metadata = None
    w.add_metadata({"/Author": "Jane Associate", "/Producer": "Acme DMS"})
    for note in ("Too low?", "Check with client", "NY or DE?"):
        annot = Text(rect=(50, 550, 200, 650), text=note)
        annot[NameObject("/T")] = TextStringObject("Jane Associate")
        w.add_annotation(page_number=0, annotation=annot)
    if encrypt:
        w.encrypt(**encrypt)
    out = BytesIO()
    w.write(out)
    return Attachment(filename="NDA.pdf", content_type="application/pdf",
                      size_bytes=len(out.getvalue()), content=out.getvalue())


def test_a_clean_copy_of_a_pdf_returns_one(captured):
    handler.handle(email("strip the metadata please", commented_pdf()))
    out = captured.sent[0]
    assert out.text_body.startswith("Clean copy attached")
    assert "3 comments from Jane Associate" in out.text_body
    assert "Author and producer names" in out.text_body
    assert "Nothing else changed; I checked the wording against yours." in out.text_body
    [copy] = out.attachments
    assert (copy.filename, copy.content_type) == ("NDA (clean).pdf", "application/pdf")
    assert b"Jane Associate" not in copy.content


def test_a_password_protected_pdf_is_declined_in_a_sentence(captured):
    handler.handle(email("clean copy please", commented_pdf(user_password="x")))
    assert ("It's password-protected; send it unlocked and I'll clean it."
            in captured.sent[0].text_body)
    assert captured.sent[0].attachments == []


# --- a reply to a comparison is about the later version -----------------------


def test_a_reply_to_a_comparison_reviews_the_later_version_with_it(captured, monkeypatch):
    """"Look harder at clause 4" used to arrive as a message about nothing and
    get "attach a .docx"."""
    assessed(monkeypatch)
    seen = {}

    def fake_review(doc, mode, instructions, **kwargs):
        seen["filename"] = doc.filename
        seen["instructions"] = instructions
        from lra.models import ReviewResult
        return ReviewResult(mode=mode, summary="Clause 4 locks you in.", findings=[])

    monkeypatch.setattr(handler.review, "review", fake_review)
    handler.handle(email("Compare these please.",
                         attach("Supply v2.docx", docx(V2)),
                         attach("Supply v1.docx", docx(V1))))
    assert len(captured.sent) == 1

    handler.handle(InboundEmail(
        message_id="v-2", from_address="jim@firm.com", to=["review@legal.firm.com"],
        subject="Re: Supply agreement", text_body="Look harder at clause 4.",
        in_reply_to="v-1", references=["v-1"], received_at=datetime.now(UTC),
    ))
    assert len(captured.sent) == 2
    assert seen["filename"] == "Supply v2.docx"
    assert "clause 4" in seen["instructions"].lower()
    assert "Clause 4 locks you in" in captured.sent[1].text_body

    # And from then on it is an ordinary review thread.
    handler.handle(InboundEmail(
        message_id="v-3", from_address="jim@firm.com", to=["review@legal.firm.com"],
        subject="Re: Supply agreement", text_body="Thanks!",
        in_reply_to="v-2", references=["v-1", "v-2"], received_at=datetime.now(UTC),
    ))
    assert len(captured.sent) == 2


def test_thanks_after_a_comparison_gets_no_reply(captured, monkeypatch):
    assessed(monkeypatch)
    handler.handle(email("Compare these please.",
                         attach("Supply v2.docx", docx(V2)),
                         attach("Supply v1.docx", docx(V1))))
    handler.handle(InboundEmail(
        message_id="v-2", from_address="jim@firm.com", to=["review@legal.firm.com"],
        subject="Re: Supply agreement", text_body="Great, thanks",
        in_reply_to="v-1", references=["v-1"], received_at=datetime.now(UTC),
    ))
    assert len(captured.sent) == 1

