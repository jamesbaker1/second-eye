"""The associate loop: answering, undoing, and acting on a reply.

Before this existed, the reply footer promised "reply in plain English and I
will make the changes" and a reply with no attachment got "I could not find a
document to review", which made the promise false on the most natural next
action a lawyer could take.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document

from lra import followup, handler, thread
from lra.mail.console import ConsoleProvider
from lra.models import Attachment, Finding, InboundEmail, ReviewResult, Severity

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return f"agent-msg-{len(self.sent)}"


@pytest.fixture
def captured(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/f.sqlite3")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    from lra.config import settings

    settings.cache_clear()
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    yield provider
    settings.cache_clear()


def contract() -> bytes:
    d = Document()
    d.add_paragraph("1. Term. The term is thirty (13) months from the Effective Date.")
    d.add_paragraph("2. Notice. The Reciever shall give notice under Section 7.")
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def first_email() -> InboundEmail:
    raw = contract()
    return InboundEmail(
        message_id="first-1", thread_id="thread-1",
        from_address="jim@firm.com", to=["review@legal.firm.com"],
        subject="Acme NDA", text_body="Quick look before I send?",
        attachments=[Attachment(filename="Acme NDA.docx", content_type=DOCX,
                                size_bytes=len(raw), content=raw)],
        received_at=datetime.now(UTC),
    )


def reply_email(body: str, mid: str, in_reply_to="agent-msg-1") -> InboundEmail:
    return InboundEmail(
        message_id=mid, thread_id="thread-1", in_reply_to=in_reply_to,
        from_address="jim@firm.com", to=["review@legal.firm.com"],
        subject="Re: Acme NDA", text_body=body, received_at=datetime.now(UTC),
    )


def stub(*findings):
    def _review(doc, mode, instructions, **kw):
        return ReviewResult(mode=mode, summary="Short NDA.", findings=list(findings))
    return _review


TYPO = Finding(severity=Severity.FORMATTING, category="typo",
               title='"Reciever" is misspelled', explanation="",
               anchor="Reciever", suggested_text="Recipient", auto_apply=True)


def run_first(captured, monkeypatch):
    monkeypatch.setattr(handler.review, "review", stub(TYPO))
    handler.handle(first_email())
    return captured


# --- the reply is recognised at all --------------------------------------

def test_a_reply_with_no_attachment_is_not_rejected(captured, monkeypatch):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("undo everything", "r2"))
    assert "could not find a document" not in captured.sent[1].text_body


def test_a_reply_resolves_through_the_agents_own_message_id(captured, monkeypatch):
    """A mail client's In-Reply-To points at the agent's reply, not the
    lawyer's original, so that id has to resolve back to the conversation."""
    run_first(captured, monkeypatch)
    handler.handle(reply_email("undo everything", "r2", in_reply_to="agent-msg-1"))
    assert "reversed" in captured.sent[1].text_body


# --- answering a question ------------------------------------------------

def test_a_one_word_answer_produces_a_tracked_change(captured, monkeypatch):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("30", "r2"))
    out = captured.sent[1]
    assert "Noted: 30" in out.text_body
    assert out.attachments, "no updated document came back"

    from lra.pipeline import redline

    ok, why = redline.verify(out.attachments[0].content, expect_revisions=True)
    assert ok, why


def test_the_answer_is_described_in_the_words_that_changed(captured, monkeypatch):
    """Anchors carry context for locating, so quoting one whole read
    'I changed "tinues for thirty (13) months from the E"'."""
    run_first(captured, monkeypatch)
    handler.handle(reply_email("30", "r2"))
    assert '\u201c13\u201d \u2192 \u201c30\u201d' in captured.sent[1].text_body


def test_an_answered_question_stops_being_asked(captured, monkeypatch):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("30", "r2"))
    body = captured.sent[1].text_body
    assert "Should this be 30 or 13?" not in body.split("Still waiting on:")[-1]


# --- undo -----------------------------------------------------------------

def test_undoing_one_change_leaves_the_others(captured, monkeypatch):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("30", "r2"))
    handler.handle(reply_email("actually undo the Reciever change", "r3"))
    body = captured.sent[2].text_body
    assert "reversed" in body
    assert '\u201c13\u201d \u2192 \u201c30\u201d' in body


def test_undoing_everything_returns_the_original(captured, monkeypatch):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("undo everything", "r2"))
    out = captured.sent[1]
    assert "this is your original" in out.text_body
    assert out.attachments == []


def test_an_undo_that_matches_nothing_asks_rather_than_guessing(captured, monkeypatch):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("undo the governing law change", "r2"))
    body = captured.sent[1].text_body
    assert body.startswith("Which one?")
    assert "I have not reversed anything" in body
    assert '  1. "Reciever" is misspelled' in body
    assert captured.sent[1].attachments == []


# --- the pieces in isolation ---------------------------------------------

def test_a_word_numeral_answer_fixes_both_halves():
    q = thread.Question(id=1, question="Should this be 30 or 13?",
                        anchor="term is thirty (13) months", options=["30", "13"],
                        answered_with=None, round=0)
    state = thread.ThreadState(id="t", owner="jim@firm.com", filename="a.docx")
    change = followup.apply_answer(state, q, "30")
    assert change.suggested_text == "term is thirty (30) months"


def test_choosing_the_numeral_rewrites_the_word_too():
    q = thread.Question(id=1, question="Should this be 30 or 13?",
                        anchor="thirty (13) days", options=["30", "13"],
                        answered_with=None, round=0)
    state = thread.ThreadState(id="t", owner="jim@firm.com", filename="a.docx")
    change = followup.apply_answer(state, q, "13")
    assert change.suggested_text == "thirteen (13) days"


@pytest.mark.parametrize("phrase", [
    "undo that", "please revert the term change", "put that back",
    "leave it as it was", "roll back your changes",
])
def test_plain_english_undo_is_recognised(phrase):
    state = thread.ThreadState(id="t", owner="j@f.com", filename="a.docx")
    kind, _ = followup.classify(
        state,
        InboundEmail(message_id="m", from_address="j@f.com", text_body=phrase,
                     received_at=datetime.now(UTC)),
    )
    assert kind == "undo"


def test_an_attachment_on_a_reply_is_a_new_version():
    state = thread.ThreadState(id="t", owner="j@f.com", filename="a.docx")
    email = InboundEmail(
        message_id="m", from_address="j@f.com", text_body="here is the revised one",
        attachments=[Attachment(filename="a.docx", content_type=DOCX,
                                size_bytes=4, content=b"abcd")],
        received_at=datetime.now(UTC),
    )
    kind, _ = followup.classify(state, email)
    assert kind == "new_version"


# --- undo must never guess --------------------------------------------------

def _ledger(*titles_and_anchors):
    changes = [
        thread.Change(id=100 + i, anchor=anchor, replacement=anchor + "!", title=title,
                      category="spelling", origin="check", round=0, active=True)
        for i, (title, anchor) in enumerate(titles_and_anchors)
    ]
    return thread.ThreadState(id="t", owner="j@f.com", filename="a.docx", changes=changes)


def _classify(state, text):
    return followup.classify(
        state,
        InboundEmail(message_id="m", from_address="j@f.com", text_body=text,
                     received_at=datetime.now(UTC)),
    )


TWO = (('"Reciever" is misspelled', "the Reciever shall"),
       ("Indemnity made mutual", "the Recipient shall indemnify"))


@pytest.mark.parametrize("text", [
    "Please don't change the governing law, but tighten the indemnity further",
    "restore the cap to 5m and tighten the indemnity",
    "Can you undo the spelling fix and also add a force majeure clause after 7",
])
def test_an_undo_verb_inside_a_wider_instruction_is_not_an_undo(text):
    """Each of these reversed the indemnity change on the spot."""
    kind, _ = _classify(_ledger(*TWO), text)
    assert kind == "instruction"


def test_undo_by_the_number_in_the_email():
    kind, payload = _classify(_ledger(*TWO), "undo 2")
    assert kind == "undo"
    assert payload["match"].ids == [101]


def test_undo_two_numbers_with_a_sign_off():
    _, payload = _classify(_ledger(*TWO), "undo 1 and 2\n\nThanks\nJ")
    assert payload["match"].ids == [100, 101]


def test_a_number_that_is_not_a_change_asks():
    _, payload = _classify(_ledger(*TWO), "undo 3")
    assert payload["match"].ids == [] and payload["match"].ask


def test_a_clause_number_is_not_read_as_a_change_number():
    """ "undo the clause 2 change" reversed change 2, whatever clause it was in."""
    state = _ledger(*TWO)
    state.original = b""  # unreadable: the clause cannot be placed
    _, payload = _classify(state, "undo the clause 2 change")
    assert payload["match"].ids == [] and payload["match"].ask


def test_undo_that_with_several_changes_asks():
    _, payload = _classify(_ledger(*TWO), "undo that")
    assert payload["match"].ids == [] and payload["match"].ask


def test_undo_that_with_one_change_is_clear():
    _, payload = _classify(_ledger(TWO[0]), "undo that")
    assert payload["match"].ids == [100]


def test_words_that_pick_out_one_change_still_work():
    _, payload = _classify(_ledger(*TWO), "undo the indemnity change")
    assert payload["match"].ids == [101]


def test_numbers_survive_an_undo():
    """The number printed on the first email must mean the same change later."""
    state = _ledger(*TWO)
    state.changes[0].active = False
    assert state.number(state.active_changes[0]) == 2


# --- clean copy on a thread -------------------------------------------------

def test_a_bare_clean_copy_reply_says_which_copy_it_cleaned(captured, monkeypatch):
    """Rebuilt from the thread, a clean copy brings back any change the lawyer
    rejected in Word. It has to say so, and how to get theirs cleaned."""
    run_first(captured, monkeypatch)
    handler.handle(reply_email("clean copy", "r2"))
    body = captured.sent[1].text_body
    assert body.startswith("Clean copy attached")
    assert "attach your copy" in body


def test_a_clean_copy_with_a_question_unanswered_says_so_first(captured, monkeypatch):
    """A clean copy reads as ready to send. One still saying "thirty (13)"
    because nobody answered the question about it must say that in its first
    line, not leave the lawyer to remember."""
    run_first(captured, monkeypatch)
    handler.handle(reply_email("clean copy", "r2"))
    body = captured.sent[1].text_body
    assert body.splitlines()[0] == "Clean copy attached, but 2 questions are still open."
    assert "Should this be 30 or 13?" in body.split("Still open, so not in this copy")[1]


def test_a_clean_copy_request_with_their_file_cleans_their_file(captured, monkeypatch):
    run_first(captured, monkeypatch)
    doc = Document()
    doc.add_paragraph("Their own final wording.")
    buf = BytesIO()
    doc.save(buf)
    e = reply_email("clean copy please", "r2")
    e.attachments = [Attachment(
        filename="Acme NDA final.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size_bytes=len(buf.getvalue()), content=buf.getvalue(),
    )]
    handler.handle(e)
    out = captured.sent[1]
    assert "attach your copy" not in out.text_body
    assert all("Acme NDA final" in a.filename for a in out.attachments)


# --- a follow-up that fails is answered once and never retried --------------

def test_an_oversized_follow_up_goes_out_without_the_file(captured, monkeypatch):
    run_first(captured, monkeypatch)
    real = captured.send

    def refuse_files(out):
        if out.attachments:
            raise RuntimeError("message too large")
        return real(out)

    monkeypatch.setattr(captured, "send", refuse_files)
    handler.handle(reply_email("30", "r2"))
    body = captured.sent[1].text_body
    assert body.startswith("I could not attach the document")
    assert "“13” → “30”" in body


def test_a_follow_up_that_blows_up_says_so_and_does_not_raise(captured, monkeypatch):
    """Raising sent the job back to the queue, which undid the change again on
    retry and, after three, left the lawyer with nothing."""
    run_first(captured, monkeypatch)

    def boom(*a, **k):
        raise RuntimeError("ledger unavailable")

    monkeypatch.setattr(handler.thread, "undo", boom)
    handler.handle(reply_email("undo 1", "r2"))     # must not raise
    assert "Something went wrong on my side" in captured.sent[1].text_body


# --- several things in one reply ----------------------------------------------
#
# The review's footer suggests answering every question in one reply and then
# asks for "clean copy". Each part is deterministic on its own; the message as
# a whole went to the instruction agent (minutes, and a failure with no model),
# or, with "clean copy" on its own line, cleaned the document and dropped the
# answer.


@pytest.fixture
def no_model(monkeypatch):
    def refuse(*a, **k):
        raise AssertionError("the model was called for a reply that needs none")
    monkeypatch.setattr(handler.instruct, "carry_out", refuse)


def _words(content: bytes) -> str:
    return " ".join(p.text for p in Document(BytesIO(content)).paragraphs)


def test_the_answers_to_every_question_in_one_reply(captured, monkeypatch, no_model):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("30; 2", "r2"))
    assert len(captured.sent) == 2
    out = captured.sent[1]
    assert out.text_body.startswith("Noted: 30 and 2. I have made those changes.")
    assert "Still waiting on" not in out.text_body

    from lra.pipeline import redline

    ok, why = redline.verify(out.attachments[0].content, expect_revisions=True)
    assert ok, why


@pytest.mark.parametrize("body", [
    "30\nclean copy",
    "30; 2; clean copy",
    "30, then clean copy",
    "30 and a clean copy please",
    "30. Clean copy please.\n\nThanks\nJim",
])
def test_answers_and_a_clean_copy_in_one_reply(captured, monkeypatch, no_model, body):
    run_first(captured, monkeypatch)
    handler.handle(reply_email(body, "r2"))
    assert len(captured.sent) == 2, "one reply, not one per request"
    out = captured.sent[1]
    assert out.text_body.startswith("Clean copy attached")
    assert "Noted: 30" in out.text_body
    [copy] = out.attachments
    assert copy.filename == "Acme NDA (clean).docx"
    words = _words(copy.content)
    assert "thirty (30)" in words and "(13)" not in words
    assert "Recipient" in words


def test_answering_everything_with_a_clean_copy_leaves_nothing_open(captured, monkeypatch,
                                                                    no_model):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("30; 2; clean copy", "r2"))
    out = captured.sent[1]
    assert out.text_body.splitlines()[0] == "Clean copy attached."
    assert "Section 2" in _words(out.attachments[0].content)


def test_accept_all_is_a_clean_copy(captured, monkeypatch, no_model):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("Accept all", "r2"))
    out = captured.sent[1]
    assert out.text_body.startswith("Clean copy attached")
    assert "Recipient" in _words(out.attachments[0].content)


def test_an_undo_with_the_rest_does_it_all_in_order(captured, monkeypatch, no_model):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("undo 1, then 30, then clean copy", "r2"))
    out = captured.sent[1]
    assert "I reversed change 1" in out.text_body
    words = _words(out.attachments[0].content)
    assert "Reciever" in words and "thirty (30)" in words


def test_an_undo_that_is_unclear_stops_the_whole_reply(captured, monkeypatch, no_model):
    """The undo keeps its rule: act only on an unambiguous reference. The rest
    is not done around it, or the clean copy would not be the one asked for."""
    run_first(captured, monkeypatch)
    handler.handle(reply_email("undo 9; 30; clean copy", "r2"))
    out = captured.sent[1]
    assert out.text_body.startswith("Which one?")
    assert "not done the rest of your reply" in out.text_body
    assert out.attachments == []
    state = thread.load(thread.resolve("jim@firm.com", "thread-1", "x", "agent-msg-1", []))
    assert len(state.active_changes) == 1
    assert len(state.open_questions) == 2


def test_a_reply_with_a_part_not_understood_goes_to_the_agent_whole(captured, monkeypatch):
    run_first(captured, monkeypatch)
    heard: list[str] = []

    def carry_out(doc, instruction, **kw):
        heard.append(instruction)
        raise RuntimeError("no model here")

    monkeypatch.setattr(handler.instruct, "carry_out", carry_out)
    handler.handle(reply_email("30; and make the cap five million", "r2"))
    assert heard == ["30; and make the cap five million"]


# The ledger-level rules, without a document.

def _asked(*option_lists):
    return thread.ThreadState(
        id="t", owner="j@f.com", filename="a.docx",
        questions=[thread.Question(id=i, question=f"q{i}", anchor=f"a{i}", options=opts,
                                   answered_with=None, round=0)
                   for i, opts in enumerate(option_lists)],
        changes=_ledger(*TWO).changes,
    )


def _plan(state, text):
    return followup.plan_reply(state, InboundEmail(
        message_id="m", from_address="j@f.com", text_body=text,
        received_at=datetime.now(UTC)))


def test_an_answer_that_fits_two_questions_is_not_guessed():
    assert _plan(_asked(["30", "13"], ["30", "60"]), "30; clean copy") is None


def test_each_question_is_answered_once_in_order():
    plan = _plan(_asked(["30", "13"], ["30", "60"]), "13; 30")
    assert [(q.id, v) for q, v in plan.answers] == [(0, "13"), (1, "30")]


def test_one_part_alone_keeps_its_own_path():
    state = _asked(["30", "13"])
    assert _plan(state, "30") is None
    assert _plan(state, "undo 2") is None
    assert _plan(state, "30 thanks") is None


def test_a_clause_abbreviation_does_not_split_an_undo():
    plan = _plan(_asked(["30", "13"]), "undo cl. 7. Clean copy")
    assert plan is not None and plan.clean
    assert plan.undo is not None and plan.undo.ask     # clause 7 cannot be placed


def test_one_request_with_a_thank_you_is_still_understood():
    """ "30" then "Thanks, Jim" read as one message was no answer at all, and
    went to the model."""
    state = _asked(["30", "13"])
    plan = _plan(state, "30\n\nThanks\nJim")
    assert [v for _, v in plan.answers] == ["30"] and not plan.clean
    assert _plan(state, "thanks; clean copy").clean
    assert _plan(state, "undo 2\n\nThanks\nJ") is None    # the undo path reads it


def test_a_sign_off_is_not_a_request_but_a_sentence_is():
    state = _asked(["30", "13"])
    assert _plan(state, "30; clean copy\n\nJim Baker").clean
    assert _plan(state, "30; clean copy\n\nsee you at the meeting") is None


def test_only_the_answers_that_reached_the_document_are_claimed():
    """ "I have made that change" over an unchanged clause is the most damaging
    sentence in the product: the lawyer stops looking."""
    def change(anchor, text):
        return Finding(severity=Severity.FORMATTING, category="answered", title="t",
                       explanation="", anchor=anchor, suggested_text=text, auto_apply=True)

    placed = thread.Change(id=1, anchor="thirty (13)", replacement="thirty (30)", title="t",
                           category="answered", origin="instruction", round=0, active=True)
    said = handler._answers_said(
        [("30", change("thirty (13)", "thirty (30)")), ("2", change("Section 7", "Section 2"))],
        [placed])
    assert said == ("Noted: 30 and 2. I have made the change for 30. I could not place "
                    "the change for 2 safely, so that one is yours to make.")


# --- a change and a clean copy in one message --------------------------------

def test_an_instruction_beside_clean_copy_is_made_not_dropped(captured, monkeypatch):
    """ "clean copy" on its own line used to win outright: the change asked
    for above it was dropped without a word, and the clean copy lacked it."""
    from lra.pipeline import instruct

    heard: list[str] = []

    def carry_out(doc, instruction, **kw):
        heard.append(instruction)
        outcome = instruct.Outcome()
        outcome.understood = "Change the start of the term to the Signing Date."
        outcome.changes = [{"anchor": "from the Effective Date",
                            "text": "from the Signing Date", "title": "Term runs from signing"}]
        return outcome

    run_first(captured, monkeypatch)
    monkeypatch.setattr(handler.instruct, "carry_out", carry_out)
    handler.handle(reply_email("Run the term from signing instead.\nclean copy\n\nJim", "r2"))

    assert heard == ["Run the term from signing instead"]
    out = captured.sent[1]
    assert [a.filename for a in out.attachments] == ["Acme NDA (redline).docx"]
    assert "You asked for a clean copy as well" in out.text_body


def test_a_clean_copy_asked_for_in_one_sentence_needs_no_model(captured, monkeypatch,
                                                               no_model):
    run_first(captured, monkeypatch)
    handler.handle(reply_email("Could you send me a clean version please?", "r2"))
    assert captured.sent[1].text_body.startswith("Clean copy attached")


def test_undoing_an_answer_reopens_its_question(captured, monkeypatch):
    """Production: "30", "undo 1", then "30" again went to the model as an
    instruction, because the question stayed answered after its change was
    undone."""
    run_first(captured, monkeypatch)
    handler.handle(reply_email("30", "r2"))
    handler.handle(reply_email("undo everything", "r3"))
    handler.handle(reply_email("30", "r4"))
    assert captured.sent[-1].text_body.startswith("Noted: 30.")


def test_an_instruction_with_no_model_says_what_works_instead(captured, monkeypatch):
    """Production without credit: every unparseable reply got "Something went
    wrong ... try me again", which no retry could fix."""
    run_first(captured, monkeypatch)
    def unset(*a, **k):
        raise handler.managed.NotConfigured("Managed Agents is not set up")

    monkeypatch.setattr(handler.instruct, "carry_out", unset)
    handler.handle(reply_email("tighten the indemnity please", "r2"))
    body = captured.sent[-1].text_body
    assert body.startswith("I can't read that one")
    assert "Should this be 30 or 13? (30 or 13)" in body
    assert "Something went wrong" not in body
