"""The deal's figures, timings and party names.

Pinned from both sides, as every check is (PRODUCT.md): the Project Falcon
defects are caught, with both figures and both clause labels, and the clean
Falcon and every clean agreement in the corpus stay silent.
"""

from __future__ import annotations

from datetime import UTC
from datetime import datetime as dt
from io import BytesIO

import pytest
from docx import Document

from evals import corpus, model_corpus, realistic
from lra.models import Attachment, InboundEmail, ReviewResult, Severity
from lra.pipeline import checks, dealmath, extract, sigpack, timeline

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
NEW = {"arithmetic", "party-name"}


def docx_doc(content: bytes, name: str = "Falcon SPA.docx"):
    return extract.extract(Attachment(filename=name, content_type=DOCX,
                                      size_bytes=len(content), content=content))


def text_doc(*paragraphs):
    raw = "\n".join(paragraphs).encode()
    return extract.extract(Attachment(filename="t.txt", content_type="text/plain",
                                      size_bytes=len(raw), content=raw))


def table_doc(rows, before=("SCHEDULE 1", "Payments"), after=()):
    d = Document()
    for text in before:
        d.add_paragraph(text)
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            t.rows[r].cells[c].text = cell
    for text in after:
        d.add_paragraph(text)
    buf = BytesIO()
    d.save(buf)
    return docx_doc(buf.getvalue(), "t.docx")


@pytest.fixture(scope="module")
def falcon():
    return docx_doc(realistic.falcon_spa(3, defects=True))


@pytest.fixture(scope="module")
def falcon_clean():
    return docx_doc(realistic.falcon_spa(3))


# --- Silence: the clean corpus ------------------------------------------------


@pytest.mark.parametrize("spec", corpus.CLEAN, ids=lambda s: s.name)
def test_no_clean_agreement_gets_a_new_finding(spec):
    content = spec.build()
    doc = docx_doc(content, f"{spec.name}.docx")
    assert [f.title for f in checks.run_all(doc, content) if f.category in NEW] == []
    assert timeline.build(doc)["flags"] == []


@pytest.mark.parametrize("spec", model_corpus.ALL, ids=lambda s: s.name)
def test_the_model_corpus_plants_nothing_these_checks_should_see(spec):
    """Its defects are the playbook's; its figures, dates and names are right,
    so a convenience notice beside a breach cure period is not a flag."""
    content = spec.build()
    doc = docx_doc(content, f"{spec.name}.docx")
    assert [f.title for f in checks.run_all(doc, content) if f.category in NEW] == []
    assert timeline.build(doc)["flags"] == []


def test_the_clean_falcon_is_silent_and_its_sums_do_add_up(falcon_clean):
    assert [f for f in checks.run_all(falcon_clean) if f.category in NEW] == []
    survey = dealmath.survey(falcon_clean)
    assert survey["findings"] == []
    (split,) = survey["splits"]
    assert split["kind"] == "instalments" and split["sum"] == 12_500_000
    schedule_2 = next(t for t in survey["tables"] if t["where"] == "Schedule 2")
    amount = schedule_2["columns"][1]
    assert amount["body_sum"] == amount["total"] == 12_500_000 and amount["plain_sum"]
    assert survey["defined_amounts"]["Purchase Price"][0]["where"] == "clause 3.1"


# --- Arithmetic ---------------------------------------------------------------


def test_schedule_2_against_the_price_names_both_figures_and_clauses(falcon):
    found = {f.title: f for f in checks.run_all(falcon) if f.category == "arithmetic"}
    price = found["The Purchase Price is £12,500,000 in clause 3.1, but Schedule 2 "
                  "totals £12,050,000"]
    assert price.severity is Severity.BLOCKER and price.where == "Schedule 2"
    assert price.options == ["£12,500,000", "£12,050,000"] and not price.auto_apply
    assert "£450,000" in price.explanation


def test_instalments_that_do_not_add_up_to_the_price(falcon):
    (f,) = [f for f in checks.run_all(falcon) if "instalments" in f.title]
    assert f.title == "Clause 3.2: the instalments add up to £12,750,000, not £12,500,000"
    assert f.where == "clause 3.2" and f.anchor in " ".join(b.text for b in falcon.blocks)


def test_a_total_row_that_is_not_the_sum_of_the_rows():
    doc = table_doc([["Shareholder", "Percentage"], ["A Limited", "60%"],
                     ["B Limited", "30%"], ["C Limited", "5%"], ["Total", "100%"]])
    (f,) = dealmath.arithmetic(doc)
    assert f.title == "Schedule 1: the Total is 100%, but the rows above add up to 95%"
    assert f.anchor == "100%" and f.severity is Severity.BLOCKER


def test_a_defined_amount_restated_differently():
    doc = text_doc(
        "3.1 The consideration is £2,500,000 (two million five hundred thousand "
        "pounds) (the “Consideration”).",
        "9.2 The Buyer shall pay the Consideration of £2,050,000 by electronic transfer.")
    (f,) = dealmath.arithmetic(doc)
    assert f.title == ("The Consideration is £2,500,000 in clause 3.1 and £2,050,000 "
                       "in clause 9.2")
    assert "£2,050,000" in f.anchor


@pytest.mark.parametrize("rows", [
    # A subtotal and a deduction: not one plain sum.
    [["Item", "£"], ["Price", "£1,000"], ["Subtotal", "£1,000"], ["Less retention", "£100"],
     ["Total", "£900"]],
    # A table in thousands, rounded.
    [["Item", "£'000"], ["Shares", "3,334"], ["Debt", "3,333"], ["Cash", "3,333"],
     ["Total", "10,001"]],
    # Rates, not parts of a whole.
    [["Band", "Rate"], ["First £1m", "5%"], ["Next £1m", "3%"], ["Total", "4%"]],
])
def test_tables_that_are_not_plain_sums_are_left_alone(rows):
    assert dealmath.arithmetic(table_doc(rows)) == []


def test_a_part_of_a_defined_amount_is_not_a_restatement():
    doc = text_doc(
        "3.1 The consideration is £2,500,000 (the “Consideration”).",
        "3.2 The portion of the Consideration of £500,000 shall be satisfied in shares.",
        "3.3 10% of the Consideration of £2,500,000 is retained.")
    assert dealmath.arithmetic(doc) == []


def test_every_amount_and_cell_carries_its_clause_label(falcon):
    survey = dealmath.survey(falcon)
    price = next(a for a in survey["amounts"] if a["defines"] == "Purchase Price")
    assert (price["where"], price["value"], price["currency"]) == ("clause 3.1", 12_500_000, "GBP")
    cells = [a for a in survey["amounts"] if a["cell"]]
    assert {a["where"] for a in cells} == {"Schedule 1", "Schedule 2"}
    assert all(b.cell is None for b in falcon.blocks if b.kind != "table-cell")


# --- Timing -------------------------------------------------------------------


def test_the_falcon_timeline_flags_claims_before_release_and_the_cure(falcon):
    built = timeline.build(falcon)
    kinds = {f["kind"]: f for f in built["flags"]}
    assert set(kinds) == {"claims-before-release", "cure-longer-than-notice"}
    assert kinds["claims-before-release"]["where"] == ["clause 6.1", "clause 5.1"]
    assert [r["date"] for r in built["dated"]] == ["2026-09-01", "2026-10-30", "2026-12-31"]
    completion = [r for r in built["relative"] if r["from"] == "Completion"]
    assert [r["days"][0] for r in completion] == sorted(r["days"][0] for r in completion)


def test_a_long_stop_before_the_completion_date_is_flagged():
    doc = text_doc("This Agreement is dated 1 September 2026.",
                   "1. “Completion Date” means 30 January 2027. “Long Stop Date” means "
                   "31 December 2026.")
    (flag,) = timeline.build(doc)["flags"]
    assert flag["kind"] == "long-stop-before-date"
    assert "31 December 2026" in flag["title"] and "30 January 2027" in flag["title"]


def test_the_conflicts_the_table_already_proves_are_flags_too():
    doc = text_doc("This Agreement is dated 1 September 2026.",
                   "1. “Long Stop Date” means 31 August 2026.")
    assert [f["kind"] for f in timeline.build(doc)["flags"]] == ["long-stop-before-signing"]


# --- Party names --------------------------------------------------------------


def test_the_buyer_renamed_in_the_parties_clause_only(falcon):
    drift = [f for f in checks.run_all(falcon) if f.category == "party-name"]
    assert [(f.where, f.title) for f in drift] == [
        ("signature page", ("The Buyer's signature block names Falcon Holdings Limited as "
                           "the Buyer, but the parties clause says Falcon Topco Limited")),
        ("clause 8.1", ("Clause 8.1 names Falcon Holdings Limited as the Buyer, but the "
                       "parties clause says Falcon Topco Limited")),
    ]
    assert all(f.severity is Severity.BLOCKER for f in drift)
    # Set the way each place sets names.
    assert drift[0].suggested_text == "FALCON TOPCO LIMITED"
    assert drift[1].suggested_text == "Falcon Topco Limited"
    # Said once, as drift: not also a stranger signing and a party with no block.
    assert not [f for f in checks.run_all(falcon) if f.category == "signature"]


def test_against_the_previous_version_it_says_what_changed(falcon):
    earlier = docx_doc(realistic.falcon_spa(2))
    found = checks.party_drift_since(earlier, falcon, "in v3")
    assert [f.title for f in found] == [
        ("Clause 8.1 still names Falcon Holdings Limited as the Buyer; the parties "
        "clause now says Falcon Topco Limited"),
        ("The Buyer's signature block still names Falcon Holdings Limited as the Buyer; "
        "the parties clause now says Falcon Topco Limited"),
    ]
    assert "changed to Falcon Topco Limited in v3" in found[1].explanation
    assert checks.party_drift_since(earlier, docx_doc(realistic.falcon_spa(3))) == []


def test_a_role_two_parties_share_ties_nobody():
    doc = text_doc(
        "(1) Alpha Limited (each a “Seller”); (2) Beta Limited (each a “Seller”); and "
        "(3) Gamma Limited (the “Buyer”).",
        "1. Notices",
        "1.1 Seller: Delta Limited, 1 High Street.",
        "SIGNED by a director for and on behalf of the Seller, DELTA LIMITED",
        "By: ________", "Name: ", "Title: ")
    assert checks.party_drift(doc) == []


def test_the_signature_pack_is_blocked_by_drift():
    blockers = sigpack.survey(realistic.falcon_spa(3, defects=True), "Falcon.docx").blockers
    assert "The Buyer's signature block names Falcon Holdings Limited as the Buyer, but " \
           "the parties clause says Falcon Topco Limited" in blockers
    assert sigpack.survey(realistic.falcon_spa(3), "Falcon.docx").blockers == []


# --- Through the handler --------------------------------------------------------


@pytest.fixture
def captured(monkeypatch, tmp_path):
    from lra import handler
    from lra.config import settings
    from lra.mail.console import ConsoleProvider

    class Captured(ConsoleProvider):
        def __init__(self):
            self.sent = []

        def send(self, email_out):
            self.sent.append(email_out)
            return f"<sent-{len(self.sent)}@test>"

    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/d.sqlite3")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    settings.cache_clear()
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    monkeypatch.setattr(handler.review, "review", lambda doc, mode, *a, **k: ReviewResult(
        mode=mode, summary="", findings=[]))
    yield provider
    settings.cache_clear()


def _email(n, content, body="Quick look?", filename="Falcon SPA.docx", **kw):
    attachments = [Attachment(filename=filename, content_type=DOCX, size_bytes=len(content),
                              content=content)] if content else []
    return InboundEmail(message_id=f"<f-{n}@firm.com>", from_address="jim@firm.com",
                        to=["review@legal.firm.com"], subject="Falcon SPA", text_body=body,
                        attachments=attachments, received_at=dt.now(UTC), **kw)


def test_a_new_version_on_the_thread_says_the_buyer_was_renamed(captured):
    from lra import handler

    handler.handle(_email(1, realistic.falcon_spa(2)))
    handler.handle(_email(2, realistic.falcon_spa(3, defects=True)))
    body = captured.sent[1].text_body
    assert "still names Falcon Holdings Limited as the Buyer; the parties clause now " \
           "says Falcon Topco Limited" in body
