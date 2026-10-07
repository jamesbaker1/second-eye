"""The tracked-changes writer.

These are the tests that decide whether the product is worth anything. Every
case here is one that breaks a naive implementation, and every one of them
occurs routinely in documents that have been through several hands.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO

import pytest
from docx import Document

from lra.models import Finding, Mode, ReviewResult, Severity
from lra.pipeline import redline
from lra.pipeline.ooxml import (
    AnchorAmbiguous,
    AnchorNotFound,
    AnchorSpansParagraphs,
    Revision,
    RevisionWriter,
)

# --- builders -------------------------------------------------------------

def save(doc) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def simple(text: str) -> bytes:
    d = Document()
    d.add_paragraph(text)
    return save(d)


def split_runs(*parts) -> bytes:
    """A paragraph whose text is split across runs, some of them bold."""
    d = Document()
    p = d.add_paragraph()
    for part in parts:
        if isinstance(part, tuple):
            p.add_run(part[0]).bold = True
        else:
            p.add_run(part)
    return save(d)


def xml_of(content: bytes) -> str:
    return zipfile.ZipFile(BytesIO(content)).read("word/document.xml").decode()


def edit(anchor, replacement, auto_apply=True) -> ReviewResult:
    return ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.FORMATTING, category="c", title="t",
                explanation="", anchor=anchor, suggested_text=replacement,
                auto_apply=auto_apply)
    ])


def write(content: bytes, anchor: str, replacement: str | None):
    doc = Document(BytesIO(content))
    RevisionWriter(doc, author="Reviewer").apply(
        Revision(anchor=anchor, replacement=replacement, author="Reviewer")
    )
    return save(doc)


# --- the output must be a real tracked change -----------------------------

def test_produces_native_revision_markup():
    out = write(simple("The term is thirty days."), "thirty", "sixty")
    x = xml_of(out)
    assert "<w:ins " in x and "<w:del " in x
    assert "<w:delText" in x
    assert 'w:author="Reviewer"' in x


def test_output_opens_and_verifies():
    out = write(simple("The term is thirty days."), "thirty", "sixty")
    ok, reason = redline.verify(out, expect_revisions=True)
    assert ok, reason


def test_deleted_text_is_preserved_not_destroyed():
    """A tracked change must be rejectable, which means the old text is still there."""
    out = write(simple("The term is thirty days."), "thirty", "sixty")
    x = xml_of(out)
    assert ">thirty<" in x
    assert ">sixty<" in x


def test_revision_ids_are_unique():
    out = write(simple("Alpha and beta and gamma."), "beta", "delta")
    ids = re.findall(r'w:id="(\d+)"', xml_of(out))
    assert len(ids) == len(set(ids))


def test_revision_ids_do_not_collide_with_existing_ones():
    """A document already under revision by someone else must stay coherent."""
    base = simple("The term is thirty days.")
    items = {n: zipfile.ZipFile(BytesIO(base)).read(n)
             for n in zipfile.ZipFile(BytesIO(base)).namelist()}
    x = items["word/document.xml"].decode().replace(
        "<w:r>",
        '<w:ins w:id="5000" w:author="A. Associate" w:date="2026-01-01T00:00:00Z"><w:r>',
        1).replace("</w:r>", "</w:r></w:ins>", 1)
    items["word/document.xml"] = x.encode()
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in items.items():
            z.writestr(n, data)

    out = write(buf.getvalue(), "thirty", "sixty")
    ids = [int(i) for i in re.findall(r'w:id="(\d+)"', xml_of(out))]
    assert len(ids) == len(set(ids))
    assert max(ids) > 5000


# --- run splitting: the case that breaks naive implementations ------------

def test_anchor_spanning_three_runs_with_bold_in_the_middle():
    content = split_runs("Notice within ", ("thirty",), " (13) days of closing.")
    out = write(content, "thirty (13) days", "thirty (30) days")
    x = xml_of(out)
    assert "<w:ins " in x and "<w:del " in x
    assert redline.verify(out, expect_revisions=True)[0]


def test_formatting_survives_a_split_run_edit():
    content = split_runs("Notice within ", ("thirty",), " (13) days of closing.")
    out = write(content, "thirty (13) days", "thirty (30) days")
    x = xml_of(out)
    ins = re.search(r"<w:ins .*?</w:ins>", x, re.DOTALL).group(0)
    assert "<w:b/>" in ins, "inserted text lost the formatting of what it replaced"


def test_text_outside_the_anchor_is_untouched():
    content = split_runs("Notice within ", ("thirty",), " (13) days of closing.")
    out = write(content, "thirty (13) days", "thirty (30) days")
    x = xml_of(out)
    assert "Notice within " in x
    assert " of closing." in x


def test_partial_word_edit_at_a_run_boundary():
    content = split_runs("The Reci", "ever shall comply.")
    out = write(content, "Reciever", "Recipient")
    assert redline.verify(out, expect_revisions=True)[0]
    assert ">Recipient<" in xml_of(out)


# --- refusing to guess ----------------------------------------------------

def test_missing_anchor_raises_rather_than_editing_something_else():
    doc = Document(BytesIO(simple("The term is thirty days.")))
    with pytest.raises(AnchorNotFound):
        RevisionWriter(doc, "A").apply(Revision("ninety days", "sixty days", "A"))


def test_ambiguous_anchor_raises_rather_than_picking_one():
    doc = Document(BytesIO(simple("Payment is due. Payment is due.")))
    with pytest.raises(AnchorAmbiguous):
        RevisionWriter(doc, "A").apply(Revision("Payment is due", "Payment is late", "A"))


def test_anchor_crossing_a_paragraph_boundary_is_refused():
    d = Document()
    d.add_paragraph("The first part of the sentence")
    d.add_paragraph("and the second part.")
    doc = Document(BytesIO(save(d)))
    with pytest.raises(AnchorSpansParagraphs):
        RevisionWriter(doc, "A").apply(
            Revision("the sentence and the second", "x", "A"))


# --- whitespace and quote tolerance --------------------------------------

def test_anchor_matches_despite_collapsed_whitespace():
    out = write(simple("The  term   is thirty days."), "The term is thirty", "The term is sixty")
    assert redline.verify(out, expect_revisions=True)[0]


def test_anchor_matches_despite_curly_quotes():
    d = Document()
    d.add_paragraph('The “Agreement” means this document.')
    out = write(save(d), '"Agreement" means', '"Agreement" shall mean')
    assert redline.verify(out, expect_revisions=True)[0]


# --- tables, headers, and other places text hides -------------------------

def test_edit_inside_a_table_cell():
    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Term"
    t.rows[0].cells[1].text = "thirty days"
    out = write(save(d), "thirty days", "sixty days")
    assert redline.verify(out, expect_revisions=True)[0]


def test_edit_inside_a_header():
    d = Document()
    d.sections[0].header.paragraphs[0].text = "DRAFT - Acme Holdings"
    d.add_paragraph("Body text.")
    doc = Document(BytesIO(save(d)))
    RevisionWriter(doc, "A").apply(Revision("DRAFT - Acme", "Acme", "A"))
    assert redline.verify(save(doc))[0]


# --- other people's tracked changes --------------------------------------

def test_another_authors_revisions_are_not_stripped():
    from pathlib import Path

    original = Path("samples/redline-sample.docx").read_bytes()
    before = redline.revision_authors(original)
    assert "A. Associate" in before

    out = redline.apply(original, edit("Mutual Non-Disclosure", "Mutual Nondisclosure"))
    after = redline.revision_authors(out.content)
    assert "A. Associate" in after, "we destroyed someone else's tracked changes"


# --- the apply() contract -------------------------------------------------

def test_judgment_calls_are_not_auto_applied():
    out = redline.apply(simple("The term is thirty days."),
                        edit("thirty", "sixty", auto_apply=False))
    assert not out.applied
    assert len(out.skipped) == 1


def test_autopilot_applies_what_conservative_mode_would_not():
    out = redline.apply(simple("The term is thirty days."),
                        edit("thirty", "sixty", auto_apply=False), autopilot=True)
    assert len(out.applied) == 1


def test_memo_only_findings_are_never_applied():
    out = redline.apply(simple("Text."), edit("Text", None))
    assert not out.applied


def test_an_unlocatable_anchor_is_reported_not_silently_dropped():
    out = redline.apply(simple("The term is thirty days."), edit("ninety", "sixty"))
    assert not out.applied
    assert len(out.skipped) == 1


def test_an_ambiguous_anchor_produces_a_note_the_lawyer_can_read():
    out = redline.apply(simple("Payment is due. Payment is due."),
                        edit("Payment is due", "Payment is late"))
    assert out.notes
    assert "more than once" in out.notes[0]


def test_several_edits_in_one_document_all_land():
    d = Document()
    d.add_paragraph("The Reciever shall pay thirty (13) days after closing.")
    result = ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.FORMATTING, category="typo", title="a", explanation="",
                anchor="Reciever", suggested_text="Recipient", auto_apply=True),
        Finding(severity=Severity.BLOCKER, category="amount", title="b", explanation="",
                anchor="thirty (13) days", suggested_text="thirty (30) days",
                auto_apply=True),
    ])
    out = redline.apply(save(d), result)
    assert len(out.applied) == 2
    x = xml_of(out.content)
    assert x.count("<w:ins ") == 2
    assert ">Recipient<" in x and ">thirty (30) days<" in x


def test_original_is_returned_unchanged_when_nothing_applies():
    original = simple("Nothing to do here.")
    out = redline.apply(original, ReviewResult(mode=Mode.REDLINE, summary=""))
    assert out.content == original


# --- verification ---------------------------------------------------------

def test_verify_rejects_a_non_docx():
    ok, reason = redline.verify(b"this is not a docx at all")
    assert not ok and "zip" in reason


def test_verify_rejects_a_package_missing_document_xml():
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
    ok, reason = redline.verify(buf.getvalue())
    assert not ok and "missing parts" in reason


def test_verify_rejects_claimed_revisions_that_are_not_there():
    ok, reason = redline.verify(simple("No revisions here."), expect_revisions=True)
    assert not ok and "no revision markup" in reason


# --- damage the audit found -----------------------------------------------

def _tiny_png(tmp_path):
    import base64
    data = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    path = tmp_path / "px.png"
    path.write_bytes(data)
    return str(path)


def test_an_inline_image_is_never_duplicated_by_an_edit(tmp_path):
    """Splitting a run that also held a drawing produced three copies of the
    image, as an untracked change the recipient cannot reject."""
    from docx.shared import Inches

    d = Document()
    p = d.add_paragraph()
    run = p.add_run("The mark ")
    run.add_picture(_tiny_png(tmp_path), width=Inches(0.1))
    run.add_text(" is licensed for the term.")
    content = save(d)
    before = xml_of(content).count("<w:drawing")

    out = redline.apply(content, edit("licensed for the term", "licensed for the Term"))
    assert xml_of(out.content).count("<w:drawing") == before
    assert not out.applied
    assert out.notes and "image" in out.notes[0]


def test_text_inside_a_hyperlink_is_visible_to_the_writer():
    """The text model must match what Word displays, or offsets are wrong and
    the ambiguity guard cannot work."""
    from docx.oxml.ns import qn as _qn
    from lxml import etree as _etree

    from lra.pipeline.ooxml import _paragraph_text

    d = Document()
    p = d.add_paragraph()
    p.add_run("The Company shall pay ")
    link = _etree.SubElement(p._p, _qn("w:hyperlink"))
    run = _etree.SubElement(link, _qn("w:r"))
    node = _etree.SubElement(run, _qn("w:t"))
    node.text = "the Fee Schedule"
    p.add_run(" within 30 days.")

    reloaded = Document(BytesIO(save(d)))
    assert _paragraph_text(reloaded.paragraphs[0]) == (
        "The Company shall pay the Fee Schedule within 30 days."
    )


def test_an_edit_overlapping_a_hyperlink_is_refused_not_botched():
    from docx.oxml.ns import qn as _qn
    from lxml import etree as _etree

    d = Document()
    p = d.add_paragraph()
    p.add_run("Pay ")
    link = _etree.SubElement(p._p, _qn("w:hyperlink"))
    run = _etree.SubElement(link, _qn("w:r"))
    node = _etree.SubElement(run, _qn("w:t"))
    node.text = "the Fee Schedule"
    p.add_run(" promptly.")

    out = redline.apply(save(d), edit("the Fee Schedule", "the Fees Annex"))
    assert not out.applied
    assert out.notes and "hyperlink" in out.notes[0]


def test_a_refused_edit_leaves_the_paragraph_untouched():
    """The writer used to mutate the paragraph and then discover the problem,
    leaving the document damaged while reporting the finding as merely skipped."""
    import base64
    import os
    import tempfile

    from docx.shared import Inches

    data = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    fd, path = tempfile.mkstemp(suffix=".png")
    os.write(fd, data)
    os.close(fd)

    d = Document()
    p = d.add_paragraph()
    run = p.add_run("The mark ")
    run.add_picture(path, width=Inches(0.1))
    run.add_text(" is licensed here.")
    content = save(d)

    out = redline.apply(content, edit("licensed here", "licensed there"))
    assert out.content == content, "the document was modified despite the edit failing"


# --- comments: asking a question where the clause is ----------------------

def test_a_question_becomes_a_word_comment_anchored_to_the_clause():
    """A question that only appears in the email makes the lawyer hunt for the
    clause. A comment sits beside it in Word's review pane."""
    import zipfile

    content = simple("1. Term. The term is thirty (13) months.")
    out = redline.apply(content, ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.BLOCKER, category="amount", title="mismatch",
                explanation="", anchor="thirty (13)",
                question="Should this be 30 or 13?", options=["30", "13"]),
    ]))
    assert len(out.commented) == 1
    names = zipfile.ZipFile(BytesIO(out.content)).namelist()
    assert "word/comments.xml" in names
    body = zipfile.ZipFile(BytesIO(out.content)).read("word/comments.xml").decode()
    assert "Should this be 30 or 13?" in body
    assert "30 / 13" in body


def test_a_document_with_only_comments_still_verifies():
    content = simple("1. Term. The term is thirty (13) months.")
    out = redline.apply(content, ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.BLOCKER, category="amount", title="t",
                explanation="", anchor="thirty (13)", question="30 or 13?"),
    ]))
    ok, why = redline.verify(out.content)
    assert ok, why


def test_comments_and_tracked_changes_coexist():
    content = simple("The Reciever pays thirty (13) days after closing.")
    out = redline.apply(content, ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.FORMATTING, category="typo", title="typo",
                explanation="", anchor="Reciever", suggested_text="Recipient",
                auto_apply=True),
        Finding(severity=Severity.BLOCKER, category="amount", title="mismatch",
                explanation="", anchor="thirty (13)", question="30 or 13?"),
    ]))
    assert len(out.applied) == 1
    assert len(out.commented) == 1
    ok, why = redline.verify(out.content, expect_revisions=True)
    assert ok, why


def test_a_comment_is_attributed_to_the_agent():
    import zipfile

    content = simple("1. Term. The term is thirty (13) months.")
    out = redline.apply(content, ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.BLOCKER, category="amount", title="t",
                explanation="", anchor="thirty (13)", question="30 or 13?"),
    ]))
    body = zipfile.ZipFile(BytesIO(out.content)).read("word/comments.xml").decode()
    assert "Reviewer" in body


def test_a_comment_with_an_unlocatable_anchor_is_skipped_quietly():
    content = simple("1. Term. The term is thirty (13) months.")
    out = redline.apply(content, ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.BLOCKER, category="x", title="t", explanation="",
                anchor="text that is not in this document", question="Which?"),
    ]))
    assert out.commented == []
    assert out.content == content


def test_commenting_is_allowed_where_editing_is_refused():
    """A comment asserts nothing and changes nothing, so it is safe beside an
    image or inside a hyperlink where an edit would damage the document."""
    from docx.oxml.ns import qn as _qn
    from lxml import etree as _etree

    d = Document()
    p = d.add_paragraph()
    p.add_run("Pay ")
    link = _etree.SubElement(p._p, _qn("w:hyperlink"))
    run = _etree.SubElement(link, _qn("w:r"))
    node = _etree.SubElement(run, _qn("w:t"))
    node.text = "the Fee Schedule"
    p.add_run(" promptly.")
    content = save(d)

    out = redline.apply(content, ReviewResult(mode=Mode.REDLINE, summary="", findings=[
        Finding(severity=Severity.SUBSTANTIVE, category="x", title="t", explanation="",
                anchor="the Fee Schedule", question="Is this the right schedule?"),
    ]))
    assert len(out.commented) == 1


def test_the_quick_normalisation_is_the_mapped_one():
    """_locate skips a paragraph when the anchor is not in its quick
    normalisation, so the two must agree character for character, every
    Unicode space and dash included, or an anchor could be missed."""
    import sys

    from lra.pipeline import ooxml

    spaces = "".join(chr(c) for c in range(sys.maxunicode + 1) if chr(c).isspace())
    samples = [
        "", "   ", "a", " lead and trail ", "one\ttwo\n\nthree",
        "x" + spaces + "y", spaces + "edge" + spaces,
        "“quoted” ‘x’ – —",
        "thirty (30) days", "\x1c\x1d\x1e\x1fseparators\x85",
    ]
    for text in samples:
        assert ooxml._normalize(text) == ooxml._normalize_with_map(text)[0], repr(text)


# --- someone else's tracked changes, decided one by one -----------------------
#
# "Accept their typo fixes, reject the cap change": each of the other side's
# changes is accepted or rejected by id, and every one not named is left
# exactly as it was, id, author and text.

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_THEM = 'w:author="Them" w:date="2026-09-01T00:00:00Z"'


def _theirs() -> bytes:
    d = Document()
    d.add_paragraph("7.1 Each party keeps the other's informaton confidential.")
    d.add_paragraph("9.4 Liability is capped at 2x the Fees in any year.")
    d.add_paragraph("12.1 English law governs.")
    w = RevisionWriter(d, "Opposing Counsel")
    w.apply(Revision("informaton", "information", "Opposing Counsel"))
    w.apply(Revision("2x the Fees", "the Fees", "Opposing Counsel"))
    w.insert_paragraph_after("English law governs.", "12.2 Disputes go to arbitration.")
    return save(d)


def _texts(content: bytes) -> list[str]:
    from lra.pipeline.extract import visible_text

    return [visible_text(p._p) for p in Document(BytesIO(content)).paragraphs]


def _changes(content: bytes):
    from lra.pipeline import ooxml

    return ooxml.tracked_changes(Document(BytesIO(content)))


def _decide(content: bytes, accept=(), reject=()) -> bytes:
    from lra.pipeline import ooxml

    d = Document(BytesIO(content))
    ooxml.decide(d, accept=accept, reject=reject)
    return save(d)


def _with_body(*paragraphs: str) -> bytes:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = Document()
    body = d.element.body
    for p in list(body.iterchildren(_W + "p")):
        body.remove(p)
    for inner in paragraphs:
        body.insert(len(body) - 1, parse_xml(f"<w:p {nsdecls('w')}>{inner}</w:p>"))
    return save(d)


def test_their_changes_are_listed_as_the_decisions_word_offers():
    changes = _changes(_theirs())
    kinds = [(c.kind, c.text) for c in changes]
    assert ("new paragraph", "12.2 Disputes go to arbitration.") in kinds, \
        "a new paragraph and its mark are one decision"
    assert ("replaced", "informaton → information") in kinds, \
        "words struck and replaced by the same hand are one decision"
    assert ("replaced", "2x the Fees → the Fees") in kinds
    assert len(kinds) == 3
    assert {c.author for c in changes} == {"Opposing Counsel"}
    assert len({c.id for c in changes}) == len(changes)


def test_accepting_and_rejecting_by_id_touches_nothing_else():
    from lra.pipeline import validate

    original = _theirs()
    changes = {(c.kind, c.text): c for c in _changes(original)}
    typo = [changes[("replaced", "informaton → information")].id]
    cap = [changes[("replaced", "2x the Fees → the Fees")].id]

    out = _decide(original, accept=typo, reject=cap)
    texts = _texts(out)
    assert texts[0] == "7.1 Each party keeps the other's information confidential."
    assert texts[1] == "9.4 Liability is capped at 2x the Fees in any year."
    new = changes[("new paragraph", "12.2 Disputes go to arbitration.")]
    assert [(c.kind, c.id, c.author) for c in _changes(out)] == [
        ("new paragraph", new.id, "Opposing Counsel")]
    assert validate.validate(out).ok
    assert redline.verify(out)[0]


def test_rejecting_a_new_paragraph_takes_the_whole_paragraph_out():
    original = _theirs()
    new = next(c for c in _changes(original) if c.kind == "new paragraph")
    out = _decide(original, reject=[new.id])
    assert "12.2 Disputes go to arbitration." not in _texts(out)
    assert len(Document(BytesIO(out)).paragraphs) == 3, "no empty paragraph is left behind"
    accepted = _decide(original, accept=[new.id])
    assert "12.2 Disputes go to arbitration." in _texts(accepted)
    assert not [c for c in _changes(accepted) if c.kind == "new paragraph"]


def test_an_unknown_id_refuses_and_changes_nothing():
    from lra.pipeline import ooxml

    d = Document(BytesIO(_theirs()))
    before = [c.as_dict() for c in ooxml.tracked_changes(d)]
    with pytest.raises(ooxml.RevisionNotFound):
        ooxml.decide(d, accept=[before[0]["id"], "9999"])
    assert [c.as_dict() for c in ooxml.tracked_changes(d)] == before
    with pytest.raises(ooxml.CannotDecide):
        ooxml.decide(d, accept=[before[0]["id"]], reject=[before[0]["id"]])
    assert [c.as_dict() for c in ooxml.tracked_changes(d)] == before


def test_rejecting_their_insertion_keeps_the_comment_on_it():
    """Taking out someone's words must never take out the comment on them."""
    from docx.text.run import Run

    from lra.pipeline import validate, wordcomments

    d = Document(BytesIO(_theirs()))
    inserted = next(ins for ins in d.element.body.iter(_W + "ins")
                    if "".join(t.text for t in ins.iter(_W + "t")) == "the Fees")
    paragraph = d.paragraphs[1]
    d.add_comment([Run(r, paragraph) for r in inserted.findall(_W + "r")],
                  text="Why remove the multiple?", author="Partner", initials="P")
    marked = save(d)
    ids = [c.id for c in _changes(marked) if c.text == "2x the Fees → the Fees"]
    out = _decide(marked, reject=ids)
    assert [c.text for c in wordcomments.read(out)] == ["Why remove the multiple?"]
    assert "capped at 2x the Fees" in _texts(out)[1]
    report = validate.validate(out)
    assert report.ok, report


def test_a_move_is_one_decision_and_leaves_no_range_markers():
    content = _with_body(
        f'<w:moveFromRangeStart w:id="1" w:name="m1" {_THEM}/>'
        f'<w:moveFrom w:id="2" {_THEM}><w:r><w:delText>Moved words.</w:delText></w:r>'
        '</w:moveFrom><w:moveFromRangeEnd w:id="1"/><w:r><w:t xml:space="preserve"> Stay.'
        '</w:t></w:r>',
        f'<w:r><w:t xml:space="preserve">End. </w:t></w:r><w:moveToRangeStart w:id="3" '
        f'w:name="m1" {_THEM}/><w:moveTo w:id="4" {_THEM}><w:r><w:t>Moved words.</w:t>'
        '</w:r></w:moveTo><w:moveToRangeEnd w:id="3"/>')
    changes = _changes(content)
    assert [(c.kind, c.text) for c in changes] == [("moved", "Moved words.")]
    accepted = _decide(content, accept=[changes[0].id])
    assert _texts(accepted) == [" Stay.", "End. Moved words."]
    rejected = _decide(content, reject=["4"])
    assert _texts(rejected) == ["Moved words. Stay.", "End. "]
    for out in (accepted, rejected):
        assert b"moveFromRange" not in out and b"moveToRange" not in out
        assert not _changes(out)


def test_rejecting_a_formatting_change_restores_the_old_formatting():
    content = _with_body(
        f'<w:r><w:rPr><w:b/><w:rPrChange w:id="5" {_THEM}><w:rPr><w:i/></w:rPr>'
        '</w:rPrChange></w:rPr><w:t>Bold now.</w:t></w:r>')
    change = _changes(content)[0]
    assert change.kind == "formatting"
    rejected = Document(BytesIO(_decide(content, reject=[change.id])))
    assert rejected.paragraphs[0].runs[0].italic and not rejected.paragraphs[0].runs[0].bold
    accepted = _decide(content, accept=[change.id])
    assert Document(BytesIO(accepted)).paragraphs[0].runs[0].bold
    assert not _changes(accepted)


def test_a_table_cell_change_is_left_for_word():
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    from lra.pipeline import ooxml

    d = Document()
    table = d.add_table(rows=1, cols=1)
    table.rows[0].cells[0]._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:cellIns {nsdecls("w")} w:id="7" {_THEM}/>'))
    content = save(d)
    change = _changes(content)[0]
    assert change.kind == "other"
    with pytest.raises(ooxml.CannotDecide, match="Word"):
        ooxml.decide(Document(BytesIO(content)), accept=[change.id])
