"""File identification and the replies it produces.

Every case here used to be an unhandled exception that sent the sender a
generic "something went wrong". A lawyer who gets that once does not send a
second document, which is how the habit dies, so each one now answers with a
sentence that says what to do next.
"""

from __future__ import annotations

import zipfile
from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document

from lra.models import Attachment, InboundEmail, Mode
from lra.pipeline import extract, intake
from lra.pipeline.filetype import Kind, identify

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def docx_bytes(text="A clause.") -> bytes:
    d = Document()
    d.add_paragraph(text)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def zip_not_docx() -> bytes:
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("notes.txt", "hello")
    return buf.getvalue()


def att(name, content, ctype=DOCX) -> Attachment:
    return Attachment(filename=name, content_type=ctype,
                      size_bytes=len(content), content=content)


def email(*attachments) -> InboundEmail:
    return InboundEmail(message_id="m", from_address="jim@firm.com", subject="s",
                        attachments=list(attachments), received_at=datetime.now(UTC))


# --- identification -------------------------------------------------------

@pytest.mark.parametrize("content,expected", [
    (docx_bytes(), Kind.DOCX),
    (b"%PDF-1.7\nstuff", Kind.PDF),
    (b"{\\rtf1\\ansi}", Kind.RTF),
    (OLE + b"\x00" * 400, Kind.LEGACY_DOC),
    (OLE + b"EncryptedPackage" + b"\x00" * 100, Kind.ENCRYPTED),
    (zip_not_docx(), Kind.ZIP_NOT_DOCX),
    (b"", Kind.EMPTY),
    (b"Plain agreement text.", Kind.TEXT),
    (docx_bytes()[:300], Kind.DAMAGED),
])
def test_identified_by_content_not_extension(content, expected):
    assert identify(content, "anything.docx") is expected


def test_a_pdf_named_docx_is_still_a_pdf():
    assert identify(b"%PDF-1.7\nx", "Contract.docx") is Kind.PDF


# --- each failure produces a sentence, never a crash ----------------------

@pytest.mark.parametrize("name,content,expect_phrase", [
    ("locked.docx", OLE + b"EncryptedPackage" + b"\x00" * 100, "password protected"),
    ("fake.docx", zip_not_docx(), "not a Word document"),
    ("empty.docx", b"", "came through empty"),
    ("broken.docx", docx_bytes()[:300], "would not open"),
])
def test_unreviewable_files_are_rejected_with_something_actionable(
        name, content, expect_phrase):
    with pytest.raises(intake.Rejection) as e:
        intake.pick_document(email(att(name, content)))
    assert expect_phrase in str(e.value)


def test_a_legacy_doc_is_refused_only_where_nothing_can_convert_it(monkeypatch):
    """With LibreOffice (or the sandbox) available, a .doc is converted and
    reviewed; the refusal "save it as .docx" is for the machine that has
    neither. CI has LibreOffice and a developer's laptop may not, so the test
    pins both answers rather than whichever this machine happens to give."""
    from lra.pipeline import reflow

    legacy = att("old.docx", OLE + b"\x00" * 400)
    monkeypatch.setattr(reflow, "conversion_available", lambda: False)
    with pytest.raises(intake.Rejection, match="save it as .docx"):
        intake.pick_document(email(legacy))

    monkeypatch.setattr(reflow, "conversion_available", lambda: True)
    assert intake.pick_document(email(legacy)) is legacy


def test_a_pdf_is_selected_for_review():
    """PDFs used to raise here and tell the lawyer to reply "review it anyway",
    which raised again. There was no sequence of replies that produced a
    review. See tests/test_pdf.py for the extraction itself."""
    picked = intake.pick_document(email(att("brief.pdf", b"%PDF-1.7\nx",
                                            "application/pdf")))
    assert picked.filename == "brief.pdf"


def test_a_pdf_keeps_redline_mode_because_a_word_copy_is_built_from_it():
    """This used to downgrade every PDF to notes only. The tracked changes now
    go into a Word copy built from the PDF's text (pipeline/reflow.py), and a
    scan is decided when that build is attempted, not here."""
    a = att("brief.pdf", b"%PDF-1.7\nx", "application/pdf")
    assert intake.mode_for(a, Mode.REDLINE) is Mode.REDLINE


def test_a_deck_is_still_downgraded_to_memo_mode():
    """A slide has no paragraph to write a tracked change into."""
    from tests.test_local_reading import pptx_bytes

    a = att("deck.pptx", pptx_bytes(), "application/octet-stream")
    assert intake.mode_for(a, Mode.REDLINE) is Mode.MEMO_ONLY


def test_a_docx_stays_in_redline_mode():
    assert intake.mode_for(att("a.docx", docx_bytes()), Mode.REDLINE) is Mode.REDLINE


# --- choosing between attachments ----------------------------------------

def test_a_signature_logo_is_never_mistaken_for_the_document():
    a = att("signature.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 50, "image/png")
    with pytest.raises(intake.Rejection) as e:
        intake.pick_document(email(a))
    assert "images" in str(e.value)


def test_the_real_document_wins_over_an_attached_logo():
    logo = att("logo.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 50, "image/png")
    doc = att("NDA.docx", docx_bytes())
    assert intake.pick_document(email(logo, doc)).filename == "NDA.docx"


def test_a_docx_wins_over_an_attached_pdf_exhibit():
    pdf = att("exhibit.pdf", b"%PDF-1.7\nx", "application/pdf")
    doc = att("SPA.docx", docx_bytes())
    assert intake.pick_document(email(pdf, doc)).filename == "SPA.docx"


def test_two_real_documents_asks_and_names_them():
    with pytest.raises(intake.Rejection) as e:
        intake.pick_document(email(att("a.docx", docx_bytes()),
                                   att("b.docx", docx_bytes("Other."))))
    assert "a.docx" in str(e.value) and "b.docx" in str(e.value)


def test_a_broken_file_alongside_a_good_one_does_not_block_the_review():
    broken = att("corrupt.docx", docx_bytes()[:200])
    good = att("NDA.docx", docx_bytes())
    assert intake.pick_document(email(broken, good)).filename == "NDA.docx"


# --- extraction resilience ------------------------------------------------

def test_a_document_with_no_text_does_not_crash():
    d = Document()
    buf = BytesIO()
    d.save(buf)
    doc = extract.extract(att("blank.docx", buf.getvalue()))
    assert doc.blocks == []


def test_a_table_only_document_is_still_extracted():
    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Term"
    t.rows[0].cells[1].text = "thirty (13) days"
    buf = BytesIO()
    d.save(buf)
    doc = extract.extract(att("table.docx", buf.getvalue()))
    assert any("thirty (13) days" in b.text for b in doc.blocks)


def test_right_to_left_and_unicode_survive_extraction():
    d = Document()
    d.add_paragraph("هذه اتفاقية سرية بين الطرفين.")
    buf = BytesIO()
    d.save(buf)
    doc = extract.extract(att("اتفاقية.docx", buf.getvalue()))
    assert "اتفاقية" in doc.blocks[0].text


# --- very long documents --------------------------------------------------

def test_a_long_document_is_scoped_rather_than_silently_truncated():
    """Reviewing the first third of a credit agreement and sounding confident
    about it is worse than saying the document is too long."""
    d = Document()
    for i in range(3000):
        d.add_paragraph(f"{i}. A clause with a reasonable amount of text in it.")
    buf = BytesIO()
    d.save(buf)
    doc = extract.extract(att("huge.docx", buf.getvalue()))

    rendered = doc.as_prompt(limit=5000)
    assert doc.truncated_at is not None
    assert len(rendered) <= 5100


def test_a_normal_document_is_not_marked_truncated():
    doc = extract.extract(att("small.docx", docx_bytes("One short clause.")))
    doc.as_prompt(limit=240_000)
    assert doc.truncated_at is None


def test_an_oversize_document_is_named_not_silently_skipped(monkeypatch):
    """Reviewing the cover note instead of the execution copy and replying
    "looks good" was the worst outcome the audit found in intake."""
    monkeypatch.setenv("MAX_ATTACHMENT_MB", "1")
    from lra.config import settings

    settings.cache_clear()
    try:
        big = att("Falcon SPA EXECUTION.docx", docx_bytes("x" * 1_200_000))
        small = att("Cover note.docx", docx_bytes("A short note."))
        with pytest.raises(intake.Rejection) as e:
            intake.pick_document(email(big, small))
        assert "Falcon SPA EXECUTION.docx" in str(e.value)
    finally:
        settings.cache_clear()


# --- decompression bombs --------------------------------------------------

def bomb(megabytes: int = 21) -> bytes:
    """A .docx-shaped package whose document part is megabytes of one byte:
    a few tens of kilobytes on the wire, a thousand times that unpacked."""
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", b"\x20" * (megabytes * 1024 * 1024))
    return buf.getvalue()


def test_a_decompression_bomb_is_identified_without_being_inflated():
    content = bomb()
    assert len(content) < 1024 * 1024
    assert identify(content, "Contract.docx") is Kind.ARCHIVE_BOMB


def test_check_archive_bounds_total_size_and_entry_count(monkeypatch):
    from lra.pipeline import filetype

    filetype.check_archive(docx_bytes())  # a real document passes
    monkeypatch.setattr(filetype, "MAX_UNPACKED_BYTES", 1000)
    with pytest.raises(filetype.ArchiveTooLarge):
        filetype.check_archive(docx_bytes())
    monkeypatch.setattr(filetype, "MAX_UNPACKED_BYTES", 10**9)
    monkeypatch.setattr(filetype, "MAX_ENTRIES", 3)
    with pytest.raises(filetype.ArchiveTooLarge):
        filetype.check_archive(docx_bytes())


def test_a_bomb_is_refused_with_a_sentence_and_never_extracted():
    with pytest.raises(extract.Unreadable) as refused:
        extract.extract(att("Contract.docx", bomb()))
    assert "unpacks to far more" in str(refused.value)
    with pytest.raises(intake.Rejection):
        intake.pick_document(email(att("Contract.docx", bomb())))


def test_an_ordinary_document_is_still_a_docx():
    assert identify(docx_bytes("x" * 200_000)) is Kind.DOCX
