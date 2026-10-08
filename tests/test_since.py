"""What changed since the version I reviewed (ROADMAP item 8).

PRODUCT.md: the most valuable review is almost never everything wrong with the
contract, it is what changed since the version the lawyer sent on Tuesday. On a
second review of a document the reply leads with that, attaches the comparison,
and tells the agent where to dwell. On a resend, a document never seen, or any
failure inside the comparison, the review is exactly what it was before.

The model is stubbed; the documents, the comparison and the proof are real.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document

from secondeye import handler
from secondeye.mail.console import ConsoleProvider
from secondeye.models import Attachment, Finding, InboundEmail, ReviewResult, Severity
from secondeye.pipeline import compare, redline

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

V1 = [
    "1. Term. This Agreement continues for twelve (12) months.",
    "2. Liability. The Supplier's liability is capped at the fees paid.",
    "3. Governing law. New York law governs this Agreement.",
]
V2 = [
    "1. Term. This Agreement continues for twenty-four (24) months.",
    "2. Liability. The Supplier's liability is capped at the fees paid.",
    "3. Governing law. New York law governs this Agreement.",
    "4. Exclusivity. The Buyer shall purchase only from the Supplier.",
]


class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return "captured"


@pytest.fixture
def captured(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/since.sqlite3")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    from secondeye.config import settings

    settings.cache_clear()
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    yield provider
    settings.cache_clear()


@pytest.fixture
def seen(monkeypatch):
    """What the review agent was told, per call."""
    calls: list[dict] = []

    def _review(doc, mode, instructions, **kwargs):
        calls.append(kwargs)
        finding = Finding(severity=Severity.SUBSTANTIVE, category="liability",
                          title="The cap is the fees paid", explanation="Low for this deal.",
                          anchor="capped at the fees paid")
        return ReviewResult(mode=mode, summary="Looked it over.", findings=[finding])

    monkeypatch.setattr(handler.review, "review", _review)
    return calls


@pytest.fixture
def assessed(monkeypatch):
    def _assess(comparison, later, instructions, house_style):
        for change in comparison.changes:
            if "Exclusivity" in change.after:
                change.risk, change.impact = "high", "Locks the Buyer to one supplier."
            else:
                change.risk, change.impact = "low", "Doubles the term."

    monkeypatch.setattr(compare, "_assess", _assess)


def docx(paragraphs) -> bytes:
    d = Document()
    for text in paragraphs:
        d.add_paragraph(text)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def attach(name: str, content: bytes) -> Attachment:
    return Attachment(filename=name, content_type=DOCX, size_bytes=len(content), content=content)


def email(body: str, *attachments: Attachment, message_id: str) -> InboundEmail:
    return InboundEmail(
        message_id=message_id, from_address="jim@firm.com", to=["review@firm.com"],
        subject="Supply agreement", text_body=body, attachments=list(attachments),
        received_at=datetime.now(UTC),
    )


def test_a_second_review_of_a_changed_document_leads_with_what_changed(captured, seen,
                                                                       assessed):
    handler.handle(email("First look.", attach("Supply.docx", docx(V1)), message_id="s-1"))
    handler.handle(email("They sent it back.", attach("Supply.docx", docx(V2)),
                         message_id="s-2"))

    first, second = captured.sent
    assert "Since the version I reviewed" not in first.text_body

    body = second.text_body
    now = datetime.now(UTC)
    today = f"{now.day} {now.strftime('%b')}"
    heading = f"Since the version I reviewed on {today}"
    assert heading in body
    # Directly under the verdict, before anything else.
    verdict, _, first_line = body.split("\n")[:3]
    assert verdict.startswith("Nearly ready")
    assert first_line == heading
    assert "2 changes, 1 worth attention." in body
    # The risky one leads and the words are quoted, shortened.
    assert body.index("Key: Added, 4") < body.index("Changed, 1")
    assert "twelve (12)" in body and "twenty-four (24)" in body
    assert "<li>" in second.html_body and heading in second.html_body

    names = [a.filename for a in second.attachments]
    assert names == ["Supply (changes since last time).docx"]
    ok, reason = redline.verify(second.attachments[0].content, expect_revisions=True)
    assert ok, reason

    # The agent was told what moved, in its own block, and still ran in full.
    assert seen[0].get("changes_block", "") == ""
    block = seen[1]["changes_block"]
    assert "inserted at 4" in block and "Exclusivity" in block
    assert "replaced at 1" in block
    assert second.to == ["jim@firm.com"]


def test_an_identical_resend_gets_neither_section_nor_comparison(captured, seen):
    handler.handle(email("First look.", attach("Supply.docx", docx(V1)), message_id="s-1"))
    handler.handle(email("Sending again.", attach("Supply.docx", docx(V1)), message_id="s-2"))
    second = captured.sent[1]
    assert "Since the version I reviewed" not in second.text_body
    assert not any("changes since" in a.filename for a in second.attachments)
    assert seen[1].get("changes_block", "") == ""


def test_the_same_words_in_a_freshly_saved_file_are_a_resend_too(captured, seen):
    """Word rewrites the package on every save, so the bytes differ while the
    document does not. Comparing the text, not the bytes, is what keeps this
    from announcing "0 changes"."""
    handler.handle(email("First look.", attach("Supply.docx", docx(V1)), message_id="s-1"))
    d = Document(BytesIO(docx(V1)))
    d.core_properties.modified = datetime(2026, 9, 25, tzinfo=UTC)
    buf = BytesIO()
    d.save(buf)
    handler.handle(email("Resaved.", attach("Supply.docx", buf.getvalue()), message_id="s-2"))
    assert "Since the version I reviewed" not in captured.sent[1].text_body


def test_a_document_never_seen_gets_neither(captured, seen):
    handler.handle(email("First look.", attach("Supply.docx", docx(V1)), message_id="s-1"))
    other = ["1. Scope. The Consultant provides the Services described in Schedule 1.",
             "2. Fees. The Client pays the Fees monthly in arrears."]
    handler.handle(email("Different deal.", attach("Consultancy.docx", docx(other)),
                         message_id="s-2"))
    second = captured.sent[1]
    assert "Since the version I reviewed" not in second.text_body
    assert not any("changes since" in a.filename for a in second.attachments)


def test_a_same_named_but_different_document_is_not_compared(captured, seen):
    """Two clients' NDAs are both "NDA.docx". A by-name match is held to the
    similarity bar, so this gets an ordinary review rather than a comparison
    reporting a rewrite that never happened."""
    handler.handle(email("First look.", attach("NDA.docx", docx(V1)), message_id="s-1"))
    other = ["1. Purpose. The Parties wish to explore a possible transaction.",
             "2. Confidentiality. Each Party keeps the other's information secret.",
             "3. Term. Two years from the Effective Date."]
    handler.handle(email("Another one.", attach("NDA.docx", docx(other)), message_id="s-2"))
    assert "Since the version I reviewed" not in captured.sent[1].text_body


def test_a_failure_inside_the_comparison_is_swallowed_and_the_review_still_replies(
        captured, seen, monkeypatch):
    handler.handle(email("First look.", attach("Supply.docx", docx(V1)), message_id="s-1"))

    def boom(*a, **k):
        raise RuntimeError("the comparison fell over")

    monkeypatch.setattr(handler.compare, "compare", boom)
    handler.handle(email("They sent it back.", attach("Supply.docx", docx(V2)),
                         message_id="s-2"))
    second = captured.sent[1]
    assert second.text_body.startswith("Nearly ready")
    assert "The cap is the fees paid" in second.text_body
    assert "Since the version I reviewed" not in second.text_body
    assert [a.filename for a in second.attachments] == []


def test_the_model_failing_to_assess_still_gives_the_count(captured, seen, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("no network")

    monkeypatch.setattr(compare, "_assess", boom)
    handler.handle(email("First look.", attach("Supply.docx", docx(V1)), message_id="s-1"))
    handler.handle(email("Back again.", attach("Supply.docx", docx(V2)), message_id="s-2"))
    body = captured.sent[1].text_body
    assert "Since the version I reviewed" in body
    assert "2 changes." in body
    assert "worth attention" not in body


def test_a_pdf_of_the_document_reviewed_as_word_is_compared_too(captured, seen, assessed):
    """The lawyer reviewed the Word draft; the other side returned a PDF. The
    Word copy built from the PDF is what the comparison runs against."""
    from tests.test_versions import as_pdf

    handler.handle(email("First look.", attach("Supply.docx", docx(V1)), message_id="s-1"))
    handler.handle(email("Their PDF.", as_pdf("Supply.pdf", V2), message_id="s-2"))
    second = captured.sent[1]
    assert "Since the version I reviewed" in second.text_body
    assert "2 changes, 1 worth attention." in second.text_body
    names = [a.filename for a in second.attachments]
    assert "Supply (changes since last time).docx" in names
    # A PDF also gets its finding as a note on the page, as in any PDF review.
    assert "Supply (annotated).pdf" in names


# --- finding the earlier version quickly ------------------------------------
#
# Every review looks for an earlier version among the lawyer's conversations,
# before the model is called. That search used to cost a quarter of a second of
# character matching per conversation on a long agreement, and two database
# round trips per conversation for ledgers nobody read.


def test_the_quick_similarity_answers_exactly_as_the_full_measure_does():
    """The cheap tests only ever answer no early. Across the eval corpus, with
    drafts of each at several depths of edit, the verdict is the one the full
    character measure gives."""
    import random
    from difflib import SequenceMatcher

    from evals import corpus
    from secondeye import reconcile

    texts = {}
    for spec in corpus.ALL[::3]:
        raw = spec.build()
        texts[spec.name] = reconcile._text_of(attach(f"{spec.name}.docx", raw))
    rnd = random.Random(7)

    def redrafted(text: str, share: float) -> str:
        words = text.split(" ")
        for _ in range(int(len(words) * share)):
            i = rnd.randrange(len(words))
            words[i:i + 1] = rnd.choice([[], ["promptly"], ["the", "Purchaser"], ["thirty"]])
        return " ".join(words)

    names = list(texts)
    pairs = [(texts[a], texts[b]) for i, a in enumerate(names) for b in names[i + 1:i + 3]]
    pairs += [(texts[a], redrafted(texts[a], share)) for a in names
              for share in (0.05, 0.2, 0.35)]
    for a, b in pairs:
        full = SequenceMatcher(None, a[:20_000], b[:20_000],
                               autojunk=False).ratio() >= reconcile.MIN_SIMILARITY
        assert reconcile._similar(a, b) is full


def test_other_conversations_are_ruled_out_without_reading_their_ledgers(
        captured, seen, monkeypatch):
    from secondeye import thread

    others = [
        ["1. Scope. The Consultant provides the Services described in Schedule 1.",
         "2. Fees. The Client pays the Fees monthly in arrears."],
        ["1. Premises. The Landlord lets the Premises to the Tenant.",
         "2. Rent. The Tenant pays the Rent quarterly in advance."],
        ["1. Employment. The Employee is employed as Head of Legal.",
         "2. Hours. The Employee works 37.5 hours a week."],
    ]
    for i, paragraphs in enumerate(others):
        handler.handle(email("Look.", attach(f"Deal {i}.docx", docx(paragraphs)),
                             message_id=f"o-{i}"))
    loaded: list[str] = []
    real_load = thread.load
    monkeypatch.setattr(thread, "load", lambda key: loaded.append(key) or real_load(key))

    handler.handle(email("New one.", attach("Supply.docx", docx(V1)), message_id="s-1"))

    new_key = thread.recent_for_owner("jim@firm.com", limit=1)[0][0]
    assert set(loaded) <= {new_key}, "a conversation about another document was loaded whole"
    assert "Since the version I reviewed" not in captured.sent[-1].text_body


def test_a_renamed_redraft_is_still_found_by_its_words(captured, seen, assessed):
    handler.handle(email("First look.", attach("Supply.docx", docx(V1)), message_id="s-1"))
    handler.handle(email("Back from them.", attach("Supply v2 (theirs).docx", docx(V2)),
                         message_id="s-2"))
    assert "Since the version I reviewed" in captured.sent[1].text_body
