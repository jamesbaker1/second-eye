"""Blacklines by email: the PDF, the change summary, and baselines named in words.

Four layers, each tested where it lives. The renderer and the counts
(pipeline/blackline_pdf.py) directly. The skill (skills/lra-blackline) built
and run from its own folder by a fresh interpreter, as the container runs it,
once with pydantic and once with the stand-in. The host (blackline.py): the
versions a conversation keeps, what is offered to the agent, and what it
checks. And the whole email, through the handler, with a detached session
whose files are made by running the skill's scripts on what the host uploaded.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import zipfile
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

from secondeye import blackline, handler, managed, negotiation, skillsync, thread
from secondeye.mail.console import ConsoleProvider
from secondeye.models import Attachment, InboundEmail
from secondeye.pipeline import blackline_pdf as bp
from secondeye.pipeline import compare
from tests import fake_sessions as fs
from tests.conftest import documents

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

V1 = ["1. Term. This agreement starts on the Effective Date and lasts for twelve (12) months.",
      "2. Fees. The Customer shall pay each invoice within thirty (30) days of receipt.",
      ("3. Liability. The Supplier's total liability under this agreement is capped at the "
       "Fees paid in the twelve months before the claim, except for fraud."),
      "4. Confidentiality. Each party shall keep the other's Confidential Information secret.",
      "5. Governing law. English law governs this agreement."]
# What we sent: the term to twenty-four months.
V2 = [V1[0].replace("twelve (12)", "twenty-four (24)"), *V1[1:]]
# What they sent back: our term kept, payment doubled, the liability clause
# moved to the end without its fraud carve-out, and an audit right added.
V3 = [V2[0], V1[1].replace("thirty (30)", "sixty (60)"), V1[3], V1[4],
      ("6. Liability. The Supplier's total liability under this agreement is capped at the "
       "Fees paid in the twelve months before the claim."),
      "7. Audit. The Customer may audit the Supplier once in each year.",
      "8. Notices. Notices must be in writing and sent by email."]


def docx(paragraphs: list[str], title: str = "Services Agreement") -> bytes:
    d = Document()
    # Fixed timestamps: python-docx stamps "now", so the same text built a
    # second apart hashed differently and the version-keeping test flaked.
    from datetime import UTC, datetime
    d.core_properties.created = d.core_properties.modified = datetime(2026, 9, 1, tzinfo=UTC)
    d.add_heading(title, 1)
    for text in paragraphs:
        d.add_paragraph(text)
    buffer = BytesIO()
    d.save(buffer)
    # The zip stamps each part with the save time to two seconds, so the
    # fixed properties alone still let a slow run straddle a tick. Re-pack
    # with a fixed time: the same text is then the same bytes.
    out = BytesIO()
    with zipfile.ZipFile(buffer) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(zipfile.ZipInfo(info.filename, (2026, 9, 1, 0, 0, 0)),
                         src.read(info.filename), zipfile.ZIP_DEFLATED)
    return out.getvalue()


def labels() -> bp.Labels:
    return bp.Labels("Services v1.docx", "Services v3.docx", "the original", "attached today")


# --------------------------------------------------------------------------
# The counts and the PDF
# --------------------------------------------------------------------------


def test_the_counts_are_by_kind_and_by_clause():
    result = compare.compare(docx(V1), docx(V3))
    summary = bp.summarise(result)
    assert summary["counts"] == {"insertions": 4, "deletions": 2, "moves": 1,
                                 "changes": 5, "clauses": 5}
    assert summary["sentence"] == ("5 changes: 4 insertions, 2 deletions, 1 move, "
                                   "in 5 clauses.")
    by_clause = {c["clause"]: c for c in summary["by_clause"]}
    assert by_clause["2"] == {"clause": "2", "insertions": 1, "deletions": 1, "moves": 0}
    assert by_clause["6"]["moves"] == 1, "a move is counted where it landed"
    fees = next(c for c in summary["changes"] if c["clause"] == "2")
    assert (fees["before"], fees["after"]) == ("… within thirty (30) days",
                                               "… within sixty (60) days")


def test_the_pdf_has_the_summary_first_and_every_change_of_both_versions():
    result = compare.compare(docx(V1), docx(V3))
    material = [bp.Material("2", "within thirty (30) days", "within sixty (60) days",
                            "Doubles the time the Customer has to pay.")]
    rendered = bp.render(result.content, result.changes, labels(), material, engine="typeset")
    texts = bp.page_texts(rendered.content)
    assert rendered.engine == "typeset" and rendered.pages == len(texts) >= 2
    assert "Change summary" in texts[0] and "Doubles the time the Customer" in texts[0]
    assert "5 changes: 4 insertions, 2 deletions, 1 move, in 5 clauses." in texts[0]
    for number, text in enumerate(texts, start=1):
        assert "Blackline: Services v3.docx against Services v1.docx" in text
        assert f"Page {number} of {len(texts)}" in text
    body = "".join(texts[1:])
    for words in ("twelve (12", "twenty-four (24", "thirty (30", "sixty (60",
                  "except for fraud", "7. Audit.", "8. Notices."):
        assert words in body
    assert bp.check(rendered.content, result.changes, labels(), bp.summarise(result)["sentence"]) == []


def test_the_check_finds_a_blackline_that_misses_a_change():
    result = compare.compare(docx(V1), docx(V3))
    rendered = bp.render(result.content, result.changes, labels(), [], engine="typeset")
    missing = compare.Change("inserted", "9", after="9. Assignment. Neither party may assign.")
    problems = bp.check(rendered.content, [*result.changes, missing], labels(), None)
    assert problems and "9. Assignment" in problems[0]
    assert bp.check(b"not a pdf", result.changes, labels(), None)[0].startswith(
        "the PDF does not open")
    other = bp.Labels("A.docx", "B.docx")
    assert any("is not headed" in p for p in bp.check(rendered.content, [], other, None))


def test_text_outside_the_standard_fonts_is_drawn_and_checked_the_same_way():
    old, new = docx(["1. Price. The price is €100 → payable now."]), docx(
        ["1. Price. The price is €200 → payable now, Zoë's rate ≤ cap."])
    result = compare.compare(old, new)
    rendered = bp.render(result.content, result.changes, labels(), [], engine="typeset")
    assert bp.check(rendered.content, result.changes, labels(), None) == []


def test_auto_prints_through_libreoffice_only_when_it_shows_the_changes(monkeypatch):
    result = compare.compare(docx(V1), docx(V3))
    # A LibreOffice that printed the document without its changes.
    blank = bp.write_pdf([bp.Canvas()])
    monkeypatch.setattr(bp, "libreoffice_pdf", lambda content: blank)
    rendered = bp.render(result.content, result.changes, labels(), [])
    assert rendered.engine == "typeset"
    with pytest.raises(RuntimeError, match="not in the blackline"):
        bp.render(result.content, result.changes, labels(), [], engine="libreoffice")

    # One that did: its pages are kept, stamped, behind our summary.
    printed = bp.render(result.content, result.changes, labels(), [], engine="typeset")
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for page in PdfReader(BytesIO(printed.content)).pages[1:]:
        writer.add_page(page)
    body = BytesIO()
    writer.write(body)
    monkeypatch.setattr(bp, "libreoffice_pdf", lambda content: body.getvalue())
    through = bp.render(result.content, result.changes, labels(), [])
    assert through.engine == "libreoffice"
    assert bp.check(through.content, result.changes, labels(),
                    bp.summarise(result)["sentence"]) == []


# --------------------------------------------------------------------------
# The skill, run from its bundle as the container runs it
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def bundle(tmp_path_factory) -> Path:
    return skillsync.build(tmp_path_factory.mktemp("skill"), "lra-blackline")


def run(bundle: Path, script: str, *args: str, shim: bool = False, ok: bool = True) -> dict:
    env = {k: v for k, v in os.environ.items() if k != "PYTHONPATH"}
    if shim:
        env["SECOND_EYE_FORCE_SHIM"] = "1"
    done = subprocess.run([sys.executable, str(bundle / "scripts" / script), *args],
                          capture_output=True, text=True, env=env, cwd=bundle,
                          timeout=120, check=False)
    assert done.stdout, done.stderr
    assert (done.returncode == 0) == ok, done.stdout
    return json.loads(done.stdout)


def test_the_bundle_has_the_shape_the_skills_api_requires(bundle):
    head = (bundle / "SKILL.md").read_text().split("---")[1]
    name = next(line for line in head.splitlines() if line.startswith("name:"))
    description = next(line for line in head.splitlines() if line.startswith("description:"))
    assert name.split(":", 1)[1].strip() == "lra-blackline"
    assert 0 < len(description.split(":", 1)[1].strip()) <= 1024
    for relative in skillsync.BLACKLINE_BUNDLED:
        assert (bundle / "lib" / "secondeye" / relative).read_bytes() == (
            skillsync.PACKAGE / relative).read_bytes()
    for script in ("pick_baseline.py", "summary.py", "render_blackline_pdf.py"):
        assert (bundle / "scripts" / script).exists()


def workspace(tmp_path: Path, today: datetime) -> Path:
    """Three versions and their list, as the host mounts them."""
    root = tmp_path / "workspace"
    root.mkdir()
    entries = []
    for n, (paras, source, days) in enumerate(
            [(V1, "received", 9), (V2, "sent", 6), (V3, "attached", 0)], start=1):
        path = root / f"V{n} Services v{n}.docx"
        path.write_bytes(docx(paras))
        when = (today - timedelta(days=days)).isoformat()
        entries.append({"id": f"V{n}", "file": str(path), "filename": f"Services v{n}.docx",
                        "label": f"v{n}", "source": source, "from": "", "when": when,
                        "day": when[:10], "ordinal": n, "version": n, "current": n == 3,
                        "detail": f"{source} {when[:10]}"})
    (root / "blackline-versions.json").write_text(json.dumps(
        {"today": today.date().isoformat(), "ask": "", "versions": entries}))
    return root


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_a_blackline_runs_from_the_bundle(bundle, tmp_path, shim):
    today = datetime(2026, 9, 28, 10, tzinfo=UTC)          # a Monday
    root = workspace(tmp_path, today)
    listing = str(root / "blackline-versions.json")

    picked = run(bundle, "pick_baseline.py", listing, "--ask",
                 "Blackline against what we sent on Tuesday", shim=shim)
    assert picked["suggested"] == {"earlier": "V2", "later": "V3"}
    assert not picked["ambiguous"]
    sent = next(v for v in picked["versions"] if v["id"] == "V2")
    assert "it is what we sent" in sent["baseline_because"]
    assert any("Tuesday" in r for r in sent["baseline_because"])

    assert run(bundle, "pick_baseline.py", listing, "--ask", "blackline v3 against v1",
               shim=shim)["suggested"] == {"earlier": "V1", "later": "V3"}
    assert run(bundle, "pick_baseline.py", listing, "--ask", "against the original",
               shim=shim)["suggested"]["earlier"] == "V1"
    vague = run(bundle, "pick_baseline.py", listing, "--ask", "blackline it", shim=shim)
    assert vague["ambiguous"] and vague["suggested"]["earlier"] is None

    summary = run(bundle, "summary.py", str(root / "V1 Services v1.docx"),
                  str(root / "V3 Services v3.docx"), shim=shim)
    assert summary["counts"]["changes"] == 5 and summary["proven"] is True
    fees = next(c for c in summary["changes"] if c["clause"] == "2")

    decision = tmp_path / "significance.json"
    decision.write_text(json.dumps({"earlier": "V1", "later": "V3", "why": "", "material": [
        {"index": fees["index"], "significance": "Doubles the time the Customer has to pay."}]}))
    outputs = tmp_path / "outputs"
    done = run(bundle, "render_blackline_pdf.py", "--versions", listing, "--significance",
               str(decision), "--outdir", str(outputs), "--engine", "typeset", shim=shim)
    assert done["pdf"] == "Services v3 (blackline against v1).pdf"
    assert done["docx"] == "Services v3 (blackline against v1).docx"
    assert done["engine"] == "typeset" and done["pages"] >= 2
    assert json.loads((outputs / "blackline.json").read_text())["material"][0]["clause"] == "2"
    pdf = (outputs / done["pdf"]).read_bytes()
    from pypdf import PdfReader

    assert len(PdfReader(BytesIO(pdf)).pages) == done["pages"]


def test_render_refuses_a_bad_significance_file_and_writes_nothing(bundle, tmp_path):
    root = workspace(tmp_path, datetime(2026, 9, 28, tzinfo=UTC))
    decision = tmp_path / "significance.json"
    decision.write_text(json.dumps({"earlier": "V1", "later": "V3", "material": [
        {"index": 99, "significance": "x"}, {"index": 0, "significance": "two\nlines"}]}))
    outputs = tmp_path / "outputs"
    failed = run(bundle, "render_blackline_pdf.py", "--versions",
                 str(root / "blackline-versions.json"), "--significance", str(decision),
                 "--outdir", str(outputs), ok=False)
    assert "material[0]: index must be" in failed["error"]
    assert "material[1]: significance must be one line" in failed["error"]
    assert not outputs.exists()
    decision.write_text(json.dumps({"earlier": "V9", "later": "V3"}))
    assert "earlier" in run(bundle, "render_blackline_pdf.py", "--versions",
                            str(root / "blackline-versions.json"), "--significance",
                            str(decision), ok=False)["error"]


# --------------------------------------------------------------------------
# The host: the versions a conversation keeps
# --------------------------------------------------------------------------


class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return f"agent-msg-{len(self.sent)}"


@pytest.fixture
def env(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/b.sqlite3")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    monkeypatch.setenv("ALLOWLIST_CONTACT", "Jim Baker <jim@firm.com>")
    fs.configure(monkeypatch, MANAGED_BLACKLINE_AGENT_ID="agent_blackline",
                 SANDBOX_BLACKLINE_SKILL_ID="skill_blackline", MANAGED_PLAYBOOK_AGENT_ID="",
                 MANAGED_COMMENTS_AGENT_ID="")
    now = [0.0]
    monkeypatch.setattr(managed.time, "monotonic", lambda: now[0])
    monkeypatch.setattr(managed.time, "sleep", lambda s: now.__setitem__(0, now[0] + s))
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    yield provider
    from secondeye.config import settings
    settings.cache_clear()


def email(body: str, attachments: list[tuple[str, bytes]] | None = None,
          reply_to: str | None = None) -> InboundEmail:
    return InboundEmail(
        message_id=f"<m-{abs(hash(body))}@firm.com>", thread_id="t-blackline",
        in_reply_to=reply_to, from_address="jim@firm.com", to=["review@legal.firm.com"],
        subject="Services Agreement", text_body=body, received_at=datetime.now(UTC),
        attachments=[Attachment(filename=n, content_type=DOCX, size_bytes=len(c), content=c)
                     for n, c in attachments or []])


def conversation() -> str:
    """v1 reviewed on Tuesday of last week, v2 sent back to them last
    Tuesday (we were copied), v3 arrived from them today."""
    key = thread.key("t-blackline", "<first@firm.com>", "jim@firm.com")
    thread.start(key, "jim@firm.com", "Services v1.docx", docx(V1))
    thread.add_alias(key, "<first@firm.com>", "t-blackline", "<our-reply@legal.firm.com>")
    negotiation.record_sent(key, "Services v2.docx", docx(V2))
    thread.supersede(key, "Services v3.docx", docx(V3))
    today = datetime.now(UTC)
    tuesday = today - timedelta(days=(today.weekday() - 1) % 7 or 7)
    from secondeye.store import connect

    with connect() as c:
        c.execute("UPDATE blackline_versions SET arrived_at = ? WHERE filename = ?",
                  ((tuesday - timedelta(days=7)).isoformat(), "Services v1.docx"))
        c.execute("UPDATE negotiation_sent SET created_at = ?", (tuesday.isoformat(),))
    return key


def test_every_version_a_conversation_holds_is_kept_and_offered(env):
    key = conversation()
    assert [name for name, *_ in blackline.stored(key)] == ["Services v1.docx"]
    thread.supersede(key, "Services v3.docx", docx(V3))       # the same bytes: one version
    assert len(blackline.stored(key)) == 1

    offered = blackline.candidates("jim@firm.com", key, email("blackline against v1"))
    assert [(v.id, v.filename, v.source) for v in offered] == [
        ("V1", "Services v1.docx", "received"), ("V2", "Services v2.docx", "sent"),
        ("V3", "Services v3.docx", "received")]
    assert offered[-1].current and offered[0].ordinal == 1
    listed = blackline.listing(offered, "blackline against v1", datetime.now(UTC))
    first = listed["versions"][0]
    assert first["file"] == "/workspace/V1 Services v1.docx"
    assert first["detail"].startswith("the first version on this conversation, received on ")
    assert listed["versions"][1]["detail"].startswith("what we sent on ")


def test_an_attached_copy_of_a_version_we_hold_is_that_version(env):
    key = conversation()
    offered = blackline.candidates("jim@firm.com", key, email(
        "blackline against the original", [("Services v3.docx", docx(V3))]))
    assert [v.id for v in offered] == ["V1", "V2", "V3"]
    assert offered[-1].current


def test_the_versions_follow_the_negotiation_ledger(env):
    key = conversation()
    negotiation.move(key, "elsewhere")
    assert blackline.stored(key) == [] and len(blackline.stored("elsewhere")) == 1


def test_versions_go_when_their_conversation_does(env):
    key = conversation()
    assert blackline.stored(key)
    thread.purge("jim@firm.com")
    assert blackline.stored(key) == []


@pytest.mark.parametrize("body", [
    "Blackline against what we sent on Tuesday",
    "Can you blackline v4 against v1?",
    "Please send me a blackline against the original.",
    "blackline please",
    "Compare this against what we sent.",
    "Redline it to v2 please",
])
def test_blackline_requests_are_recognised(env, body):
    assert blackline.intent(email(body))


@pytest.mark.parametrize("body", [
    "Their blackline is attached, please review.",
    "Blackline attached for your review.",
    "Please review the attached draft.",
    "Compare these two.",
])
def test_other_requests_are_not_blacklines(env, body):
    assert not blackline.intent(email(body))


def test_without_the_agent_a_blackline_is_the_comparison_it_was(env, monkeypatch):
    monkeypatch.setenv("MANAGED_BLACKLINE_AGENT_ID", "")
    from secondeye.config import settings

    settings.cache_clear()
    assert not blackline.intent(email("Blackline against what we sent on Tuesday"))


# --------------------------------------------------------------------------
# The whole email
# --------------------------------------------------------------------------


class SkillSessions(fs.DetachedSessions):
    """A detached session whose agent runs the real skill on what the host
    uploaded: pick_baseline.py on the lawyer's words, a decision, then
    render_blackline_pdf.py. `decide` stands in for the model's choices."""

    def __init__(self, bundle: Path, work: Path, decide, spoil=None, turns: int = 1):
        super().__init__([fs.Turn(fs.message("Done."), fs.idle()) for _ in range(turns)])
        self.bundle, self.work, self.decide, self.spoil = bundle, work, decide, spoil

    def create(self, **kwargs):
        made = super().create(**kwargs)
        root = self.work / "workspace"
        root.mkdir()
        for name, data in self.uploaded:
            (root / managed.mount_name(name)).write_bytes(data)
        listing = json.loads((root / blackline.VERSIONS_FILE).read_text())
        for entry in listing["versions"]:
            entry["file"] = entry["file"].replace(managed.WORKSPACE, str(root))
        (root / blackline.VERSIONS_FILE).write_text(json.dumps(listing))
        picked = run(self.bundle, "pick_baseline.py", str(root / blackline.VERSIONS_FILE))
        decision = self.work / "significance.json"
        decision.write_text(json.dumps(self.decide(picked, listing)))
        outdir = self.work / "outputs"
        run(self.bundle, "render_blackline_pdf.py", "--versions",
            str(root / blackline.VERSIONS_FILE), "--significance", str(decision),
            "--outdir", str(outdir), "--engine", "typeset")
        outputs = {p.name: p.read_bytes() for p in outdir.iterdir()}
        if self.spoil:
            first = dict(outputs)
            self.spoil(first)
            self._turn.outputs = first
            if self.turns:                  # the turn after our report: the fix
                self.turns[0].outputs = outputs
        else:
            self._turn.outputs = outputs
        return made


def material_fees(picked, listing):
    """The model's part: the suggested versions, and clause 2 as material."""
    earlier, later = picked["suggested"]["earlier"], picked["suggested"]["later"]
    files = {v["id"]: Path(v["file"]).read_bytes() for v in listing["versions"]}
    changes = bp.summarise(compare.compare(files[earlier], files[later]))["changes"]
    fees = next(c["index"] for c in changes if c["clause"] == "2")
    return {"earlier": earlier, "later": later, "why": "", "material": [
        {"index": fees, "significance": "Doubles the time the Customer has to pay."}]}


def test_blackline_against_what_we_sent_on_tuesday_end_to_end(env, bundle, tmp_path):
    key = conversation()
    fake = SkillSessions(bundle, tmp_path, material_fees)
    ask = "Blackline against what we sent on Tuesday please."
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email(ask, reply_to="<our-reply@legal.firm.com>"))

    out = env.sent[-1]
    first = out.text_body.splitlines()[0]
    assert first == ("Blackline: Services v3.docx against Services v2.docx. "
                     "4 changes, 1 material.")
    assert "Earlier: Services v2.docx, what we sent on " in out.text_body
    assert "Later: Services v3.docx, received today." in out.text_body
    assert "4 changes: 3 insertions, 1 deletion, 1 move, in 4 clauses." in out.text_body
    assert ("  - cl. 2: was “… within thirty (30) days”, now “… within sixty (60) days”. "
            "Doubles the time the Customer has to pay.") in out.text_body
    assert "<table" in out.html_body and "Why it matters" in out.html_body

    attached = {a.filename: a for a in documents(out)}
    pdf = attached["Services v3 (blackline against v2).pdf"]
    assert pdf.content_type == "application/pdf"
    word = attached["Services v3 (blackline against v2).docx"]
    earlier, later = docx(V2), docx(V3)
    assert blackline.proves(word.content, earlier, later, complete=True)
    result = compare.compare(earlier, later)
    labels_ = bp.Labels("Services v2.docx", "Services v3.docx")
    assert not [p for p in bp.check(pdf.content, result.changes, None, None)]
    assert "Blackline: " + labels_.later in bp.page_texts(pdf.content)[0]

    created = fake.created[0]
    assert created["agent"] == "agent_blackline"
    opening = created["initial_events"][0]["content"][0]["text"]
    assert ask in opening and opening.rstrip().endswith("two or three lines.")
    assert sorted(name for name, _ in fake.uploaded) == [
        "V1 Services v1.docx", "V2 Services v2.docx", "V3 Services v3.docx",
        blackline.VERSIONS_FILE]
    assert fake.deleted_sessions == ["sesn_fake"]
    assert thread.load(key).filename == "Services v3.docx", "the conversation is unchanged"


def test_cumulative_blackline_of_an_attached_version_against_the_original(env, bundle, tmp_path):
    conversation()
    v4 = docx([*V3, "9. Assignment. Neither party may assign this agreement."])

    def decide(picked, listing):
        return {"earlier": "V1", "later": picked["suggested"]["later"], "material": []}

    fake = SkillSessions(bundle, tmp_path, decide)
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Blackline v4 against the original",
                             [("Services v4.docx", v4)]))
    out = env.sent[-1]
    assert out.text_body.startswith("Blackline: Services v4.docx against Services v1.docx.")
    assert "Later: Services v4.docx, attached to your email today" in out.text_body
    assert len(documents(out)) == 2


def test_a_spoiled_pdf_goes_back_once_and_the_fixed_one_is_sent(env, bundle, tmp_path):
    conversation()

    def spoil(outputs):
        name = next(n for n in outputs if n.endswith(".pdf"))
        outputs[name] = bp.write_pdf([bp.Canvas()])

    fake = SkillSessions(bundle, tmp_path, material_fees, spoil=spoil, turns=2)
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Blackline against v1 please", reply_to="<our-reply@legal.firm.com>"))
    correction = fake.sent[0][0]["content"][0]["text"]
    assert "PDF: " in correction and "render_blackline_pdf.py" in correction
    out = env.sent[-1]
    assert any(a.filename.endswith(".pdf") for a in documents(out))
    assert "did not pass" not in out.text_body


def test_a_pdf_that_never_passes_is_not_attached_but_the_word_file_is(env, bundle, tmp_path):
    conversation()

    def spoil(outputs):
        name = next(n for n in outputs if n.endswith(".pdf"))
        outputs[name] = b"%PDF-1.4 broken"

    fake = SkillSessions(bundle, tmp_path, material_fees, spoil=spoil, turns=1)
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Blackline against v1 please", reply_to="<our-reply@legal.firm.com>"))
    out = env.sent[-1]
    assert [a.filename.rsplit(".", 1)[1] for a in documents(out)] == ["docx"]
    assert "The PDF blackline did not pass my checks" in out.text_body


def test_one_version_is_not_enough(env):
    fake = fs.DetachedSessions([])
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(email("Blackline against the original", [("Services v1.docx",
                                                                  docx(V1))]))
    out = env.sent[-1]
    assert out.text_body.startswith("I can't make a blackline: I have only Services v1.docx")
    assert fake.created == []


def test_the_blackline_skill_is_attached_to_its_own_agent_only(monkeypatch):
    fs.configure(monkeypatch, SANDBOX_BLACKLINE_SKILL_ID="skill_blackline")
    ref = {"type": "custom", "skill_id": "skill_blackline", "version": "latest"}
    own = managed.agent_body(managed.load_manifest(managed.AGENTS / "blackline.agent.yaml"), [])
    review = managed.agent_body(managed.load_manifest(managed.AGENTS / "review.agent.yaml"), [])
    assert ref in own["skills"] and ref not in review["skills"]
    assert own["system"].startswith("You are a junior associate preparing a blackline")
    assert not [t for t in own["tools"] if t.get("type") == "custom"]
    from secondeye.config import settings

    settings.cache_clear()
