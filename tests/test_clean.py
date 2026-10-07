"""The clean copy: nothing in the file but the document, and the words intact."""

from __future__ import annotations

import zipfile
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from lra.pipeline import checks, clean, compare, redline
from lra.pipeline.ooxml import Revision, RevisionWriter


def save(document) -> bytes:
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def marked_up() -> bytes:
    """Tracked changes, a comment, hidden text and an author, all at once."""
    d = Document()
    d.add_paragraph("1. The term is thirty days from signing.")
    d.add_paragraph("2. The Buyer shall pay the Price on completion.")
    secret = d.add_paragraph("3. Governing law is New York.").add_run(" ASK JANE ABOUT THIS")
    secret.font.hidden = True
    d.core_properties.author = "Jane Associate"
    d.core_properties.last_modified_by = "Peter Partner"
    d.core_properties.title = "Acme NDA built from the Beta precedent"

    writer = RevisionWriter(d, "Peter Partner")
    writer.apply(Revision("thirty days", "sixty days", "Peter Partner"))
    writer.insert_paragraph_after("2. The Buyer", "2A. Interest accrues at 8%.")
    writer.comment("on completion", "Is this what the client agreed?")
    return save(d)


def text_of(content: bytes) -> list[str]:
    return [p.text for p in Document(BytesIO(content)).paragraphs if p.text.strip()]


def test_everything_the_leftovers_check_finds_is_gone():
    raw = marked_up()
    assert len(checks.leftovers(raw)) >= 4
    out = clean.clean(raw)
    assert checks.leftovers(out.content) == []


def test_the_words_are_the_accepted_document():
    out = clean.clean(marked_up())
    assert text_of(out.content) == [
        "1. The term is sixty days from signing.",
        "2. The Buyer shall pay the Price on completion.",
        "2A. Interest accrues at 8%.",
        "3. Governing law is New York.",
    ]


def test_the_reply_can_say_exactly_what_came_out():
    told = " ".join(clean.clean(marked_up()).removed)
    assert "Peter Partner" in told and "Jane Associate" in told
    assert "1 comment" in told
    assert "hidden text" in told
    assert "Acme NDA built from the Beta precedent" in told


def test_comment_parts_leave_no_dangling_reference():
    out = clean.clean(marked_up())
    package = zipfile.ZipFile(BytesIO(out.content))
    assert "word/comments.xml" not in package.namelist()
    assert b"comments" not in package.read("word/_rels/document.xml.rels")
    assert b"comments" not in package.read("[Content_Types].xml")
    ok, reason = redline.verify(out.content)
    assert ok, reason


def test_accepting_a_deleted_paragraph_removes_it_whole():
    d = Document()
    for text in ("1. Keep this.", "2. Strike this clause entirely.", "3. And keep this."):
        d.add_paragraph(text)
    later = Document()
    for text in ("1. Keep this.", "3. And keep this."):
        later.add_paragraph(text)
    compared = compare.compare(save(d), save(later))

    out = clean.clean(compared.content)
    assert text_of(out.content) == ["1. Keep this.", "3. And keep this."]
    # No empty clause left behind where the deleted one was.
    assert len(Document(BytesIO(out.content)).paragraphs) == 2


def test_an_already_clean_document_reports_nothing():
    d = Document()
    d.add_paragraph("Nothing to see here.")
    d.core_properties.author = ""
    d.core_properties.comments = ""
    out = clean.clean(save(d))
    assert not [r for r in out.removed if "Accepted" in r or "comment" in r]


def test_a_tracked_cell_deletion_is_declined_with_a_sentence():
    d = Document()
    table = d.add_table(rows=1, cols=2)
    cell_props = table.cell(0, 0)._tc.get_or_add_tcPr()
    marker = OxmlElement("w:cellDel")
    marker.set(qn("w:id"), "9")
    marker.set(qn("w:author"), "Someone")
    cell_props.append(marker)
    with pytest.raises(clean.CannotClean, match="table"):
        clean.clean(save(d))


def test_not_a_word_file():
    with pytest.raises(clean.CannotClean):
        clean.clean(b"%PDF-1.7 not a zip")
