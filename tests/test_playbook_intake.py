"""The firm's playbook, taught by email.

The session is faked (tests/fake_sessions.py) and so is the memory store, so
these run offline. What is exercised is everything around the model: which
emails are playbook emails and whose, that a report becomes position files in
the lra-playbook format marked as the firm's, that every quote is checked
against the material, that nothing takes effect before "approve playbook",
that reviews then mount the firm's store, and that undo walks back through
the versions to the starters.
"""

from __future__ import annotations

import re
from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace as NS
from unittest.mock import patch

import pytest
from docx import Document

import secondeye.config
from secondeye import handler, managed, memory, playbook
from secondeye.mail.console import ConsoleProvider
from secondeye.models import Attachment, InboundEmail
from secondeye.pipeline import redline
from tests import api_contract as contract
from tests import fake_sessions as fs

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

MEMO_LINES = [
    "Positions memo 2026",
    "Limitation of liability: we ask for a cap of 150% of annual fees.",
    "We will accept a cap of 100% of annual fees if pushed.",
    "We never accept a cap below 50% of annual fees.",
    "Indemnity: the supplier indemnifies for third party IP claims.",
]
PRECEDENT_LINES = [
    "Master Services Agreement",
    ("12.1 Each party's total liability in any Contract Year shall not exceed the Charges "
     "paid in that Contract Year."),
]


class Captured(ConsoleProvider):
    def __init__(self):
        self.sent = []

    def send(self, email_out):
        self.sent.append(email_out)
        return f"agent-msg-{len(self.sent)}"


class FakeStore:
    """client.beta.memory_stores: files by store and path, with deletes."""

    def __init__(self):
        self.stores: dict[str, dict[str, str]] = {}
        self.created: list[dict] = []

    def create_store(self, **kw):
        self.created.append(kw)
        sid = f"memstore_{len(self.created)}"
        self.stores[sid] = {}
        return NS(id=sid, name=kw["name"])

    def create(self, store_id, *, path, content, **kw):
        assert path not in self.stores[store_id]
        self.stores[store_id][path] = content
        return NS(type="memory", id=f"mem|{path}", path=path, memory_store_id=store_id)

    def list(self, store_id, **kw):
        return [NS(type="memory", id=f"mem|{p}", path=p, memory_store_id=store_id)
                for p in self.stores[store_id]]

    def update(self, memory_id, *, memory_store_id, content, **kw):
        path = memory_id.split("|", 1)[1]
        self.stores[memory_store_id][path] = content
        return NS(type="memory", id=memory_id, path=path, memory_store_id=memory_store_id)

    def delete(self, memory_id, *, memory_store_id, **kw):
        del self.stores[memory_store_id][memory_id.split("|", 1)[1]]
        return NS(type="memory_deleted", id=memory_id)


@pytest.fixture
def env(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/p.sqlite3")
    monkeypatch.setenv("MAIL_AGENT_ADDRESS", "review@legal.firm.com")
    monkeypatch.setenv("FIRM_DOMAINS", "firm.com")
    monkeypatch.setenv("ALLOWLIST_CONTACT", "Jim Baker <jim@firm.com>")
    monkeypatch.setenv("PLAYBOOK_ADMINS", "")
    fs.configure(monkeypatch, MANAGED_PLAYBOOK_AGENT_ID="agent_playbook")
    store = FakeStore()
    client = contract.strict(NS(beta=NS(memory_stores=NS(create=store.create_store,
                                                         memories=store))))
    monkeypatch.setattr(secondeye.config, "anthropic_client", lambda: client)
    provider = Captured()
    monkeypatch.setattr(handler, "get_provider", lambda: provider)
    yield NS(provider=provider, store=store)
    secondeye.config.settings.cache_clear()


def docx(lines: list[str]) -> bytes:
    d = Document()
    for line in lines:
        d.add_paragraph(line)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def attachment(name: str, lines: list[str]) -> Attachment:
    raw = docx(lines)
    return Attachment(filename=name, content_type=DOCX, size_bytes=len(raw), content=raw)


def email(body: str, mid: str, subject: str = "Playbook", sender: str = "jim@firm.com",
          attachments: list[Attachment] | None = None, to=("review@legal.firm.com",)):
    return InboundEmail(
        message_id=mid, thread_id=f"t-{mid}", from_address=sender, to=list(to),
        subject=subject, text_body=body, attachments=attachments or [],
        received_at=datetime.now(UTC))


def playbook_email(mid: str = "pb-1") -> InboundEmail:
    return email("Here's our playbook. The memo and our MSA precedent are attached.", mid,
                 attachments=[attachment("Positions memo.docx", MEMO_LINES),
                              attachment("MSA precedent.docx", PRECEDENT_LINES)])


LOL = {
    "clause": "Limitation of liability", "side": "customer",
    "position": "A cap of 150% of annual fees.",
    "fallback": "A cap of 100% of annual fees.",
    "walk_away": "A cap below 50% of annual fees.",
    "watch_for": [],
    "model_wording": ("Each party's total liability in any Contract Year shall not exceed "
                      "the Charges paid in that Contract Year."),
    "sources": [
        {"document": "Positions memo.docx",
         "quote": "we ask for a cap of 150% of annual fees"},
        {"document": "Positions memo.docx",
         "quote": "We will accept a cap of 100% of annual fees if pushed."},
        {"document": "Positions memo.docx",
         "quote": "We never accept a cap below 50% of annual fees."},
        {"document": "MSA precedent.docx",
         "quote": "Each party’s total liability in any Contract Year shall not exceed"},
    ],
}
INDEMNITY = {
    "clause": "Indemnity", "side": "customer",
    "position": "The supplier indemnifies for third-party IP claims.",
    "sources": [{"document": "Positions memo.docx",
                 "quote": "the supplier indemnifies for third party IP claims"}],
}


def run(env, inbound: InboundEmail, *reports: dict) -> list[fs.FakeSessions]:
    """Handle one email with the agent making these record_playbook calls in
    turn, each after the previous one's result."""
    script: list = []
    for i, report in enumerate(reports):
        call = fs.tool_use("record_playbook", **report)
        if i == 0:
            script.append(call)
        else:
            script.append(fs.after_tool_result(call))
    script += [fs.after_tool_result(fs.idle())]
    fake = fs.FakeSessions(script)
    with patch.object(managed, "anthropic_client", lambda: fake.client):
        handler.handle(inbound)
    return [fake]


def reply(env):
    return env.provider.sent[-1]


def store_files(env) -> dict[str, str]:
    sid = playbook.store_id()
    return env.store.stores.get(sid, {}) if sid else {}


# --- which emails are playbook emails -----------------------------------------


@pytest.mark.parametrize("body,subject,expected", [
    ("Here's our playbook.", "", playbook.UPDATE),
    ("Attached are our negotiating positions.", "", playbook.UPDATE),
    ("Update the playbook: we now accept a 2x cap", "", playbook.UPDATE),
    ("Please add this to the playbook: no uncapped indemnities.", "", playbook.UPDATE),
    ("", "Playbook update", playbook.UPDATE),
    ("Thanks", "Re: Playbook update", None),
    ("Looks right to me, will read it tonight.", "Re: Playbook update", None),
    ("approve playbook", "Re: Playbook", playbook.APPROVE),
    ("Approve the playbook.", "", playbook.APPROVE),
    ("undo the last playbook change", "", playbook.UNDO),
    ("discard playbook", "", playbook.DISCARD),
    ("Please review this against our playbook.", "Supply agreement", None),
    ("Is the cap on our playbook?", "", None),
    ("Can you check the indemnity?", "Supply agreement", None),
    ("approve", "", None),
])
def test_a_playbook_email_is_told_from_a_review(env, body, subject, expected):
    assert playbook.intent(email(body, "m", subject=subject)) == expected


@pytest.mark.parametrize("body,subject", [
    # Outlook, Apple Mail and phones type the curly apostrophe.
    ("Here\u2019s our playbook.", ""),
    ("Here\u2018s the firm\u2019s playbook.", ""),
    ("", "Here\u2019s our playbook"),
])
def test_a_curly_apostrophe_is_still_a_playbook_email(env, body, subject):
    assert playbook.intent(email(body, "m", subject=subject)) == playbook.UPDATE


def test_a_playbook_forwarded_to_someone_else_is_not_an_update(env):
    bcc = email("Here's our playbook, welcome aboard.", "m", to=("newjoiner@firm.com",))
    assert playbook.intent(bcc) is None


def test_admins_default_to_the_allowlist_contact_then_the_firm(env, monkeypatch):
    assert playbook.is_admin("Jim@Firm.com")
    assert not playbook.is_admin("amy@firm.com")
    monkeypatch.setenv("ALLOWLIST_CONTACT", "")
    secondeye.config.settings.cache_clear()
    assert playbook.is_admin("amy@firm.com")
    assert not playbook.is_admin("counsel@other.com")
    monkeypatch.setenv("PLAYBOOK_ADMINS", "kl@firm.com")
    secondeye.config.settings.cache_clear()
    assert playbook.is_admin("kl@firm.com") and not playbook.is_admin("amy@firm.com")


# --- a new playbook, drafted and approved --------------------------------------


def test_positions_arrive_as_a_draft_with_a_summary_to_check(env):
    [fake] = run(env, playbook_email(), {
        "summary": "Two positions.", "positions": [LOL, INDEMNITY],
        "unsettled": ["Indemnity: the memo gives no fallback."]})

    created = fake.created[0]
    assert created["agent"] == "agent_playbook"
    [outcome] = created["initial_events"]
    assert outcome["type"] == "user.define_outcome"
    assert "word for word" in outcome["rubric"]["content"]
    assert "Positions memo 2026" in outcome["description"]
    assert {name for name, _ in fake.uploaded} == {"Positions memo.docx", "MSA precedent.docx"}
    assert fake.sent[0][0]["content"][0]["text"] == "Recorded 2 positions."

    out = reply(env)
    first = out.text_body.splitlines()[0]
    assert first == ("Playbook updated: 2 positions from 2 documents "
                     "(Limitation of liability, Indemnity).")
    assert "1 thing I couldn't settle:" in out.text_body
    assert "- Indemnity: the memo gives no fallback." in out.text_body
    assert 'reply "approve playbook"' in out.text_body
    [summary] = out.attachments
    assert summary.filename == playbook.SUMMARY_NAME
    assert redline.verify(summary.content)[0]
    text = "\n".join(p.text for p in Document(BytesIO(summary.content)).paragraphs)
    assert "Limitation of liability (New)" in text and "150% of annual fees" in text

    # A draft only: reviews still use the starters.
    assert playbook.live() is None
    assert "playbook" not in memory.stores_for("amy@firm.com")
    assert store_files(env) == {}


def test_approval_publishes_the_firm_files_that_reviews_then_mount(env):
    run(env, playbook_email(), {"summary": "s", "positions": [LOL, INDEMNITY]})
    handler.handle(email("approve playbook", "ok-1", subject="Re: Playbook"))

    assert reply(env).text_body.startswith("Approved. From the next review")
    files = store_files(env)
    assert set(files) == {"/README.md", "/positions/limitation-of-liability.md",
                          "/positions/indemnity.md"}
    assert memory.stores_for("amy@firm.com")["playbook"] == playbook.store_id()
    assert env.store.created[0]["name"] == playbook.STORE_NAME
    labels = managed.memory_resources({"playbook": "memstore_1"})
    assert "replace" in labels[0]["instructions"] and labels[0]["access"] == "read_only"


def test_a_position_file_is_in_the_playbook_format_marked_as_the_firms(env):
    rendered = playbook.render({**LOL, "updated": "2026-09-28"})
    front, body = rendered.split("---\n", 2)[1:]
    meta = dict(line.split(": ", 1) for line in front.strip().splitlines())
    assert meta["status"] == "firm"
    assert meta["clause"] == "Limitation of liability"
    assert meta["source"] == "Positions memo.docx; MSA precedent.docx"
    starter = (playbook.STARTERS / "limitation-of-liability.md").read_text()
    starter_sections = re.findall(r"^## .+$", starter, re.MULTILINE)
    assert re.findall(r"^## .+$", body, re.MULTILINE) == starter_sections + ["## Source"]
    assert '- MSA precedent.docx: "Each party’s total liability' in body
    # A part the material did not give is said to be missing, never filled.
    sparse = playbook.render(INDEMNITY)
    assert "## Fallback\nNot stated in the firm's documents." in sparse
    assert "## Model wording (fallback)\nNone in the firm's documents." in sparse


# --- nothing invented ------------------------------------------------------------


def test_a_quote_not_in_the_material_is_handed_back_once(env):
    invented = {**INDEMNITY, "sources": [{"document": "Positions memo.docx",
                                          "quote": "the supplier gives an uncapped indemnity"}]}
    [fake] = run(env, playbook_email(), {"summary": "s", "positions": [invented]},
                 {"summary": "s", "positions": [INDEMNITY]})

    first_result = fake.sent[0][0]["content"][0]["text"]
    assert first_result.startswith("Not recorded.")
    assert "the supplier gives an uncapped indemnity" in first_result
    assert fake.sent[1][0]["content"][0]["text"] == "Recorded 1 position."
    assert "Check before approving" not in reply(env).text_body


def test_a_quote_still_missing_stands_but_is_named_for_the_partner(env):
    invented = {**INDEMNITY, "model_wording": "The Supplier shall indemnify on demand.",
                "sources": [{"document": "Positions memo.docx", "quote": "made up words"}]}
    run(env, playbook_email(), {"summary": "s", "positions": [invented]},
        {"summary": "s", "positions": [invented]})

    body = reply(env).text_body
    assert "Check before approving: Indemnity: \"made up words\" is not word for word" in body
    assert "the model wording is not word for word" in body
    assert playbook.pending("jim@firm.com").positions["indemnity"]["sources"]


def test_a_quote_found_in_another_document_is_cited_to_that_one(env):
    misfiled = {**INDEMNITY, "sources": [{"document": "MSA precedent.docx",
                                          "quote": "the supplier indemnifies for third "
                                                   "party IP claims"}]}
    run(env, playbook_email(), {"summary": "s", "positions": [misfiled]})
    source = playbook.pending("jim@firm.com").positions["indemnity"]["sources"][0]
    assert source["document"] == "Positions memo.docx"


# --- updates, history and undo ---------------------------------------------------


def two_x() -> dict:
    return {**LOL, "fallback": "A cap of 200% of annual fees.",
            "sources": LOL["sources"][:1] + [
                {"document": "email", "quote": "we now accept a 2x cap"}]}


def test_a_one_line_update_changes_one_position_and_keeps_the_rest(env):
    run(env, playbook_email(), {"summary": "s", "positions": [LOL, INDEMNITY]})
    handler.handle(email("approve playbook", "ok-1"))
    [fake] = run(env, email("Update the playbook: we now accept a 2x cap", "pb-2"),
                 {"summary": "s", "positions": [two_x()]})

    # The agent was shown the current playbook to update.
    assert "<current-playbook-" in fake.created[0]["initial_events"][0]["description"]
    body = reply(env).text_body
    assert body.splitlines()[0] == ("Playbook updated: 1 position from 1 document and your "
                                    "email (Limitation of liability).")
    assert "The other position stays as it was." in body
    draft = playbook.pending("jim@firm.com")
    assert set(draft.positions) == {"limitation-of-liability", "indemnity"}
    # Not live yet: the store still says 100%.
    assert "100% of annual fees" in store_files(env)["/positions/limitation-of-liability.md"]

    handler.handle(email("approve playbook", "ok-2"))
    assert "200% of annual fees" in store_files(env)["/positions/limitation-of-liability.md"]


def test_undo_walks_back_through_the_versions_to_the_starters(env):
    run(env, playbook_email(), {"summary": "s", "positions": [LOL, INDEMNITY]})
    handler.handle(email("approve playbook", "ok-1"))
    run(env, email("Update the playbook: we now accept a 2x cap", "pb-2"),
        {"summary": "s", "positions": [two_x()], "removed": ["Indemnity"]})
    handler.handle(email("approve playbook", "ok-2"))
    assert "/positions/indemnity.md" not in store_files(env)

    handler.handle(email("undo the last playbook change", "u-1"))
    assert reply(env).text_body.startswith("Undone. The playbook is back to the version "
                                           "approved on")
    files = store_files(env)
    assert "/positions/indemnity.md" in files
    assert "100% of annual fees" in files["/positions/limitation-of-liability.md"]

    handler.handle(email("undo the last playbook change", "u-2"))
    assert reply(env).text_body == ("Undone. Reviews are back to the starter positions.\n" + "\nPrivileged & Confidential — Attorney Work Product\n")
    assert store_files(env) == {}
    assert "playbook" not in memory.stores_for("amy@firm.com")

    statuses = [v.status for v in playbook.history()]
    assert statuses == ["undone", "undone"]


def test_approving_a_draft_built_on_a_playbook_since_changed_is_refused(env, monkeypatch):
    monkeypatch.setenv("PLAYBOOK_ADMINS", "jim@firm.com,kl@firm.com")
    secondeye.config.settings.cache_clear()
    run(env, playbook_email("a"), {"summary": "s", "positions": [LOL]})
    run(env, email("Here's our playbook", "b", sender="kl@firm.com",
                   attachments=[attachment("Positions memo.docx", MEMO_LINES)]),
        {"summary": "s", "positions": [INDEMNITY]})
    handler.handle(email("approve playbook", "ok-kl", sender="kl@firm.com"))
    handler.handle(email("approve playbook", "ok-jim"))

    assert "someone else's change" in reply(env).text_body
    assert set(playbook.live().positions) == {"indemnity"}


# --- who may, and when it cannot run ---------------------------------------------


def test_only_an_admin_can_change_the_playbook(env):
    handler.handle(email("Here's our playbook.", "x-1", sender="amy@firm.com",
                         attachments=[attachment("Positions memo.docx", MEMO_LINES)]))
    body = reply(env).text_body
    assert body.startswith("Only the firm's playbook admins (jim@firm.com)")
    assert playbook.history() == []

    handler.handle(email("undo the last playbook change", "x-2", sender="amy@firm.com"))
    assert reply(env).text_body.startswith("Only the firm's playbook admins")


def test_approving_with_nothing_waiting_says_so(env):
    handler.handle(email("approve playbook", "ok-0"))
    assert reply(env).text_body.startswith("There is no playbook change of yours waiting")


def test_without_the_playbook_agent_nothing_changes_and_it_says_why(env, monkeypatch):
    monkeypatch.setenv("MANAGED_PLAYBOOK_AGENT_ID", "")
    secondeye.config.settings.cache_clear()
    handler.handle(playbook_email())
    assert "isn't set up yet" in reply(env).text_body
    assert playbook.history() == []


def test_material_with_no_positions_leaves_the_playbook_alone(env):
    run(env, playbook_email(), {"summary": "Nothing here.", "positions": [],
                                "unsettled": ["The memo is a cover note only."]})
    body = reply(env).text_body
    assert body.startswith("I couldn't find any negotiating positions in that")
    assert "- The memo is a cover note only." in body
    assert playbook.history() == []


def test_the_playbook_agent_is_applied_with_its_own_tool(env):
    body = managed.agent_body(managed.load_manifest(managed.AGENTS / "playbook.agent.yaml"),
                              playbook.CUSTOM_TOOLS)
    assert body["system"] == playbook.SYSTEM
    names = [t.get("name") for t in body["tools"]]
    assert "record_playbook" in names and "make_changes" not in names
