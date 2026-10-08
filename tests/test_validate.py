"""The OOXML validator.

Its job is to refuse documents Word would refuse or silently mangle. A validator
that never rejects anything is worthless, so most of these tests build
deliberately broken revision markup and assert it is caught.
"""

from __future__ import annotations

import zipfile
from io import BytesIO

import pytest
from docx import Document

from secondeye.pipeline import redline
from secondeye.pipeline.ooxml import Revision, RevisionWriter
from secondeye.pipeline.validate import validate


def save(doc) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def simple(text="The term is thirty days.") -> bytes:
    d = Document()
    d.add_paragraph(text)
    return save(d)


def edited() -> bytes:
    doc = Document(BytesIO(simple()))
    RevisionWriter(doc, "Legal Review Agent").apply(
        Revision("thirty", "sixty", "Legal Review Agent")
    )
    return save(doc)


def rewrite_xml(content: bytes, transform) -> bytes:
    """Rebuild a package with document.xml passed through `transform`."""
    src = zipfile.ZipFile(BytesIO(content))
    items = {n: src.read(n) for n in src.namelist()}
    items["word/document.xml"] = transform(
        items["word/document.xml"].decode()
    ).encode()
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in items.items():
            out.writestr(name, data)
    return buf.getvalue()


# --- valid documents must pass -------------------------------------------

def test_an_untouched_document_is_valid():
    assert validate(simple()).ok


def test_our_own_output_is_valid():
    report = validate(edited())
    assert report.ok, str(report)
    assert report.revisions == 2                 # one insertion, one deletion
    assert report.authors == {"Legal Review Agent"}


def test_a_document_with_several_edits_stays_valid():
    doc = Document(BytesIO(simple("The Reciever pays thirty (13) days after closing.")))
    writer = RevisionWriter(doc, "Legal Review Agent")
    writer.apply(Revision("Reciever", "Recipient", "Legal Review Agent"))
    writer.apply(Revision("thirty (13) days", "thirty (30) days", "Legal Review Agent"))
    assert validate(save(doc)).ok


# --- the rules Word enforces ---------------------------------------------

def test_deleted_text_using_w_t_is_rejected():
    """The single most common way to write tracked changes wrongly. Word either
    refuses the file or silently drops the text."""
    broken = rewrite_xml(edited(), lambda x: x.replace("<w:delText", "<w:t").replace(
        "</w:delText>", "</w:t>"))
    report = validate(broken)
    assert not report.ok
    assert any("delText" in e for e in report.errors)


def test_duplicate_revision_ids_are_rejected():
    broken = rewrite_xml(edited(), lambda x: __import__("re").sub(
        r'w:id="\d+"', 'w:id="7"', x))
    report = validate(broken)
    assert not report.ok
    assert any("duplicate revision id" in e for e in report.errors)


def test_a_revision_without_an_author_is_rejected():
    broken = rewrite_xml(edited(), lambda x: __import__("re").sub(
        r'\sw:author="[^"]*"', "", x))
    report = validate(broken)
    assert not report.ok
    assert any("author" in e for e in report.errors)


def test_a_malformed_date_is_rejected():
    broken = rewrite_xml(edited(), lambda x: __import__("re").sub(
        r'w:date="[^"]*"', 'w:date="last Tuesday"', x))
    report = validate(broken)
    assert not report.ok
    assert any("w:date" in e for e in report.errors)


def test_a_dangling_relationship_reference_is_rejected():
    broken = rewrite_xml(
        edited(),
        lambda x: x.replace(
            "<w:p>",
            '<w:p><w:hyperlink r:id="rIdDoesNotExist"><w:r><w:t>link</w:t></w:r>'
            "</w:hyperlink>",
            1,
        ),
    )
    report = validate(broken)
    assert not report.ok
    assert any("rIdDoesNotExist" in e for e in report.errors)


def test_malformed_xml_is_rejected():
    broken = rewrite_xml(edited(), lambda x: x.replace("</w:body>", ""))
    report = validate(broken)
    assert not report.ok
    assert any("well-formed" in e for e in report.errors)


def test_a_missing_document_part_is_rejected():
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
    report = validate(buf.getvalue())
    assert not report.ok
    assert any("word/document.xml" in e for e in report.errors)


def test_something_that_is_not_a_zip_is_rejected():
    assert not validate(b"this is not a docx").ok


# --- empty deleted text means the original content was lost --------------

def test_empty_deleted_text_is_warned_about():
    broken = rewrite_xml(
        edited(),
        lambda x: __import__("re").sub(
            r"<w:delText([^>]*)>[^<]*</w:delText>", r"<w:delText\1></w:delText>", x
        ),
    )
    report = validate(broken)
    assert any("may have been lost" in w for w in report.warnings)


# --- the gate actually uses it -------------------------------------------

def test_verify_refuses_a_document_the_validator_rejects():
    broken = rewrite_xml(edited(), lambda x: x.replace("<w:delText", "<w:t").replace(
        "</w:delText>", "</w:t>"))
    ok, reason = redline.verify(broken, expect_revisions=True)
    assert not ok
    assert "delText" in reason


def test_another_authors_revisions_are_reported_not_rejected():
    """A document already under revision by a colleague is normal, not broken."""
    two_authors = rewrite_xml(
        edited(),
        lambda x: x.replace(
            "<w:p>",
            '<w:p><w:ins w:id="9001" w:author="A. Associate" '
            'w:date="2026-01-01T00:00:00Z"><w:r><w:t>prior</w:t></w:r></w:ins>',
            1,
        ),
    )
    report = validate(two_authors)
    assert report.ok, str(report)
    assert "A. Associate" in report.authors
    assert "Legal Review Agent" in report.authors


@pytest.mark.parametrize("text", [
    "A paragraph with a table reference.",
    "Text with “curly quotes” and an em dash — here.",
    "هذه اتفاقية سرية بين الطرفين.",
])
def test_unusual_content_still_validates(text):
    assert validate(simple(text)).ok
