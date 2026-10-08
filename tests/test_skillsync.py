"""The custom skill: what is uploaded has to run on its own.

The container has no network, no configuration and possibly no pydantic. So the
built bundle is executed here as the container would execute it: from its own
folder, by a fresh interpreter, once with pydantic and once with the stand-in.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from secondeye import managed, skillsync
from secondeye.pipeline.ooxml import Revision, RevisionWriter


def docx(path: Path, paragraphs, author: str = "") -> Path:
    d = Document()
    for text in paragraphs:
        d.add_paragraph(text)
    if author:
        d.core_properties.author = author
        RevisionWriter(d, author).apply(Revision("thirty", "sixty", author))
    d.save(path)
    return path


@pytest.fixture(scope="module")
def bundle(tmp_path_factory) -> Path:
    return skillsync.build(tmp_path_factory.mktemp("skill"))


def run(bundle: Path, script: str, *args: str, shim: bool) -> dict:
    env = {k: v for k, v in os.environ.items() if k != "PYTHONPATH"}
    if shim:
        env["SECOND_EYE_FORCE_SHIM"] = "1"
    done = subprocess.run(
        [sys.executable, str(bundle / "scripts" / script), *args],
        capture_output=True, text=True, env=env, cwd=bundle, timeout=60, check=False,
    )
    assert done.stdout, done.stderr
    return json.loads(done.stdout)


def test_the_bundle_has_the_shape_the_skills_api_requires(bundle):
    assert bundle.name == "lra-document-tools"
    head = (bundle / "SKILL.md").read_text().split("---")[1]
    name = next(line for line in head.splitlines() if line.startswith("name:"))
    description = next(line for line in head.splitlines() if line.startswith("description:"))
    assert name.split(":", 1)[1].strip() == "lra-document-tools"
    assert 0 < len(description.split(":", 1)[1].strip()) <= 1024
    assert "claude" not in name and "anthropic" not in name
    size = sum(f.stat().st_size for f in bundle.rglob("*") if f.is_file())
    assert size < 30 * 1024 * 1024


def test_the_bundle_carries_the_real_modules_not_a_fork(bundle):
    for relative in skillsync.BUNDLED:
        assert (bundle / "lib" / "secondeye" / relative).read_bytes() == (
            skillsync.PACKAGE / relative).read_bytes()


def test_nothing_bundled_reaches_for_configuration_at_import(bundle):
    for relative in skillsync.BUNDLED:
        top_level = [
            line for line in (bundle / "lib" / "secondeye" / relative).read_text().splitlines()
            if line.startswith(("import ", "from "))
        ]
        assert not [line for line in top_level if "secondeye.config" in line
                    or "secondeye.managed" in line or "anthropic" in line], relative


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_compare_runs_from_the_bundle(bundle, tmp_path, shim):
    earlier = docx(tmp_path / "v1.docx", ["1. The term is thirty days.", "2. Law: New York."])
    later = docx(tmp_path / "v2.docx", ["1. The term is sixty days.", "2. Law: New York."])
    out = tmp_path / "out" / "v2 (comparison).docx"

    result = run(bundle, "compare.py", str(earlier), str(later), "--out", str(out), shim=shim)
    assert result["proven"] is True
    assert [c["kind"] for c in result["changes"]] == ["replaced"]
    assert out.exists() and result["file"] == str(out)


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_clean_and_check_run_from_the_bundle(bundle, tmp_path, shim):
    dirty = docx(tmp_path / "nda.docx",
                 ["1. The term is thirty days from [DATE]."], author="Jane Associate")
    out = tmp_path / "nda (clean).docx"

    cleaned = run(bundle, "clean.py", str(dirty), "--out", str(out), shim=shim)
    assert any("Jane Associate" in line for line in cleaned["removed"])
    assert "sixty days" in Document(BytesIO(out.read_bytes())).paragraphs[0].text

    checked = run(bundle, "check.py", str(dirty), shim=shim)
    categories = {f["category"] for f in checked["findings"]}
    assert "placeholder" in categories and "leftovers" in categories


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_the_deal_math_scripts_run_from_the_bundle(bundle, tmp_path, shim):
    """Figures, timings and party names, as the review agent runs them."""
    from evals import realistic

    v3 = tmp_path / "Falcon SPA v3.docx"
    v2 = tmp_path / "Falcon SPA v2.docx"
    v3.write_bytes(realistic.falcon_spa(3, defects=True))
    v2.write_bytes(realistic.falcon_spa(2))

    figures = run(bundle, "deal_math.py", str(v3), shim=shim)
    assert {f["category"] for f in figures["findings"]} == {"arithmetic"}
    assert figures["splits"][0]["sum"] == 12_750_000
    assert any(t["where"] == "Schedule 2" for t in figures["tables"])

    timing = run(bundle, "timeline.py", str(v3), shim=shim)
    assert {f["kind"] for f in timing["flags"]} == {"claims-before-release",
                                                   "cure-longer-than-notice"}

    parties = run(bundle, "parties.py", str(v3), "--previous", str(v2), shim=shim)
    assert [p["role"] for p in parties["parties"]] == ["Seller", "Buyer"]
    assert len(parties["findings"]) == len(parties["since_previous"]) == 2

    assert "error" in run(bundle, "deal_math.py", str(tmp_path / "missing.docx"), shim=shim)


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_the_writer_runs_from_the_bundle_and_verifies_its_own_output(bundle, tmp_path, shim):
    """The redline is written in the container now (DECISIONS 28), so the
    writer has to stand on its own there like the other scripts."""
    from secondeye.pipeline import redline

    source = docx(tmp_path / "spa.docx", ["1. The term is thirty (13) days.",
                                          "2. Law: New York. Law: New York."])
    findings = tmp_path / "findings.json"
    findings.write_text(json.dumps([
        {"severity": "blocker", "category": "amount", "title": "Mismatch",
         "explanation": "The words say thirty and the figure says 13. Clause 4 "
                        "counts thirty days.",
         "anchor": "thirty (13) days", "suggested_text": "thirty (30) days",
         "auto_apply": True},
        {"severity": "substantive", "category": "law", "title": "Duplicate",
         "explanation": "", "anchor": "Law: New York.", "suggested_text": "Law: Delaware.",
         "auto_apply": True},
        {"severity": "substantive", "category": "indemnity", "title": "Cap?",
         "explanation": "", "anchor": "1. The term", "question": "Is the cap intended?"},
    ]))
    out = tmp_path / "outputs" / "spa (redline).docx"

    result = run(bundle, "redline.py", str(source), "--findings", str(findings),
                 "--out", str(out), shim=shim)
    assert result["verified"] is True and result["file"] == str(out)
    assert result["applied"] == ["Mismatch"]
    assert "Duplicate" in result["skipped"], "an ambiguous anchor must be refused"
    assert result["commented"] == ["Cap?"]
    assert result["explained"] == ["Mismatch"], "the reason goes beside the change"
    ok, why = redline.verify(out.read_bytes(), expect_revisions=True)
    assert ok, why
    # The lawyer's file is untouched.
    assert "thirty (13)" in Document(BytesIO(source.read_bytes())).paragraphs[0].text


def test_the_writer_refuses_findings_it_cannot_read(bundle, tmp_path):
    source = docx(tmp_path / "a.docx", ["1. Text."])
    findings = tmp_path / "f.json"
    findings.write_text(json.dumps([{"severity": "critical", "title": "x"}]))
    result = run(bundle, "redline.py", str(source), "--findings", str(findings),
                 "--out", str(tmp_path / "o.docx"), shim=False)
    assert "could not be read" in result["error"]
    assert not (tmp_path / "o.docx").exists()


def test_a_failure_is_json_not_a_traceback(bundle, tmp_path):
    bad = tmp_path / "not-word.docx"
    bad.write_bytes(b"plain text")
    result = run(bundle, "clean.py", str(bad), "--out", str(tmp_path / "x.docx"), shim=False)
    assert "error" in result


def test_the_custom_skill_is_loaded_only_once_it_has_an_id(monkeypatch):
    from secondeye.config import settings

    monkeypatch.setenv("SANDBOX_SKILL_ID", "")
    monkeypatch.setenv("SANDBOX_PLAYBOOK_SKILL_ID", "")
    settings.cache_clear()
    assert {s["type"] for s in managed.skill_refs()} == {"anthropic"}

    monkeypatch.setenv("SANDBOX_SKILL_ID", "skill_01Example")
    settings.cache_clear()
    try:
        assert managed.skill_refs()[-1] == {
            "type": "custom", "skill_id": "skill_01Example", "version": "latest"}
    finally:
        settings.cache_clear()


# --- the playbook skill ------------------------------------------------------------


def test_the_playbook_skill_builds_without_code_and_has_a_file_per_clause(tmp_path):
    from secondeye import skillsync

    bundle = skillsync.build(tmp_path, "lra-playbook")
    assert (bundle / "SKILL.md").exists()
    assert not (bundle / "lib").exists(), "the playbook carries instructions, not modules"
    positions = sorted(p.name for p in (bundle / "positions").glob("*.md") if p.name != "README.md")
    assert "limitation-of-liability.md" in positions and len(positions) >= 10
    for p in positions:
        text = (bundle / "positions" / p).read_text()
        assert "status: starter" in text, f"{p} must say it is a starter until the firm replaces it"
        for heading in ("## Position", "## Fallback", "## Walk-away", "## Watch for",
                        "## Model wording"):
            assert heading in text, f"{p} lacks {heading}"


def test_the_playbook_skill_is_loaded_alongside_the_document_tools(monkeypatch):
    from secondeye.config import settings

    monkeypatch.setenv("SANDBOX_SKILL_ID", "skill_tools")
    monkeypatch.setenv("SANDBOX_PLAYBOOK_SKILL_ID", "skill_playbook")
    settings.cache_clear()
    try:
        ids = [s["skill_id"] for s in managed.skill_refs()]
        assert ids == ["docx", "xlsx", "pdf", "pptx", "skill_tools", "skill_playbook"]
    finally:
        settings.cache_clear()


def test_an_unknown_skill_name_is_refused(tmp_path):
    import pytest

    from secondeye import skillsync

    with pytest.raises(KeyError, match="no such skill"):
        skillsync.build(tmp_path, "lra-nonsense")


def test_the_key_terms_skill_builds_and_quotes_rather_than_infers(tmp_path):
    from secondeye import skillsync

    bundle = skillsync.build(tmp_path, "lra-key-terms")
    text = (bundle / "SKILL.md").read_text()
    assert not (bundle / "lib").exists()
    assert "Not addressed" in text and "never infer" in text.lower()
    for term in ("Liability cap", "Change of control", "Governing law", "Signatories"):
        assert term in text


# --- the findings, recorded in the sandbox (docs/migration.md, phase 1) ------------


def _json_type(annotation):
    """A pydantic field's annotation as the JSON Schema type it should have,
    with Optional unwrapped: null is how a finding says it has no value."""
    import typing
    from enum import Enum

    args = typing.get_args(annotation)
    if type(None) in args:
        annotation = next(a for a in args if a is not type(None))
    origin = typing.get_origin(annotation)
    if annotation is str:
        return {"type": "string"}
    if annotation is float:
        return {"type": "number"}
    if annotation is bool:
        return {"type": "boolean"}
    if isinstance(annotation, type) and issubclass(annotation, Enum):
        return {"type": "string", "enum": [e.value for e in annotation]}
    if origin is typing.Literal:
        return {"type": "string", "enum": list(typing.get_args(annotation))}
    if origin is list:
        return {"type": "array", "items": _json_type(typing.get_args(annotation)[0])}
    raise AssertionError(f"no JSON type for {annotation!r}")


def _schema() -> dict:
    return json.loads((skillsync.SKILL_SOURCE / "findings.schema.json").read_text())


def test_the_findings_schema_is_models_finding_field_for_field():
    """The schema the sandbox checks against and the model the server builds
    must agree, or a finding the agent was told is recorded fails on arrival,
    or one the server would take is refused in the sandbox."""
    from secondeye.models import Finding

    items = _schema()["properties"]["findings"]["items"]
    assert set(items["properties"]) == set(Finding.model_fields)
    assert set(items["required"]) == {name for name, field in Finding.model_fields.items()
                                      if field.is_required()}
    assert items["additionalProperties"] is False
    for name, field in Finding.model_fields.items():
        prop = {k: v for k, v in items["properties"][name].items() if k != "description"}
        assert prop == _json_type(field.annotation), name


def test_the_host_reads_findings_with_the_schema_the_skill_ships():
    """One schema: the sandbox validates with it, the finish checks with it,
    and it is closed, because a list of free-form objects is where a bad
    severity gets through."""
    from secondeye.pipeline import review

    assert review.findings_schema() == _schema()
    items = _schema()["properties"]["findings"]["items"]
    assert items["additionalProperties"] is False
    assert set(items["required"]) >= {"severity", "anchor", "title"}


def _draft(tmp_path: Path, report, name: str = "draft.json") -> Path:
    path = tmp_path / name
    path.write_text(json.dumps(report))
    return path


GOOD = {"severity": "blocker", "category": "amount", "title": "Mismatch",
        "explanation": "Words and figure disagree.", "anchor": "thirty (13) days",
        "suggested_text": "thirty (30) days", "auto_apply": True, "question": None}


def _run_raw(bundle: Path, *args: str, shim: bool) -> subprocess.CompletedProcess:
    env = {k: v for k, v in os.environ.items() if k != "PYTHONPATH"}
    if shim:
        env["SECOND_EYE_FORCE_SHIM"] = "1"
    return subprocess.run([sys.executable, str(bundle / "scripts" / "validate_findings.py"),
                           *args], capture_output=True, text=True, env=env, cwd=bundle,
                          timeout=60, check=False)


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_the_validator_records_findings_that_all_pass(bundle, tmp_path, shim):
    out = tmp_path / "outputs" / "findings.json"
    draft = _draft(tmp_path, {"summary": "", "findings": [GOOD]})
    result = run(bundle, "validate_findings.py", str(draft), "--out", str(out), shim=shim)
    assert result["recorded"] == 1 and result["file"] == str(out)
    written = json.loads(out.read_text())
    assert written["findings"][0]["anchor"] == "thirty (13) days"
    assert "question" not in written["findings"][0], "a null is an absent field"


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_the_validator_records_nothing_when_one_finding_fails(bundle, tmp_path, shim):
    """The stand-in pydantic validates nothing, so the check cannot lean on it:
    a bad severity has to be refused in a container without pydantic too."""
    out = tmp_path / "outputs" / "findings.json"
    bad = dict(GOOD, severity="critical", title="Wrong name")
    missing = {k: v for k, v in GOOD.items() if k != "anchor"}
    draft = _draft(tmp_path, {"summary": "", "findings": [GOOD, bad, missing]})
    done = _run_raw(bundle, str(draft), "--out", str(out), shim=shim)
    assert done.returncode != 0
    error = json.loads(done.stdout)["error"]
    assert error.startswith("NOTHING WAS RECORDED")
    assert "#2 (Wrong name)" in error and "critical" in error
    assert "#3 (Mismatch): anchor is missing" in error
    assert "validate_findings.py" in error, "it says how to try again"
    assert not out.exists(), "nothing is written unless every finding passes"


def test_a_refused_draft_leaves_the_earlier_good_one_in_place(bundle, tmp_path):
    out = tmp_path / "findings.json"
    run(bundle, "validate_findings.py",
        str(_draft(tmp_path, {"summary": "", "findings": [GOOD]})), "--out", str(out),
        shim=False)
    before = out.read_text()
    done = _run_raw(bundle, str(_draft(tmp_path, {"summary": "", "findings": [{"title": "x"}]},
                                       "second.json")), "--out", str(out), shim=False)
    assert done.returncode != 0
    assert out.read_text() == before


def test_the_validator_refuses_what_is_not_json(bundle, tmp_path):
    draft = tmp_path / "draft.json"
    draft.write_text("{not json")
    done = _run_raw(bundle, str(draft), "--out", str(tmp_path / "f.json"), shim=False)
    assert json.loads(done.stdout)["error"].startswith("NOTHING WAS RECORDED")


def test_the_writer_takes_the_recorded_file_as_it_stands(bundle, tmp_path):
    """The prompt tells the agent to run redline.py on findings.json, so the
    file the validator writes must be one the writer reads."""
    source = docx(tmp_path / "spa.docx", ["1. The term is thirty (13) days."])
    recorded = tmp_path / "findings.json"
    run(bundle, "validate_findings.py",
        str(_draft(tmp_path, {"summary": "", "findings": [GOOD]})), "--out", str(recorded),
        shim=False)
    result = run(bundle, "redline.py", str(source), "--findings", str(recorded),
                 "--out", str(tmp_path / "spa (redline).docx"), shim=False)
    assert result["applied"] == ["Mismatch"] and result["verified"] is True


def test_the_server_and_the_sandbox_refuse_in_the_same_words():
    from secondeye.pipeline import review

    answer = review._Report().record("", [dict(GOOD, severity="critical")])
    assert answer.startswith("NOTHING WAS RECORDED")
    assert "run validate_findings.py on it with the COMPLETE list" in answer


# --- the closing skill -------------------------------------------------------------


@pytest.mark.parametrize("shim", [False, True], ids=["pydantic", "stand-in"])
def test_the_closing_skill_validates_its_own_state_from_the_bundle(tmp_path, shim):
    """lra-closing carries code, so it gets its modules and the stand-in
    pydantic like the document tools, and its validator refuses an update
    before anything is written for the host."""
    bundle = skillsync.build(tmp_path / "b", "lra-closing")
    assert (bundle / "shims" / "pydantic" / "__init__.py").exists()
    for relative in skillsync.CLOSING_BUNDLED:
        assert (bundle / "lib" / "secondeye" / relative).exists(), relative
    update = tmp_path / "update.json"
    update.write_text(json.dumps({"schema": "lra-closing-update/1", "message": ["x"],
                                  "attach": ["missing.pdf"], "state": {}}))
    out = tmp_path / "outputs"
    out.mkdir()
    result = run(bundle, "validate_state.py", "--update", str(update),
                 "--previous", str(tmp_path / "none.json"), "--outputs", str(out),
                 "--held", str(tmp_path), "--write", str(out / "closing.json"), shim=shim)
    assert "NOTHING WAS WRITTEN" in result["error"]
    assert "attach: 'missing.pdf' is not in the outputs" in result["error"]
    assert not (out / "closing.json").exists()
